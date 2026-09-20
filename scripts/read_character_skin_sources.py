"""Read the approved workbook source rows for the server-side importer.

This helper is intentionally read-only. It emits a compact JSON document to
stdout; it never writes the workbook, public data, or any database state.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")


CHARACTER_COLUMNS = [
    "character_id",
    "name_cn",
    "name_vi",
    "fullname_cn",
    "fullname_vi",
    "nickname_vi",
    "tags_cn",
    "tags_vi",
]

SKIN_COLUMNS = [
    "skin_id",
    "character_id",
    "skin_name_cn",
    "skin_name_vi",
    "desc_cn",
    "desc_vi",
    "obtain_cn",
    "obtain_vi",
    "is_base_skin",
    "skin_type",
    "unlock_date",
    "price",
    "currency",
    "is_high_skin",
    "skin_rare",
    "cv_name",
    "drawing_path",
    "card_path",
    "series_id",
    "series_name_cn",
    "series_name_vi",
    "avatar_path",
    "item_id",
    "goods_id",
    "discount_goods_id",
    "discount_price",
    "discount_start",
    "discount_end",
]


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def rows_for(sheet, columns):
    header = [clean(cell.value) for cell in sheet[1]]
    missing = [column for column in columns if column not in header]
    if missing:
        raise ValueError(f"{sheet.title} missing required columns: {', '.join(missing)}")
    positions = {column: header.index(column) for column in columns}
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        record = {column: clean(values[index]) for column, index in positions.items()}
        if record[columns[0]] is not None:
            rows.append(record)
    ids = [record[columns[0]] for record in rows]
    if len(ids) != len(set(ids)):
        raise ValueError(f"{sheet.title} contains duplicate {columns[0]} values")
    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    args = parser.parse_args()
    workbook_path = args.workbook.resolve()
    if workbook_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("workbook must be .xlsx or .xlsm")
    if not workbook_path.is_file():
        raise FileNotFoundError(workbook_path)

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        required_sheets = {"CHARACTER", "SKIN"}
        missing = required_sheets - set(workbook.sheetnames)
        if missing:
            raise ValueError(f"workbook missing sheets: {', '.join(sorted(missing))}")
        payload = {
            "workbook": str(workbook_path),
            "characters": rows_for(workbook["CHARACTER"], CHARACTER_COLUMNS),
            "skins": rows_for(workbook["SKIN"], SKIN_COLUMNS),
        }
    finally:
        workbook.close()

    json.dump(payload, sys.stdout, ensure_ascii=False, separators=(",", ":"))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"SOURCE_READER_ERROR: {exc}", file=sys.stderr)
        raise
