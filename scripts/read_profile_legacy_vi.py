# scripts/read_profile_legacy_vi.py
"""Read the workbook PROFILE sheet (read-only) and print rows with a VI cell as JSON.

Never writes the workbook. Used only by import-character-profile.mjs --seed-legacy-workbook.
"""
from __future__ import annotations

import argparse
import json
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
COLUMNS = ["profile_id", "character_id", "category", "text_cn", "text_vi"]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    args = parser.parse_args()
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    rows = workbook["PROFILE"].iter_rows(values_only=True)
    header = list(next(rows))
    missing = [name for name in COLUMNS if name not in header]
    if missing:
        raise SystemExit(f"PROFILE sheet is missing columns: {missing}")
    out = []
    for values in rows:
        record = dict(zip(header, values))
        if str(record.get("text_vi") or "").strip():
            out.append({name: (None if record.get(name) is None else str(record[name])) for name in COLUMNS})
    json.dump(out, sys.stdout, ensure_ascii=False)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
