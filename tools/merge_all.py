import os
import sys
import shutil
from datetime import datetime
import pandas as pd
import openpyxl

# Xác định đường dẫn thư mục gốc
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)

MASTER_FILE = os.path.join(PROJECT_ROOT, "localization", "localization_master.xlsx")
BACKUP_DIR = os.path.join(PROJECT_ROOT, "localization", "backups")

# 1. Khóa chính (Primary Key) để merge từng dòng cho dữ liệu nhân vật
PRIMARY_KEYS = {
    "CHARACTER": "character_id",
    "SKILL": "skill_id",
    "BUFF_STATUS": "buff_id",
    "ZHIZHI": ("character_id", "star"),
    "HUANZHANG": "brilliant_id",
    "PROFILE": "profile_id",
    "SKIN": "skin_id"
}

# 2. Bảng mapping tên sheet linh hoạt (hỗ trợ cả chữ hoa lẫn chữ thường số nhiều)
BATCH_SHEET_MAP = {
    "characters": "CHARACTER",
    "character": "CHARACTER",
    "skills": "SKILL",
    "skill": "SKILL",
    "buffs": "BUFF_STATUS",
    "buff_status": "BUFF_STATUS",
    "zhizhi": "ZHIZHI",
    "brilliant": "HUANZHANG",
    "huanzhang": "HUANZHANG",
    "profile": "PROFILE",
    "skins": "SKIN",
    "skin": "SKIN"
}

# 3. Các sheet thay thế nguyên khối từ file terms
TERMS_SHEETS = ["GLOSSARY", "TALENT", "REVIEW", "REVIEW_CANDIDATES"]

def merge_all(terms_file=None, batch_file=None):
    # Tự động tìm đường dẫn file master nếu file nằm ở thư mục gốc
    master_path = MASTER_FILE
    if not os.path.exists(master_path):
        alt_master = os.path.join(PROJECT_ROOT, "localization_master.xlsx")
        if os.path.exists(alt_master):
            master_path = alt_master
        else:
            print(f"[-] Không tìm thấy file master tại: {master_path}")
            return

    # Tự động nhận diện đường dẫn 2 file con nếu chưa truyền vào
    if not terms_file:
        for fname in ["terms_reference_export_3.xlsx", "terms_reference_export.xlsx", "localization/terms_reference_export.xlsx"]:
            p = os.path.join(PROJECT_ROOT, fname)
            if os.path.exists(p):
                terms_file = p
                break

    if not batch_file:
        for fname in ["batch_5_characters_export_4.xlsx", "batch_5_characters_export.xlsx", "localization/batch_5_characters_export.xlsx"]:
            p = os.path.join(PROJECT_ROOT, fname)
            if os.path.exists(p):
                batch_file = p
                break

    # Tự động tạo bản sao lưu (Backup an toàn tuyệt đối)
    os.makedirs(BACKUP_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(BACKUP_DIR, f"localization_master_merge_all_{ts}.xlsx")
    shutil.copyfile(master_path, backup_path)
    print(f"[+] Đã tạo bản sao lưu an toàn tại: {backup_path}")

    wb_master = openpyxl.load_workbook(master_path)

    # ==========================================
    # BƯỚC 1: GỘP CÁC SHEET QUY ƯỚC TỪ FILE TERMS
    # ==========================================
    if terms_file and os.path.exists(terms_file):
        print(f"\n[+] Đang gộp file quy ước thuật ngữ: {terms_file}")
        xl_terms = pd.ExcelFile(terms_file)
        for s in TERMS_SHEETS:
            if s in xl_terms.sheet_names:
                df_term = pd.read_excel(terms_file, sheet_name=s)
                if s in wb_master.sheetnames:
                    del wb_master[s]
                ws_new = wb_master.create_sheet(title=s)
                ws_new.append(list(df_term.columns))
                for row in df_term.itertuples(index=False):
                    ws_new.append(list(row))
                print(f" -> Đã đồng bộ nguyên khối sheet [{s}] ({len(df_term)} dòng)")
    else:
        print(" [!] Bỏ qua gộp file terms (không tìm thấy tệp).")

    # ==========================================
    # BƯỚC 2: GỘP DỮ LIỆU NHÂN VẬT TỪ FILE BATCH
    # ==========================================
    if batch_file and os.path.exists(batch_file):
        print(f"\n[+] Đang gộp file batch nhân vật: {batch_file}")
        xl_batch = pd.ExcelFile(batch_file)
        total_batch_updated = 0

        for b_sheet in xl_batch.sheet_names:
            target_sheet = BATCH_SHEET_MAP.get(b_sheet.lower(), b_sheet.upper())
            if target_sheet not in wb_master.sheetnames:
                continue

            pk = PRIMARY_KEYS.get(target_sheet)
            if not pk:
                continue

            df_b = pd.read_excel(batch_file, sheet_name=b_sheet)
            if df_b.empty:
                continue

            ws_m = wb_master[target_sheet]
            master_rows = list(ws_m.iter_rows(values_only=True))
            master_header = list(master_rows[0])
            batch_cols = list(df_b.columns)

            # Lập index vị trí các dòng trên Master
            master_index = {}
            for row_idx, r in enumerate(master_rows[1:], start=2):
                if isinstance(pk, tuple):
                    idx1 = master_header.index(pk[0])
                    idx2 = master_header.index(pk[1])
                    k = (str(r[idx1]).strip(), str(r[idx2]).strip())
                else:
                    idx = master_header.index(pk)
                    k = str(r[idx]).strip()
                master_index[k] = row_idx

            # Chỉ cập nhật các cột dịch thuật và trạng thái, không sửa ID
            update_count = 0
            for _, b_row in df_b.iterrows():
                if isinstance(pk, tuple):
                    b_k = (str(b_row[pk[0]]).strip(), str(b_row[pk[1]]).strip())
                else:
                    b_k = str(b_row[pk]).strip()

                if b_k in master_index:
                    target_row_idx = master_index[b_k]
                    for col in batch_cols:
                        if col.endswith("_vi") or col in ["confidence", "status", "notes"]:
                            if col in master_header:
                                val = b_row[col]
                                if pd.isna(val):
                                    val = None
                                m_col_idx = master_header.index(col) + 1
                                ws_m.cell(row=target_row_idx, column=m_col_idx, value=val)
                    update_count += 1

            print(f" -> Sheet [{target_sheet}]: Đã cập nhật thành công {update_count} dòng")
            total_batch_updated += update_count

        print(f"\n[V] Đã hoàn tất cập nhật {total_batch_updated} mục nhân vật!")
    else:
        print(" [!] Bỏ qua gộp file batch (không tìm thấy tệp).")

    # Lưu thay đổi vào file mẹ
    wb_master.save(master_path)
    print(f"\n[SUCCESS] Toàn bộ dữ liệu đã được gộp thành công vào file mẹ: {master_path}")

if __name__ == "__main__":
    t_file = sys.argv[1] if len(sys.argv) > 1 else None
    b_file = sys.argv[2] if len(sys.argv) > 2 else None
    merge_all(t_file, b_file)