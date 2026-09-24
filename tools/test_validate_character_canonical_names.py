import unittest
from pathlib import Path
import sys
import uuid

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

from validate_character_canonical_names import validate_workbook


def make_workbook(target_text):
    output_dir = Path(__file__).resolve().parent / "test_fixtures" / "canonical_validator" / uuid.uuid4().hex
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "canonical.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "CHARACTER"
    ws.append(["character_id", "name_cn", "fullname_cn", "name_vi", "fullname_vi"])
    ws.append(["A0100", "愿望杯", "愿望杯", "Lotus Chalice", "Lotus Chalice"])
    hz = wb.create_sheet("HUANZHANG")
    hz.append(["brilliant_id", "character_id", "icon_name_cn", "icon_name_vi", "icon_info_cn", "icon_info_vi", "buff_show_cn", "buff_show_vi", "translation_review_status"])
    hz.append(["S01554", "S0155", "x", "x", "向愿望杯许愿。", target_text, "", "", "CHATGPT_REVIEWED"])
    skill = wb.create_sheet("SKILL")
    skill.append(["skill_id", "character_id", "skill_name_cn", "skill_name_vi", "desc_cn", "desc_vi", "translation_review_status"])
    buff = wb.create_sheet("BUFF_STATUS")
    buff.append(["buff_id", "buff_name_cn", "buff_name_vi", "buff_desc_cn", "buff_desc_vi"])
    wb.save(path)
    return path


class CharacterCanonicalValidatorTests(unittest.TestCase):
    def test_catches_cross_character_legacy_alias(self):
        path = make_workbook("Nguyện Vọng Bôi lắng nghe lời nguyện.")
        _, _, issues = validate_workbook(path)
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]["canonical_cn"], "愿望杯")
        self.assertEqual(issues[0]["canonical_vi"], "Lotus Chalice")
        self.assertEqual(issues[0]["legacy_aliases_present"], ["Nguyện Vọng Bôi"])

    def test_passes_cross_character_canonical_name(self):
        path = make_workbook("Lotus Chalice lắng nghe lời nguyện.")
        _, _, issues = validate_workbook(path)
        self.assertEqual(issues, [])


if __name__ == "__main__":
    unittest.main()
