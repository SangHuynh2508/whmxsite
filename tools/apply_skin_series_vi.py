"""
Safe workbook mutation: Apply owner-approved Series Vietnamese names to SKIN sheet.
Canonical mapping for 19 Series:
202: Tân Xuân, 203: Hoa Triêu, 204: Tiết Khí, 205: Phi Di, 206: Nhàn Thú,
207: Trường An, 208: Ỷ Mộng, 209: Hạnh Thực, 210: Kỷ Niệm, 211: Dị Tượng,
212: Huyễn Cảnh, 213: Hành Giả, 214: Tài Dạng, 215: Linh Luật, 216: Tần Âm,
217: Chức Thải, 218: Dị Thế, 219: Tiêu Thử, 220: Vân Tưởng Tân Thường.
"""
from __future__ import annotations

import hashlib
import json
import os
import sys

from safe_workbook_mutation import safe_mutate_workbook

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass


SERIES_MAP_VI = {
    202: "Tân Xuân",
    203: "Hoa Triêu",
    204: "Tiết Khí",
    205: "Phi Di",
    206: "Nhàn Thú",
    207: "Trường An",
    208: "Ỷ Mộng",
    209: "Hạnh Thực",
    210: "Kỷ Niệm",
    211: "Dị Tượng",
    212: "Huyễn Cảnh",
    213: "Hành Giả",
    214: "Tài Dạng",
    215: "Linh Luật",
    216: "Tần Âm",
    217: "Chức Thải",
    218: "Dị Thế",
    219: "Tiêu Thử",
    220: "Vân Tưởng Tân Thường",
}


def get_sha256(filepath: str) -> str:
    h = hashlib.sha256()
    with open(filepath, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


def main():
    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    master_path = os.path.join(base_dir, "localization", "localization_master.xlsx")

    print("=== APPLY SERIES VI TO SKIN SHEET ===")
    before_sha = get_sha256(master_path)
    print(f"Master workbook path: {master_path}")
    print(f"Master SHA-256 before: {before_sha}")

    mutated_skin_ids = []

    def mutator(wb):
        ws = wb["SKIN"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        s_id_idx = headers.index("series_id") + 1
        s_vi_idx = headers.index("series_name_vi") + 1

        applied_count = 0
        unassigned_count = 0

        for r in range(2, ws.max_row + 1):
            sid = ws.cell(r, 1).value
            if not sid:
                continue

            mutated_skin_ids.append(str(sid))
            s_id = ws.cell(r, s_id_idx).value
            if s_id and int(s_id) in SERIES_MAP_VI:
                vi_name = SERIES_MAP_VI[int(s_id)]
                ws.cell(r, s_vi_idx, vi_name)
                applied_count += 1
            else:
                ws.cell(r, s_vi_idx, None)
                unassigned_count += 1

        print(f"Mutator summary: {applied_count} rows assigned VI names, {unassigned_count} unassigned (null)")
        assert applied_count == 144, f"Expected 144 assigned, got {applied_count}"
        assert unassigned_count == 1, f"Expected 1 unassigned, got {unassigned_count}"

    # Load skin IDs for authorized_deps
    import openpyxl
    wb_pre = openpyxl.load_workbook(master_path, read_only=True)
    all_sids = [str(r[0]) for r in wb_pre["SKIN"].iter_rows(min_row=2, max_col=1, values_only=True) if r[0]]
    wb_pre.close()

    authorized_deps = {
        "skin_ids": all_sids
    }
    authorized_deleted_rows = {
        "SKIN": {""}
    }

    safe_mutate_workbook(
        base_dir=base_dir,
        mutator_fn=mutator,
        authorized_deps=authorized_deps,
        authorized_deleted_rows=authorized_deleted_rows,
    )

    after_sha = get_sha256(master_path)
    print(f"Master SHA-256 after: {after_sha}")
    print("Safe mutation completed successfully!")


if __name__ == "__main__":
    main()
