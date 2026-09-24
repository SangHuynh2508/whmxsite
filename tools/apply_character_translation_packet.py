"""Guarded importer for a reviewed character translation packet."""
from __future__ import annotations

import argparse
import hashlib
import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

from safe_workbook_mutation import AtomicReplaceLockError, safe_mutate_workbook
from workbook_semantic_fingerprint import workbook_semantic_fingerprint


ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "localization/localization_master.xlsx"
DEFAULT_PACKET = ROOT / "localization/reviews/character_translation_packet_20260913174045_TRANSLATED.xlsx"
DEFAULT_EXPECTED = "65255FDD5BF0977A3DE418A5A7CE87C92DE50C7AE3F16C34421EEF33AF35FF36"
COLOR = re.compile(r"^<color=#[0-9A-Fa-f]{6}>([^<]+)</color>$")


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest().upper()


def rows(sheet):
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if any(value not in (None, "") for value in row)]


def fail(kind, identifier, field, master_value, packet_value):
    raise SystemExit(
        f"{kind}:{identifier}:{field}:master={master_value!r}:packet={packet_value!r}"
    )


def same_name(left, right):
    if left == right:
        return True
    left_color, right_color = COLOR.fullmatch(str(left or "")), COLOR.fullmatch(str(right or ""))
    return bool(left_color and left_color.group(1) == right or right_color and right_color.group(1) == left)


def same_optional_text(left, right):
    return left == right or (left in (None, "") and right in (None, ""))


def _action(row):
    return row.get("translation_action") or row.get("proposal_action")


SUPPORTED_ACTIONS = {
    "SKILLS": {"REVIEWED_UNCHANGED", "REVISED", "NEW_TRANSLATION"},
    "BUFFS": {"LOCKED_REFERENCE", "REVIEWED_REFERENCE", "REVISED", "REVIEWED_UNCHANGED", "NEW_TRANSLATION"},
    "HUANZHANG": {"INTERNAL_REFERENCE", "REVIEWED_UNCHANGED", "REVISED", "NEW_TRANSLATION"},
}
REQUIRED_COLUMNS = {
    "SKILLS": {"character_id", "skill_id", "skill_name_cn", "desc_cn", "proposed_vi_name", "proposed_vi_desc"},
    "BUFFS": {"buff_id", "buff_name_cn", "buff_desc_cn", "proposed_vi_name", "proposed_vi_desc"},
    "HUANZHANG": {"character_id", "huanzhang_id", "title_name_cn", "full_cn_lore_story", "proposed_vi_title", "proposed_vi_lore"},
}
ID_FIELDS = {"SKILLS": "skill_id", "BUFFS": "buff_id", "HUANZHANG": "huanzhang_id"}


def _require_columns(sheet_name, records):
    if not records:
        return
    headers = set(records[0])
    if sheet_name in ("SKILLS", "BUFFS"):
        action_col = "translation_action" if "translation_action" in headers else "proposal_action"
        if action_col not in headers:
            raise SystemExit(f"PACKET_SHAPE_MISMATCH:{sheet_name}:missing_action_column")
    missing = sorted(REQUIRED_COLUMNS[sheet_name] - headers)
    if missing:
        raise SystemExit(f"PACKET_SHAPE_MISMATCH:{sheet_name}:missing_columns={missing!r}")


def _require_unique_ids(sheet_name, records):
    id_field = ID_FIELDS[sheet_name]
    ids = [row.get(id_field) for row in records]
    duplicates = sorted({identifier for identifier in ids if identifier and ids.count(identifier) > 1})
    if duplicates:
        raise SystemExit(f"DUPLICATE_{id_field.upper()}:{duplicates!r}")


def _require_supported_actions(sheet_name, records):
    supported = SUPPORTED_ACTIONS[sheet_name]
    bad = sorted({_action(row) for row in records if _action(row) and _action(row) not in supported})
    if bad:
        raise SystemExit(f"UNSUPPORTED_PROPOSAL_ACTION:{sheet_name}:{bad!r}")


def _summary_value(packet_rows, field):
    summary = packet_rows.get("BATCH_SUMMARY") or []
    if not summary or field not in summary[0] or summary[0].get(field) in (None, ""):
        return None
    try:
        return int(summary[0][field])
    except (TypeError, ValueError):
        raise SystemExit(f"PACKET_SUMMARY_NON_NUMERIC:{field}:{summary[0].get(field)!r}")


def require_packet_shape(packet_rows):
    for sheet_name in ("SKILLS", "BUFFS", "HUANZHANG"):
        if sheet_name not in packet_rows:
            raise SystemExit(f"PACKET_SHAPE_MISMATCH:missing_sheet={sheet_name}")
        _require_columns(sheet_name, packet_rows[sheet_name])
        _require_unique_ids(sheet_name, packet_rows[sheet_name])
        _require_supported_actions(sheet_name, packet_rows[sheet_name])

    metadata_checks = {
        "exact_skill_rows": len(packet_rows["SKILLS"]),
        "unique_buff_count": len(packet_rows["BUFFS"]),
        "locked_buff_count": sum(_action(row) == "LOCKED_REFERENCE" for row in packet_rows["BUFFS"]),
        "reviewed_reference_buff_count": sum(_action(row) == "REVIEWED_REFERENCE" for row in packet_rows["BUFFS"]),
        "buff_rows_reviewed": sum(_action(row) == "REVISED" for row in packet_rows["BUFFS"]),
        "review_existing_buff_count": sum(_action(row) in {"REVISED", "REVIEWED_UNCHANGED"} for row in packet_rows["BUFFS"]),
        "huanzhang_player_rows_translated": sum(
            any(row.get(field) not in (None, "") for field in ("proposed_vi_title", "proposed_vi_lore", "proposed_vi_buff_show"))
            for row in packet_rows["HUANZHANG"]
        ),
        "huanzhang_character_count": len({
            row["character_id"]
            for row in packet_rows["HUANZHANG"]
            if any(row.get(field) not in (None, "") for field in ("proposed_vi_title", "proposed_vi_lore", "proposed_vi_buff_show"))
        }),
    }
    for field, actual in metadata_checks.items():
        expected = _summary_value(packet_rows, field)
        if expected is not None and actual != expected:
            raise SystemExit(f"PACKET_SUMMARY_MISMATCH:{field}:actual={actual}:expected={expected}")

    qa = packet_rows.get("TRANSLATION_QA") or []
    failing_qa = [row for row in qa if str(row.get("result", "")).strip().upper() not in {"", "PASS", "OK", "INFO"}]
    if failing_qa:
        raise SystemExit(f"PACKET_TRANSLATION_QA_NOT_PASS:{failing_qa!r}")


def canonical_lotus_guard(master_rows, allow_repair=False):
    character = master_rows["CHARACTER"]
    records = [row for row in character if row.get("character_id") == "A0100"]
    expected = {"name_cn": "愿望杯", "fullname_cn": "愿望杯", "name_vi": "Lotus Chalice", "fullname_vi": "Lotus Chalice"}
    if len(records) != 1 or {field: records[0].get(field) for field in expected} != expected:
        actual = None if not records else {field: records[0].get(field) for field in expected}
        raise SystemExit(f"CANONICAL_CHARACTER_A0100_MISMATCH:{actual!r}")

    all_wrong_references = []
    for sheetname, records_for_sheet in master_rows.items():
        for record in records_for_sheet:
            for field, value in record.items():
                if isinstance(value, str) and "Nguyện Vọng Bôi" in value:
                    all_wrong_references.append((sheetname, record.get("brilliant_id") or record.get("skill_id") or record.get("character_id") or record.get("buff_id"), field, value.count("Nguyện Vọng Bôi")))
    if not allow_repair:
        if all_wrong_references:
            raise SystemExit(f"CANONICAL_CHARACTER_UNEXPECTED_OLD_REFERENCE:{all_wrong_references!r}")
        return None
    if all_wrong_references != [("HUANZHANG", "S01554", "icon_info_vi", 1)]:
        raise SystemExit(f"CANONICAL_CHARACTER_UNEXPECTED_OLD_REFERENCE:{all_wrong_references!r}")
    repair = next(row for row in master_rows["HUANZHANG"] if row.get("brilliant_id") == "S01554")
    if "愿望杯" not in (repair.get("icon_info_cn") or ""):
        raise SystemExit("CANONICAL_CHARACTER_S0155_SOURCE_MISSING_CN")
    if (repair.get("icon_info_vi") or "").count("Nguyện Vọng Bôi") != 1:
        raise SystemExit("CANONICAL_CHARACTER_S0155_EXPECTED_ONE_OLD_REFERENCE")
    return repair


def guarded_in_place_fallback(master_path: Path, candidate_path: Path, backup_path: Path) -> None:
    current_sha = sha(master_path)
    current_fp = workbook_semantic_fingerprint(master_path)
    candidate_fp = workbook_semantic_fingerprint(candidate_path)
    if candidate_fp == current_fp:
        raise RuntimeError("FALLBACK_CANDIDATE_HAS_NO_DELTA")

    if sha(backup_path) != current_sha or workbook_semantic_fingerprint(backup_path) != current_fp:
        raise RuntimeError("FALLBACK_BACKUP_VERIFICATION_FAILED")

    candidate_bytes = candidate_path.read_bytes()

    import ctypes
    from ctypes import wintypes
    import msvcrt

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    create_file = kernel32.CreateFileW
    create_file.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    create_file.restype = wintypes.HANDLE
    invalid_handle = ctypes.c_void_p(-1).value
    write_handle = create_file(str(master_path), 0xC0000000, 0, None, 3, 0x80, None)
    if write_handle == invalid_handle:
        raise RuntimeError(f"FALLBACK_EXCLUSIVE_WRITE_FAILED winerror={ctypes.get_last_error()}")

    def overwrite(handle: int, bytes_to_write: bytes) -> None:
        fd = msvcrt.open_osfhandle(handle, os.O_RDWR | getattr(os, "O_BINARY", 0))
        with os.fdopen(fd, "r+b", closefd=True) as destination:
            destination.seek(0)
            destination.truncate(0)
            destination.write(bytes_to_write)
            destination.flush()
            os.fsync(destination.fileno())

    try:
        overwrite(write_handle, candidate_bytes)
        if sha(master_path) != sha(candidate_path):
            raise RuntimeError("FALLBACK_SHA_VERIFICATION_FAILED")
        if workbook_semantic_fingerprint(master_path) != candidate_fp:
            raise RuntimeError("FALLBACK_SEMANTIC_VERIFICATION_FAILED")
        try:
            candidate_path.unlink(missing_ok=True)
        except Exception:
            pass
    except Exception as exc:
        restore_handle = create_file(str(master_path), 0xC0000000, 0, None, 3, 0x80, None)
        if restore_handle != invalid_handle:
            try:
                overwrite(restore_handle, backup_path.read_bytes())
            except Exception:
                pass
        raise RuntimeError(f"FALLBACK_FAILED_RESTORED cause={exc!r}") from exc


def validate_source_guards(packet_rows, master_rows):
    skills = {row["skill_id"]: row for row in master_rows["SKILL"]}
    buffs = {row["buff_id"]: row for row in master_rows["BUFF_STATUS"]}
    huanzhang = {row["brilliant_id"]: row for row in master_rows["HUANZHANG"]}
    for row in packet_rows["SKILLS"]:
        current = skills.get(row["skill_id"])
        if current is None:
            fail("SOURCE_CHANGED_SKILL", row["skill_id"], "missing_id", None, row["skill_id"])
        for field in ("character_id", "skill_name_cn", "desc_cn"):
            if current[field] != row[field]:
                fail("SOURCE_CHANGED_SKILL", row["skill_id"], field, current[field], row[field])
        if "skill_name_vi_current" in row and not same_optional_text(current.get("skill_name_vi"), row["skill_name_vi_current"]):
            fail("SOURCE_CHANGED_SKILL", row["skill_id"], "skill_name_vi", current.get("skill_name_vi"), row["skill_name_vi_current"])
        if "desc_vi_current" in row and not same_optional_text(current.get("desc_vi"), row["desc_vi_current"]):
            fail("SOURCE_CHANGED_SKILL", row["skill_id"], "desc_vi", current.get("desc_vi"), row["desc_vi_current"])
    for row in packet_rows["BUFFS"]:
        current = buffs.get(row["buff_id"])
        if current is None:
            fail("SOURCE_CHANGED_BUFF", row["buff_id"], "missing_id", None, row["buff_id"])
        if current["buff_desc_cn"] != row["buff_desc_cn"]:
            fail("SOURCE_CHANGED_BUFF", row["buff_id"], "buff_desc_cn", current["buff_desc_cn"], row["buff_desc_cn"])
        if not same_name(current["buff_name_cn"], row["buff_name_cn"]):
            fail("SOURCE_CHANGED_BUFF", row["buff_id"], "buff_name_cn", current["buff_name_cn"], row["buff_name_cn"])
    for row in packet_rows["HUANZHANG"]:
        current = huanzhang.get(row["huanzhang_id"])
        if current is None:
            fail("SOURCE_CHANGED_HUANZHANG", row["huanzhang_id"], "missing_id", None, row["huanzhang_id"])
        if current.get("character_id") != row.get("character_id"):
            fail("SOURCE_CHANGED_HUANZHANG", row["huanzhang_id"], "character_id", current.get("character_id"), row.get("character_id"))
        for master_field, packet_field in (("icon_name_cn", "title_name_cn"), ("icon_info_cn", "full_cn_lore_story"), ("buff_show_cn", "buff_show_cn")):
            if not same_optional_text(current.get(master_field), row.get(packet_field)):
                fail("SOURCE_CHANGED_HUANZHANG", row["huanzhang_id"], master_field, current.get(master_field), row.get(packet_field))
        if "title_name_vi_current" in row and not same_optional_text(current.get("icon_name_vi"), row.get("title_name_vi_current")):
            fail("SOURCE_CHANGED_HUANZHANG", row["huanzhang_id"], "icon_name_vi", current.get("icon_name_vi"), row.get("title_name_vi_current"))
        if "current_vi_lore_story" in row and not same_optional_text(current.get("icon_info_vi"), row.get("current_vi_lore_story")):
            fail("SOURCE_CHANGED_HUANZHANG", row["huanzhang_id"], "icon_info_vi", current.get("icon_info_vi"), row.get("current_vi_lore_story"))
        if "buff_show_vi_current" in row and not same_optional_text(current.get("buff_show_vi"), row.get("buff_show_vi_current")):
            fail("SOURCE_CHANGED_HUANZHANG", row["huanzhang_id"], "buff_show_vi", current.get("buff_show_vi"), row.get("buff_show_vi_current"))


def plan_packet_apply(packet_rows, master_rows, allow_repair=False):
    locked = [row for row in packet_rows["BUFFS"] if _action(row) in {"LOCKED_REFERENCE", "REVIEWED_REFERENCE"}]
    reviewed = [row for row in packet_rows["BUFFS"] if row not in locked and (row.get("proposed_vi_name") not in (None, "") or row.get("proposed_vi_desc") not in (None, ""))]
    hz_apply = [
        row for row in packet_rows["HUANZHANG"]
        if _action(row) != "INTERNAL_REFERENCE"
        and any(row.get(field) not in (None, "") for field in ("proposed_vi_title", "proposed_vi_lore", "proposed_vi_buff_show"))
    ]
    repair = canonical_lotus_guard(master_rows, allow_repair=allow_repair)
    dependencies = {
        "skill_ids": {row["skill_id"] for row in packet_rows["SKILLS"]},
        "player_facing_buff_ids": {row["buff_id"] for row in reviewed},
        "huanzhang_ids": {row["huanzhang_id"] for row in hz_apply},
    }
    return locked, reviewed, hz_apply, repair, dependencies


def build_mutator(packet_name, packet_rows, reviewed_buffs, hz_apply, repair, character_updates=None):
    def mutate(workbook):
        for name in ("SKILL", "HUANZHANG"):
            worksheet = workbook[name]
            headers = [cell.value for cell in worksheet[1]]
            for column in ("translation_review_status", "translation_review_source"):
                if column not in headers:
                    worksheet.cell(1, len(headers) + 1).value = column
                    headers.append(column)

        skill_sheet = workbook["SKILL"]
        skill_headers = {cell.value: position + 1 for position, cell in enumerate(skill_sheet[1])}
        skill_rows = {str(skill_sheet.cell(row, skill_headers["skill_id"]).value): row for row in range(2, skill_sheet.max_row + 1)}
        for row in packet_rows["SKILLS"]:
            target = skill_rows[row["skill_id"]]
            skill_sheet.cell(target, skill_headers["skill_name_vi"]).value = row["proposed_vi_name"]
            skill_sheet.cell(target, skill_headers["desc_vi"]).value = row["proposed_vi_desc"]
            skill_sheet.cell(target, skill_headers["translation_review_status"]).value = "CHATGPT_REVIEWED"
            skill_sheet.cell(target, skill_headers["translation_review_source"]).value = packet_name

        buff_sheet = workbook["BUFF_STATUS"]
        buff_headers = {cell.value: position + 1 for position, cell in enumerate(buff_sheet[1])}
        buff_rows = {str(buff_sheet.cell(row, buff_headers["buff_id"]).value): row for row in range(2, buff_sheet.max_row + 1)}
        for row in reviewed_buffs:
            target = buff_rows[row["buff_id"]]
            buff_sheet.cell(target, buff_headers["buff_name_vi"]).value = row["proposed_vi_name"]
            buff_sheet.cell(target, buff_headers["buff_desc_vi"]).value = row["proposed_vi_desc"]
            current_name_auth = buff_sheet.cell(target, buff_headers["name_authority"]).value
            if current_name_auth != "OWNER_APPROVED":
                buff_sheet.cell(target, buff_headers["name_authority"]).value = "CHATGPT_REVIEWED"
            buff_sheet.cell(target, buff_headers["desc_translation_status"]).value = "CHATGPT_REVIEWED"
            buff_sheet.cell(target, buff_headers["desc_translation_source"]).value = packet_name

        hz_sheet = workbook["HUANZHANG"]
        hz_headers = {cell.value: position + 1 for position, cell in enumerate(hz_sheet[1])}
        hz_rows = {str(hz_sheet.cell(row, hz_headers["brilliant_id"]).value): row for row in range(2, hz_sheet.max_row + 1)}
        for row in hz_apply:
            target = hz_rows[row["huanzhang_id"]]
            if row.get("proposed_vi_title") not in (None, ""):
                hz_sheet.cell(target, hz_headers["icon_name_vi"]).value = row["proposed_vi_title"]
            if row.get("proposed_vi_lore") not in (None, ""):
                hz_sheet.cell(target, hz_headers["icon_info_vi"]).value = row["proposed_vi_lore"]
            if row.get("proposed_vi_buff_show") not in (None, "") and "buff_show_vi" in hz_headers:
                hz_sheet.cell(target, hz_headers["buff_show_vi"]).value = row["proposed_vi_buff_show"]
            if "translation_review_status" in hz_headers:
                hz_sheet.cell(target, hz_headers["translation_review_status"]).value = "CHATGPT_REVIEWED"
            if "translation_review_source" in hz_headers:
                hz_sheet.cell(target, hz_headers["translation_review_source"]).value = packet_name
        if repair is not None:
            target = hz_rows["S01554"]
            current = hz_sheet.cell(target, hz_headers["icon_info_vi"]).value or ""
            if current.count("Nguyện Vọng Bôi") != 1:
                raise RuntimeError("S0155 repair precondition changed during mutation")
            hz_sheet.cell(target, hz_headers["icon_info_vi"]).value = current.replace("Nguyện Vọng Bôi", "Lotus Chalice", 1)
        if character_updates:
            char_sheet = workbook["CHARACTER"]
            char_headers = {cell.value: position + 1 for position, cell in enumerate(char_sheet[1])}
            char_rows = {str(char_sheet.cell(row, char_headers["character_id"]).value): row for row in range(2, char_sheet.max_row + 1)}
            for cid, spec in character_updates.items():
                if cid not in char_rows:
                    raise RuntimeError(f"CHARACTER_NOT_FOUND:{cid}")
                target = char_rows[cid]
                for field, exp in spec.get("expected", {}).items():
                    actual = char_sheet.cell(target, char_headers[field]).value
                    actual_s = "" if actual is None else str(actual).strip()
                    exp_s = "" if exp is None else str(exp).strip()
                    if actual_s != exp_s:
                        raise RuntimeError(f"CHARACTER_GUARD_FAIL:{cid}:{field}:actual={actual!r}:expected={exp!r}")
                for field, val in spec.get("updates", {}).items():
                    char_sheet.cell(target, char_headers[field]).value = val
    return mutate


D0183_METADATA_SPEC = {
    "D0183": {
        "expected": {
            "character_id": "D0183",
            "name_cn": "幻戏图",
            "name_vi": None,
            "fullname_cn": "李嵩《骷髅幻戏图》页",
            "fullname_vi": None,
            "confidence": "LOW",
            "status": "PENDING",
        },
        "updates": {
            "name_vi": "Huyễn Hý Đồ",
            "fullname_vi": "Lý Tung 《Khô Lâu Huyễn Hý Đồ》 Hiệt",
            "confidence": "HIGH",
            "status": "APPROVED",
        },
    }
}


def apply_character_packet(root_path: Path, packet_path: Path, expected_sha: str | None = None, allow_repair: bool = False, character_updates: dict | None = None):
    root = Path(root_path).resolve()
    master = root / "localization/localization_master.xlsx"
    packet = Path(packet_path).resolve()
    actual_packet_hash = sha(packet)
    print(f"PACKET_SHA256={actual_packet_hash}")
    if expected_sha and actual_packet_hash != expected_sha.upper():
        raise SystemExit("PACKET_HASH_MISMATCH")

    packet_workbook = load_workbook(packet, read_only=True, data_only=True)
    master_workbook = load_workbook(master, read_only=True, data_only=True)
    try:
        packet_rows = {name: rows(packet_workbook[name]) for name in ("SKILLS", "BUFFS", "HUANZHANG")}
        for optional in ("BATCH_SUMMARY", "TRANSLATION_QA"):
            if optional in packet_workbook.sheetnames:
                packet_rows[optional] = rows(packet_workbook[optional])
        master_rows = {name: rows(master_workbook[name]) for name in ("CHARACTER", "SKILL", "BUFF_STATUS", "HUANZHANG")}
    finally:
        packet_workbook.close()
        master_workbook.close()

    require_packet_shape(packet_rows)
    validate_source_guards(packet_rows, master_rows)
    print(f"FULL_SOURCE_GUARD=PASS skills={len(packet_rows['SKILLS'])} buffs={len(packet_rows['BUFFS'])} huanzhang={len(packet_rows['HUANZHANG'])}")

    locked, reviewed, hz_apply, repair, dependencies = plan_packet_apply(packet_rows, master_rows, allow_repair=allow_repair)
    mutate = build_mutator(packet.name, packet_rows, reviewed, hz_apply, repair, character_updates=character_updates)

    authorized_cells = set()
    if repair is not None:
        authorized_cells.add(("HUANZHANG", "S01554", "icon_info_vi"))
    if character_updates:
        for cid, spec in character_updates.items():
            for field in spec.get("updates", {}):
                authorized_cells.add(("CHARACTER", cid, field))

    try:
        backup = safe_mutate_workbook(
            str(root),
            mutate,
            dependencies,
            authorized_new_columns={"SKILL": {"translation_review_status", "translation_review_source"}, "HUANZHANG": {"translation_review_status", "translation_review_source"}},
            authorized_cells=authorized_cells,
        )
        write_method = "ATOMIC_REPLACE"
    except AtomicReplaceLockError as lock_err:
        print("[INFO] Atomic replace hit Windows lock (WinError 32); engaging GUARDED_IN_PLACE_FALLBACK...", file=sys.stderr)
        guarded_in_place_fallback(Path(lock_err.master_path), Path(lock_err.temp_path), Path(lock_err.backup_path))
        write_method = "GUARDED_IN_PLACE_FALLBACK"
        backup = lock_err.backup_path

    print(f"WRITE_METHOD={write_method}")
    print(f"BACKUP={backup}")
    print(f"MASTER_BEFORE_AFTER={sha(backup)} {sha(master)}")
    print(f"MASTER_SEMANTIC_FINGERPRINT_AFTER={workbook_semantic_fingerprint(master)}")
    char_cell_count = sum(len(s.get("updates", {})) for s in (character_updates or {}).values())
    print(f"EXACT_APPLY_SCOPE=skills={len(dependencies['skill_ids'])} buffs={len(reviewed)} huanzhang={len(hz_apply)} explicit_cells={1 if repair else 0} character_cells={char_cell_count}")
    return {
        "write_method": write_method,
        "backup": str(backup),
        "master_sha_before": sha(backup),
        "master_sha_after": sha(master),
        "master_fp_after": workbook_semantic_fingerprint(master),
        "skills_applied": len(dependencies["skill_ids"]),
        "buffs_applied": len(reviewed),
        "buffs_locked": len(locked),
        "huanzhang_applied": len(hz_apply),
        "character_cells_applied": char_cell_count,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--packet", type=Path, default=DEFAULT_PACKET)
    parser.add_argument("--expected-sha256", default=DEFAULT_EXPECTED)
    parser.add_argument("--repair-s0155-lotus-chalice", action="store_true")
    parser.add_argument("--d0183-character-metadata", action="store_true", help="Apply reviewed D0183 character metadata")
    parser.add_argument("--character-metadata-file", type=Path, default=None, help="Path to JSON file with character metadata updates")
    args = parser.parse_args()

    char_updates = None
    if args.character_metadata_file:
        import json
        char_updates = json.loads(args.character_metadata_file.read_text(encoding="utf-8"))
    elif args.d0183_character_metadata:
        char_updates = D0183_METADATA_SPEC

    apply_character_packet(ROOT, args.packet, expected_sha=args.expected_sha256, allow_repair=args.repair_s0155_lotus_chalice, character_updates=char_updates)


if __name__ == "__main__":
    main()
