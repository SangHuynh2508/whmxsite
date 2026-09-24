"""Stable semantic fingerprints for XLSX workbooks.

The digest deliberately represents workbook content rather than XLSX ZIP
members, compression, timestamps, or document properties.  ``None`` and an
empty string are distinct values so callers can use it as a content guard.
"""
from __future__ import annotations

import argparse
from datetime import date, datetime, time
import hashlib
import json
from pathlib import Path
from typing import Any

import openpyxl


def _value_record(value: Any, data_type: str) -> dict[str, str]:
    if value is None:
        return {"kind": "none", "data_type": data_type}
    if isinstance(value, bool):
        return {"kind": "bool", "data_type": data_type, "value": "true" if value else "false"}
    if isinstance(value, (datetime, date, time)):
        return {"kind": type(value).__name__, "data_type": data_type, "value": value.isoformat()}
    if isinstance(value, float):
        return {"kind": "float", "data_type": data_type, "value": value.hex()}
    return {"kind": type(value).__name__, "data_type": data_type, "value": str(value)}


def workbook_semantic_fingerprint(path: str | Path) -> str:
    """Return SHA-256 over ordered sheets, cells, formulas, and values."""
    workbook = openpyxl.load_workbook(Path(path), read_only=True, data_only=False)
    digest = hashlib.sha256()
    try:
        def write(record: Any) -> None:
            digest.update(json.dumps(record, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8"))
            digest.update(b"\n")

        write({"format": "whmx-workbook-semantic-v1", "sheet_order": workbook.sheetnames})
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            write({"sheet": sheet_name, "max_row": worksheet.max_row, "max_column": worksheet.max_column})
            for row in worksheet.iter_rows():
                write(
                    {
                        "row": row[0].row if row else None,
                        "cells": [_value_record(cell.value, cell.data_type) for cell in row],
                    }
                )
    finally:
        workbook.close()
    return digest.hexdigest().upper()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    print(workbook_semantic_fingerprint(args.workbook))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
