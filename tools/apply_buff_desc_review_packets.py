#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
apply_buff_desc_review_packets.py

Imports the 74 reviewed buff descriptions from the 5 FILLED packets into
localization_master.xlsx (sheet BUFF_STATUS) using the safe_mutate_workbook guard.

Rules:
- Match ONLY by exact buff_id.
- Update ONLY buff_desc_vi.
- Do NOT touch buff_name_vi, CN fields, or any other sheet.
- Fail closed on any conflict, unmatched ID, or duplicate ID.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

from safe_workbook_mutation import safe_mutate_workbook, create_workbook_backup
from workbook_semantic_fingerprint import workbook_semantic_fingerprint

PACKET_DIR = ROOT / "localization" / "reviews" / "buff_desc_packets"
DEFAULT_MASTER = ROOT / "localization" / "localization_master.xlsx"

FILLED_PACKETS = [
    "buff_desc_review_packet_01_20260917_FILLED.xlsx",
    "buff_desc_review_packet_02_20260917_FILLED.xlsx",
    "buff_desc_review_packet_03_20260917_FILLED.xlsx",
    "buff_desc_review_packet_04_20260917_FILLED.xlsx",
    "buff_desc_review_packet_05_20260917_FILLED.xlsx",
]

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()

def load_all_packet_rows() -> list[dict[str, Any]]:
    rows = []
    for fn in FILLED_PACKETS:
        p = PACKET_DIR / fn
        if not p.is_file():
            raise FileNotFoundError(f"Packet not found: {p}")
        wb = openpyxl.load_workbook(p, data_only=True)
        if "TO_TRANSLATE" not in wb.sheetnames:
            raise ValueError(f"Sheet TO_TRANSLATE missing from {fn}")
        ws = wb["TO_TRANSLATE"]
        header = [c for c in next(ws.iter_rows(values_only=True))]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if any(r):
                rows.append(dict(zip(header, r)))
        wb.close()
    return rows

def build_merge_plan(packet_rows: list[dict[str, Any]], master_path: Path) -> dict[str, Any]:
    wb_m = openpyxl.load_workbook(master_path, data_only=True)
    if "BUFF_STATUS" not in wb_m.sheetnames:
        raise ValueError(f"Master missing BUFF_STATUS sheet")
    ws_m = wb_m["BUFF_STATUS"]
    m_headers = [c for c in next(ws_m.iter_rows(values_only=True))]
    hm_map = {c: i for i, c in enumerate(m_headers)}

    master_buffs = {}
    duplicate_master_ids = []
    for r_idx, r in enumerate(ws_m.iter_rows(min_row=2, values_only=True), start=2):
        bid = str(r[hm_map["buff_id"]] or "").strip()
        if bid:
            if bid in master_buffs:
                duplicate_master_ids.append(bid)
            master_buffs[bid] = (r_idx, dict(zip(m_headers, r)))
    wb_m.close()

    if duplicate_master_ids:
        raise ValueError(f"Duplicate buff_id in master BUFF_STATUS: {duplicate_master_ids}")

    # Check packet uniqueness
    seen_ids = set()
    dup_packet_ids = []
    for r in packet_rows:
        bid = str(r.get("buff_id") or "").strip()
        if not bid:
            raise ValueError(f"Empty buff_id in packet row: {r}")
        if bid in seen_ids:
            dup_packet_ids.append(bid)
        seen_ids.add(bid)

    if dup_packet_ids:
        raise ValueError(f"Duplicate buff_id in packets: {dup_packet_ids}")

    matched = []
    unmatched = []
    already_identical = []
    updates = []
    conflicts = []

    for r in packet_rows:
        bid = str(r.get("buff_id") or "").strip()
        n_cn_p = str(r.get("buff_name_cn") or "").strip()
        n_vi_p = str(r.get("buff_name_vi") or "").strip()
        d_cn_p = str(r.get("buff_desc_cn") or "").strip()
        d_vi_p = str(r.get("buff_desc_vi") or "").strip()
        tf = str(r.get("translation_field") or "").strip()

        if tf != "DESCRIPTION_ONLY":
            conflicts.append({"buff_id": bid, "reason": f"UNEXPECTED_TRANSLATION_FIELD_{tf}"})
            continue

        if not d_vi_p:
            conflicts.append({"buff_id": bid, "reason": "EMPTY_BUFF_DESC_VI_IN_PACKET"})
            continue

        if bid not in master_buffs:
            unmatched.append({"buff_id": bid, "buff_desc_cn": d_cn_p})
            continue

        r_idx, mr = master_buffs[bid]
        m_name_cn = str(mr.get("buff_name_cn") or "").strip()
        m_name_vi = str(mr.get("buff_name_vi") or "").strip()
        m_desc_cn = str(mr.get("buff_desc_cn") or "").strip()
        m_desc_vi = str(mr.get("buff_desc_vi") or "").strip()

        matched.append(bid)

        # Invariant checks
        if n_vi_p != m_name_vi:
            conflicts.append({
                "buff_id": bid,
                "reason": "BUFF_NAME_VI_MISMATCH",
                "packet": n_vi_p,
                "master": m_name_vi,
            })
            continue

        if d_cn_p != m_desc_cn:
            conflicts.append({
                "buff_id": bid,
                "reason": "BUFF_DESC_CN_MISMATCH",
                "packet": d_cn_p,
                "master": m_desc_cn,
            })
            continue

        if m_desc_vi == d_vi_p:
            already_identical.append({"buff_id": bid, "desc_vi": d_vi_p})
        elif m_desc_vi and m_desc_vi != d_vi_p:
            conflicts.append({
                "buff_id": bid,
                "reason": "PREEXISTING_DESC_VI_CONFLICT",
                "master": m_desc_vi,
                "packet": d_vi_p,
            })
        else:
            updates.append({
                "buff_id": bid,
                "master_row_index": r_idx,
                "new_desc_vi": d_vi_p,
                "old_desc_vi": m_desc_vi,
                "name_vi": m_name_vi,
            })

    return {
        "packet_rows_read": len(packet_rows),
        "unique_buff_ids": len(seen_ids),
        "rows_matched": len(matched),
        "rows_unmatched": len(unmatched),
        "rows_already_identical": len(already_identical),
        "rows_updated": len(updates),
        "conflicts_count": len(conflicts),
        "unmatched_list": unmatched,
        "already_identical_list": already_identical,
        "conflicts_list": conflicts,
        "updates_list": updates,
    }

def apply_merge(plan: dict[str, Any], master_path: Path, base_dir: Path) -> dict[str, Any]:
    if plan["conflicts_count"] > 0:
        raise RuntimeError(f"Cannot apply merge with {plan['conflicts_count']} conflicts: {plan['conflicts_list']}")
    if plan["rows_unmatched"] > 0:
        raise RuntimeError(f"Cannot apply merge with {plan['rows_unmatched']} unmatched rows: {plan['unmatched_list']}")

    update_map = {u["buff_id"]: u["new_desc_vi"] for u in plan["updates_list"]}
    authorized_buff_ids = set(update_map.keys())

    def mutator(wb: openpyxl.Workbook):
        ws = wb["BUFF_STATUS"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        h_map = {c: i + 1 for i, c in enumerate(headers)}
        id_col = h_map["buff_id"]
        desc_col = h_map["buff_desc_vi"]

        applied = 0
        for r in range(2, ws.max_row + 1):
            bid = str(ws.cell(r, id_col).value or "").strip()
            if bid in update_map:
                ws.cell(r, desc_col).value = update_map[bid]
                applied += 1

        print(f"[MUTATOR] Successfully applied {applied} buff_desc_vi updates into BUFF_STATUS")
        if applied != len(update_map):
            raise ValueError(f"Applied count {applied} != update_map count {len(update_map)}")

    authorized_deps = {"player_facing_buff_ids": list(authorized_buff_ids)}

    print("[APPLY] Invoking safe_mutate_workbook...")
    safe_mutate_workbook(
        base_dir=str(base_dir),
        mutator_fn=mutator,
        authorized_deps=authorized_deps,
    )

    return {"applied_count": len(update_map)}

def main():
    print("=== APPLYING 74 REVIEWED BUFF DESCRIPTIONS ===")
    pre_sha = sha256_file(DEFAULT_MASTER)
    pre_fp = workbook_semantic_fingerprint(DEFAULT_MASTER)
    print(f"Authoritative master: {DEFAULT_MASTER}")
    print(f"Pre-mutation SHA-256: {pre_sha}")
    print(f"Pre-mutation Semantic Fingerprint: {pre_fp}")

    # Load packet rows
    packet_rows = load_all_packet_rows()
    print(f"Loaded {len(packet_rows)} rows across {len(FILLED_PACKETS)} packet files")

    # Build merge plan
    plan = build_merge_plan(packet_rows, DEFAULT_MASTER)
    print(f"\n--- Merge Plan Summary ---")
    print(f"  Packet rows read: {plan['packet_rows_read']}")
    print(f"  Unique Buff IDs: {plan['unique_buff_ids']}")
    print(f"  Rows matched: {plan['rows_matched']}")
    print(f"  Rows to update: {plan['rows_updated']}")
    print(f"  Rows already identical: {plan['rows_already_identical']}")
    print(f"  Rows unmatched: {plan['rows_unmatched']}")
    print(f"  Conflicts: {plan['conflicts_count']}")

    if plan["conflicts_count"] > 0 or plan["rows_unmatched"] > 0:
        print("ERROR: Conflicts or unmatched rows detected! Aborting.")
        sys.exit(1)

    # Apply merge
    apply_merge(plan, DEFAULT_MASTER, ROOT)

    # Post-mutation verification
    post_sha = sha256_file(DEFAULT_MASTER)
    post_fp = workbook_semantic_fingerprint(DEFAULT_MASTER)
    print(f"\n--- Post-Mutation Verification ---")
    print(f"Post-mutation SHA-256: {post_sha}")
    print(f"Post-mutation Semantic Fingerprint: {post_fp}")

    # Detailed cell check
    wb_after = openpyxl.load_workbook(DEFAULT_MASTER, data_only=True)
    ws_after = wb_after["BUFF_STATUS"]
    headers = [c for c in next(ws_after.iter_rows(values_only=True))]
    hm_map = {c: i for i, c in enumerate(headers)}
    
    update_map = {u["buff_id"]: u["new_desc_vi"] for u in plan["updates_list"]}
    verified_updates = 0
    for r in ws_after.iter_rows(min_row=2, values_only=True):
        bid = str(r[hm_map["buff_id"]] or "").strip()
        if bid in update_map:
            actual_desc = str(r[hm_map["buff_desc_vi"]] or "").strip()
            expected_desc = update_map[bid]
            if actual_desc != expected_desc:
                raise ValueError(f"Post-check mismatch for {bid}!")
            verified_updates += 1
    wb_after.close()

    print(f"Verified {verified_updates}/{len(update_map)} updated cells match packet values exactly!")
    print("\nMERGE COMPLETED SUCCESSFULLY!")

if __name__ == "__main__":
    main()
