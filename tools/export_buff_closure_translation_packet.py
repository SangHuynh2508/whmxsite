#!/usr/bin/env python3
"""Export character-centric buff closure translation packet.

Strictly read-only:
- Never mutates localization_master.xlsx
- Never mutates NeoArtifacts
- Reuses snapshot authority, transitive popup-graph traversal, and write_character_translation_packet styling
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from collections import defaultdict, deque
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

PROJECT = Path(__file__).resolve().parent.parent
REPO = PROJECT.parent
NEO = REPO / "NeoArtifacts"
MASTER = PROJECT / "localization" / "localization_master.xlsx"
EXPORTS_DIR = PROJECT / "localization" / "exports"

sys.path.insert(0, str(Path(__file__).resolve().parent))
from audit_incremental_character_coverage import (
    load_authority, workbook_rows, explicit_list, skill_levels, strings, refs
)
from write_character_translation_packet import write_packet_from_dict

TAG_RE = re.compile(r"<[^>]+>")
COLOR_TERM_RE = re.compile(r'<color=[^>]+>(.*?)</color>', re.IGNORECASE | re.DOTALL)
NON_APPLY_OPS = {"cleanbuff", "checkbufflayers", "checkbuff", "immunizebuff", "checknotbuff", "cleanbuffgroup"}


def clean_tag(t: Any) -> str:
    return TAG_RE.sub('', str(t or '')).strip()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def extract_applied_buffs(attr_list: list) -> list[tuple[str, str]]:
    """Return list of (buff_id, raw_evidence) applied by an Attr list."""
    applied = []
    for item in attr_list or []:
        if isinstance(item, list) and len(item) > 0:
            val = item[0]
            parts = [p.strip() for p in val.split(',')]
            key = parts[0]
            if key.startswith("Effect") and not key.endswith("Para") and not key.endswith("Tips"):
                tokens = parts[1:]
                if tokens and tokens[0].lower() in NON_APPLY_OPS:
                    continue
                for i, token in enumerate(tokens):
                    if token.startswith("Buff_"):
                        prev = tokens[i-1].lower() if i > 0 else ""
                        if prev not in NON_APPLY_OPS:
                            applied.append((token, val))
    return applied


def run_export(output_path: Path | None = None) -> tuple[Path, str, dict[str, int]]:
    initial_master_hash = sha256_file(MASTER)

    tables, auth = load_authority()
    wb_data = workbook_rows()

    chars = {cid: value for cid, value in tables["characterTable"].items()
             if isinstance(value, dict) and str(value.get("Switch")) == "0"}

    buff_rows = {str(row["buff_id"]): row for row in wb_data["BUFF_STATUS"] if row.get("buff_id")}
    skill_rows = {str(row["skill_id"]): row for row in wb_data["SKILL"] if row.get("skill_id")}
    hz_rows = {str(row["brilliant_id"]): row for row in wb_data["HUANZHANG"] if row.get("brilliant_id")}
    char_rows = {str(row["character_id"]): row for row in wb_data["CHARACTER"] if row.get("character_id")}

    # Index buff names for popup term matching
    buff_ids_by_cn_name = defaultdict(set)
    for bid, b in tables["buffMap"].items():
        if isinstance(b, dict):
            nm = clean_tag(b.get("NameLanText"))
            if nm:
                buff_ids_by_cn_name[nm].add(bid)

    # Index group owners
    owners_by_group = defaultdict(set)
    for source_name in ("characterSkillMap", "characterPassiveSkillMap"):
        for record in tables[source_name].values():
            if isinstance(record, dict) and record.get("HeroId") and record.get("GroupId"):
                owners_by_group[str(record["GroupId"])].add(str(record["HeroId"]))

    # Index Hoan Chuong
    hz_by_owner = defaultdict(list)
    for brilliant_id, record in tables["BrilliantMap"].items():
        if not isinstance(record, dict): continue
        groups = sorted(set(explicit_list(record, "Buff") + explicit_list(record, "Skill1") + explicit_list(record, "Skill2")))
        owners = sorted({owner for group in groups for owner in owners_by_group.get(group, set()) if owner in chars})
        if len(owners) == 1:
            hz_by_owner[owners[0]].append((brilliant_id, groups, record))

    # Historical early baseline check for missing buffs
    early_wb_path = PROJECT / "localization/backups/localization_master_20260907_125847.xlsx"
    early_buffs = set()
    if early_wb_path.exists():
        ewb = openpyxl.load_workbook(early_wb_path, read_only=True)
        ews = ewb['BUFF_STATUS']
        eh = [c.value for c in next(ews.iter_rows(max_row=1))]
        idx_eid = eh.index('buff_id')
        for r in ews.iter_rows(min_row=2, values_only=True):
            if r[idx_eid]:
                early_buffs.add(str(r[idx_eid]))
        ewb.close()

    results_by_char = {}
    all_source_issues = []
    seen_source_issues = set()

    for cid, char in sorted(chars.items()):
        base_groups = [explicit_list(char, f"skill{slot}")[0] for slot in range(1, 7) if explicit_list(char, f"skill{slot}")]
        ex_groups = []
        for star in range(1, 7):
            role = tables["roleattrMap"].get(f"{cid}{star}", {})
            if isinstance(role, dict):
                for attr in strings(role.get("starUpAttr", [])):
                    match = re.search(r"(?:^|,)SkillUP,([A-Za-z0-9_]+)(?:,|$)", attr)
                    if match:
                        base = match.group(1)
                        enhanced = f"{base}ex"
                        if cid in owners_by_group.get(enhanced, set()):
                            if enhanced not in ex_groups:
                                ex_groups.append(enhanced)
        hz_groups = []
        for b_id, groups, r in hz_by_owner.get(cid, []):
            for g in groups:
                if g not in hz_groups:
                    hz_groups.append(g)

        # Collect skills for card
        card_skills = []
        for g in base_groups:
            for raw_key, level, record in skill_levels(tables["skillMap"], g):
                card_skills.append((g, record, "BASE"))
            for sname in ("characterSkillMap", "characterPassiveSkillMap"):
                for record in tables[sname].values():
                    if isinstance(record, dict) and str(record.get("GroupId")) == g:
                        card_skills.append((g, record, "BASE_MAP"))
        for g in ex_groups:
            for raw_key, level, record in skill_levels(tables["skillMap"], g):
                card_skills.append((g, record, "EX"))
            for sname in ("characterSkillMap", "characterPassiveSkillMap"):
                for record in tables[sname].values():
                    if isinstance(record, dict) and str(record.get("GroupId")) == g:
                        card_skills.append((g, record, "EX_MAP"))
        for g in hz_groups:
            for raw_key, level, record in skill_levels(tables["skillMap"], g):
                card_skills.append((g, record, "HZ"))
            for sname in ("characterSkillMap", "characterPassiveSkillMap"):
                for record in tables[sname].values():
                    if isinstance(record, dict) and str(record.get("GroupId")) == g:
                        card_skills.append((g, record, "HZ_MAP"))
        for b_id, groups, r in hz_by_owner.get(cid, []):
            card_skills.append((b_id, r, "HZ_BRILLIANT"))

        # Direct buff edges
        direct_buffs = set()
        direct_evidence = defaultdict(list)
        buff_to_skills = defaultdict(set)

        for card_id, rec, cat in card_skills:
            desc = str(rec.get("DescriptionLanText") or rec.get("BuffShowLanText") or "")
            for b in re.findall(r'\{(Buff_[^}\s]+)\}', desc):
                direct_buffs.add(b)
                direct_evidence[b].append((card_id, f"{cat}_TEXT_TAG", desc))
                buff_to_skills[b].add(card_id)
            for b, ev in extract_applied_buffs(rec.get("Attr", [])):
                direct_buffs.add(b)
                direct_evidence[b].append((card_id, f"{cat}_ATTR", ev))
                buff_to_skills[b].add(card_id)
            if cat == "HZ_BRILLIANT":
                for field in ("Buff", "BuffUp"):
                    for val in explicit_list(rec, field):
                        if val.startswith("Buff_"):
                            direct_buffs.add(val)
                            direct_evidence[val].append((card_id, f"HZ_DIRECT_{field}", f"{field}={val}"))
                            buff_to_skills[val].add(card_id)

        # Transitive closure
        closure = set()
        relation_path = {}
        parent_map = defaultdict(set)
        children_map = defaultdict(set)
        depth_map = {}
        raw_evidence_map = {}
        queue = deque()

        for b in sorted(direct_buffs):
            depth_map[b] = 0
            srcs = [f"{e[0]}:{e[1]}" for e in direct_evidence[b]]
            ev_str = "; ".join(f"{e[1]}: {e[2]}" for e in direct_evidence[b])
            raw_evidence_map[b] = ev_str
            relation_path[b] = f"{cid} -> {direct_evidence[b][0][0]} -> {b}"
            queue.append(b)

        while queue:
            bid = queue.popleft()
            if bid in closure: continue
            closure.add(bid)

            b_rec = tables["buffMap"].get(bid)
            if not isinstance(b_rec, dict):
                issue_key = (cid, bid)
                if issue_key not in seen_source_issues:
                    seen_source_issues.add(issue_key)
                    all_source_issues.append({
                        "character_id": cid,
                        "issue_type": "ABSENT_FROM_BUFFMAP",
                        "item_id": bid,
                        "raw_reference": raw_evidence_map.get(bid, ""),
                        "relation_path": relation_path.get(bid, ""),
                        "description": f"Referenced Buff_ID {bid} is absent from authoritative buffMap.json"
                    })
                continue

            child_refs = set()
            for child, ev in extract_applied_buffs(b_rec.get("Attr", [])):
                if child != bid:
                    child_refs.add((child, f"ATTR_EFFECT: {ev}"))
            desc = str(b_rec.get("DescriptionLanText") or "")
            for child in re.findall(r'\{(Buff_[^}\s]+)\}', desc):
                if child != bid:
                    child_refs.add((child, f"TEXT_TAG: {{{child}}}"))
            for m in COLOR_TERM_RE.finditer(desc):
                term = clean_tag(m.group(1))
                matching_ids = buff_ids_by_cn_name.get(term, set())
                char_family_matches = [x for x in matching_ids if cid in x or x.startswith(f"Buff_{cid}")]
                if len(matching_ids) == 1:
                    child_refs.add((list(matching_ids)[0], f"POPUP_COLOR_TERM: {term}"))
                elif len(char_family_matches) == 1:
                    child_refs.add((char_family_matches[0], f"POPUP_COLOR_TERM: {term}"))
                elif matching_ids:
                    for x in matching_ids:
                        if cid in x or x.startswith(f"Buff_{cid}"):
                            child_refs.add((x, f"POPUP_COLOR_TERM: {term}"))

            cur_depth = depth_map.get(bid, 0)
            for child, ev in sorted(child_refs):
                parent_map[child].add(bid)
                children_map[bid].add(child)
                for s in buff_to_skills[bid]:
                    buff_to_skills[child].add(s)
                if child not in depth_map or cur_depth + 1 < depth_map[child]:
                    depth_map[child] = cur_depth + 1
                if child not in relation_path:
                    relation_path[child] = f"{relation_path.get(bid, bid)} -> {child}"
                if child not in raw_evidence_map:
                    raw_evidence_map[child] = ev
                if child not in closure:
                    queue.append(child)

        # Graph node classification
        graph_nodes = []
        missing_units = []
        approved_units = []

        for bid in sorted(closure):
            b_rec = tables["buffMap"].get(bid)
            depth = depth_map.get(bid, 0)
            parents = sorted(parent_map.get(bid, set()))
            children = sorted(children_map.get(bid, set()))
            parent_id = parents[0] if parents else "ROOT"
            path_str = relation_path.get(bid, "")
            ev_str = raw_evidence_map.get(bid, "")

            if not isinstance(b_rec, dict):
                graph_nodes.append({
                    "character_id": cid,
                    "skill_id": ", ".join(sorted(buff_to_skills.get(bid, []))),
                    "root_buff_id": path_str.split(" -> ")[1] if " -> " in path_str else bid,
                    "buff_id": bid,
                    "parent_buff_id": parent_id,
                    "depth": depth,
                    "relation_path": path_str,
                    "semantic_type": "SOURCE_BLOCKER",
                    "player_facing": "NO",
                    "name_cn": "",
                    "name_vi": "",
                    "desc_cn": "",
                    "desc_vi": "",
                    "translation_state": "SOURCE_BLOCKER",
                    "raw_relation_evidence": ev_str,
                })
                continue

            name_cn = clean_tag(b_rec.get("NameLanText"))
            desc_cn = str(b_rec.get("DescriptionLanText") or "").strip()
            is_player_facing = bool(name_cn or desc_cn)

            wb_row = buff_rows.get(bid)
            name_vi = clean_tag((wb_row.get("buff_name_vi") if wb_row else "") or "")
            desc_vi = str((wb_row.get("buff_desc_vi") if wb_row else "") or "").strip()
            status_val = str((wb_row.get("status") if wb_row else "") or "").strip()
            authority_val = str((wb_row.get("name_authority") if wb_row else "") or "").strip()

            if not is_player_facing:
                semantic_type = "CONTROLLER_ONLY"
                translation_state = "CONTROLLER_CONTEXT_ONLY"
                graph_nodes.append({
                    "character_id": cid,
                    "skill_id": ", ".join(sorted(buff_to_skills.get(bid, []))),
                    "root_buff_id": path_str.split(" -> ")[1] if " -> " in path_str else bid,
                    "buff_id": bid,
                    "parent_buff_id": parent_id,
                    "depth": depth,
                    "relation_path": path_str,
                    "semantic_type": semantic_type,
                    "player_facing": "NO",
                    "name_cn": name_cn,
                    "name_vi": name_vi,
                    "desc_cn": desc_cn,
                    "desc_vi": desc_vi,
                    "translation_state": translation_state,
                    "raw_relation_evidence": ev_str,
                })
                continue

            semantic_type = "PLAYER_FACING_STATUS" if "status" in bid.lower() or "state" in bid.lower() else "PLAYER_FACING_BUFF"
            need_name = bool(name_cn) and not bool(name_vi)
            need_desc = bool(desc_cn) and not bool(desc_vi)

            if need_name or need_desc:
                translation_state = "MISSING_TRANSLATION"
                trans_field = "NAME_AND_DESC" if (need_name and need_desc) else ("NAME_ONLY" if need_name else "DESC_ONLY")

                in_early = bid in early_buffs
                if not in_early:
                    hist_class = "NEW_GAME_CONTENT"
                    why_missed = "Absent from 2026-09-07 early baseline, added in later game/master sync"
                else:
                    hist_class = "MISSED_BY_OLD_DIRECT_ONLY_TRAVERSAL"
                    why_missed = "Present in early masterdata but missed by old exporter which only checked direct {Buff_...} text tags in skill descriptions"

                src_skills = sorted(buff_to_skills.get(bid, []))
                src_skill_names_cn = [clean_tag(skill_rows.get(s, {}).get("skill_name_cn") or tables["skillMap"].get(s, {}).get("NameLanText") or s) for s in src_skills]
                src_skill_names_vi = [clean_tag(skill_rows.get(s, {}).get("skill_name_vi") or "") for s in src_skills]

                missing_units.append({
                    "character_id": cid,
                    "character_name_cn": char_rows.get(cid, {}).get("name_cn", ""),
                    "character_name_vi": char_rows.get(cid, {}).get("name_vi", ""),
                    "buff_id": bid,
                    "group_root_id": src_skills[0] if src_skills else "",
                    "buff_name_cn": name_cn,
                    "buff_name_vi": name_vi,
                    "buff_desc_cn": desc_cn,
                    "buff_desc_vi": desc_vi,
                    "translation_field": trans_field,
                    "semantic_type": semantic_type,
                    "source_skill_ids": ", ".join(src_skills),
                    "source_skill_names_cn": ", ".join(src_skill_names_cn),
                    "source_skill_names_vi": ", ".join(src_skill_names_vi),
                    "relation_path": path_str,
                    "parent_buff_id": parent_id,
                    "child_buff_ids": ", ".join(children),
                    "context_type": "SHARED_GENERIC" if not (cid in bid or bid.startswith(f"Buff_{cid}")) else "CHARACTER_PRIVATE",
                    "notes": "Keep Hán-Việt for buff identity; natural Vietnamese for description; preserve all placeholders [EffectParam,...]",
                    "provenance": f"MasterData:buffMap.json; Snapshot:{auth['selected_snapshot']}",
                    "missing_reason": hist_class,
                    "why_missed": why_missed,
                    "raw_evidence": ev_str,
                    "is_nested_only": bid not in direct_buffs,
                })
            else:
                translation_state = "ALREADY_TRANSLATED"
                approved_units.append({
                    "character_id": cid,
                    "character_name_vi": char_rows.get(cid, {}).get("name_vi", ""),
                    "buff_id": bid,
                    "buff_name_cn": name_cn,
                    "buff_name_vi": name_vi,
                    "buff_desc_cn": desc_cn,
                    "buff_desc_vi": desc_vi,
                    "relation_to_target": "PARENT_BUFF" if bid in parent_map else ("CHILD_BUFF" if bid in children_map else "SAME_FAMILY"),
                    "approval_status": authority_val or status_val or "EXISTING_VI",
                })

            graph_nodes.append({
                "character_id": cid,
                "skill_id": ", ".join(sorted(buff_to_skills.get(bid, []))),
                "root_buff_id": path_str.split(" -> ")[1] if " -> " in path_str else bid,
                "buff_id": bid,
                "parent_buff_id": parent_id,
                "depth": depth,
                "relation_path": path_str,
                "semantic_type": semantic_type,
                "player_facing": "YES",
                "name_cn": name_cn,
                "name_vi": name_vi,
                "desc_cn": desc_cn,
                "desc_vi": desc_vi,
                "translation_state": translation_state,
                "raw_relation_evidence": ev_str,
            })

        if missing_units:
            results_by_char[cid] = {
                "name_cn": char_rows.get(cid, {}).get("name_cn", ""),
                "name_vi": char_rows.get(cid, {}).get("name_vi", ""),
                "missing_units": missing_units,
                "approved_units": approved_units,
                "graph_nodes": graph_nodes,
                "affected_skills": sorted({s for u in missing_units for s in u["source_skill_ids"].split(", ") if s}),
                "source_blockers_count": sum(1 for n in graph_nodes if n["semantic_type"] == "SOURCE_BLOCKER"),
            }

    # 1. SUMMARY
    summary_rows = []
    for cid, info in sorted(results_by_char.items()):
        summary_rows.append({
            "character_id": cid,
            "name_cn": info["name_cn"],
            "name_vi": info["name_vi"],
            "missing_buff_count": len(info["missing_units"]),
            "source_changed_count": 0,
            "source_blocker_count": info["source_blockers_count"],
            "affected_skill_count": len(info["affected_skills"]),
            "selection_reason": f"Has {len(info['missing_units'])} untranslated player-facing buff(s) reachable via transitive closure"
        })

    # 2. TO_TRANSLATE & 8. AUDIT
    to_translate_rows = []
    audit_rows = []
    for cid, info in sorted(results_by_char.items()):
        for u in info["missing_units"]:
            to_translate_rows.append({
                "character_id": u["character_id"],
                "character_name_cn": u["character_name_cn"],
                "character_name_vi": u["character_name_vi"],
                "buff_id": u["buff_id"],
                "group_root_id": u["group_root_id"],
                "buff_name_cn": u["buff_name_cn"],
                "buff_name_vi": u["buff_name_vi"],
                "buff_desc_cn": u["buff_desc_cn"],
                "buff_desc_vi": u["buff_desc_vi"],
                "translation_field": u["translation_field"],
                "semantic_type": u["semantic_type"],
                "source_skill_ids": u["source_skill_ids"],
                "source_skill_names_cn": u["source_skill_names_cn"],
                "source_skill_names_vi": u["source_skill_names_vi"],
                "relation_path": u["relation_path"],
                "parent_buff_id": u["parent_buff_id"],
                "child_buff_ids": u["child_buff_ids"],
                "context_type": u["context_type"],
                "notes": u["notes"],
                "provenance": u["provenance"],
                "missing_reason": u["missing_reason"],
            })
            audit_rows.append({
                "character_id": u["character_id"],
                "buff_id": u["buff_id"],
                "discovery_method": "TRANSITIVE_POPUP_GRAPH_TRAVERSAL" if u["is_nested_only"] else "DIRECT_CARD_RELATION",
                "why_missed_historically": u["why_missed"],
                "historical_classification": u["missing_reason"],
                "graph_traversal_evidence": f"Path: {u['relation_path']}; Raw: {u['raw_evidence']}",
            })

    # 3. CHARACTER_CONTEXT
    character_context_rows = []
    target_cids = set(results_by_char.keys())
    for cid in sorted(target_cids):
        crow = char_rows.get(cid, {})
        character_context_rows.append({
            "character_id": cid,
            "character_name_vi": crow.get("name_vi", ""),
            "category": "CHARACTER",
            "item_id": cid,
            "slot_or_type": "BASE_INFO",
            "name_cn": crow.get("name_cn", ""),
            "name_vi": crow.get("name_vi", ""),
            "desc_cn": crow.get("fullname_cn", ""),
            "desc_vi": crow.get("fullname_vi", ""),
            "tags_or_notes": crow.get("tag_lan_text", ""),
        })
        c_skills = [row for row in wb_data["SKILL"] if str(row.get("character_id")) == cid]
        for s in c_skills:
            character_context_rows.append({
                "character_id": cid,
                "character_name_vi": crow.get("name_vi", ""),
                "category": "SKILL",
                "item_id": str(s.get("skill_id", "")),
                "slot_or_type": str(s.get("type_label") or s.get("skill_slot") or "SKILL"),
                "name_cn": str(s.get("skill_name_cn", "")),
                "name_vi": str(s.get("skill_name_vi", "")),
                "desc_cn": str(s.get("desc_cn", "")),
                "desc_vi": str(s.get("desc_vi", "")),
                "tags_or_notes": f"Group: {s.get('skill_group_id', '')}",
            })
        c_hz = [row for row in wb_data["HUANZHANG"] if str(row.get("character_id")) == cid]
        for hz in c_hz:
            character_context_rows.append({
                "character_id": cid,
                "character_name_vi": crow.get("name_vi", ""),
                "category": "HUANZHANG",
                "item_id": str(hz.get("brilliant_id", "")),
                "slot_or_type": "HUANZHANG_LORE",
                "name_cn": str(hz.get("icon_name_cn", "")),
                "name_vi": str(hz.get("icon_name_vi", "")),
                "desc_cn": str(hz.get("icon_info_cn", "")),
                "desc_vi": str(hz.get("icon_info_vi", "")),
                "tags_or_notes": str(hz.get("buff_show_cn", "")),
            })

    # 4. BUFF_GRAPH
    buff_graph_rows = []
    for cid, info in sorted(results_by_char.items()):
        for node in info["graph_nodes"]:
            buff_graph_rows.append(node)

    # 5. APPROVED_CONTEXT
    approved_context_rows = []
    seen_approved = set()
    for cid, info in sorted(results_by_char.items()):
        for a in info["approved_units"]:
            key_app = (cid, a["buff_id"])
            if key_app not in seen_approved:
                seen_approved.add(key_app)
                approved_context_rows.append(a)

    # 6. GLOSSARY_CONTEXT
    wb_master_read = openpyxl.load_workbook(MASTER, read_only=True)
    glossary_sheet = wb_master_read['GLOSSARY']
    g_headers = [c.value for c in next(glossary_sheet.iter_rows(max_row=1))]
    raw_glossary_rows = [dict(zip(g_headers, r)) for r in glossary_sheet.iter_rows(min_row=2, values_only=True)]
    wb_master_read.close()

    canonical_overrides = {
        "瞄准": ("Nhắm Bắn", "CANONICAL LOCK: 瞄准 -> Nhắm Bắn (overrides stale Miêu Chuẩn per owner rule)"),
        "流失": ("Mất Máu", "OWNER_CANONICAL_RULE: Common gameplay mechanic naturalization"),
        "眩晕": ("Choáng", "OWNER_CANONICAL_RULE: Common gameplay mechanic naturalization"),
        "沉睡": ("Ngủ Say", "OWNER_CANONICAL_RULE: Common gameplay mechanic naturalization"),
        "援护": ("Hộ Vệ", "OWNER_CANONICAL_RULE: Common gameplay mechanic naturalization"),
        "移动力": ("Điểm Di Chuyển", "OWNER_CANONICAL_RULE: Common mechanic naturalization; never use literal 'Di Động Lực'"),
    }

    target_cn_corpus = []
    for r in character_context_rows:
        target_cn_corpus.append(str(r["name_cn"] or ""))
        target_cn_corpus.append(str(r["desc_cn"] or ""))
    for r in to_translate_rows:
        target_cn_corpus.append(str(r["buff_name_cn"] or ""))
        target_cn_corpus.append(str(r["buff_desc_cn"] or ""))
    for r in buff_graph_rows:
        target_cn_corpus.append(str(r["name_cn"] or ""))
        target_cn_corpus.append(str(r["desc_cn"] or ""))
    full_target_cn_text = "\n".join(target_cn_corpus)

    glossary_context_rows = []
    seen_terms = set()
    for g in raw_glossary_rows:
        term_cn = str(g.get("term_cn") or "").strip()
        if not term_cn or term_cn in seen_terms: continue
        if term_cn in full_target_cn_text:
            seen_terms.add(term_cn)
            term_vi = str(g.get("term_vi") or "").strip()
            notes = str(g.get("notes") or "").strip()
            status_val = str(g.get("status") or "")
            if term_cn in canonical_overrides:
                term_vi, override_notes = canonical_overrides[term_cn]
                notes = override_notes
                status_val = "OWNER_APPROVED"
            glossary_context_rows.append({
                "term_id": str(g.get("term_id") or f"GLOSSARY_{len(glossary_context_rows)+1:03d}"),
                "category": str(g.get("category") or "gameplay_terminology"),
                "term_cn": term_cn,
                "term_vi": term_vi,
                "han_viet": str(g.get("han_viet") or ""),
                "status": status_val,
                "notes": notes,
            })

    for term_cn, (term_vi, note) in canonical_overrides.items():
        if term_cn not in seen_terms and term_cn in full_target_cn_text:
            seen_terms.add(term_cn)
            glossary_context_rows.append({
                "term_id": f"OWNER_RULE_{len(glossary_context_rows)+1:03d}",
                "category": "owner_canonical_rule",
                "term_cn": term_cn,
                "term_vi": term_vi,
                "han_viet": "",
                "status": "OWNER_APPROVED",
                "notes": note,
            })

    # 7. SOURCE_ISSUES
    source_issues_rows = all_source_issues

    sheets_payload = {
        "SUMMARY": summary_rows,
        "TO_TRANSLATE": to_translate_rows,
        "CHARACTER_CONTEXT": character_context_rows,
        "BUFF_GRAPH": buff_graph_rows,
        "APPROVED_CONTEXT": approved_context_rows,
        "GLOSSARY_CONTEXT": glossary_context_rows,
        "SOURCE_ISSUES": source_issues_rows,
        "AUDIT": audit_rows,
    }

    if output_path is None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
        output_path = EXPORTS_DIR / f"buff_closure_translation_packet_{stamp}.xlsx"

    write_packet_from_dict(sheets_payload, output_path)

    out_hash = sha256_file(output_path)
    post_master_hash = sha256_file(MASTER)
    assert initial_master_hash == post_master_hash, "Master mutation detected!"

    counts = {s: len(r) for s, r in sheets_payload.items()}
    return output_path, out_hash, counts


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Export character buff-closure translation packet")
    parser.add_argument("--output", type=str, default=None, help="Custom output XLSX path")
    args = parser.parse_args()

    out_p = Path(args.output) if args.output else None
    out_file, out_sha, counts = run_export(out_p)
    print(json.dumps({
        "output_path": str(out_file),
        "sha256": out_sha,
        "counts": counts,
        "master_unchanged": True
    }, indent=2, ensure_ascii=False))
