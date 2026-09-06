import openpyxl
import json
import os
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

def extract_character_batch(character_id):
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    master_path = os.path.join(base_dir, 'localization', 'localization_master.xlsx')
    gameplay_path = os.path.join(base_dir, 'localization', 'translation_gameplay.xlsx')
    out_dir = os.path.join(base_dir, 'localization', 'batch_context')

    if not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    print(f"Loading master workbook: {master_path}...")
    wb_master = openpyxl.load_workbook(master_path, data_only=True)
    
    print(f"Loading relational links: {gameplay_path}...")
    wb_gameplay = openpyxl.load_workbook(gameplay_path, data_only=True) if os.path.exists(gameplay_path) else None

    # 1. Read CHARACTER row (READ-ONLY context)
    ws_char = wb_master['CHARACTER']
    char_rows = list(ws_char.iter_rows(values_only=True))[1:]
    char_rec = None
    for r in char_rows:
        if r[0] == character_id:
            char_rec = {
                'character_id': r[0],
                'name_cn': r[1],
                'name_vi': r[2],
                'fullname_cn': r[3],
                'fullname_vi': r[4],
                'nickname_vi': r[5],
                'tags_cn': r[6],
                'tags_vi': r[7],
                'rare': r[8],
                'status': r[10]
            }
            break

    if not char_rec:
        print(f"ERROR: Character ID '{character_id}' not found in CHARACTER sheet.")
        sys.exit(1)

    # 2. Read SKILL raw rows
    ws_skill = wb_master['SKILL']
    skill_rows = list(ws_skill.iter_rows(values_only=True))[1:]
    character_skills = []
    skill_ids_set = set()

    for r in skill_rows:
        if r[1] == character_id:
            sid = r[0]
            skill_ids_set.add(sid)
            desc_cn = str(r[7]) if len(r) > 7 and r[7] is not None else ''
            
            # extract placeholders and rich-text tags
            placeholders = re.findall(r'\[Effect[^\]]+\]', desc_cn)
            color_tags = re.findall(r'</?color[^>]*>', desc_cn)

            character_skills.append({
                'skill_id': sid,
                'character_id': r[1],
                'skill_group_id': r[2],
                'skill_slot': r[3],
                'type_label': r[4],
                'skill_name_cn': r[5],
                'skill_name_vi': r[6],
                'desc_cn': desc_cn,
                'desc_vi': r[8],
                'placeholders': placeholders,
                'color_tags': color_tags,
                'status': r[10]
            })

    # 3. Read SKILL_BUFF_LINKS (Primary deterministic relational edge source)
    verified_buff_ids = set()
    link_evidence = []
    if wb_gameplay and 'SKILL_BUFF_LINKS' in wb_gameplay.sheetnames:
        ws_links = wb_gameplay['SKILL_BUFF_LINKS']
        for r in list(ws_links.iter_rows(values_only=True))[1:]:
            if r[0] == character_id or r[2] in skill_ids_set:
                bid = str(r[4]).strip() if r[4] else None
                if bid and bid != 'None':
                    verified_buff_ids.add(bid)
                    link_evidence.append({
                        'skill_id': r[2],
                        'buff_id': bid,
                        'buff_name_cn': r[5],
                        'source': r[6],
                        'confidence': r[7],
                        'notes': r[8]
                    })

    # 4. Read BUFF_STATUS rows and categorize
    ws_buff = wb_master['BUFF_STATUS']
    buff_rows = list(ws_buff.iter_rows(values_only=True))[1:]
    
    verified_buffs = []
    context_only_buffs = []
    unresolved_buffs = []

    # Map of all buff_id to row
    all_buff_map = {str(b[0]): b for b in buff_rows}

    for bid in sorted(verified_buff_ids):
        if bid in all_buff_map:
            b = all_buff_map[bid]
            verified_buffs.append({
                'buff_id': b[0],
                'group_root_id': b[1],
                'buff_name_cn': b[2],
                'buff_name_vi': b[3],
                'buff_desc_cn': b[4],
                'buff_desc_vi': b[5],
                'classification_scope': b[6],
                'status': b[8],
                'provenance': 'SKILL_BUFF_LINKS Verified Edge'
            })
        else:
            unresolved_buffs.append({
                'buff_id': bid,
                'reason': 'Linked in SKILL_BUFF_LINKS but missing from BUFF_STATUS sheet'
            })

    # Check for character-specific buffs that match character ID
    for b in buff_rows:
        bid = str(b[0])
        if (character_id in bid) and bid not in verified_buff_ids:
            context_only_buffs.append({
                'buff_id': b[0],
                'group_root_id': b[1],
                'buff_name_cn': b[2],
                'buff_name_vi': b[3],
                'buff_desc_cn': b[4],
                'buff_desc_vi': b[5],
                'classification_scope': b[6],
                'status': b[8],
                'provenance': 'Character-ID Match (Context Only)'
            })

    # 5. Read ZHIZHI rows
    ws_zhizhi = wb_master['ZHIZHI']
    zhizhi_rows = [r for r in list(ws_zhizhi.iter_rows(values_only=True))[1:] if r[0] == character_id]
    character_zhizhi = []
    for z in zhizhi_rows:
        character_zhizhi.append({
            'character_id': z[0],
            'star': z[1],
            'rank_numeral_cn': z[2],
            'rank_numeral_vi': z[3],
            'effect_type': z[4],
            'effect_summary_cn': z[5],
            'effect_summary_vi': z[6],
            'skill_up_base_id': z[7],
            'skill_up_ex_id': z[8],
            'status': z[10]
        })

    # 6. Read HUANZHANG rows
    ws_hz = wb_master['HUANZHANG']
    hz_rows = [r for r in list(ws_hz.iter_rows(values_only=True))[1:] if r[1] == character_id]
    character_hz = []
    for h in hz_rows:
        character_hz.append({
            'brilliant_id': h[0],
            'character_id': h[1],
            'icon_name_cn': h[2],
            'icon_name_vi': h[3],
            'icon_info_cn': h[4],
            'icon_info_vi': h[5],
            'buff_show_cn': h[6],
            'buff_show_vi': h[7],
            'status': h[9]
        })

    # 7. Read GLOSSARY context (Approved vs Review)
    ws_gl = wb_master['GLOSSARY']
    gl_rows = list(ws_gl.iter_rows(values_only=True))[1:]
    glossary_approved = []
    glossary_pending_review = []

    for g in gl_rows:
        gitem = {
            'term_id': g[0],
            'category': g[1],
            'term_cn': g[2],
            'term_vi': g[3],
            'han_viet': g[4],
            'status': g[6]
        }
        if g[6] == 'APPROVED':
            glossary_approved.append(gitem)
        else:
            glossary_pending_review.append(gitem)

    # Compile Context Bundle
    bundle = {
        'metadata': {
            'character_id': character_id,
            'character_name_vi': char_rec['name_vi'],
            'character_name_cn': char_rec['name_cn'],
            'extracted_at': 'Phase 3 Context Pipeline',
            'primary_link_source': 'localization/translation_gameplay.xlsx (SKILL_BUFF_LINKS sheet)'
        },
        'character_context': char_rec,
        'skill_translation_targets': character_skills,
        'buff_references': {
            'verified_references': verified_buffs,
            'context_only_references': context_only_buffs,
            'unresolved_references': unresolved_buffs,
            'link_evidence_graph': link_evidence
        },
        'zhizhi_translation_targets': character_zhizhi,
        'huanzhang_translation_targets': character_hz,
        'glossary_context': {
            'approved_terms': glossary_approved,
            'review_pending_terms': glossary_pending_review
        }
    }

    out_file = os.path.join(out_dir, f"{character_id}_context.json")
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump(bundle, f, ensure_ascii=False, indent=2)

    print(f"\n[OK] Extracted context bundle for {character_id} ({char_rec['name_vi']})")
    print(f"     Saved to: {out_file}")
    print(f"     SKILL targets: {len(character_skills)}")
    print(f"     BUFF VERIFIED_REFERENCE: {len(verified_buffs)}")
    print(f"     BUFF CONTEXT_ONLY: {len(context_only_buffs)}")
    print(f"     BUFF UNRESOLVED: {len(unresolved_buffs)}")
    print(f"     ZHIZHI targets: {len(character_zhizhi)}")
    print(f"     HUANZHANG targets: {len(character_hz)}")

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python tools/extract_character_batch.py <CHARACTER_ID>")
        sys.exit(1)
    extract_character_batch(sys.argv[1])
