"""Export the 89 gameplay Han-leak review packet for ChatGPT/owner review.

This tool is strictly an exporter. It performs NO translation.
It produces the review artifact:
    localization/reviews/han_leak_gameplay_cleanup_<timestamp>.xlsx
"""
from __future__ import annotations

from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import re
import sys
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
MASTER_PATH = BASE_DIR / "localization" / "localization_master.xlsx"
REVIEWS_DIR = BASE_DIR / "localization" / "reviews"

HAN_RE = re.compile(r"[\u4e00-\u9fff]")
WHITELIST = ["《鹿王本生图》", "雷威"]
BUFF_MARKER_RE = re.compile(r"\{?(Buff_[A-Za-z0-9_]+)\}?")
COLOR_TAG_RE = re.compile(r"</?color[^>]*>")
PLACEHOLDER_RE = re.compile(r"\[Effect[^\]]+\]")
SPAN_RE = re.compile(r"<color=[^>]*>.*?</color>", re.DOTALL)
TAG_RE = re.compile(r"<[^>]+>")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def visible_text(value: Any) -> str:
    return TAG_RE.sub("", str(value or "")).strip()


def extract_character_id_from_buff(buff_id: str) -> str:
    """Extract character ID from buff IDs like Buff_A0162_18_1ex -> A0162."""
    match = re.search(r"Buff_([A-Z]\d{4})", buff_id)
    return match.group(1) if match else ""


def build_packet() -> tuple[Path, str, int, int, int]:
    wb_master = openpyxl.load_workbook(MASTER_PATH, data_only=True)

    # 1. Build CHARACTER lookup
    characters: dict[str, dict[str, str]] = {}
    ws_char = wb_master["CHARACTER"]
    char_headers = [c for c in next(ws_char.iter_rows(min_row=1, max_row=1, values_only=True))]
    for row in ws_char.iter_rows(min_row=2, values_only=True):
        item = dict(zip(char_headers, row))
        cid = str(item.get("character_id") or "").strip()
        if cid:
            characters[cid] = {
                "name_cn": str(item.get("name_cn") or "").strip(),
                "name_vi": str(item.get("name_vi") or "").strip(),
                "fullname_cn": str(item.get("fullname_cn") or "").strip(),
                "fullname_vi": str(item.get("fullname_vi") or "").strip(),
            }

    # 2. Build BUFF registry
    buffs: dict[str, dict[str, str]] = {}
    ws_buff = wb_master["BUFF_STATUS"]
    buff_headers = [c for c in next(ws_buff.iter_rows(min_row=1, max_row=1, values_only=True))]
    for row in ws_buff.iter_rows(min_row=2, values_only=True):
        item = dict(zip(buff_headers, row))
        bid = str(item.get("buff_id") or "").strip()
        if bid:
            b_cn = visible_text(item.get("buff_name_cn"))
            b_vi = visible_text(item.get("buff_name_vi"))
            auth = str(item.get("name_authority") or item.get("status") or "").strip()
            buffs[bid] = {
                "name_cn": b_cn,
                "name_vi": b_vi,
                "auth": auth,
            }

    # 3. Find 83 SKILL Han leaks
    ws_skill = wb_master["SKILL"]
    skill_headers = [c for c in next(ws_skill.iter_rows(min_row=1, max_row=1, values_only=True))]
    skill_rows: list[dict[str, Any]] = []
    for row in ws_skill.iter_rows(min_row=2, values_only=True):
        item = dict(zip(skill_headers, row))
        vi_val = str(item.get("desc_vi") or "")
        if vi_val:
            check_val = vi_val
            for w in WHITELIST:
                check_val = check_val.replace(w, "")
            if HAN_RE.search(check_val):
                skill_rows.append(item)

    # 4. Find 6 BUFF_STATUS Han leaks
    buff_rows: list[dict[str, Any]] = []
    for row in ws_buff.iter_rows(min_row=2, values_only=True):
        item = dict(zip(buff_headers, row))
        vi_val = str(item.get("buff_desc_vi") or "")
        if vi_val:
            check_val = vi_val
            for w in WHITELIST:
                check_val = check_val.replace(w, "")
            if HAN_RE.search(check_val):
                buff_rows.append(item)

    wb_master.close()

    if len(skill_rows) != 83 or len(buff_rows) != 6:
        raise RuntimeError(
            f"UNEXPECTED_GAMEPLAY_HAN_LEAK_COUNT skill={len(skill_rows)} (expected 83) "
            f"buff={len(buff_rows)} (expected 6)"
        )

    # 5. Format the rows for export
    output_rows: list[dict[str, Any]] = []

    # Process SKILL rows
    for item in skill_rows:
        cid = str(item.get("character_id") or "").strip()
        source_cn = str(item.get("desc_cn") or "")
        current_vi = str(item.get("desc_vi") or "")

        # Canonical character context
        char_info = characters.get(cid, {})
        char_ctx = f"{cid}: {char_info.get('name_cn', '')} ({char_info.get('name_vi', '')})"
        if char_info.get("fullname_vi"):
            char_ctx += f" | {char_info.get('fullname_vi')}"

        # Resolve buff references
        markers = set(BUFF_MARKER_RE.findall(source_cn) + BUFF_MARKER_RE.findall(current_vi))
        resolved_buffs = []
        for bid in sorted(markers):
            if bid in buffs:
                b = buffs[bid]
                resolved_buffs.append(f"{{{bid}}}: {b['name_cn']} -> {b['name_vi']} [{b['auth']}]")
        # Also check for character-owned buffs
        if cid:
            char_buff_prefix = f"Buff_{cid}_"
            for bid, b in buffs.items():
                if bid.startswith(char_buff_prefix) and b["name_cn"] and b["name_cn"] in source_cn:
                    entry = f"{{{bid}}}: {b['name_cn']} -> {b['name_vi']} [{b['auth']}]"
                    if entry not in resolved_buffs:
                        resolved_buffs.append(entry)

        output_rows.append({
            "sheet": "SKILL",
            "record_id": str(item.get("skill_id") or "").strip(),
            "character_id": cid,
            "field": "desc_vi",
            "source_cn": source_cn,
            "current_vi": current_vi,
            "skill_or_buff_name_cn": str(item.get("skill_name_cn") or "").strip(),
            "skill_or_buff_name_vi": str(item.get("skill_name_vi") or "").strip(),
            "status": str(item.get("status") or "").strip(),
            "translation_review_status": str(item.get("translation_review_status") or "").strip(),
            "name_authority": "",
            "desc_translation_status": "",
            "source_placeholders": ", ".join(sorted(PLACEHOLDER_RE.findall(source_cn))),
            "current_vi_placeholders": ", ".join(sorted(PLACEHOLDER_RE.findall(current_vi))),
            "source_color_tokens": ", ".join(COLOR_TAG_RE.findall(source_cn)),
            "current_vi_color_tokens": ", ".join(COLOR_TAG_RE.findall(current_vi)),
            "source_tagged_spans": " | ".join(SPAN_RE.findall(source_cn)),
            "current_vi_tagged_spans": " | ".join(SPAN_RE.findall(current_vi)),
            "canonical_character_context": char_ctx,
            "resolved_canonical_buff_references": " ; ".join(resolved_buffs),
            "translation_proposed_vi": "",
            "review_notes": "",
        })

    # Process BUFF_STATUS rows
    for item in buff_rows:
        bid = str(item.get("buff_id") or "").strip()
        cid = extract_character_id_from_buff(bid)
        source_cn = str(item.get("buff_desc_cn") or "")
        current_vi = str(item.get("buff_desc_vi") or "")

        char_info = characters.get(cid, {})
        char_ctx = f"{cid}: {char_info.get('name_cn', '')} ({char_info.get('name_vi', '')})" if cid else ""
        if char_info.get("fullname_vi"):
            char_ctx += f" | {char_info.get('fullname_vi')}"

        markers = set(BUFF_MARKER_RE.findall(source_cn) + BUFF_MARKER_RE.findall(current_vi))
        resolved_buffs = []
        for ref_bid in sorted(markers):
            if ref_bid in buffs:
                b = buffs[ref_bid]
                resolved_buffs.append(f"{{{ref_bid}}}: {b['name_cn']} -> {b['name_vi']} [{b['auth']}]")

        output_rows.append({
            "sheet": "BUFF_STATUS",
            "record_id": bid,
            "character_id": cid,
            "field": "buff_desc_vi",
            "source_cn": source_cn,
            "current_vi": current_vi,
            "skill_or_buff_name_cn": str(item.get("buff_name_cn") or "").strip(),
            "skill_or_buff_name_vi": str(item.get("buff_name_vi") or "").strip(),
            "status": str(item.get("status") or "").strip(),
            "translation_review_status": "",
            "name_authority": str(item.get("name_authority") or "").strip(),
            "desc_translation_status": str(item.get("desc_translation_status") or "").strip(),
            "source_placeholders": ", ".join(sorted(PLACEHOLDER_RE.findall(source_cn))),
            "current_vi_placeholders": ", ".join(sorted(PLACEHOLDER_RE.findall(current_vi))),
            "source_color_tokens": ", ".join(COLOR_TAG_RE.findall(source_cn)),
            "current_vi_color_tokens": ", ".join(COLOR_TAG_RE.findall(current_vi)),
            "source_tagged_spans": " | ".join(SPAN_RE.findall(source_cn)),
            "current_vi_tagged_spans": " | ".join(SPAN_RE.findall(current_vi)),
            "canonical_character_context": char_ctx,
            "resolved_canonical_buff_references": " ; ".join(resolved_buffs),
            "translation_proposed_vi": "",
            "review_notes": "",
        })

    # 6. Create Excel workbook
    wb_out = openpyxl.Workbook()
    # Main sheet
    ws_main = wb_out.active
    ws_main.title = "GAMEPLAY_HAN_LEAKS"

    columns = [
        "sheet",
        "record_id",
        "character_id",
        "field",
        "source_cn",
        "current_vi",
        "skill_or_buff_name_cn",
        "skill_or_buff_name_vi",
        "status",
        "translation_review_status",
        "name_authority",
        "desc_translation_status",
        "source_placeholders",
        "current_vi_placeholders",
        "source_color_tokens",
        "current_vi_color_tokens",
        "source_tagged_spans",
        "current_vi_tagged_spans",
        "canonical_character_context",
        "resolved_canonical_buff_references",
        "translation_proposed_vi",
        "review_notes",
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws_main.append(columns)
    for col_idx in range(1, len(columns) + 1):
        cell = ws_main.cell(1, col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_data in output_rows:
        ws_main.append([row_data[col] for col in columns])

    for row_idx in range(2, len(output_rows) + 2):
        for col_idx in range(1, len(columns) + 1):
            cell = ws_main.cell(row_idx, col_idx)
            cell.font = Font(name="Calibri", size=10)
            cell.border = border
            # Highlight editable empty review columns in light yellow
            col_name = columns[col_idx - 1]
            if col_name in ("translation_proposed_vi", "review_notes"):
                cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws_main.row_dimensions[1].height = 28
    for col_idx, col_name in enumerate(columns, 1):
        letter = get_column_letter(col_idx)
        if col_name in ("source_cn", "current_vi", "translation_proposed_vi"):
            ws_main.column_dimensions[letter].width = 45
        elif col_name in ("source_tagged_spans", "current_vi_tagged_spans", "resolved_canonical_buff_references"):
            ws_main.column_dimensions[letter].width = 35
        elif col_name in ("canonical_character_context", "review_notes"):
            ws_main.column_dimensions[letter].width = 30
        elif col_name in ("skill_or_buff_name_cn", "skill_or_buff_name_vi"):
            ws_main.column_dimensions[letter].width = 20
        elif col_name in ("record_id", "source_placeholders", "current_vi_placeholders"):
            ws_main.column_dimensions[letter].width = 20
        else:
            ws_main.column_dimensions[letter].width = 15

    # QA Sheet
    ws_qa = wb_out.create_sheet(title="QA")
    ws_qa.append(["check", "result", "details"])
    ws_qa.append(["total_gameplay_han_leaks", "PASS", str(len(output_rows))])
    ws_qa.append(["skill_desc_vi_count", "PASS", str(len(skill_rows))])
    ws_qa.append(["buff_desc_vi_count", "PASS", str(len(buff_rows))])
    ws_qa.append(["profile_rows_excluded", "PASS", "0 PROFILE rows included (backlog deferred)"])
    ws_qa.append(["untranslated_proposals_blank", "PASS", "All 89 translation_proposed_vi left blank"])
    ws_qa.append(["exported_at", "INFO", datetime.now().isoformat()])

    for cell in ws_qa[1]:
        cell.font = header_font
        cell.fill = header_fill

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = REVIEWS_DIR / f"han_leak_gameplay_cleanup_{timestamp}.xlsx"
    wb_out.save(out_path)
    wb_out.close()

    file_sha = sha256(out_path)
    return out_path, file_sha, len(skill_rows), len(buff_rows), len(output_rows)


def main() -> int:
    if getattr(sys.stdout, "encoding", "").lower() != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8")
    REVIEWS_DIR.mkdir(parents=True, exist_ok=True)
    path, file_sha, skill_cnt, buff_cnt, total_cnt = build_packet()
    summary = {
        "packet_path": str(path),
        "packet_sha256": file_sha,
        "skill_count": skill_cnt,
        "buff_status_count": buff_cnt,
        "total_count": total_cnt,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
