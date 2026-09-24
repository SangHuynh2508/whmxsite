#!/usr/bin/env python3
"""WHMX - Apply All 5 Reviewed Skin Lore Translation Packets.

Safely applies 145 manually reviewed skin lore translations into
localization/localization_master.xlsx (SKIN sheet -> desc_vi).

Safety guarantees:
- Computes SHA-256 before & after
- Creates timestamped backup under localization/backups/ before modifying
- Validates all 5 packets (145 rows, exact source_cn match) before writing
- Verifies post-apply integrity
"""
from __future__ import annotations

import hashlib
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

PROJECT_ROOT = Path(__file__).resolve().parent.parent
MASTER_PATH = PROJECT_ROOT / "localization" / "localization_master.xlsx"
BACKUPS_DIR = PROJECT_ROOT / "localization" / "backups"
EXPORTS_DIR = PROJECT_ROOT / "localization" / "exports"

PACKET_FILENAMES = [
    "skin_lore_translation_packet_01_20260918_215226_translated_vi.xlsx",
    "skin_lore_translation_packet_02_20260918_215226_translated_vi.xlsx",
    "skin_lore_translation_packet_03_20260918_215226_translated_vi.xlsx",
    "skin_lore_translation_packet_04_20260918_215226_translated_vi.xlsx",
    "skin_lore_translation_packet_05_20260918_215226_translated_vi.xlsx",
]


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def main():
    print("=== STEP 1: PRE-APPLY HASH & BACKUP ===")
    assert MASTER_PATH.exists(), f"Missing master file: {MASTER_PATH}"
    pre_sha = sha256_file(MASTER_PATH)
    print(f"localization_master.xlsx SHA-256 (Pre-apply): {pre_sha}")

    BACKUPS_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = BACKUPS_DIR / f"localization_master_before_skin_lore_apply_{ts}.xlsx"
    assert not backup_path.exists(), f"Backup already exists: {backup_path}"
    shutil.copy2(MASTER_PATH, backup_path)
    backup_sha = sha256_file(backup_path)
    assert pre_sha == backup_sha, "Backup SHA-256 mismatch!"
    print(f"Created backup at: {backup_path}")
    print(f"Backup SHA-256:    {backup_sha}")

    print("\n=== STEP 2: LOAD & VALIDATE ALL 5 PACKETS ===")
    all_packet_rows: list[dict[str, str]] = []
    packet_counts: dict[str, int] = {}

    for fn in PACKET_FILENAMES:
        fp = EXPORTS_DIR / fn
        assert fp.exists(), f"Missing packet file: {fp}"
        wb_p = openpyxl.load_workbook(fp, data_only=True, read_only=True)
        ws_p = wb_p.active
        rows = list(ws_p.iter_rows(values_only=True))
        header = rows[0]
        data_rows = rows[1:]

        assert len(data_rows) == 29, f"{fn}: expected 29 rows, got {len(data_rows)}"
        packet_counts[fn] = len(data_rows)

        cols = {str(name): idx for idx, name in enumerate(header)}
        for req_col in ["skin_id", "field_type", "source_cn", "translation_vi"]:
            assert req_col in cols, f"{fn} missing column {req_col}"

        for idx, r in enumerate(data_rows, start=2):
            sid = str(r[cols["skin_id"]]).strip() if r[cols["skin_id"]] is not None else ""
            ft = str(r[cols["field_type"]]).strip() if r[cols["field_type"]] is not None else ""
            scn = str(r[cols["source_cn"]]).strip() if r[cols["source_cn"]] is not None else ""
            tvi = str(r[cols["translation_vi"]]).strip() if r[cols["translation_vi"]] is not None else ""

            assert sid, f"{fn} row {idx}: empty skin_id"
            assert ft == "desc", f"{fn} row {idx}: invalid field_type '{ft}' (expected 'desc')"
            assert scn, f"{fn} row {idx}: empty source_cn"
            assert tvi, f"{fn} row {idx}: empty translation_vi for skin {sid}"

            all_packet_rows.append({
                "skin_id": sid,
                "field_type": ft,
                "source_cn": scn,
                "translation_vi": tvi,
                "packet": fn,
            })

    total_packet_rows = len(all_packet_rows)
    print(f"Loaded {len(PACKET_FILENAMES)} packets, total rows: {total_packet_rows}")
    assert total_packet_rows == 145, f"Expected 145 rows, got {total_packet_rows}"

    unique_sids = set(r["skin_id"] for r in all_packet_rows)
    assert len(unique_sids) == 145, f"Duplicate skin_ids found! {len(unique_sids)} unique"

    print("\n=== STEP 3: CROSS-VALIDATE PACKETS WITH MASTER SKIN SHEET ===")
    master_ro_wb = openpyxl.load_workbook(MASTER_PATH, data_only=True, read_only=True)
    skin_ws_ro = master_ro_wb["SKIN"]
    master_rows = list(skin_ws_ro.iter_rows(values_only=True))
    m_header = master_rows[0]
    m_cols = {str(name): idx for idx, name in enumerate(m_header)}
    master_skins = {}
    for r in master_rows[1:]:
        sid = str(r[m_cols["skin_id"]]).strip() if r[m_cols["skin_id"]] is not None else ""
        d_cn = str(r[m_cols["desc_cn"]]).strip() if r[m_cols["desc_cn"]] is not None else ""
        master_skins[sid] = d_cn

    print(f"Master SKIN rows: {len(master_skins)}")
    assert len(master_skins) == 145, f"Expected 145 skins in master, found {len(master_skins)}"

    mismatches = []
    for r in all_packet_rows:
        sid = r["skin_id"]
        if sid not in master_skins:
            mismatches.append(f"{sid}: Not found in master SKIN sheet")
        elif r["source_cn"] != master_skins[sid]:
            mismatches.append(f"{sid}: source_cn mismatch!\n  Master: {master_skins[sid]}\n  Packet: {r['source_cn']}")

    if mismatches:
        print(f"FATAL: Found {len(mismatches)} mismatches! Aborting.")
        for m in mismatches:
            print(" -", m)
        sys.exit(1)
    print("[VALIDATION PASSED] All 145 packet rows match current master source_cn exactly!")

    print("\n=== STEP 4: APPLY TRANSLATIONS TO LOCALIZATION MASTER ===")
    # Load editable workbook
    master_wb = openpyxl.load_workbook(MASTER_PATH)
    skin_ws = master_wb["SKIN"]

    header_row = [cell.value for cell in skin_ws[1]]
    skin_id_col_idx = header_row.index("skin_id") + 1
    desc_vi_col_idx = header_row.index("desc_vi") + 1

    packet_map = {r["skin_id"]: r["translation_vi"] for r in all_packet_rows}
    applied_count = 0

    for row_idx in range(2, skin_ws.max_row + 1):
        cell_sid = skin_ws.cell(row=row_idx, column=skin_id_col_idx).value
        if cell_sid is not None:
            sid_str = str(cell_sid).strip()
            if sid_str in packet_map:
                target_cell = skin_ws.cell(row=row_idx, column=desc_vi_col_idx)
                target_cell.value = packet_map[sid_str]
                applied_count += 1

    print(f"Applied translations to {applied_count} rows in SKIN sheet.")
    assert applied_count == 145, f"Expected 145 applied translations, got {applied_count}"

    master_wb.save(MASTER_PATH)
    print("Saved modified master workbook successfully.")

    print("\n=== STEP 5: POST-APPLY AUDIT ===")
    post_sha = sha256_file(MASTER_PATH)
    print(f"localization_master.xlsx SHA-256 (Post-apply): {post_sha}")
    assert post_sha != pre_sha, "File hash did not change after save!"

    # Reload read-only to audit
    audit_wb = openpyxl.load_workbook(MASTER_PATH, data_only=True, read_only=True)
    audit_ws = audit_wb["SKIN"]
    a_rows = list(audit_ws.iter_rows(values_only=True))
    a_header = a_rows[0]
    a_cols = {str(name): idx for idx, name in enumerate(a_header)}

    missing_desc_vi = 0
    suspicious_copied = 0
    verified_matches = 0

    for r in a_rows[1:]:
        sid = str(r[a_cols["skin_id"]]).strip()
        d_vi = str(r[a_cols["desc_vi"]]).strip() if r[a_cols["desc_vi"]] is not None else ""
        d_cn = str(r[a_cols["desc_cn"]]).strip() if r[a_cols["desc_cn"]] is not None else ""
        o_vi = str(r[a_cols["obtain_vi"]]).strip() if r[a_cols["obtain_vi"]] is not None else ""
        o_cn = str(r[a_cols["obtain_cn"]]).strip() if r[a_cols["obtain_cn"]] is not None else ""

        if not d_vi:
            missing_desc_vi += 1

        expected_vi = packet_map.get(sid, "")
        if d_vi == expected_vi:
            verified_matches += 1
        else:
            print(f"[AUDIT FAIL] Mismatch for {sid}!")

        if d_vi and o_vi and d_vi == o_vi and d_cn != o_cn:
            suspicious_copied += 1
            print(f"[SUSPICIOUS] {sid}: desc_vi == obtain_vi ('{d_vi}')")

    print(f"Total skin rows: {len(a_rows) - 1}")
    print(f"Verified matches with packet translations: {verified_matches} / 145")
    print(f"Remaining missing desc_vi: {missing_desc_vi}")
    print(f"Suspicious desc_vi == obtain_vi cases: {suspicious_copied}")
    assert verified_matches == 145, "Post-apply verification failed: not all 145 matched"
    assert missing_desc_vi == 0, f"Remaining missing desc_vi: {missing_desc_vi}"
    assert suspicious_copied == 0, f"Suspicious desc_vi == obtain_vi found: {suspicious_copied}"

    print("\n[SUCCESS] ALL 145 SKIN LORE TRANSLATIONS APPLIED AND VERIFIED PERFECTLY!")


if __name__ == "__main__":
    main()
