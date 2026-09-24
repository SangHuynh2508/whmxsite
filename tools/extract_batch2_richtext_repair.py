"""Extract Batch #2 rich-text repair diagnostics without mutating translations."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "localization/localization_master.xlsx"
DEFAULT_BEFORE = ROOT / "localization/backups/localization_master_20260913_153106_372204.xlsx"
PLACEHOLDER_RE = re.compile(r"\[Effect[^\]]+\]")
COLOR_TOKEN_RE = re.compile(r"</?color[^>]*>")
COLOR_OPEN_RE = re.compile(r"<color=(#[0-9A-Fa-f]{6})>")
COLOR_SPAN_RE = re.compile(r"(<color=(#[0-9A-Fa-f]{6})>)(.*?)(</color>)", re.S)


def text(value):
    return "" if value is None else str(value)


def sha256(path: Path):
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def load_sheet_rows(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        result = {}
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [text(value).strip() for value in rows[0]]
            records = []
            for row_number, values in enumerate(rows[1:], start=2):
                if not any(value not in (None, "") for value in values):
                    continue
                record = dict(zip(headers, values))
                record["_row_number"] = row_number
                records.append(record)
            result[sheet.title] = {"headers": headers, "records": records}
        return result
    finally:
        workbook.close()


def cn_vi_pairs(headers):
    pairs = []
    for header in headers:
        lowered = header.lower()
        if "cn" not in lowered and "zh" not in lowered:
            continue
        vi_header = lowered.replace("cn", "vi").replace("zh", "vi")
        for candidate in headers:
            if candidate.lower() == vi_header and candidate.lower() not in {"status", "notes", "provenance_vi"}:
                pairs.append((header, candidate))
                break
    return pairs


def primary_id(sheet_name, record):
    for field in ("skill_id", "buff_id", "brilliant_id", "character_id"):
        if record.get(field) not in (None, ""):
            return text(record.get(field))
    return f"Row_{record.get('_row_number')}"


def indexed_records(data):
    output = {}
    for sheet_name, sheet_data in data.items():
        output[sheet_name] = {primary_id(sheet_name, record): record for record in sheet_data["records"]}
    return output


def placeholders(value):
    return PLACEHOLDER_RE.findall(text(value))


def color_tokens(value):
    return COLOR_TOKEN_RE.findall(text(value))


def color_open_colors(value):
    return [match.group(1).lower() for match in COLOR_OPEN_RE.finditer(text(value))]


def color_spans(value):
    spans = []
    for index, match in enumerate(COLOR_SPAN_RE.finditer(text(value)), start=1):
        spans.append(
            {
                "index": index,
                "color": match.group(2).lower(),
                "visible_text": match.group(3),
                "opening_tag": match.group(1),
                "closing_tag": match.group(4),
            }
        )
    return spans


def token_mismatch(source, target):
    return placeholders(source) != placeholders(target) or color_tokens(source) != color_tokens(target)


def classify(source, target):
    source_ph = placeholders(source)
    target_ph = placeholders(target)
    if Counter(source_ph) != Counter(target_ph):
        return "PLACEHOLDER_MISSING" if len(source_ph) > len(target_ph) else "PLACEHOLDER_EXTRA"
    if source_ph != target_ph:
        return "PLACEHOLDER_ORDER_CHANGED"

    source_colors = color_open_colors(source)
    target_colors = color_open_colors(target)
    if Counter(source_colors) != Counter(target_colors):
        return "MISSING_COLOR_TAG" if len(source_colors) > len(target_colors) else "EXTRA_COLOR_TAG"
    if source_colors != target_colors:
        return "COLOR_TAG_ORDER_CHANGED"
    source_spans = color_spans(source)
    target_spans = color_spans(target)
    if len(source_spans) == len(target_spans) and any(a["visible_text"] != b["visible_text"] for a, b in zip(source_spans, target_spans)):
        return "COLOR_SPAN_CONTENT_CHANGED"
    if color_tokens(source) != color_tokens(target):
        return "COLOR_SPAN_MOVED"
    return "OTHER_RICH_TEXT_MISMATCH"


def official_findings(data):
    findings = set()
    for sheet_name, sheet_data in data.items():
        for source_field, target_field in cn_vi_pairs(sheet_data["headers"]):
            for record in sheet_data["records"]:
                target = text(record.get(target_field)).strip()
                if not target or target == "NONE":
                    continue
                source = text(record.get(source_field)).strip()
                if sorted(placeholders(source)) != sorted(placeholders(target)):
                    findings.add((sheet_name, primary_id(sheet_name, record), target_field, "placeholder"))
                if sorted(color_tokens(source)) != sorted(color_tokens(target)):
                    findings.add((sheet_name, primary_id(sheet_name, record), target_field, "rich"))
    return findings


def validator_error(source, target):
    errors = []
    if placeholders(source) != placeholders(target):
        errors.append(f"Placeholder sequence mismatch: CN has {placeholders(source)!r}, VI has {placeholders(target)!r}")
    if color_tokens(source) != color_tokens(target):
        errors.append(f"Rich-text tag sequence mismatch: CN has {color_tokens(source)!r}, VI has {color_tokens(target)!r}")
    return " | ".join(errors)


def repair_rows(master: Path = MASTER, before: Path = DEFAULT_BEFORE):
    current_data = load_sheet_rows(master)
    before_data = load_sheet_rows(before)
    current_index = indexed_records(current_data)
    before_index = indexed_records(before_data)
    rows = []

    for sheet_name, sheet_data in current_data.items():
        if sheet_name not in before_index:
            continue
        for source_field, target_field in cn_vi_pairs(sheet_data["headers"]):
            for record in sheet_data["records"]:
                record_id = primary_id(sheet_name, record)
                before_record = before_index[sheet_name].get(record_id)
                if not before_record:
                    continue
                source = text(record.get(source_field)).strip()
                target = text(record.get(target_field)).strip()
                old_target = text(before_record.get(target_field)).strip()
                if not target or target == "NONE" or target == old_target or not token_mismatch(source, target):
                    continue
                base = {
                    "sheet": sheet_name,
                    "record_id": record_id,
                    "character_id": text(record.get("character_id")),
                    "field": target_field,
                    "source_cn": source,
                    "current_vi": target,
                    "source_placeholders": json.dumps(placeholders(source), ensure_ascii=False),
                    "vi_placeholders": json.dumps(placeholders(target), ensure_ascii=False),
                    "source_color_tokens": json.dumps(color_tokens(source), ensure_ascii=False),
                    "vi_color_tokens": json.dumps(color_tokens(target), ensure_ascii=False),
                    "source_tagged_spans": json.dumps(color_spans(source), ensure_ascii=False),
                    "vi_tagged_spans": json.dumps(color_spans(target), ensure_ascii=False),
                    "validator_error": validator_error(source, target),
                    "error_type": classify(source, target),
                    "translation_review_status": text(record.get("translation_review_status")),
                    "repair_proposed_vi": "",
                    "repair_notes": "",
                    "skill_id": text(record.get("skill_id")),
                    "skill_name_cn": text(record.get("skill_name_cn")),
                    "skill_name_vi": text(record.get("skill_name_vi")),
                    "buff_id": text(record.get("buff_id")),
                    "buff_name_cn": text(record.get("buff_name_cn")),
                    "buff_name_vi": text(record.get("buff_name_vi")),
                    "brilliant_id": text(record.get("brilliant_id")),
                    "icon_name_cn": text(record.get("icon_name_cn")),
                }
                rows.append(base)
    return rows, current_data, before_data


def write_json(rows, output: Path, master: Path, before: Path, current_data, before_data):
    official_current = official_findings(current_data)
    official_before = official_findings(before_data)
    payload = {
        "metadata": {
            "master": str(master.resolve()),
            "master_sha256": sha256(master),
            "before": str(before.resolve()),
            "before_sha256": sha256(before),
            "official_current_count": len(official_current),
            "official_before_count": len(official_before),
            "official_new_count": len(official_current - official_before),
            "official_resolved_count": len(official_before - official_current),
            "detailed_changed_token_mismatch_count": len(rows),
            "breakdown": dict(Counter(row["sheet"] for row in rows)),
            "error_type_counts": dict(Counter(row["error_type"] for row in rows)),
        },
        "rows": rows,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--master", type=Path, default=MASTER)
    parser.add_argument("--before", type=Path, default=DEFAULT_BEFORE)
    parser.add_argument("--output-json", type=Path, required=True)
    args = parser.parse_args()
    rows, current_data, before_data = repair_rows(args.master, args.before)
    payload = write_json(rows, args.output_json, args.master, args.before, current_data, before_data)
    print(json.dumps(payload["metadata"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
