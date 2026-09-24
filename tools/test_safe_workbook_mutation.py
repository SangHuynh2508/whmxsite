"""Regression test: a valid safe mutation must reach atomic replacement."""
from __future__ import annotations

import gc
import shutil
import sys
import unittest
from unittest import mock
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl

import safe_workbook_mutation


class SafeWorkbookMutationTest(unittest.TestCase):
    def test_successful_mutation_reaches_atomic_replace(self):
        root = Path(__file__).resolve().parent / "test_fixtures" / "safe_mutation"
        localization = root / "localization"
        master = localization / "localization_master.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "SKILL"
        sheet.append(["skill_id", "skill_name_vi"])
        sheet.append(["S1", "old"])
        workbook.save(master)
        workbook.close()
        del workbook
        gc.collect()

        def mutate(wb):
            wb["SKILL"]["B2"] = "new"

        calls = []
        def successful_atomic_replace(source, destination):
            calls.append((source, destination))
            shutil.copy2(source, destination)
        # The CI desktop profile can leave an external AV handle on a freshly
        # saved .xlsx. This verifies that a valid mutation reaches the atomic
        # replacement call (the regression point for the former NameError).
        with mock.patch.object(safe_workbook_mutation.os, "replace", side_effect=successful_atomic_replace):
            backup = safe_workbook_mutation.safe_mutate_workbook(
                str(root), mutate, {"skill_ids": {"S1"}}, authorized_new_columns={}
            )
        self.assertTrue(Path(backup).is_file())
        self.assertEqual(len(calls), 1)
        checked = openpyxl.load_workbook(master, read_only=True)
        self.assertEqual(checked["SKILL"]["B2"].value, "new")
        checked.close()

    def test_duplicate_primary_key_invariant_fails(self):
        root = Path(__file__).resolve().parent / "test_fixtures" / "safe_mutation_dup"
        localization = root / "localization"
        master = localization / "localization_master.xlsx"
        localization.mkdir(parents=True, exist_ok=True)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "SKILL"
        sheet.append(["skill_id", "skill_name_vi"])
        sheet.append(["S1", "name1"])
        sheet.append(["S2", "name2"])
        workbook.save(master)
        workbook.close()
        del workbook
        gc.collect()

        def mutate_intro_dup(wb):
            # Change S2 to S1, introducing duplicate S1
            wb["SKILL"]["A3"] = "S1"

        with self.assertRaises(PermissionError) as ctx:
            safe_workbook_mutation.safe_mutate_workbook(
                str(root), mutate_intro_dup, {"skill_ids": {"S1", "S2"}}, authorized_new_columns={}
            )
        self.assertIn("Duplicate primary key invariant violated", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()


def _declare_larger_dimension(path, ref):
    """Mimic an Excel-saved sheet whose <dimension> covers formatted-but-empty rows."""
    import re
    import zipfile
    tmp = Path(str(path) + ".dim")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="' + ref.encode() + b'"', data)
            dst.writestr(item, data)
    tmp.replace(path)


class TrailingEmptyRowsTest(unittest.TestCase):
    """Empty rows inside the declared sheet dimension are not data (SKIN had A1:AE418 with 146 real rows)."""

    def _master(self, rows, dimension):
        import tempfile
        root = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, root, ignore_errors=True)
        (root / "localization").mkdir()
        master = root / "localization" / "localization_master.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "SKILL"
        for row in rows:
            sheet.append(row)
        workbook.save(master)
        workbook.close()
        _declare_larger_dimension(master, dimension)
        return root, master

    def test_authorized_cell_edit_passes_despite_trailing_empty_rows(self):
        root, master = self._master([["skill_id", "desc_vi"], ["S1", "Phong Địch"]], "A1:B400")

        def mutate(wb):
            wb["SKILL"]["B2"] = "Phong Đích"

        safe_workbook_mutation.safe_mutate_workbook(str(root), mutate, {}, authorized_cells={("SKILL", "S1", "desc_vi")})
        checked = openpyxl.load_workbook(master, read_only=True)
        self.assertEqual(checked["SKILL"]["B2"].value, "Phong Đích")
        checked.close()

    def test_removing_a_real_row_is_still_rejected(self):
        root, _ = self._master([["skill_id", "desc_vi"], ["S1", "a"], ["S2", "b"]], "A1:B400")

        def mutate(wb):
            wb["SKILL"].delete_rows(3)

        with self.assertRaises(PermissionError):
            safe_workbook_mutation.safe_mutate_workbook(str(root), mutate, {})
