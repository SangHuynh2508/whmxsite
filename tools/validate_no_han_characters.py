import openpyxl
import re
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

# Whitelist of approved Chinese text / titles in Vietnamese player-facing fields
APPROVED_WHITELIST = [
    "《鹿王本生图》",  # Book/artwork title in character full name
    "雷威",            # Historical Tang Dynasty guqin maker name in A01604 story
]

def validate_no_han_characters():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    master_path = os.path.join(base_dir, 'localization', 'localization_master.xlsx')

    if not os.path.exists(master_path):
        print(f"[ERROR] Master workbook not found: {master_path}")
        sys.exit(1)

    print(f"Loading master workbook for Han-character validation: {master_path}...")
    wb = openpyxl.load_workbook(master_path, data_only=True)

    # CJK Unified Ideographs block: U+4E00 to U+9FFF
    cjk_regex = re.compile(r'[\u4e00-\u9fff]')

    total_audited = 0
    errors = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        header = rows[0]
        vi_indices = [i for i, h in enumerate(header) if h and '_vi' in str(h).lower()]

        for row_idx, r in enumerate(rows[1:], start=2):
            rec_id = str(r[0]) if r[0] is not None else f"Row_{row_idx}"
            for idx in vi_indices:
                val = str(r[idx]) if r[idx] is not None else ''
                if val:
                    total_audited += 1
                    
                    # Clean out whitelisted phrases before checking
                    check_val = val
                    for w in APPROVED_WHITELIST:
                        check_val = check_val.replace(w, '')
                        
                    han_matches = cjk_regex.findall(check_val)
                    if han_matches:
                        errors.append({
                            'sheet': sheet_name,
                            'row': row_idx,
                            'record_id': rec_id,
                            'field': header[idx],
                            'value': val,
                            'offending_chars': list(set(han_matches))
                        })

    print(f"\n=== ACCIDENTAL HAN-CHARACTER VALIDATION REPORT ===")
    print(f"Total Populated '_vi' Fields Audited: {total_audited}")
    print(f"Total Accidental Han-Character Leaks Found: {len(errors)}")

    if errors:
        print("\n[ERROR LIST]")
        for err in errors:
            print(f"  - Sheet: {err['sheet']}, Record: {err['record_id']}, Field: {err['field']}")
            print(f"    Offending Chars: {err['offending_chars']}")
            print(f"    Text: {err['value']}\n")
        return False
    else:
        print("[SUCCESS] 0 ACCIDENTAL HAN CHARACTERS FOUND IN POPULATED VI FIELDS!\n")
        return True

if __name__ == '__main__':
    success = validate_no_han_characters()
    if not success:
        sys.exit(1)
