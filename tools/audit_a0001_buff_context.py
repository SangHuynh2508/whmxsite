"""Emit exact A0001 raw skill/buff context without inventing popup identities."""
from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT.parent / "NeoArtifacts" / "MasterData" / "json"
MASTER = ROOT / "localization" / "localization_master.xlsx"
MARKER = re.compile(r"\{([^{}]+)\}")


def s(value): return "" if value is None else str(value)


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--output", type=Path, required=True); args = parser.parse_args()
    gameplay = json.loads((ROOT / "translation_gameplay.json").read_text(encoding="utf-8"))
    raw_buffs = json.loads((RAW / "buffMap.json").read_text(encoding="utf-8"))
    wb = load_workbook(MASTER, read_only=True, data_only=False)
    skill_ws, buff_ws = wb["SKILL"], wb["BUFF_STATUS"]
    sh = {value: index for index, value in enumerate(next(skill_ws.iter_rows(values_only=True)))}
    bh = {value: index for index, value in enumerate(next(buff_ws.iter_rows(values_only=True)))}
    skill_rows = [dict(zip(sh, row)) for row in skill_ws.iter_rows(min_row=2, values_only=True) if s(row[sh["character_id"]]) == "A0001"]
    buffs = {s(row[bh["buff_id"]]): dict(zip(bh, row)) for row in buff_ws.iter_rows(min_row=2, values_only=True)}
    links = [link for link in gameplay["skill_buff_links"] if link.get("character_id") == "A0001"]
    paths = defaultdict(list)
    for link in links:
        paths[link["buff_id"]].append({"skill_id": link["skill_id"], "relationship_source": link["relationship_source"], "evidence": link["evidence_notes"]})
    dependency_rows = []
    for bid in sorted(paths):
        raw = raw_buffs.get(bid, {})
        local = buffs.get(bid, {})
        name_cn, desc_cn = s(raw.get("NameLanText")), s(raw.get("DescriptionLanText"))
        # A named raw controller can carry a template while only feeding a
        # terminal display buff through Attr parameters.  Treat it as
        # player-facing only when the card itself contains its exact marker;
        # otherwise it must never become a popup display authority.
        has_direct_marker = any(
            path.get("relationship_source") == "skill_desc_tag"
            for path in paths[bid]
        )
        player_facing = bool(name_cn and desc_cn and has_direct_marker)
        term_class = (
            "CHARACTER_NAMED" if player_facing
            else "RAW_NAMED_CONTROLLER" if name_cn and desc_cn
            else "INTERNAL_CONTROLLER"
        )
        dependency_rows.append({
            "buff_id": bid, "name_cn": name_cn, "canonical_name_vi": s(local.get("buff_name_vi")),
            "desc_cn": desc_cn, "canonical_desc_vi": s(local.get("buff_desc_vi")),
            "term_class": term_class,
            "scope": "CHARACTER_SPECIFIC" if bid.startswith("Buff_A0001") else "SHARED",
            "raw_relation_path": paths[bid], "translation_status": s(local.get("status")) or "INTERNAL_NO_LOCALIZATION_ROW",
            "player_facing": player_facing,
        })
    groups = defaultdict(list)
    for row in skill_rows:
        groups[s(row["skill_group_id"])].append({"skill_id": s(row["skill_id"]), "skill_name_cn": s(row["skill_name_cn"]), "status": s(row["status"]), "ordered_markers": MARKER.findall(s(row["desc_cn"]))})
    payload = {"character_id": "A0001", "skill_groups": dict(groups), "buff_dependencies": dependency_rows}
    args.output.parent.mkdir(parents=True, exist_ok=True); args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(args.output), "skill_rows": len(skill_rows), "dependencies": len(dependency_rows)}, ensure_ascii=False))


if __name__ == "__main__": main()
