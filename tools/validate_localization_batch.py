import openpyxl
import json
import os
import sys

from batch_manifest_helper import load_batch_manifest, collect_batch_dependencies
from zhizhi_resolver import contains_raw_dev_codes

def validate_localization_batch(manifest_path, base_dir=None):
    if base_dir is None:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        
    manifest = load_batch_manifest(manifest_path)
    deps = collect_batch_dependencies(base_dir, manifest)
    
    master_path = os.path.join(base_dir, 'localization', 'localization_master.xlsx')
    wb = openpyxl.load_workbook(master_path, data_only=True)
    
    batch_chars = set(deps['characters'])
    player_buffs = set(deps['player_facing_buff_ids'])
    huanzhang_ids = set(deps['huanzhang_ids'])
    
    errors = []

    # 1. SKILL Coverage
    ws_skill = wb['SKILL']
    headers_sk = list(next(ws_skill.iter_rows(values_only=True)))
    cid_idx_sk = headers_sk.index('character_id')
    sid_idx_sk = headers_sk.index('skill_id')
    name_vi_idx_sk = headers_sk.index('skill_name_vi')
    desc_vi_idx_sk = headers_sk.index('desc_vi')

    sk_total = 0
    sk_populated = 0

    for r in ws_skill.iter_rows(min_row=2, values_only=True):
        cid = str(r[cid_idx_sk] or '').strip()
        if cid in batch_chars:
            sk_total += 1
            sid = str(r[sid_idx_sk] or '').strip()
            name_vi = str(r[name_vi_idx_sk] or '').strip()
            desc_vi = str(r[desc_vi_idx_sk] or '').strip()

            if not name_vi or name_vi == 'None':
                errors.append(f"SKILL {sid} ({cid}): skill_name_vi is empty/None")
            if not desc_vi or desc_vi == 'None':
                errors.append(f"SKILL {sid} ({cid}): desc_vi is empty/None")

            if name_vi and name_vi != 'None' and desc_vi and desc_vi != 'None':
                sk_populated += 1

    # 2. BUFF_STATUS Coverage
    ws_buff = wb['BUFF_STATUS']
    headers_bf = list(next(ws_buff.iter_rows(values_only=True)))
    bid_idx_bf = headers_bf.index('buff_id')
    name_vi_idx_bf = headers_bf.index('buff_name_vi')
    desc_vi_idx_bf = headers_bf.index('buff_desc_vi')

    bf_total = len(player_buffs)
    bf_populated = 0

    for r in ws_buff.iter_rows(min_row=2, values_only=True):
        bid = str(r[bid_idx_bf] or '').strip()
        if bid in player_buffs:
            name_vi = str(r[name_vi_idx_bf] or '').strip()
            desc_vi = str(r[desc_vi_idx_bf] or '').strip()

            if not name_vi or name_vi == 'None':
                errors.append(f"BUFF_STATUS {bid}: buff_name_vi is empty/None")
            if not desc_vi or desc_vi == 'None':
                errors.append(f"BUFF_STATUS {bid}: buff_desc_vi is empty/None")

            if name_vi and name_vi != 'None' and desc_vi and desc_vi != 'None':
                bf_populated += 1

    # 3. ZHIZHI Coverage
    ws_zhizhi = wb['ZHIZHI']
    headers_zh = list(next(ws_zhizhi.iter_rows(values_only=True)))
    cid_idx_zh = headers_zh.index('character_id')
    star_idx_zh = headers_zh.index('star')
    vi_idx_zh = headers_zh.index('effect_summary_vi')

    zh_total = len(batch_chars) * 6
    zh_populated = 0

    zh_found_ranks = {}

    for r in ws_zhizhi.iter_rows(min_row=2, values_only=True):
        cid = str(r[cid_idx_zh] or '').strip()
        if cid in batch_chars:
            star = r[star_idx_zh]
            zh_found_ranks[(cid, star)] = True
            vi_summary = str(r[vi_idx_zh] or '').strip()

            if not vi_summary or vi_summary == 'None':
                errors.append(f"ZHIZHI {cid} Rank {star}: effect_summary_vi is empty/None")
            elif contains_raw_dev_codes(vi_summary):
                errors.append(f"ZHIZHI {cid} Rank {star}: contains raw dev code '{vi_summary}'")
            else:
                zh_populated += 1

    if len(zh_found_ranks) < zh_total:
        errors.append(f"ZHIZHI total ranks found ({len(zh_found_ranks)}) < expected ({zh_total})")

    # 4. HUANZHANG Coverage
    ws_hz = wb['HUANZHANG']
    headers_hz = list(next(ws_hz.iter_rows(values_only=True)))
    hid_idx_hz = headers_hz.index('brilliant_id')
    info_vi_idx_hz = headers_hz.index('icon_info_vi')

    hz_total = len(huanzhang_ids)
    hz_populated = 0

    for r in ws_hz.iter_rows(min_row=2, values_only=True):
        hid = str(r[hid_idx_hz] or '').strip()
        if hid in huanzhang_ids:
            info_vi = str(r[info_vi_idx_hz] or '').strip()

            if not info_vi or info_vi == 'None':
                errors.append(f"HUANZHANG {hid}: icon_info_vi is empty/None")
            else:
                hz_populated += 1

    counts = {
        'SKILL': f"{sk_populated}/{sk_total}",
        'BUFF_STATUS': f"{bf_populated}/{bf_total}",
        'ZHIZHI': f"{zh_populated}/{zh_total}",
        'HUANZHANG': f"{hz_populated}/{hz_total}"
    }

    print(f"=== BATCH LOCALIZATION COVERAGE REPORT ({manifest['batch_id']}) ===")
    print(f"SKILL: {counts['SKILL']}")
    print(f"BUFF_STATUS: {counts['BUFF_STATUS']}")
    print(f"ZHIZHI: {counts['ZHIZHI']}")
    print(f"HUANZHANG: {counts['HUANZHANG']}")

    if errors:
        print(f"\n[ERROR] Found {len(errors)} coverage failure(s):")
        for err in errors:
            print(f"  - {err}")
        return False, counts, errors
    else:
        print("\n[SUCCESS] Batch coverage is 100% complete and validated!")
        return True, counts, []

if __name__ == '__main__':
    manifest_file = sys.argv[1] if len(sys.argv) > 1 else r"d:\BaiTapCode\WHMX\WhmxCalc\localization\batches\phase3_batch1.json"
    success = validate_localization_batch(manifest_file)
    if not success:
        sys.exit(1)
