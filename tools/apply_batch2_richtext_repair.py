"""Apply the reviewed Batch #2 rich-text repair workbook with exact guards.

This importer intentionally accepts only the reviewed repair artifact described
in the Batch #2 repair apply request.  It never derives new text: the artifact
is the sole source of every replacement string.
"""
from __future__ import annotations

import argparse
from collections import Counter
import gc
import hashlib
import json
import os
from pathlib import Path
import re
import sys
from typing import Any

import openpyxl

from safe_workbook_mutation import AtomicReplaceLockError, create_workbook_backup, safe_mutate_workbook
from workbook_semantic_fingerprint import workbook_semantic_fingerprint


BASE_DIR = Path(__file__).resolve().parent.parent
DEFAULT_ARTIFACT = (
    BASE_DIR
    / "localization"
    / "reviews"
    / "batch2_richtext_repair_20260913_180005_REPAIRED.xlsx"
)
DEFAULT_MASTER = BASE_DIR / "localization" / "localization_master.xlsx"
EXPECTED_ARTIFACT_SHA256 = "CB00D174B970E8AD4145AC07F22085FD65A16F643B393EA43196A525B747466E"
EXPECTED_MASTER_SHA256 = "FA8F78913E498A75853B1911986476C086FB45A4770EB3580BEA92501388F0E7"
PRE_REPAIR_BACKUP = BASE_DIR / "localization" / "backups" / "localization_master_20260913_183816_592020.xlsx"
EXPECTED_PRE_REPAIR_SHA256 = "8F55A0BB78E4E47BE8422904E6B69FE5D1B0BB69315C1B441A8CB83BA204AB1C"
DEFAULT_FALLBACK_CANDIDATE = (
    BASE_DIR / "localization" / "localization_master.xlsx.mutation_20260913_185238_131760.tmp.xlsx"
)

REQUIRED_REPAIR_HEADERS = {
    "sheet",
    "record_id",
    "field",
    "source_cn",
    "current_vi",
    "repair_proposed_vi",
    "repair_notes",
}
SOURCE_FIELD_BY_TARGET = {
    ("SKILL", "desc_vi"): "desc_cn",
    ("BUFF_STATUS", "buff_name_vi"): "buff_name_cn",
    ("HUANZHANG", "buff_show_vi"): "buff_show_cn",
}
EXPECTED_COUNTS = {"SKILL": 42, "BUFF_STATUS": 2, "HUANZHANG": 1}
EXPECTED_BUFF_IDS = {"Buff_W0150_M_1_ex", "Buff_Weakness"}
EXPECTED_HUANZHANG_IDS = {"W01344"}
COLOR_TAG_RE = re.compile(r"</?color[^>]*>")
PLACEHOLDER_RE = re.compile(r"\[Effect[^\]]+\]")
BUFF_MARKER_RE = re.compile(r"\{Buff_[^}]+\}")
HAN_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")


class RepairGuardError(RuntimeError):
    """Raised when an exact repair precondition does not hold."""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def ordered_tokens(pattern: re.Pattern[str], value: Any) -> list[str]:
    return pattern.findall(str(value or ""))


def headers_for(ws: Any) -> dict[str, int]:
    headers = [cell.value for cell in ws[1]]
    return {str(name): index + 1 for index, name in enumerate(headers) if name is not None}


def validate_hashes(artifact_path: Path, master_path: Path) -> dict[str, str]:
    artifact_sha = sha256(artifact_path)
    master_sha = sha256(master_path)
    if artifact_sha != EXPECTED_ARTIFACT_SHA256:
        raise RepairGuardError(
            f"WRONG_REPAIR_ARTIFACT_RUNTIME_FILE expected={EXPECTED_ARTIFACT_SHA256} actual={artifact_sha}"
        )
    if master_sha != EXPECTED_MASTER_SHA256:
        raise RepairGuardError(
            f"WRONG_MASTER_RUNTIME_FILE expected={EXPECTED_MASTER_SHA256} actual={master_sha}"
        )
    return {"artifact_sha256": artifact_sha, "master_sha256_before": master_sha}


def validate_pre_repair_semantics(master_path: Path) -> dict[str, str]:
    if sha256(PRE_REPAIR_BACKUP) != EXPECTED_PRE_REPAIR_SHA256:
        raise RepairGuardError("WRONG_PRE_REPAIR_BACKUP_RUNTIME_FILE")
    current_fingerprint = workbook_semantic_fingerprint(master_path)
    backup_fingerprint = workbook_semantic_fingerprint(PRE_REPAIR_BACKUP)
    if current_fingerprint != backup_fingerprint:
        raise RepairGuardError(
            f"PRE_APPLY_SEMANTIC_GUARD_FAILED current={current_fingerprint} backup={backup_fingerprint}"
        )
    return {
        "current_semantic_fingerprint": current_fingerprint,
        "pre_repair_backup_semantic_fingerprint": backup_fingerprint,
    }


def load_repairs(artifact_path: Path) -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(artifact_path, data_only=False, read_only=False)
    try:
        if set(workbook.sheetnames) != {"RichText Repair", "Repair QA"}:
            raise RepairGuardError(f"REPAIR_ARTIFACT_SHEETS_INVALID actual={workbook.sheetnames!r}")
        ws = workbook["RichText Repair"]
        headers = headers_for(ws)
        missing = REQUIRED_REPAIR_HEADERS - set(headers)
        if missing:
            raise RepairGuardError(f"REPAIR_ARTIFACT_HEADERS_MISSING {sorted(missing)!r}")
        repairs: list[dict[str, str]] = []
        for row_number in range(2, ws.max_row + 1):
            row = {
                header: "" if ws.cell(row_number, column).value is None else str(ws.cell(row_number, column).value)
                for header, column in headers.items()
            }
            if not any(row.values()):
                continue
            repairs.append(row)
        validate_artifact_shape(repairs)
        validate_repair_qa(workbook["Repair QA"])
        return repairs
    finally:
        workbook.close()


def validate_artifact_shape(repairs: list[dict[str, str]]) -> None:
    if len(repairs) != 45:
        raise RepairGuardError(f"REPAIR_ARTIFACT_ROW_COUNT_INVALID expected=45 actual={len(repairs)}")
    counts = Counter(repair["sheet"] for repair in repairs)
    if dict(counts) != EXPECTED_COUNTS:
        raise RepairGuardError(f"REPAIR_ARTIFACT_BREAKDOWN_INVALID expected={EXPECTED_COUNTS!r} actual={dict(counts)!r}")
    for repair in repairs:
        if not repair["repair_proposed_vi"]:
            raise RepairGuardError(
                f"REPAIR_PROPOSAL_EMPTY sheet={repair['sheet']} record_id={repair['record_id']} field={repair['field']}"
            )
        key = (repair["sheet"], repair["field"])
        if key not in SOURCE_FIELD_BY_TARGET:
            raise RepairGuardError(f"REPAIR_FIELD_NOT_AUTHORIZED repair={key!r} id={repair['record_id']!r}")
    duplicate_keys = [key for key, count in Counter((r["sheet"], r["record_id"], r["field"]) for r in repairs).items() if count > 1]
    if duplicate_keys:
        raise RepairGuardError(f"REPAIR_ARTIFACT_DUPLICATE_CELLS {duplicate_keys!r}")
    if {r["record_id"] for r in repairs if r["sheet"] == "BUFF_STATUS"} != EXPECTED_BUFF_IDS:
        raise RepairGuardError("REPAIR_ARTIFACT_BUFF_SCOPE_INVALID")
    if {r["record_id"] for r in repairs if r["sheet"] == "HUANZHANG"} != EXPECTED_HUANZHANG_IDS:
        raise RepairGuardError("REPAIR_ARTIFACT_HUANZHANG_SCOPE_INVALID")
    if any(r["field"] != "desc_vi" for r in repairs if r["sheet"] == "SKILL"):
        raise RepairGuardError("REPAIR_ARTIFACT_SKILL_SCOPE_INVALID")
    if len({r["repair_proposed_vi"] for r in repairs}) != 25:
        raise RepairGuardError("REPAIR_ARTIFACT_PATTERN_COUNT_INVALID expected=25")


def validate_repair_qa(ws: Any) -> None:
    headers = headers_for(ws)
    required = {"check", "result", "details"}
    if not required.issubset(headers):
        raise RepairGuardError("REPAIR_QA_HEADERS_MISSING")
    qa = {
        str(ws.cell(row, headers["check"]).value or ""): {
            "result": str(ws.cell(row, headers["result"]).value or ""),
            "details": str(ws.cell(row, headers["details"]).value or ""),
        }
        for row in range(2, ws.max_row + 1)
    }
    for check in ("color_tag_sequence", "placeholder_sequence"):
        if qa.get(check, {}).get("result") != "PASS":
            raise RepairGuardError(f"REPAIR_QA_FAILED check={check!r} actual={qa.get(check)!r}")
    han = qa.get("han_leaks_in_repairs", {})
    if han.get("result") != "PASS" or not han.get("details", "").startswith("0"):
        raise RepairGuardError(f"REPAIR_QA_FAILED check='han_leaks_in_repairs' actual={han!r}")


def master_rows(master_path: Path) -> dict[str, tuple[dict[str, int], dict[str, dict[str, Any]]]]:
    workbook = openpyxl.load_workbook(master_path, data_only=False, read_only=True)
    try:
        result: dict[str, tuple[dict[str, int], dict[str, dict[str, Any]]]] = {}
        for sheet in SOURCE_FIELD_BY_TARGET:
            sheet_name = sheet[0]
            if sheet_name in result:
                continue
            ws = workbook[sheet_name]
            headers = headers_for(ws)
            records = {}
            for values in ws.iter_rows(min_row=2, values_only=True):
                record_id = "" if not values or values[0] is None else str(values[0])
                if record_id:
                    records[record_id] = {
                        header: values[column - 1] if column - 1 < len(values) else None
                        for header, column in headers.items()
                    }
            result[sheet_name] = (headers, records)
        return result
    finally:
        workbook.close()


def validate_current_master(master_path: Path, repairs: list[dict[str, str]]) -> None:
    workbook_data = master_rows(master_path)
    failures: list[str] = []
    for repair in repairs:
        sheet, record_id, target_field = repair["sheet"], repair["record_id"], repair["field"]
        headers, records = workbook_data[sheet]
        source_field = SOURCE_FIELD_BY_TARGET[(sheet, target_field)]
        if target_field not in headers or source_field not in headers:
            failures.append(f"RICH_TEXT_REPAIR_SCHEMA_INVALID sheet={sheet} target={target_field} source={source_field}")
            continue
        current = records.get(record_id)
        if current is None:
            failures.append(f"RICH_TEXT_REPAIR_RECORD_MISSING sheet={sheet} record_id={record_id}")
            continue
        if current[source_field] != repair["source_cn"]:
            failures.append(
                f"RICH_TEXT_REPAIR_SOURCE_CHANGED sheet={sheet} record_id={record_id} field={source_field} "
                f"master={current[source_field]!r} artifact={repair['source_cn']!r}"
            )
        if current[target_field] != repair["current_vi"]:
            failures.append(
                f"RICH_TEXT_REPAIR_TARGET_CHANGED sheet={sheet} record_id={record_id} field={target_field} "
                f"master={current[target_field]!r} artifact_expected_current_vi={repair['current_vi']!r}"
            )
    if failures:
        raise RepairGuardError("\n".join(failures))


def validate_repair_parity(repairs: list[dict[str, str]]) -> None:
    failures: list[str] = []
    for repair in repairs:
        source = repair["source_cn"]
        target = repair["repair_proposed_vi"]
        for label, pattern in (("placeholder", PLACEHOLDER_RE), ("color_tag", COLOR_TAG_RE)):
            if ordered_tokens(pattern, source) != ordered_tokens(pattern, target):
                failures.append(
                    f"REPAIR_PROPOSAL_{label.upper()}_MISMATCH sheet={repair['sheet']} "
                    f"record_id={repair['record_id']} field={repair['field']}"
                )
        if HAN_RE.search(target):
            failures.append(
                f"REPAIR_PROPOSAL_HAN_LEAK sheet={repair['sheet']} record_id={repair['record_id']} field={repair['field']}"
            )
    if failures:
        raise RepairGuardError("\n".join(failures))


def mutate_master(workbook: Any, repairs: list[dict[str, str]]) -> None:
    by_sheet: dict[str, dict[str, dict[str, str]]] = {}
    for repair in repairs:
        by_sheet.setdefault(repair["sheet"], {})[repair["record_id"]] = repair
    for sheet, by_id in by_sheet.items():
        ws = workbook[sheet]
        headers = headers_for(ws)
        for row_number in range(2, ws.max_row + 1):
            record_id = "" if ws.cell(row_number, 1).value is None else str(ws.cell(row_number, 1).value)
            repair = by_id.get(record_id)
            if repair is None:
                continue
            target_field = repair["field"]
            source_field = SOURCE_FIELD_BY_TARGET[(sheet, target_field)]
            if ws.cell(row_number, headers[source_field]).value != repair["source_cn"]:
                raise RepairGuardError(f"RICH_TEXT_REPAIR_SOURCE_CHANGED_DURING_MUTATION {sheet} {record_id} {source_field}")
            if ws.cell(row_number, headers[target_field]).value != repair["current_vi"]:
                raise RepairGuardError(f"RICH_TEXT_REPAIR_TARGET_CHANGED_DURING_MUTATION {sheet} {record_id} {target_field}")
            ws.cell(row_number, headers[target_field]).value = repair["repair_proposed_vi"]


def validate_applied_master(master_path: Path, repairs: list[dict[str, str]]) -> None:
    workbook_data = master_rows(master_path)
    failures: list[str] = []
    for repair in repairs:
        sheet, record_id, target_field = repair["sheet"], repair["record_id"], repair["field"]
        _, records = workbook_data[sheet]
        actual = records[record_id][target_field]
        if actual != repair["repair_proposed_vi"]:
            failures.append(f"REPAIR_APPLY_VALUE_MISMATCH {sheet} {record_id} {target_field} actual={actual!r}")
            continue
        source_field = SOURCE_FIELD_BY_TARGET[(sheet, target_field)]
        source = records[record_id][source_field]
        for label, pattern in (("placeholder", PLACEHOLDER_RE), ("color_tag", COLOR_TAG_RE)):
            if ordered_tokens(pattern, source) != ordered_tokens(pattern, actual):
                failures.append(f"REPAIR_APPLY_{label.upper()}_MISMATCH {sheet} {record_id} {target_field}")
        if HAN_RE.search(str(actual)):
            failures.append(f"REPAIR_APPLY_HAN_LEAK {sheet} {record_id} {target_field}")
        if ordered_tokens(BUFF_MARKER_RE, repair["current_vi"]) != ordered_tokens(BUFF_MARKER_RE, actual):
            failures.append(f"REPAIR_APPLY_BUFF_MARKER_CHANGED {sheet} {record_id} {target_field}")
    if failures:
        raise RepairGuardError("\n".join(failures))


def validate_candidate_exact_delta(master_path: Path, candidate_path: Path, repairs: list[dict[str, str]]) -> None:
    """Require the candidate to differ from current master in exactly 45 cells."""
    authorized = {(r["sheet"], r["record_id"], r["field"]) for r in repairs}
    current = openpyxl.load_workbook(master_path, data_only=False, read_only=True)
    candidate = openpyxl.load_workbook(candidate_path, data_only=False, read_only=True)
    try:
        if current.sheetnames != candidate.sheetnames:
            raise RepairGuardError("FALLBACK_CANDIDATE_SHEET_TOPOLOGY_CHANGED")
        changes: set[tuple[str, str, str]] = set()
        for sheet_name in current.sheetnames:
            source_ws, candidate_ws = current[sheet_name], candidate[sheet_name]
            source_headers = [cell.value for cell in next(source_ws.iter_rows(min_row=1, max_row=1))]
            candidate_headers = [cell.value for cell in next(candidate_ws.iter_rows(min_row=1, max_row=1))]
            if source_headers != candidate_headers or source_ws.max_row != candidate_ws.max_row:
                raise RepairGuardError(f"FALLBACK_CANDIDATE_TOPOLOGY_CHANGED sheet={sheet_name}")
            for source_row, candidate_row in zip(
                source_ws.iter_rows(min_row=2, values_only=True),
                candidate_ws.iter_rows(min_row=2, values_only=True),
            ):
                record_id = str((source_row[0] if source_row else None) or "")
                for index, (source_value, candidate_value) in enumerate(zip(source_row, candidate_row)):
                    if source_value == candidate_value:
                        continue
                    key = (sheet_name, record_id, str(source_headers[index]))
                    if key not in authorized:
                        raise RepairGuardError(f"FALLBACK_CANDIDATE_UNAUTHORIZED_DELTA {key!r}")
                    changes.add(key)
        if changes != authorized:
            raise RepairGuardError(
                f"FALLBACK_CANDIDATE_DELTA_COUNT_INVALID expected={len(authorized)} actual={len(changes)}"
            )
    finally:
        current.close()
        candidate.close()


def guarded_in_place_fallback(master_path: Path, candidate_path: Path, repairs: list[dict[str, str]]) -> Path:
    """Use one verified non-atomic overwrite only after a bounded lock failure."""
    if workbook_semantic_fingerprint(master_path) != workbook_semantic_fingerprint(PRE_REPAIR_BACKUP):
        raise RepairGuardError("FALLBACK_CURRENT_MASTER_SEMANTIC_GUARD_FAILED")
    candidate_fingerprint = workbook_semantic_fingerprint(candidate_path)
    if candidate_fingerprint == workbook_semantic_fingerprint(master_path):
        raise RepairGuardError("FALLBACK_CANDIDATE_HAS_NO_REPAIR_DELTA")
    validate_candidate_exact_delta(master_path, candidate_path, repairs)
    validate_applied_master(candidate_path, repairs)
    current_sha = sha256(master_path)
    current_fingerprint = workbook_semantic_fingerprint(master_path)
    fallback_backup = Path(create_workbook_backup(str(BASE_DIR)))
    if sha256(fallback_backup) != current_sha or workbook_semantic_fingerprint(fallback_backup) != current_fingerprint:
        raise RepairGuardError("FALLBACK_BACKUP_VERIFICATION_FAILED")
    candidate_bytes = candidate_path.read_bytes()

    # Request an exclusive Win32 handle.  It is held only for the byte copy,
    # and immediately closed before reopen/fingerprint verification.
    import ctypes
    from ctypes import wintypes
    import msvcrt

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    create_file = kernel32.CreateFileW
    create_file.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    create_file.restype = wintypes.HANDLE
    invalid_handle = ctypes.c_void_p(-1).value
    write_handle = create_file(str(master_path), 0xC0000000, 0, None, 3, 0x80, None)
    if write_handle == invalid_handle:
        raise RepairGuardError(f"FALLBACK_EXCLUSIVE_WRITE_FAILED winerror={ctypes.get_last_error()}")

    def overwrite(handle: int, bytes_to_write: bytes) -> None:
        # ``open_osfhandle`` transfers ownership.  The ``with`` block closes
        # it whether the copy succeeds or fails, so a restore can reacquire
        # exclusivity without inheriting a stale handle from the failed write.
        fd = msvcrt.open_osfhandle(handle, os.O_RDWR | getattr(os, "O_BINARY", 0))
        with os.fdopen(fd, "r+b", closefd=True) as destination:
            destination.seek(0)
            destination.truncate(0)
            destination.write(bytes_to_write)
            destination.flush()
            os.fsync(destination.fileno())

    try:
        overwrite(write_handle, candidate_bytes)
        if sha256(master_path) != sha256(candidate_path):
            raise RepairGuardError("FALLBACK_SHA_VERIFICATION_FAILED")
        # This reopens as XLSX and confirms semantic equality to the candidate.
        if workbook_semantic_fingerprint(master_path) != candidate_fingerprint:
            raise RepairGuardError("FALLBACK_SEMANTIC_VERIFICATION_FAILED")
        validate_applied_master(master_path, repairs)
        validate_candidate_exact_delta(PRE_REPAIR_BACKUP, master_path, repairs)
    except Exception as exc:
        # The exclusive handle was consumed by os.fdopen and is now closed; a
        # fresh exclusive handle is required for an attempted guarded restore.
        restore_handle = create_file(str(master_path), 0xC0000000, 0, None, 3, 0x80, None)
        if restore_handle == invalid_handle:
            raise RepairGuardError(
                f"FALLBACK_FAILED_AND_RESTORE_HANDLE_UNAVAILABLE winerror={ctypes.get_last_error()}"
            ) from exc
        try:
            overwrite(restore_handle, fallback_backup.read_bytes())
            if sha256(master_path) != sha256(fallback_backup):
                raise RepairGuardError("FALLBACK_RESTORE_SHA_VERIFICATION_FAILED")
            if workbook_semantic_fingerprint(master_path) != workbook_semantic_fingerprint(fallback_backup):
                raise RepairGuardError("FALLBACK_RESTORE_SEMANTIC_VERIFICATION_FAILED")
        except Exception as restore_exc:
            raise RepairGuardError(f"FALLBACK_FAILED_AND_RESTORE_FAILED {restore_exc!r}") from exc
        raise RepairGuardError(f"FALLBACK_FAILED_RESTORED_FROM={fallback_backup} cause={exc!r}") from exc
    return fallback_backup


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--artifact", type=Path, default=DEFAULT_ARTIFACT)
    parser.add_argument("--master", type=Path, default=DEFAULT_MASTER)
    parser.add_argument("--dry-run", action="store_true", help="Validate without creating a backup or mutating the master.")
    parser.add_argument("--fallback-only", action="store_true", help="Skip atomic replace and run only the guarded in-place fallback.")
    parser.add_argument("--candidate", type=Path, default=DEFAULT_FALLBACK_CANDIDATE)
    args = parser.parse_args()
    artifact_path, master_path = args.artifact.resolve(), args.master.resolve()
    try:
        result = validate_hashes(artifact_path, master_path)
        result.update(validate_pre_repair_semantics(master_path))
        repairs = load_repairs(artifact_path)
        validate_repair_parity(repairs)
        validate_current_master(master_path, repairs)
        result.update(
            repair_rows=len(repairs),
            repair_breakdown=dict(Counter(repair["sheet"] for repair in repairs)),
            unique_repair_patterns=len({repair["repair_proposed_vi"] for repair in repairs}),
            source_current_guard="PASS",
        )
        if args.dry_run:
            result["dry_run"] = True
            print(json.dumps(result, ensure_ascii=False, sort_keys=True))
            return 0
        # ``validate_current_master`` reads all needed values into ordinary
        # dicts and closes its read-only workbook before returning.  Force a
        # collection at this boundary so no source-guard object can survive
        # into the atomic-replace section of the safe mutator.
        gc.collect()
        authorized_cells = {(r["sheet"], r["record_id"], r["field"]) for r in repairs}
        if args.fallback_only:
            candidate_path = args.candidate.resolve()
            if not candidate_path.is_file():
                raise RepairGuardError(f"FALLBACK_CANDIDATE_MISSING {candidate_path}")
            validate_candidate_exact_delta(master_path, candidate_path, repairs)
            validate_applied_master(candidate_path, repairs)
            fallback_backup = guarded_in_place_fallback(master_path, candidate_path, repairs)
            backup_path = fallback_backup
            result["write_method"] = "GUARDED_IN_PLACE_FALLBACK"
        else:
            try:
                backup_path = safe_mutate_workbook(
                    str(BASE_DIR),
                    lambda workbook: mutate_master(workbook, repairs),
                    authorized_deps={},
                    authorized_cells=authorized_cells,
                )
                result["write_method"] = "ATOMIC_REPLACE"
            except AtomicReplaceLockError as lock_error:
                fallback_backup = guarded_in_place_fallback(master_path, Path(lock_error.temp_path), repairs)
                backup_path = fallback_backup
                result["write_method"] = "GUARDED_IN_PLACE_FALLBACK"
        validate_applied_master(master_path, repairs)
        validate_candidate_exact_delta(PRE_REPAIR_BACKUP, master_path, repairs)
        result.update(
            backup_path=str(backup_path),
            backup_sha256=sha256(Path(backup_path)),
            master_sha256_after=sha256(master_path),
            target_richtext_parity="PASS 45/45",
        )
        print(json.dumps(result, ensure_ascii=False, sort_keys=True))
        return 0
    except RepairGuardError as exc:
        print(str(exc), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
