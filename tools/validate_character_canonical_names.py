"""Validate exact CHARACTER-derived canonical references in reviewed VI fields."""
from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path
import sys
from typing import Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
REVIEWED_STATUSES = {"CHATGPT_REVIEWED", "OWNER_APPROVED", "APPROVED", "HUMAN_REVIEWED"}
LEGACY_ALIASES = {
    "愿望杯": {"Nguyện Vọng Bôi"},
}
PLAYER_FACING_PAIRS = {
    "SKILL": (
        ("skill_id", "skill_name_cn", "skill_name_vi"),
        ("skill_id", "desc_cn", "desc_vi"),
    ),
    "HUANZHANG": (
        ("brilliant_id", "icon_name_cn", "icon_name_vi"),
        ("brilliant_id", "icon_info_cn", "icon_info_vi"),
        ("brilliant_id", "buff_show_cn", "buff_show_vi"),
    ),
}


def text(value):
    return value if isinstance(value, str) else "" if value is None else str(value)


def normalized_status(value):
    return text(value).strip().upper()


def row_dicts(sheet):
    headers = [cell.value for cell in sheet[1]]
    return [
        dict(zip(headers, values))
        for values in sheet.iter_rows(min_row=2, values_only=True)
        if any(value not in (None, "") for value in values)
    ]


def build_character_registry(character_rows: Iterable[dict]):
    seen = defaultdict(set)
    character_variants = defaultdict(set)
    for row in character_rows:
        vi_candidates = {text(row.get("name_vi")).strip(), text(row.get("fullname_vi")).strip()} - {""}
        for cn_field in ("name_cn", "fullname_cn"):
            cn = text(row.get(cn_field)).strip()
            if cn:
                character_variants[cn].update(vi_candidates)
        for cn_field, vi_field in (("name_cn", "name_vi"), ("fullname_cn", "fullname_vi")):
            cn = text(row.get(cn_field)).strip()
            vi = text(row.get(vi_field)).strip()
            if cn and vi:
                seen[cn].add(vi)
    registry = {}
    ambiguities = {}
    for cn, vi_values in seen.items():
        if len(vi_values) == 1:
            registry[cn] = next(iter(vi_values))
        else:
            ambiguities[cn] = sorted(vi_values)
    return registry, ambiguities, character_variants


def reviewed_row(row):
    if "translation_review_status" not in row:
        return True
    return normalized_status(row.get("translation_review_status")) in REVIEWED_STATUSES


def validate_rows(rows_by_sheet, registry, character_variants=None):
    character_variants = character_variants or {}
    issues = []
    for sheet_name, pairs in PLAYER_FACING_PAIRS.items():
        for row in rows_by_sheet.get(sheet_name, []):
            if not reviewed_row(row):
                continue
            for id_field, source_field, target_field in pairs:
                source = text(row.get(source_field))
                target = text(row.get(target_field))
                if not source or not target:
                    continue
                for cn, vi in registry.items():
                    if cn not in source:
                        continue
                    allowed = character_variants.get(cn, {vi})
                    if any(cand in target for cand in allowed):
                        continue
                    aliases_present = sorted(alias for alias in LEGACY_ALIASES.get(cn, set()) if alias in target)
                    issues.append(
                        {
                            "sheet": sheet_name,
                            "record_id": text(row.get(id_field)),
                            "field": target_field,
                            "canonical_cn": cn,
                            "canonical_vi": vi,
                            "legacy_aliases_present": aliases_present,
                        }
                    )
    return issues


def validate_workbook(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows_by_sheet = {sheet.title: row_dicts(sheet) for sheet in workbook.worksheets}
    finally:
        workbook.close()
    registry, ambiguities, variants = build_character_registry(rows_by_sheet.get("CHARACTER", []))
    issues = validate_rows(rows_by_sheet, registry, variants)
    return registry, ambiguities, issues


def main():
    if getattr(sys.stdout, "encoding", "").lower() != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--master", type=Path, default=ROOT / "localization/localization_master.xlsx")
    args = parser.parse_args()
    registry, ambiguities, issues = validate_workbook(args.master)
    if "愿望杯" not in registry or registry["愿望杯"] != "Lotus Chalice":
        raise SystemExit(f"CANONICAL_CHARACTER_REGRESSION:愿望杯:{registry.get('愿望杯')!r}")
    if ambiguities:
        print(
            "CHARACTER_CANONICAL_AMBIGUITY="
            + repr({cn: values for cn, values in sorted(ambiguities.items())})
        )
    if issues:
        raise SystemExit(f"CANONICAL_CHARACTER_NAME_MISSING:{issues!r}")
    print(f"CANONICAL_CHARACTER_REGISTRY=PASS mappings={len(registry)} ambiguities={len(ambiguities)}")
    print("CANONICAL_CHARACTER_REGRESSION=愿望杯 → Lotus Chalice")


if __name__ == "__main__":
    main()
