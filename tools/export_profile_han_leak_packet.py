"""Export the 22 PROFILE Han-leak review packet for ChatGPT/owner review.

This tool produces:
    localization/reviews/profile_han_cleanup_<timestamp>.xlsx

It splits/classifies findings into:
- PROFILE_SCHEMA_SHIFT (20 rows)
- PROFILE_PARTIAL_HAN (2 rows)
All translation and repair proposal fields are left strictly blank.
"""
from __future__ import annotations

from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import re
import sys
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
MASTER_PATH = BASE_DIR / "localization" / "localization_master.xlsx"
REVIEWS_DIR = BASE_DIR / "localization" / "reviews"

HAN_RE = re.compile(r"[\u4e00-\u9fff]")
WHITELIST = ["《鹿王本生图》", "雷威"]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def build_profile_packet() -> tuple[Path, str, int, int, int]:
    wb = openpyxl.load_workbook(MASTER_PATH, data_only=True)
    ws = wb["PROFILE"]
    headers = [cell.value for cell in ws[1]]
    header_map = {name: idx + 1 for idx, name in enumerate(headers) if name is not None}

    schema_shift_rows = []
    partial_han_rows = []

    for row_idx in range(2, ws.max_row + 1):
        item = {h: ws.cell(row_idx, col).value for h, col in header_map.items()}
        pid = str(item.get("profile_id") or "").strip()
        title_vi = str(item.get("title_vi") or "")
        text_vi = str(item.get("text_vi") or "")

        has_han_title = bool(HAN_RE.search(title_vi))
        has_han_text = bool(HAN_RE.search(text_vi))

        if not (has_han_title or has_han_text):
            continue

        if has_han_title and text_vi == "HIGH":
            # Schema shift row
            item["problem_class"] = "PROFILE_SCHEMA_SHIFT"
            item["expected_field_mapping"] = "title_vi = translation of title_cn; text_vi = translation of text_cn; confidence = HIGH"
            item["source_cn_body"] = item.get("text_cn")
            item["current_misaligned_values"] = "title_vi contains CN body text; text_vi contains 'HIGH'"
            item["repair_proposed_title_vi"] = ""
            item["translation_proposed_text_vi"] = ""
            item["review_notes"] = ""
            schema_shift_rows.append(item)
        else:
            # Partial Han row
            item["problem_class"] = "PROFILE_PARTIAL_HAN"
            item["expected_field_mapping"] = "text_vi should be clean Vietnamese translation without residual Han fragments"
            item["source_cn_body"] = item.get("text_cn")
            item["current_misaligned_values"] = "text_vi contains partial Vietnamese translation with residual Han fragments"
            item["repair_proposed_title_vi"] = ""
            item["translation_proposed_text_vi"] = ""
            item["review_notes"] = ""
            partial_han_rows.append(item)

    wb.close()

    if len(schema_shift_rows) != 20:
        raise RuntimeError(f"EXPECTED 20 SCHEMA_SHIFT ROWS, GOT {len(schema_shift_rows)}")
    if len(partial_han_rows) != 2:
        raise RuntimeError(f"EXPECTED 2 PARTIAL_HAN ROWS, GOT {len(partial_han_rows)}")

    columns = [
        "problem_class",
        "profile_id",
        "character_id",
        "category",
        "title_cn",
        "title_vi",
        "text_cn",
        "text_vi",
        "confidence",
        "status",
        "notes",
        "expected_field_mapping",
        "source_cn_body",
        "current_misaligned_values",
        "repair_proposed_title_vi",
        "translation_proposed_text_vi",
        "review_notes",
    ]

    wb_out = openpyxl.Workbook()

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    def write_sheet(ws_target, rows_data):
        ws_target.append(columns)
        for col_idx in range(1, len(columns) + 1):
            cell = ws_target.cell(1, col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_data in rows_data:
            ws_target.append([row_data.get(col) for col in columns])

        for r_idx in range(2, len(rows_data) + 2):
            for c_idx in range(1, len(columns) + 1):
                cell = ws_target.cell(r_idx, c_idx)
                cell.font = Font(name="Calibri", size=10)
                cell.border = border
                c_name = columns[c_idx - 1]
                if c_name in ("repair_proposed_title_vi", "translation_proposed_text_vi", "review_notes"):
                    cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        ws_target.row_dimensions[1].height = 28
        for c_idx, c_name in enumerate(columns, 1):
            letter = get_column_letter(c_idx)
            if c_name in ("text_cn", "source_cn_body", "translation_proposed_text_vi", "title_vi"):
                ws_target.column_dimensions[letter].width = 45
            elif c_name in ("expected_field_mapping", "current_misaligned_values", "review_notes"):
                ws_target.column_dimensions[letter].width = 35
            elif c_name in ("title_cn", "repair_proposed_title_vi"):
                ws_target.column_dimensions[letter].width = 25
            else:
                ws_target.column_dimensions[letter].width = 15

    # Sheet 1: ALL 22 ROWS
    ws_all = wb_out.active
    ws_all.title = "PROFILE_HAN_BACKLOG"
    all_rows = schema_shift_rows + partial_han_rows
    write_sheet(ws_all, all_rows)

    # Sheet 2: SCHEMA SHIFT (20 rows)
    ws_shift = wb_out.create_sheet(title="SCHEMA_SHIFT")
    write_sheet(ws_shift, schema_shift_rows)

    # Sheet 3: PARTIAL HAN (2 rows)
    ws_partial = wb_out.create_sheet(title="PARTIAL_HAN")
    write_sheet(ws_partial, partial_han_rows)

    # Sheet 4: QA
    ws_qa = wb_out.create_sheet(title="QA")
    ws_qa.append(["check", "result", "details"])
    ws_qa.append(["total_profile_han_backlog", "PASS", str(len(all_rows))])
    ws_qa.append(["schema_shift_count", "PASS", str(len(schema_shift_rows))])
    ws_qa.append(["partial_han_count", "PASS", str(len(partial_han_rows))])
    ws_qa.append(["proposal_fields_left_blank", "PASS", "repair_proposed_title_vi and translation_proposed_text_vi are 100% blank"])
    ws_qa.append(["created_at", "INFO", datetime.now().isoformat()])

    for cell in ws_qa[1]:
        cell.font = header_font
        cell.fill = header_fill

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = REVIEWS_DIR / f"profile_han_cleanup_{timestamp}.xlsx"
    wb_out.save(out_path)
    wb_out.close()

    return out_path, sha256(out_path), len(schema_shift_rows), len(partial_han_rows), len(all_rows)


def main() -> int:
    if getattr(sys.stdout, "encoding", "").lower() != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8")
    REVIEWS_DIR.mkdir(parents=True, exist_ok=True)
    out_path, out_sha, shift_cnt, part_cnt, total_cnt = build_profile_packet()
    summary = {
        "packet_path": str(out_path),
        "packet_sha256": out_sha,
        "schema_shift_count": shift_cnt,
        "partial_han_count": part_cnt,
        "total_count": total_cnt,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
