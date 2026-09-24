"""Create and safely round-trip compact localization batches.

The master workbook deliberately remains row-complete.  This tool writes a
translation-facing workbook plus a per-target manifest; it never modifies the
real master unless the caller explicitly supplies another workbook to --apply.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import tempfile
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "localization" / "localization_master.xlsx"
BATCH_DIR = ROOT / "localization" / "batches"
REPORT_DIR = ROOT / "localization" / "batch_reports"
CANONICAL = {
    "万籁沉寂": "Vạn Âm Trầm Tịch", "瞄准": "Miêu Chuẩn", "脆弱": "Thúy Nhược",
    "蓄势": "Súc Thế", "滞缓": "Trệ Hoãn", "萧瑟": "Tiêu Sắt", "截招": "Tiệt Chiêu",
}
QUEUE = {"PENDING", "NEEDS_TRANSLATION", "REVIEW", "SOURCE_CHANGED"}
OWNER_WORDS = ("OWNER", "CANONICAL", "HUMAN_REVIEW", "APPROVED")
PLACEHOLDER_RE = re.compile(r"\[(?:Effect|Condition)[^\]]+\]|\{\d+\}")
BUFF_RE = re.compile(r"\{Buff_[^}\s]+\}", re.I)
COLOR_RE = re.compile(r"</?color(?:=[^>]+)?>", re.I)

PRIMARY = {
    "CHARACTER": ("character_id",), "SKILL": ("skill_id",),
    "BUFF_STATUS": ("buff_id",), "ZHIZHI": ("character_id", "star"),
    "HUANZHANG": ("brilliant_id",), "ITEM": ("item_id",),
    "PROFILE": ("profile_id",), "SKIN": ("skin_id",),
    "UI_SYSTEM": ("ui_key",), "GLOSSARY": ("term_id",), "TALENT": ("talent_bank_id",),
}


def s(value):
    return "" if value is None else str(value)


def header(ws):
    return {s(cell.value): cell.column for cell in ws[1] if cell.value is not None}


def row_data(ws, row, h):
    return {name: ws.cell(row, col).value for name, col in h.items()}


def norm_key(value):
    return s(value).strip().upper()


def key_for(sheet, data):
    return tuple(norm_key(data.get(col)) for col in PRIMARY[sheet])


def record_id(sheet, data):
    return ":".join(key_for(sheet, data))


def pairs(data):
    return [(c[:-3] + "_cn", c) for c in data if c.endswith("_vi") and c[:-3] + "_cn" in data]


def editable_pairs(data):
    return [(cn, vi) for cn, vi in pairs(data) if s(data.get(cn)).strip()]


def source_signature(sheet, data):
    payload = {"sheet": sheet, "key": key_for(sheet, data),
               "sources": {cn: s(data.get(cn)) for cn, _ in editable_pairs(data)}}
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def text_signatures(data):
    text = "\n".join(s(data.get(cn)) for cn, _ in editable_pairs(data))
    return {
        "placeholder_signature": json.dumps(PLACEHOLDER_RE.findall(text), ensure_ascii=False),
        "ordered_buff_markers": json.dumps(BUFF_RE.findall(text), ensure_ascii=False),
        "color_signature": json.dumps(COLOR_RE.findall(text), ensure_ascii=False),
        "newline_count": text.count("\n"),
    }


def status(data): return s(data.get("status")).strip().upper()


def owner_locked(data):
    note = s(data.get("notes")).upper()
    return status(data) in {"APPROVED", "OWNER_APPROVED", "OWNER_CORRECTED"} or any(x in note for x in OWNER_WORDS)


def needs_work(data):
    # Existing VI is retained as editable context when row needs review/source work.
    return status(data) in QUEUE and bool(editable_pairs(data)) and not owner_locked(data)


def owner_variant(data):
    return json.dumps({"status": status(data), "notes": s(data.get("notes")),
                       "vi": {vi: s(data.get(vi)) for _, vi in editable_pairs(data)}},
                      ensure_ascii=False, sort_keys=True)


def ex_class(data):
    ident = norm_key(data.get("skill_id"))
    group = norm_key(data.get("skill_group_id"))
    return "EX" if ident.endswith("EX") or group.endswith("EX") else "BASE"


def first_pair(data):
    ps = editable_pairs(data)
    return ps[0] if ps else ("", "")


def display_values(data):
    # Generic fields let all workbook domains use a single import surface.
    name_pair = next(((cn, vi) for cn, vi in editable_pairs(data) if "name" in cn or "title" in cn or "term" in cn), None)
    desc_pair = next(((cn, vi) for cn, vi in editable_pairs(data) if "desc" in cn or "text" in cn or "info" in cn or "summary" in cn or "show" in cn), None)
    if not name_pair: name_pair = first_pair(data)
    if not desc_pair: desc_pair = ("", "")
    return (s(data.get(name_pair[0])), s(data.get(name_pair[1])), s(data.get(desc_pair[0])), s(data.get(desc_pair[1])))


def group_signature(data):
    """Strictly preserve every source-relevant field; no whitespace normalization."""
    return json.dumps({
        "cid": norm_key(data.get("character_id")), "group": norm_key(data.get("skill_group_id")),
        "class": ex_class(data), "sources": {cn: s(data.get(cn)) for cn, _ in editable_pairs(data)},
        "template": text_signatures(data), "owner_variant": owner_variant(data),
        "source_version": s(data.get("source_version")),
    }, ensure_ascii=False, sort_keys=True)


def source_version(master: Path):
    return hashlib.sha256(master.read_bytes()).hexdigest()


def index_workbook(wb):
    indexed = {}
    for ws in wb.worksheets:
        if ws.title not in PRIMARY:
            continue
        h = header(ws)
        if not all(k in h for k in PRIMARY[ws.title]):
            continue
        rows, dupe = {}, []
        for r in range(2, ws.max_row + 1):
            data = row_data(ws, r, h); key = key_for(ws.title, data)
            if not any(key): continue
            if key in rows: dupe.append(key)
            rows[key] = (r, data)
        indexed[ws.title] = (ws, h, rows, dupe)
    return indexed


def parameter_source_for(skill_id, data):
    # Exact display identity is never inferred from this provenance.
    return norm_key(data.get("parameter_source_id") or data.get("skill_id") or skill_id)


def level_for(data):
    explicit = s(data.get("skill_slot") or data.get("star")).strip()
    if explicit:
        return explicit
    match = re.search(r"_(\d+)$", s(data.get("skill_id")))
    return match.group(1) if match else ""


def build_payload(master: Path):
    # We perform exact indexed lookups below.  In read-only mode ``ws.cell``
    # re-streams worksheet XML, turning a harmless audit into a quadratic scan.
    # Normal mode remains read-only in intent: this workbook object is never saved.
    wb = load_workbook(master, read_only=False, data_only=False)
    indexed = index_workbook(wb)
    version = source_version(master)
    work = defaultdict(list); manifest = []; ambiguous = []; excluded = []
    source_counts = Counter(); split_candidates = defaultdict(set)
    for sheet, (ws, h, rows, _) in indexed.items():
        for row, data in (v for v in rows.values()):
            if sheet == "BUFF_STATUS":
                scope = norm_key(data.get("classification_scope")); required = norm_key(data.get("translation_required"))
                if scope == "INTERNAL_CONTROLLER" or required == "NO":
                    excluded.append({"sheet": sheet, "id": record_id(sheet, data), "reason": "INTERNAL_CONTROLLER_OR_NO"}); continue
                if scope == "AMBIGUOUS":
                    ambiguous.append({"sheet": sheet, "id": record_id(sheet, data), "reason": "AMBIGUOUS"}); continue
            if not needs_work(data): continue
            source_counts[sheet] += 1
            sig = group_signature(data) if sheet == "SKILL" else record_id(sheet, data)
            if sheet == "SKILL":
                broad = (norm_key(data.get("character_id")), norm_key(data.get("skill_group_id")), ex_class(data), s(data.get("skill_name_cn")))
                split_candidates[broad].add(sig)
            work[(sheet, sig)].append((row, data))

    groups = []
    for (sheet, signature), targets in sorted(work.items()):
        # All non-SKILL domains intentionally stay atomic in this release.
        first = targets[0][1]
        if sheet != "SKILL" or len(targets) == 1:
            buckets = [[item] for item in targets]
        else:
            buckets = [targets]
        for bucket in buckets:
            data = bucket[0][1]; rid = record_id(sheet, data)
            key_seed = json.dumps({"sheet": sheet, "signature": signature, "targets": [record_id(sheet, x[1]) for x in bucket]}, ensure_ascii=False, sort_keys=True)
            tkey = f"{sheet[:4]}_{hashlib.sha256(key_seed.encode()).hexdigest()[:20]}"
            ncn, nvi, dcn, dvi = display_values(data)
            editable = {vi: s(data.get(vi)) for _, vi in editable_pairs(data)}
            groups.append({
                "translation_key": tkey, "representative_id": rid, "target_count": len(bucket),
                "target_ids": "See MANIFEST", "character_id": s(data.get("character_id")), "record_type": sheet,
                "skill_group_id": s(data.get("skill_group_id")), "base_or_ex": ex_class(data) if sheet == "SKILL" else "N/A",
                "source_version": version, "source_hash": source_signature(sheet, data),
                "name_cn": ncn, "name_vi": nvi, "desc_cn": dcn, "desc_vi": dvi, "status": s(data.get("status")),
                "notes": s(data.get("notes")), "editable_fields_json": json.dumps(editable, ensure_ascii=False, sort_keys=True),
                "original_editable_fields_json": json.dumps(editable, ensure_ascii=False, sort_keys=True),
                "source_sheet": sheet,
            })
            for row, target in bucket:
                extra = text_signatures(target)
                manifest.append({
                    "translation_key": tkey, "target_sheet": sheet, "target_primary_id": record_id(sheet, target),
                    "character_id": s(target.get("character_id")), "skill_group_id": s(target.get("skill_group_id")),
                    "level": level_for(target),
                    "display_source_id": s(target.get("skill_id") or record_id(sheet, target)),
                    "parameter_source_id": parameter_source_for(rid, target), "source_hash": source_signature(sheet, target),
                    "source_version": version, "grouping_reason": "strict exact source/template/provenance signature" if len(bucket) > 1 else "atomic exact record",
                    "editable_fields_json": json.dumps({vi: s(target.get(vi)) for _, vi in editable_pairs(target)}, ensure_ascii=False, sort_keys=True),
                    **extra,
                })
    split = []
    for broad, variants in split_candidates.items():
        if len(variants) > 1:
            split.append({"character_id": broad[0], "skill_group_id": broad[1], "base_or_ex": broad[2], "skill_name_cn": broad[3], "strict_variants": len(variants), "reason": "template/marker/placeholder/owner/source signature differs"})
    # D0183 remains outside the current master/public roster.  Preserve its raw
    # classifier evidence in the audit instead of manufacturing master targets.
    d0183 = ROOT / "localization" / "manifests" / "D0183_dependencies.json"
    if d0183.exists():
        raw = json.loads(d0183.read_text(encoding="utf-8"))
        for node in raw.get("nodes", []):
            if node.get("classification") == "INTERNAL_CONTROLLER":
                excluded.append({"sheet": "RAW_D0183", "id": s(node.get("id")), "reason": "INTERNAL_CONTROLLER_RAW_AUDIT_OUTSIDE_MASTER"})
    return {"source_version": version, "groups": groups, "manifest": manifest, "ambiguous": ambiguous, "excluded": excluded,
            "source_counts": dict(source_counts), "split": split}


def style_sheet(ws):
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # A manifest can contain thousands of targets.  Style only stable headers
    # and bounded column widths; per-cell style assignment is needlessly slow
    # and has no bearing on import semantics.
    for cell in ws[1]:
        ws.column_dimensions[cell.column_letter].width = min(42, max(14, len(s(cell.value)) + 4))


def write_rows(ws, rows):
    columns = list(rows[0]) if rows else ["note"]
    ws.append(columns)
    for r in rows: ws.append([r.get(c, "") for c in columns])
    style_sheet(ws)


def export(master: Path, output: Path, manifest_path: Path | None = None):
    payload = build_payload(master); output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); wb.remove(wb.active)
    grouped = defaultdict(list)
    for row in payload["groups"]: grouped[row["source_sheet"]].append(row)
    for sheet, rows in grouped.items(): write_rows(wb.create_sheet(sheet[:31]), rows)
    write_rows(wb.create_sheet("MANIFEST"), payload["manifest"])
    write_rows(wb.create_sheet("AMBIGUOUS_REVIEW"), payload["ambiguous"])
    write_rows(wb.create_sheet("EXCLUDED_CONTROLLERS"), payload["excluded"])
    write_rows(wb.create_sheet("SPLIT_GROUPS"), payload["split"])
    wb.save(output)
    manifest_path = manifest_path or output.with_suffix(".manifest.json")
    manifest_path.write_text(json.dumps({k: payload[k] for k in ("source_version", "manifest", "ambiguous", "excluded", "split", "source_counts")}, ensure_ascii=False, indent=2), encoding="utf-8")
    by_sheet = Counter(x["source_sheet"] for x in payload["groups"])
    sizes = Counter(int(x["target_count"]) for x in payload["groups"])
    dry = {
        "exact_source_rows_by_sheet": payload["source_counts"],
        "translation_groups_by_sheet": dict(by_sheet),
        "row_reduction_by_sheet": {sheet: payload["source_counts"].get(sheet, 0) - by_sheet.get(sheet, 0) for sheet in payload["source_counts"]},
        "group_size_distribution": {str(k): v for k, v in sorted(sizes.items())},
        "split_groups": len(payload["split"]), "internal_controller_excluded": len(payload["excluded"]), "ambiguous": len(payload["ambiguous"]),
        "master_mutation": False,
    }
    output.with_suffix(".dry-run.json").write_text(json.dumps(dry, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload, manifest_path


def read_batch(batch: Path):
    wb = load_workbook(batch, data_only=False)
    groups = []
    for ws in wb.worksheets:
        if ws.title in {"MANIFEST", "AMBIGUOUS_REVIEW", "EXCLUDED_CONTROLLERS", "SPLIT_GROUPS"}: continue
        h = header(ws)
        if "translation_key" not in h: continue
        for r in range(2, ws.max_row + 1): groups.append(row_data(ws, r, h))
    mw = wb["MANIFEST"]; mh = header(mw)
    manifest = [row_data(mw, r, mh) for r in range(2, mw.max_row + 1)]
    return groups, manifest


def import_batch(batch: Path, master: Path, apply=False):
    groups, manifest = read_batch(batch); by_key = defaultdict(list)
    for row in manifest: by_key[s(row["translation_key"])].append(row)
    wb = load_workbook(master, data_only=False); indexed = index_workbook(wb)
    report = {"targets": [], "accepted": 0, "source_changed": 0, "keep_master": 0, "missing_target": 0, "unexpected_changes": []}
    before = {sheet: {key: [ws.cell(row, c).value for c in range(1, ws.max_column+1)] for key,(row,_) in rows.items()} for sheet,(ws,h,rows,_) in indexed.items()}
    allowed = set()
    for group in groups:
        tkey = s(group.get("translation_key")); incoming = json.loads(s(group.get("editable_fields_json")) or "{}")
        original = json.loads(s(group.get("original_editable_fields_json")) or "{}")
        if incoming == original: continue
        for target in by_key.get(tkey, []):
            sheet = s(target["target_sheet"]); key = tuple(norm_key(x) for x in s(target["target_primary_id"]).split(":"))
            if sheet not in indexed or key not in indexed[sheet][2]:
                report["targets"].append({"translation_key":tkey,"target":s(target.get("target_primary_id")),"action":"MISSING_TARGET"}); report["missing_target"] += 1; continue
            ws,h,rows,_ = indexed[sheet]; row,data = rows[key]
            if source_signature(sheet,data) != s(target["source_hash"]) or text_signatures(data)["placeholder_signature"] != s(target["placeholder_signature"]) or text_signatures(data)["ordered_buff_markers"] != s(target["ordered_buff_markers"]):
                report["targets"].append({"translation_key":tkey,"target":record_id(sheet,data),"action":"SOURCE_CHANGED"}); report["source_changed"] += 1; continue
            if owner_locked(data):
                report["targets"].append({"translation_key":tkey,"target":record_id(sheet,data),"action":"KEEP_MASTER_OWNER"}); report["keep_master"] += 1; continue
            current = {vi:s(data.get(vi)) for _,vi in editable_pairs(data)}
            target_original = json.loads(s(target.get("editable_fields_json")) or "{}")
            if any(current.get(vi, "") not in {s(target_original.get(vi)), ""} for vi in incoming):
                report["targets"].append({"translation_key":tkey,"target":record_id(sheet,data),"action":"KEEP_MASTER_NEWER"}); report["keep_master"] += 1; continue
            changes=[]
            for vi,value in incoming.items():
                if vi not in h or vi not in current: continue
                ws.cell(row,h[vi]).value=value; changes.append(vi); allowed.add((sheet,key,vi))
            if changes:
                report["targets"].append({"translation_key":tkey,"target":record_id(sheet,data),"action":"ACCEPT","fields":changes}); report["accepted"] += 1
    if apply: wb.save(master)
    # Check only allowed localization cell values differ.
    after_index=index_workbook(wb)
    for sheet,(ws,h,rows,_) in after_index.items():
        for key,(row,_) in rows.items():
            for c,name in enumerate(h,1):
                if before[sheet][key][c-1] != ws.cell(row,c).value and (sheet,key,name) not in allowed:
                    report["unexpected_changes"].append([sheet,key,name])
    report["integrity_pass"] = not report["unexpected_changes"]
    return report


def conflict_report(master: Path, old_batch: Path, output: Path):
    old = REPORT_DIR / "batch_import_20260910_132534_dry_run.json"
    data = json.loads(old.read_text(encoding="utf-8")); conflicts=[]
    current=load_workbook(master,read_only=False,data_only=False); oldwb=load_workbook(old_batch,read_only=False,data_only=False)
    for sheet in ("SKILL","BUFF_STATUS"):
        ws=current[sheet]; h=header(ws); lookup={key_for(sheet,row_data(ws,r,h)):row_data(ws,r,h) for r in range(2,ws.max_row+1)}
        bw=oldwb[sheet]; bh=header(bw)
        for decision in data["by_sheet"][sheet].get("decisions",[]):
            if decision.get("action") != "CONFLICT": continue
            key=tuple(norm_key(x) for x in decision["key"]); master_row=lookup.get(key,{})
            batch_row={}
            for r in range(2,bw.max_row+1):
                candidate=row_data(bw,r,bh)
                if key_for(sheet,candidate)==key: batch_row=candidate; break
            changed_fields=[x.get("field") for x in decision.get("fields",[]) if x.get("field")]
            cn_fields=[field[:-3]+"_cn" for field in changed_fields if field.endswith("_vi") and field[:-3]+"_cn" in master_row]
            master_vi={field:s(master_row.get(field)) for field in changed_fields}
            batch_vi={field:s(batch_row.get(field)) for field in changed_fields}
            violations=[]
            for term,canon in CANONICAL.items():
                if any(term in s(master_row.get(cn)) for cn in cn_fields) and any(batch_vi.values()) and not any(canon in value for value in batch_vi.values()): violations.append(f"{term} must remain {canon}")
            conflicts.append({"exact_id":": ".join(key),"sheet":sheet,"cn_current":json.dumps({field:s(master_row.get(field)) for field in cn_fields},ensure_ascii=False),"master_vi":json.dumps(master_vi,ensure_ascii=False),"batch_vi":json.dumps(batch_vi,ensure_ascii=False),"master_provenance_status":f"{s(master_row.get('status'))}; {s(master_row.get('notes'))}","batch_provenance_status":f"{s(batch_row.get('status'))}; {s(batch_row.get('notes'))}","canonical_term_violations":" | ".join(violations),"recommended_action":"KEEP_MASTER_CANONICAL" if violations else "","owner_decision":""})
    out=Workbook(); ws=out.active; ws.title="CONFLICT_REVIEW"; write_rows(ws,conflicts); output.parent.mkdir(parents=True,exist_ok=True); out.save(output)
    return conflicts


def self_test(batch: Path, master: Path, report_path: Path):
    groups, manifest=read_batch(batch); by_key=defaultdict(list)
    for m in manifest: by_key[s(m["translation_key"])].append(m)
    candidates={5:next((g for g in groups if int(g.get("target_count") or 0)==5 and s(g.get("source_sheet"))=="SKILL"),None),
                3:next((g for g in groups if int(g.get("target_count") or 0)==3 and s(g.get("source_sheet"))=="SKILL" and s(g.get("representative_id")).endswith("04_1")),None)}
    if not all(candidates.values()): raise RuntimeError("Need both 5-level and 3-level SKILL groups for round-trip test")
    td = ROOT / "localization" / "test_artifacts" / f"compact_round_trip_{datetime.now():%Y%m%d_%H%M%S}"
    td.mkdir(parents=True, exist_ok=False)
    test_master=td/"master.xlsx"; test_batch=td/"batch.xlsx"; shutil.copy2(master,test_master); shutil.copy2(batch,test_batch)
    # Keep this disposable copy for inspection; it is never the real master.
    
    bw=load_workbook(test_batch); changed=[]
    for count,g in candidates.items():
        ws=bw[s(g["source_sheet"])]; h=header(ws)
        for r in range(2,ws.max_row+1):
            if s(ws.cell(r,h["translation_key"]).value)==s(g["translation_key"]):
                values=json.loads(s(ws.cell(r,h["editable_fields_json"]).value)); field=next((x for x in values if "desc" in x),next(iter(values)))
                values[field]=f"__ROUNDTRIP_{count}_LEVEL_TEST__"; ws.cell(r,h["editable_fields_json"]).value=json.dumps(values,ensure_ascii=False); changed.append((count,g,field)); break
    bw.save(test_batch)
    result=import_batch(test_batch,test_master,apply=True)
    # Alter one exact source field after the first import then re-import modified batch: it must block.
    altered=changed[0][1]; target=by_key[s(altered["translation_key"])][0]; mw=load_workbook(test_master); ws=mw[s(target["target_sheet"])]; h=header(ws); target_id=s(target["target_primary_id"]).split(":");
    for r in range(2,ws.max_row+1):
        data=row_data(ws,r,h)
        if list(key_for(ws.title,data))==[norm_key(x) for x in target_id]:
            cn,_=first_pair(data); ws.cell(r,h[cn]).value=s(ws.cell(r,h[cn]).value)+" [source-change-test]"; break
    mw.save(test_master); blocked=import_batch(test_batch,test_master,apply=False)
    expected={n:len(by_key[s(g["translation_key"])]) for n,g in candidates.items()}
    report={"normal_5_expected_targets":expected[5],"passive_3_expected_targets":expected[3],"first_import_accepted":result["accepted"],"first_import_integrity":result["integrity_pass"],"source_changed_blocked":blocked["source_changed"]>0,"no_unexpected_cells":not result["unexpected_changes"],"test_copy":str(td),"pass":result["integrity_pass"] and blocked["source_changed"]>0 and all(expected.values())}
    report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8"); return report


def identity_audit(master: Path, output: Path):
    """Evidence that exact display records stay separate from raw parameters."""
    wb=load_workbook(master,read_only=False,data_only=False); ws=wb["SKILL"]; h=header(ws)
    wanted={"A016004_1","A016004_2","A016004_3","A016004ex"}; records=[]
    for r in range(2,ws.max_row+1):
        data=row_data(ws,r,h)
        if s(data.get("skill_id")) in wanted:
            records.append({"skill_id":s(data.get("skill_id")),"skill_name_cn":s(data.get("skill_name_cn")),"skill_name_vi":s(data.get("skill_name_vi")),"desc_vi":s(data.get("desc_vi")),"display_source_id":s(data.get("skill_id")),"parameter_source_id":parameter_source_for(s(data.get("skill_id")),data)})
    raw_dir=ROOT.parent/"NeoArtifacts"/"MasterData"/"json"
    skills=json.loads((raw_dir/"skillMap.json").read_text(encoding="utf-8")); roles=json.loads((raw_dir/"roleattrMap.json").read_text(encoding="utf-8"))
    raw_levels={k:{"group_id":s(v.get("GroupId")),"level":v.get("level"),"name_cn":s(v.get("NameLanText"))} for k,v in skills.items() if k.startswith("A016004")}
    relations=[]
    for key,value in roles.items():
        if not key.startswith("A0160") or not isinstance(value,dict): continue
        for field in ("starUpAttr","starUpAttr2"):
            for entry in value.get(field,[]) if isinstance(value.get(field),list) else [value.get(field)]:
                if entry == "SkillUP,A016004": relations.append({"roleattr_id":key,"relation":entry,"enhanced_display_id":"A016004ex"})
    report={"base_records":[x for x in records if x["skill_id"]!="A016004ex"],"ex_record":[x for x in records if x["skill_id"]=="A016004ex"],"raw_level_records":raw_levels,"skillup_relations":relations,"canonical_terms":{term:canon for term,canon in CANONICAL.items() if term in {"万籁沉寂","瞄准","蓄势"}},"pass":bool(records) and all(x["skill_name_vi"]=="Vạn Âm Trầm Tịch" for x in records) and any("Súc Thế" in x["desc_vi"] for x in records if x["skill_id"]=="A016004ex")}
    output.parent.mkdir(parents=True,exist_ok=True); output.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8"); return report


def main():
    p=argparse.ArgumentParser(); sub=p.add_subparsers(dest="cmd",required=True)
    e=sub.add_parser("export"); e.add_argument("--master",type=Path,default=MASTER); e.add_argument("--output",type=Path); e.add_argument("--manifest",type=Path)
    i=sub.add_parser("import"); i.add_argument("--batch",type=Path,required=True); i.add_argument("--master",type=Path,required=True); i.add_argument("--apply",action="store_true"); i.add_argument("--report",type=Path)
    c=sub.add_parser("conflicts"); c.add_argument("--master",type=Path,default=MASTER); c.add_argument("--old-batch",type=Path,default=BATCH_DIR/"character_translation_batch_skill_huanzhang_completed.xlsx"); c.add_argument("--output",type=Path,required=True)
    t=sub.add_parser("self-test"); t.add_argument("--batch",type=Path,required=True); t.add_argument("--master",type=Path,default=MASTER); t.add_argument("--report",type=Path,required=True)
    q=sub.add_parser("identity-audit"); q.add_argument("--master",type=Path,default=MASTER); q.add_argument("--output",type=Path,required=True)
    a=p.parse_args();
    if a.cmd=="export":
        out=a.output or BATCH_DIR/f"compact_translation_batch_{datetime.now():%Y%m%d_%H%M%S}.xlsx"; payload,manifest=export(a.master,out,a.manifest); print(json.dumps({"batch":str(out),"manifest":str(manifest),"source_counts":payload["source_counts"],"groups":len(payload["groups"]),"targets":len(payload["manifest"]),"excluded":len(payload["excluded"]),"ambiguous":len(payload["ambiguous"])},ensure_ascii=False))
    elif a.cmd=="import":
        report=import_batch(a.batch,a.master,a.apply); out=a.report or REPORT_DIR/f"compact_import_{datetime.now():%Y%m%d_%H%M%S}.json"; out.parent.mkdir(parents=True,exist_ok=True); out.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8"); print(json.dumps({"report":str(out),**{k:report[k] for k in ("accepted","source_changed","keep_master","missing_target","integrity_pass")}},ensure_ascii=False))
    elif a.cmd=="conflicts": print(json.dumps({"count":len(conflict_report(a.master,a.old_batch,a.output)),"output":str(a.output)},ensure_ascii=False))
    elif a.cmd=="self-test": print(json.dumps(self_test(a.batch,a.master,a.report),ensure_ascii=False))
    else: print(json.dumps(identity_audit(a.master,a.output),ensure_ascii=False))

if __name__ == "__main__": main()
