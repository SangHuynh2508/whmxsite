"""Atomic, scope-checked mutation for the localization master workbook."""
from __future__ import annotations

from datetime import datetime
import gc
import openpyxl
import os
import shutil
import time


class AtomicReplaceLockError(PermissionError):
    """A bounded atomic-replace attempt was blocked by a Windows file lock."""

    def __init__(self, master_path, temp_path, backup_path, cause):
        super().__init__(*getattr(cause, "args", (str(cause),)))
        self.master_path = master_path
        self.temp_path = temp_path
        self.backup_path = backup_path
        self.cause = cause


def create_workbook_backup(base_dir):
    """Create a timestamped safety backup of ``localization_master.xlsx``."""
    master_path = os.path.join(base_dir, "localization", "localization_master.xlsx")
    backup_dir = os.path.join(base_dir, "localization", "backups")
    os.makedirs(backup_dir, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_path = os.path.join(backup_dir, f"localization_master_{timestamp}.xlsx")
    shutil.copy2(master_path, backup_path)
    return backup_path


def _row_is_authorized(sheetname, row, authorized_deps):
    record_id = str((row[0] if row else None) or "").strip()
    if sheetname == "SKILL":
        return record_id in set(authorized_deps.get("skill_ids", ()))
    if sheetname == "BUFF_STATUS":
        return record_id in set(authorized_deps.get("player_facing_buff_ids", ()))
    if sheetname == "HUANZHANG":
        return record_id in set(authorized_deps.get("huanzhang_ids", ()))
    if sheetname == "SKIN":
        return record_id in set(authorized_deps.get("skin_ids", ()))
    if sheetname == "ITEM":
        return record_id in set(authorized_deps.get("item_ids", ()))
    return False


def _report_current_process_open_files(master_path):
    """Emit an exact self-lock diagnostic immediately before ``os.replace``.

    This is deliberately best-effort: production mutation must remain usable
    where ``psutil`` is not installed.  The report is useful when Windows has
    an enduring sharing violation rather than the usual short Defender handle.
    """
    normalized_master = os.path.normcase(os.path.abspath(master_path))
    print(f"[SAFE MUTATION LOCK DIAG] CURRENT_PYTHON_PID={os.getpid()}")
    try:
        import psutil
    except ImportError:
        print("[SAFE MUTATION LOCK DIAG] PSUTIL_AVAILABLE=NO")
        return
    process = psutil.Process()
    open_files = process.open_files()
    print(f"[SAFE MUTATION LOCK DIAG] PSUTIL_AVAILABLE=YES OPEN_FILES_COUNT={len(open_files)}")
    master_is_open = False
    for opened in open_files:
        print(f"[SAFE MUTATION LOCK DIAG] OPEN_FILE={opened.path!r}")
        if os.path.normcase(os.path.abspath(opened.path)) == normalized_master:
            master_is_open = True
    print(f"[SAFE MUTATION LOCK DIAG] MASTER_IN_CURRENT_PROCESS_OPEN_FILES={master_is_open}")


def safe_mutate_workbook(
    base_dir,
    mutator_fn,
    authorized_deps,
    authorized_new_columns=None,
    authorized_cells=None,
    authorized_new_rows=None,
    authorized_deleted_rows=None,
):
    """Mutate master only after topology and exact-ID scope verification.

    ``authorized_cells`` contains ``(sheet, record_id, column_header)`` tuples
    for the rare explicit cell-level repair on an otherwise unauthorized row.
    No character-level authorization exists in this guard.
    """
    master_path = os.path.join(base_dir, "localization", "localization_master.xlsx")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_dir = os.path.join(base_dir, "localization", "backups")
    os.makedirs(backup_dir, exist_ok=True)
    backup_path = os.path.join(backup_dir, f"localization_master_{timestamp}.xlsx")
    temp_path = master_path + f".mutation_{timestamp}.tmp.xlsx"
    if os.path.exists(temp_path):
        raise FileExistsError(f"Refusing to overwrite existing mutation file: {temp_path}")

    shutil.copy2(master_path, backup_path)
    print(f"[SAFE MUTATION] Created timestamped backup: {backup_path}")
    shutil.copy2(master_path, temp_path)

    wb_temp = openpyxl.load_workbook(temp_path)
    try:
        mutator_fn(wb_temp)
        wb_temp.save(temp_path)
    finally:
        wb_temp.close()
        # On Windows, keeping the Workbook object alive can retain the ZIP
        # archive handle even after ``close``.  The temporary path becomes the
        # source of an atomic replace below, so release it deterministically.
        del wb_temp

    # Reopen the saved file before comparing it. ``read_only`` keeps the full
    # workbook check fast without weakening row or appended-column checks.
    try:
        wb_check = openpyxl.load_workbook(temp_path, data_only=False, read_only=True)
    except Exception as exc:
        raise RuntimeError(f"Mutation failed validation: {exc}") from exc
    wb_backup = openpyxl.load_workbook(backup_path, data_only=False, read_only=True)

    authorized_new_columns = authorized_new_columns or {}
    authorized_new_rows = {sheet: set(ids) for sheet, ids in (authorized_new_rows or {}).items()}
    authorized_deleted_rows = {sheet: set(ids) for sheet, ids in (authorized_deleted_rows or {}).items()}
    authorized_cells = set(authorized_cells or ())
    unrelated_diffs = []
    ws_before = None
    ws_after = None
    before_rows = None
    after_rows = None
    try:
        if wb_backup.sheetnames != wb_check.sheetnames:
            raise PermissionError(
                f"Sheet topology changed: before={wb_backup.sheetnames}, after={wb_check.sheetnames}"
            )

        for sheetname in wb_backup.sheetnames:
            ws_before = wb_backup[sheetname]
            ws_after = wb_check[sheetname]
            before_headers = next(ws_before.iter_rows(min_row=1, max_row=1, values_only=True))
            after_headers = next(ws_after.iter_rows(min_row=1, max_row=1, values_only=True))
            before_width, after_width = len(before_headers), len(after_headers)
            if after_headers[:before_width] != before_headers:
                raise PermissionError(f"Header topology changed in {sheetname}")
            appended_headers = after_headers[before_width:]
            if not set(appended_headers).issubset(set(authorized_new_columns.get(sheetname, ()))):
                raise PermissionError(f"Unauthorized appended headers in {sheetname}: {appended_headers}")
            if ws_after.max_row < ws_before.max_row:
                allowed_deleted_ids = authorized_deleted_rows.get(sheetname, set())
                if not allowed_deleted_ids:
                    raise PermissionError(f"Rows removed in {sheetname}: {ws_before.max_row} -> {ws_after.max_row}")
                before_id_col = [str(r[0] or "").strip() for r in ws_before.iter_rows(min_row=2, max_col=1, values_only=True)]
                after_id_col = [str(r[0] or "").strip() for r in ws_after.iter_rows(min_row=2, max_col=1, values_only=True)]
                from collections import Counter
                deleted_counts = Counter(before_id_col) - Counter(after_id_col)
                if set(deleted_counts.keys()) != allowed_deleted_ids:
                    raise PermissionError(
                        f"Deleted rows in {sheetname} do not exactly match authorized deleted IDs: "
                        f"{set(deleted_counts.keys())!r} vs {sorted(allowed_deleted_ids)!r}"
                    )
            if ws_after.max_row > ws_before.max_row:
                allowed_ids = authorized_new_rows.get(sheetname, set())
                if not allowed_ids:
                    raise PermissionError(
                        f"Unauthorized rows appended in {sheetname}: {ws_before.max_row} -> {ws_after.max_row}"
                    )
                existing_ids = {
                    str(row[0] or "").strip()
                    for row in ws_before.iter_rows(min_row=2, max_col=1, values_only=True)
                }
                appended_ids = []
                for row in ws_after.iter_rows(min_row=ws_before.max_row + 1, max_col=after_width, values_only=True):
                    record_id = str((row[0] if row else None) or "").strip()
                    appended_ids.append(record_id)
                    if not record_id or record_id in existing_ids or record_id not in allowed_ids:
                        raise PermissionError(
                            f"Unauthorized appended row in {sheetname}: {record_id!r}"
                        )
                if set(appended_ids) != allowed_ids or len(appended_ids) != len(allowed_ids):
                    raise PermissionError(
                        f"Appended rows in {sheetname} do not exactly match authorized IDs: "
                        f"{appended_ids!r} vs {sorted(allowed_ids)!r}"
                    )

            # Invariant: verify that no new duplicate primary keys were introduced
            from collections import Counter
            after_id_col = [str(row[0] or "").strip() for row in ws_after.iter_rows(min_row=2, max_col=1, values_only=True) if row and str(row[0] or "").strip()]
            before_id_col = [str(row[0] or "").strip() for row in ws_before.iter_rows(min_row=2, max_col=1, values_only=True) if row and str(row[0] or "").strip()]
            after_id_counts = Counter(after_id_col)
            before_id_counts = Counter(before_id_col)
            new_duplicate_ids = {k for k, count in after_id_counts.items() if count > 1 and count > before_id_counts.get(k, 0)}
            if new_duplicate_ids:
                raise PermissionError(
                    f"Duplicate primary key invariant violated in {sheetname}: {new_duplicate_ids}"
                )

            before_rows = ws_before.iter_rows(min_row=2, max_col=before_width, values_only=True)
            after_rows = ws_after.iter_rows(min_row=2, max_col=after_width, values_only=True)
            allowed_deleted_ids = authorized_deleted_rows.get(sheetname, set())
            if allowed_deleted_ids:
                before_it = iter(before_rows)
                for row_number, after_row in enumerate(after_rows, start=2):
                    while True:
                        before_row = next(before_it, None)
                        if before_row is None:
                            raise PermissionError(f"Row count mismatch during deletion alignment in {sheetname}")
                        padded_before = before_row + (None,) * (after_width - before_width)
                        if padded_before == after_row:
                            break
                        rec_id = str((before_row[0] if before_row else None) or "").strip()
                        if rec_id in allowed_deleted_ids:
                            continue
                        if _row_is_authorized(sheetname, before_row, authorized_deps):
                            break
                        for col_idx, (b_val, a_val) in enumerate(zip(padded_before, after_row)):
                            if b_val != a_val:
                                unrelated_diffs.append((sheetname, row_number, after_headers[col_idx], b_val, a_val))
                        break
                for remaining_before in before_it:
                    rec_id = str((remaining_before[0] if remaining_before else None) or "").strip()
                    if rec_id not in allowed_deleted_ids:
                        unrelated_diffs.append((sheetname, 0, "row_deletion", rec_id, None))
            else:
                for row_number, (before_row, after_row) in enumerate(zip(before_rows, after_rows), start=2):
                    padded_before = before_row + (None,) * (after_width - before_width)
                    if padded_before == after_row:
                        continue
                    if _row_is_authorized(sheetname, before_row, authorized_deps):
                        continue
                    record_id = str((before_row[0] if before_row else None) or "").strip()
                    for column_index, (before_value, after_value) in enumerate(zip(padded_before, after_row)):
                        if before_value == after_value:
                            continue
                        column_name = after_headers[column_index]
                        if (sheetname, record_id, column_name) not in authorized_cells:
                            unrelated_diffs.append(
                                (sheetname, row_number, column_name, before_value, after_value)
                            )
    finally:
        wb_backup.close()
        wb_check.close()
        # The worksheet variables created while validating are children of
        # these workbooks.  Dropping all workbook references before the
        # replace prevents this process from self-locking the temp xlsx.
        ws_before = ws_after = None
        before_rows = after_rows = None
        del wb_backup
        del wb_check

    if unrelated_diffs:
        print(f"[SCOPE GUARD FAIL] Found {len(unrelated_diffs)} unauthorized cell mutations!")
        for diff in unrelated_diffs[:10]:
            print(f"  {diff[0]} row={diff[1]} column={diff[2]!r}: {diff[3]!r} -> {diff[4]!r}")
        raise PermissionError("Mutation violated exact ID/cell scope guard")

    # Excel/Defender can retain a short-lived handle after an xlsx save. Retry
    # only that transient Windows sharing violation; any lasting lock keeps the
    # checked temporary workbook for forensic recovery.
    # All read-only source guards must be materialized and closed before this
    # point.  Collecting here also releases any ref-count cycle left by a ZIP
    # reader before checking our own handles immediately before replacement.
    gc.collect()
    for attempt in range(20):
        try:
            _report_current_process_open_files(master_path)
            os.replace(temp_path, master_path)
            break
        except PermissionError as exc:
            if getattr(exc, "winerror", None) == 32 and attempt == 19:
                raise AtomicReplaceLockError(master_path, temp_path, backup_path, exc) from exc
            if getattr(exc, "winerror", None) != 32:
                raise
            time.sleep(0.5)
    print(f"[SAFE MUTATION SUCCESS] Atomic replace completed on {master_path}")
    return backup_path
