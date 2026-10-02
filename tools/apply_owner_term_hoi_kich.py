"""Owner terminology 2026-10-02: 回击 (the game's counterattack) reads "Phản kích", not "Hồi kích".

Vietnamese players know the mechanic as "phản kích" (GLOSSARY already has 反击 → Phản kích, OWNER_GAMEPLAY_009; the
game's text uses 回击, which had been rendered literally). Only VI cells whose row's CN source contains 回击 are touched,
case is kept ("Hồi Kích" → "Phản Kích", "hồi kích" → "phản kích"), statuses are left as they are, and the decision is
recorded as a GLOSSARY row. Runs through tools/safe_workbook_mutation.py (backup, exact-cell scope check).

    python tools/apply_owner_term_hoi_kich.py            # dry run: list the cells
    python tools/apply_owner_term_hoi_kich.py --apply    # write
"""
from __future__ import annotations

import argparse
import os
import re
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(__file__))
from safe_workbook_mutation import safe_mutate_workbook  # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MASTER = os.path.join(BASE, "localization", "localization_master.xlsx")
TERM = re.compile(r"(?<![Tt]hu )([Hh])ồi ([Kk])ích(?! hoạt)")  # not "thu hồi" + "kích hoạt"
GLOSSARY_ID = "OWNER_GAMEPLAY_032"
GLOSSARY_ROW = {
    "term_id": GLOSSARY_ID, "category": "gameplay_terminology", "term_cn": "回击", "term_vi": "Phản kích",
    "han_viet": "Hồi Kích", "confidence": "HIGH", "status": "OWNER_APPROVED",
    "notes": "Owner 2026-10-02: 回击 shown as Phản kích (familiar to Vietnamese players), like 反击 (OWNER_GAMEPLAY_009).",
}


def swap(text: str) -> str:
    return TERM.sub(lambda m: f"{'P' if m.group(1) == 'H' else 'p'}hản {m.group(2)}ích", text)


def targets(wb):
    """(sheet, record_id, column) of every VI cell holding the term whose row's CN source contains 回击."""
    out = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "") for h in next(rows, ())]
        for row in rows:
            if not row or not row[0]:
                continue
            cn = " ".join(v for h, v in zip(header, row) if h.endswith("_cn") and isinstance(v, str))
            for h, v in zip(header, row):
                if h.endswith("_vi") and isinstance(v, str) and TERM.search(v):
                    if "回击" not in cn:
                        raise SystemExit(f"Stop: {ws.title} {row[0]} {h} has the term but no 回击 in its CN source")
                    out.append((ws.title, str(row[0]).strip(), h))
    return out


def mutate(cells):
    wanted = {}
    for sheet, rid, col in cells:
        wanted.setdefault(sheet, {}).setdefault(rid, set()).add(col)

    def run(wb):
        for sheet, rows in wanted.items():
            ws = wb[sheet]
            header = [str(c.value or "") for c in ws[1]]
            for row in ws.iter_rows(min_row=2):
                rid = str(row[0].value or "").strip()
                for col in rows.get(rid, ()):
                    cell = row[header.index(col)]
                    cell.value = swap(cell.value)
        ws = wb["GLOSSARY"]
        header = [str(c.value or "") for c in ws[1]]
        if any(str(r[0].value or "").strip() == GLOSSARY_ID for r in ws.iter_rows(min_row=2)):
            raise SystemExit(f"{GLOSSARY_ID} already exists")
        ws.append([GLOSSARY_ROW.get(h, None) for h in header])

    return run


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()
    wb = load_workbook(MASTER, read_only=True)
    cells = targets(wb)
    wb.close()  # a read-only workbook keeps the file open; Windows then refuses the atomic replace
    del wb
    by_sheet = {}
    for s, _, c in cells:
        by_sheet[f"{s}.{c}"] = by_sheet.get(f"{s}.{c}", 0) + 1
    print(f"{len(cells)} cells: {by_sheet}")
    if not args.apply:
        for s, rid, c in cells[:5]:
            print(" ", s, rid, c)
        return
    safe_mutate_workbook(BASE, mutate(cells), authorized_deps={}, authorized_cells=set(cells),
                         authorized_new_rows={"GLOSSARY": {GLOSSARY_ID}})
    check = load_workbook(MASTER, read_only=True)
    left = targets(check)
    check.close()
    print(f"applied; cells still holding the old term: {len(left)}")


if __name__ == "__main__":
    main()
