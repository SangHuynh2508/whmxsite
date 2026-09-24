"""Apply a tiny source-verified localization transaction with a safe backup."""
from __future__ import annotations
import argparse, hashlib, json, shutil
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
MASTER=ROOT/'localization'/'localization_master.xlsx'

def s(v): return '' if v is None else str(v)
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def header(ws): return {s(c.value):c.column for c in ws[1] if c.value is not None}

def main():
 p=argparse.ArgumentParser();p.add_argument('--transaction',type=Path,required=True);p.add_argument('--apply',action='store_true');p.add_argument('--report',type=Path,required=True);a=p.parse_args()
 locks=list(MASTER.parent.glob(f'~${MASTER.name}'))
 if locks: raise RuntimeError(f'Excel lock: {locks}')
 tx=json.loads(a.transaction.read_text(encoding='utf8'));wb=load_workbook(MASTER,data_only=False);ws=wb['BUFF_STATUS'];h=header(ws)
 rows={s(ws.cell(r,h['buff_id']).value):(r,{x:ws.cell(r,c).value for x,c in h.items()}) for r in range(2,ws.max_row+1)}
 plan=[]; failures=[]
 for target in tx['targets']:
  bid=target['buff_id']
  if bid not in rows: failures.append({'buff_id':bid,'reason':'MISSING_TARGET'});continue
  r,data=rows[bid]
  if s(data.get('buff_name_cn'))!=target['expected_name_cn'] or s(data.get('buff_desc_cn'))!=target['expected_desc_cn']: failures.append({'buff_id':bid,'reason':'SOURCE_CHANGED'});continue
  plan.append((r,target,data))
 report={'transaction_id':tx['transaction_id'],'planned':len(plan),'failures':failures,'applied':[]}
 if a.apply:
  if failures: raise RuntimeError('Refusing partial transaction')
  backup=ROOT/'localization'/'backups'/f'localization_master_pre_{tx["transaction_id"].lower()}_{datetime.now():%Y%m%d_%H%M%S}.xlsx';backup.parent.mkdir(parents=True,exist_ok=True);before=sha(MASTER);shutil.copy2(MASTER,backup)
  if sha(backup)!=before: raise RuntimeError('backup hash mismatch')
  for r,target,data in plan:
   changes=[]
   for field in ('buff_name_vi','buff_desc_vi'):
    old=s(ws.cell(r,h[field]).value);new=target[field];ws.cell(r,h[field]).value=new;changes.append({'field':field,'old':old,'new':new})
   ws.cell(r,h['status']).value='TRANSLATED';note=s(ws.cell(r,h['notes']).value);marker=f"Imported {tx['transaction_id']}; requires review";ws.cell(r,h['notes']).value=(note.rstrip('; ')+'; '+marker).strip('; ')
   report['applied'].append({'buff_id':target['buff_id'],'changes':changes})
  wb.save(MASTER);report.update({'backup':str(backup),'backup_sha256':before})
 a.report.parent.mkdir(parents=True,exist_ok=True);a.report.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps({'report':str(a.report),'planned':len(plan),'applied':len(report['applied']),'failures':len(failures)},ensure_ascii=False))
if __name__=='__main__':main()
