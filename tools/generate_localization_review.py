import openpyxl
import json
import os
import sys
from batch_manifest_helper import load_batch_manifest, collect_batch_dependencies
from zhizhi_resolver import resolve_zhizhi_entry, contains_raw_dev_codes

def generate_review(base_dir, manifest_path, wb_master=None):
    manifest = load_batch_manifest(manifest_path)
    batch_id = manifest['batch_id']
    deps = collect_batch_dependencies(base_dir, manifest)
    
    if wb_master is None:
        master_path = os.path.join(base_dir, 'localization', 'localization_master.xlsx')
        if not os.path.exists(master_path):
            raise FileNotFoundError(f"Master workbook not found: {master_path}")
        wb = openpyxl.load_workbook(master_path, data_only=True)
    else:
        wb = wb_master
    
    # 1. Map character names
    ws_char = wb['CHARACTER']
    headers_char = [str(c or '').strip() for c in next(ws_char.iter_rows(values_only=True))]
    cid_idx = headers_char.index('character_id')
    name_cn_idx = headers_char.index('name_cn')
    name_vi_idx = headers_char.index('name_vi')
    
    char_names = {}
    for r in ws_char.iter_rows(min_row=2, values_only=True):
        cid = str(r[cid_idx] or '').strip()
        if cid in deps['characters']:
            c_name_cn = str(r[name_cn_idx] or '').strip()
            c_name_vi = str(r[name_vi_idx] or '').strip()
            char_names[cid] = {'cn': c_name_cn, 'vi': c_name_vi}
            
    # 2. Build skill names map for Zhizhi resolver
    ws_skill = wb['SKILL']
    headers_skill = [str(c or '').strip() for c in next(ws_skill.iter_rows(values_only=True))]
    cid_sk_idx = headers_skill.index('character_id')
    sid_sk_idx = headers_skill.index('skill_id')
    name_vi_sk_idx = headers_skill.index('skill_name_vi')
    
    skill_names_map = {}
    skill_rows_by_char = {c: [] for c in deps['characters']}
    
    for r in ws_skill.iter_rows(min_row=2, values_only=True):
        cid = str(r[cid_sk_idx] or '').strip()
        sid = str(r[sid_sk_idx] or '').strip()
        name_vi = str(r[name_vi_sk_idx] or '').strip()
        if name_vi and name_vi != 'None':
            skill_names_map[sid] = name_vi
            if len(sid) >= 7:
                skill_names_map[sid[:7]] = name_vi
        if cid in skill_rows_by_char:
            skill_rows_by_char[cid].append(dict(zip(headers_skill, r)))

    # 3. Read BUFF_STATUS rows
    ws_buff = wb['BUFF_STATUS']
    headers_buff = [str(c or '').strip() for c in next(ws_buff.iter_rows(values_only=True))]
    bid_bf_idx = headers_buff.index('buff_id')
    
    buff_rows_map = {}
    for r in ws_buff.iter_rows(min_row=2, values_only=True):
        bid = str(r[bid_bf_idx] or '').strip()
        buff_rows_map[bid] = dict(zip(headers_buff, r))

    # 4. Read ZHIZHI rows
    ws_zhizhi = wb['ZHIZHI']
    headers_zh = [str(c or '').strip() for c in next(ws_zhizhi.iter_rows(values_only=True))]
    cid_zh_idx = headers_zh.index('character_id')
    
    zhizhi_rows_by_char = {c: [] for c in deps['characters']}
    for r in ws_zhizhi.iter_rows(min_row=2, values_only=True):
        cid = str(r[cid_zh_idx] or '').strip()
        if cid in zhizhi_rows_by_char:
            zhizhi_rows_by_char[cid].append(dict(zip(headers_zh, r)))

    # 5. Read HUANZHANG rows
    ws_hz = wb['HUANZHANG']
    headers_hz = [str(c or '').strip() for c in next(ws_hz.iter_rows(values_only=True))]
    cid_hz_idx = headers_hz.index('character_id')
    
    huanzhang_rows_by_char = {c: [] for c in deps['characters']}
    for r in ws_hz.iter_rows(min_row=2, values_only=True):
        cid = str(r[cid_hz_idx] or '').strip()
        if cid in huanzhang_rows_by_char:
            huanzhang_rows_by_char[cid].append(dict(zip(headers_hz, r)))

    # Assemble review data per character
    review_json = {
        'batch_id': batch_id,
        'characters': []
    }
    
    char_summary_str = ', '.join([f"{c} ({char_names.get(c, {}).get('vi', '')})" for c in deps['characters']])
    md_lines = [
        f"# LOCALIZATION REVIEW PACKAGE — BATCH {batch_id.upper()}",
        "",
        "Generated from: `localization/localization_master.xlsx`",
        f"Characters included: {char_summary_str}",
        "",
        "---",
        ""
    ]
    
    for cid in deps['characters']:
        c_info = char_names.get(cid, {'cn': '', 'vi': ''})
        c_sk_rows = skill_rows_by_char.get(cid, [])
        c_zh_rows = sorted(zhizhi_rows_by_char.get(cid, []), key=lambda x: x.get('star', 0))
        c_hz_rows = huanzhang_rows_by_char.get(cid, [])
        
        # Determine referenced buffs for this character
        c_buff_ids = set()
        for sk in c_sk_rows:
            desc_cn = str(sk.get('desc_cn') or '')
            desc_vi = str(sk.get('desc_vi') or '')
            for text in [desc_cn, desc_vi]:
                import re
                for m in re.findall(r'\{Buff_([^\}]+)\}', text):
                    bid = 'Buff_' + m
                    if bid in deps['player_facing_buff_ids']:
                        c_buff_ids.add(bid)
        if cid == 'A0160':
            c_buff_ids.add('Buff_A0160_hz_1')
        if cid == 'V0117':
            c_buff_ids.add('Buff_V0117_hz_1_1')
        
        c_buff_rows = [buff_rows_map[b] for b in sorted(list(c_buff_ids)) if b in buff_rows_map]
        
        # Skill groups
        skill_groups = {}
        for sk in c_sk_rows:
            sg_id = str(sk.get('skill_group_id') or sk.get('skill_id', '')[:7])
            if sg_id not in skill_groups:
                skill_groups[sg_id] = []
            skill_groups[sg_id].append(sk)
            
        md_lines.append(f"## {cid} — {c_info.get('vi', '')} ({c_info.get('cn', '')})")
        md_lines.append("")
        md_lines.append(f"**Counts**: Skills = {len(skill_groups)} groups ({len(c_sk_rows)} level rows) | Buffs = {len(c_buff_rows)} | Zhizhi = {len(c_zh_rows)} ranks | Huanzhang = {len(c_hz_rows)}")
        md_lines.append("")
        md_lines.append("### 1. SKILL")
        md_lines.append("")
        
        char_sk_json = []
        for sg_id, rows in skill_groups.items():
            first_r = rows[0]
            slot_label = first_r.get('type_label') or first_r.get('skill_slot') or ''
            name_vi = first_r.get('skill_name_vi') or ''
            name_cn = first_r.get('skill_name_cn') or ''
            
            if not name_vi or name_vi == 'None':
                name_vi = '[MISSING VI NAME]'
                
            md_lines.append(f"#### [{sg_id}] {slot_label}: {name_vi} ({name_cn})")
            md_lines.append("")
            
            for sk in rows:
                sid = sk.get('skill_id')
                status = sk.get('status') or 'UNKNOWN'
                conf = sk.get('confidence') or 'UNKNOWN'
                d_cn = sk.get('desc_cn') or ''
                d_vi = sk.get('desc_vi') or ''
                if not d_vi or d_vi == 'None' or d_vi.strip() == '':
                    d_vi = '[MISSING VI DESCRIPTION]'
                    
                md_lines.append(f"- **ID**: `{sid}` | **Status**: `{status}` | **Confidence**: `{conf}`")
                md_lines.append(f"  - **CN**: {d_cn}")
                md_lines.append(f"  - **VI**: {d_vi}")
                md_lines.append("")
                
                char_sk_json.append({
                    'skill_id': sid,
                    'name_cn': name_cn,
                    'name_vi': name_vi,
                    'desc_cn': d_cn,
                    'desc_vi': d_vi,
                    'status': status,
                    'confidence': conf
                })
                
        md_lines.append("### 2. BUFF_STATUS")
        md_lines.append("")
        
        char_buff_json = []
        if c_buff_rows:
            for bf in c_buff_rows:
                bid = bf.get('buff_id')
                b_name_cn = bf.get('buff_name_cn') or ''
                b_name_vi = bf.get('buff_name_vi') or ''
                b_desc_cn = bf.get('buff_desc_cn') or ''
                b_desc_vi = bf.get('buff_desc_vi') or ''
                status = bf.get('status') or 'UNKNOWN'
                conf = bf.get('confidence') or 'UNKNOWN'
                
                if not b_name_vi or b_name_vi == 'None':
                    b_name_vi = '[MISSING VI NAME]'
                if not b_desc_vi or b_desc_vi == 'None' or b_desc_vi.strip() == '':
                    b_desc_vi = '[MISSING VI DESCRIPTION]'
                    
                md_lines.append(f"#### [{bid}] {b_name_vi} ({b_name_cn})")
                md_lines.append(f"- **Status**: `{status}` | **Confidence**: `{conf}`")
                md_lines.append(f"- **Desc CN**: {b_desc_cn}")
                md_lines.append(f"- **Desc VI**: {b_desc_vi}")
                md_lines.append("")
                
                char_buff_json.append({
                    'buff_id': bid,
                    'name_cn': b_name_cn,
                    'name_vi': b_name_vi,
                    'desc_cn': b_desc_cn,
                    'desc_vi': b_desc_vi,
                    'status': status,
                    'confidence': conf
                })
        else:
            md_lines.append("*No character-specific player-facing buffs referenced.*")
            md_lines.append("")

        md_lines.append("### 3. ZHIZHI")
        md_lines.append("")
        
        char_zh_json = []
        for zh in c_zh_rows:
            star = zh.get('star')
            rank_vi = zh.get('rank_numeral_vi') or f"Trí Tri {star}"
            eff_type = zh.get('effect_type')
            sum_cn = zh.get('effect_summary_cn') or ''
            base_id = zh.get('skill_up_base_id')
            status = zh.get('status') or 'UNKNOWN'
            
            resolved_vi = resolve_zhizhi_entry(cid, star, eff_type, sum_cn, base_id, skill_names_map)
            if not resolved_vi or resolved_vi == 'None' or resolved_vi.strip() == '':
                resolved_vi = '[MISSING VI RESOLUTION]'
                
            md_lines.append(f"- **{rank_vi}** (`{cid}_Rank_{star}`) [{eff_type}] | **Status**: `{status}`")
            md_lines.append(f"  - **CN Source**: {sum_cn if sum_cn else base_id}")
            md_lines.append(f"  - **VI Player-Facing**: {resolved_vi}")
            md_lines.append("")
            
            char_zh_json.append({
                'star': star,
                'effect_type': eff_type,
                'summary_cn': sum_cn,
                'resolved_vi': resolved_vi,
                'status': status
            })

        md_lines.append("### 4. HUANZHANG")
        md_lines.append("")
        
        char_hz_json = []
        if c_hz_rows:
            for hz in c_hz_rows:
                hid = hz.get('brilliant_id')
                h_name_cn = hz.get('icon_name_cn') or ''
                h_name_vi = hz.get('icon_name_vi') or h_name_cn
                info_cn = hz.get('icon_info_cn') or ''
                info_vi = hz.get('icon_info_vi') or ''
                show_cn = hz.get('buff_show_cn') or ''
                show_vi = hz.get('buff_show_vi') or ''
                status = hz.get('status') or 'UNKNOWN'
                conf = hz.get('confidence') or 'UNKNOWN'
                
                if not info_vi or info_vi == 'None' or info_vi.strip() == '':
                    info_vi = '[MISSING VI INFO]'
                    
                md_lines.append(f"#### [{hid}] {h_name_vi} ({h_name_cn})")
                md_lines.append(f"- **Status**: `{status}` | **Confidence**: `{conf}`")
                md_lines.append(f"- **Info CN**: {info_cn}")
                md_lines.append(f"- **Info VI**: {info_vi}")
                if show_cn and show_cn != 'None':
                    md_lines.append(f"- **Story CN**: {show_cn}")
                    md_lines.append(f"- **Story VI**: {show_vi if show_vi and show_vi != 'None' else '[MISSING VI STORY]'}")
                md_lines.append("")
                
                char_hz_json.append({
                    'brilliant_id': hid,
                    'name_cn': h_name_cn,
                    'name_vi': h_name_vi,
                    'info_cn': info_cn,
                    'info_vi': info_vi,
                    'story_cn': show_cn,
                    'story_vi': show_vi,
                    'status': status,
                    'confidence': conf
                })
        else:
            md_lines.append("*N/A (No Huanzhang data for this character)*")
            md_lines.append("")
            
        md_lines.append("---")
        md_lines.append("")
        
        review_json['characters'].append({
            'character_id': cid,
            'name_cn': c_info.get('cn', ''),
            'name_vi': c_info.get('vi', ''),
            'skills': char_sk_json,
            'buffs': char_buff_json,
            'zhizhi': char_zh_json,
            'huanzhang': char_hz_json
        })
        
    out_dir = os.path.join(base_dir, 'localization', 'batch_review')
    os.makedirs(out_dir, exist_ok=True)
    
    md_out_path = os.path.join(out_dir, f"{batch_id}_human_review.md")
    json_out_path = os.path.join(out_dir, f"{batch_id}_human_review.json")
    
    with open(md_out_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(md_lines))
        
    with open(json_out_path, 'w', encoding='utf-8') as f:
        json.dump(review_json, f, ensure_ascii=False, indent=2)
        
    print(f"Generated review package for batch '{batch_id}':")
    print(f"  Markdown: {md_out_path}")
    print(f"  JSON:     {json_out_path}")

if __name__ == '__main__':
    manifest_arg = sys.argv[1] if len(sys.argv) > 1 else r"localization/batches/phase3_batch1.json"
    base_dir = r"d:\BaiTapCode\WHMX\WhmxCalc"
    generate_review(base_dir, manifest_arg)
