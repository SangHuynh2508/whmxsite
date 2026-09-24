"""Apply one complete, source-verified SKILL translation transaction."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "localization" / "localization_master.xlsx"


def s(value): return "" if value is None else str(value)
def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def token_signature(value): return re.findall(r"<[^>]+>|\{[^{}]+\}|\[(?:Effect|Condition)\w*,\d+\]", value)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--transaction", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    locks = list(MASTER.parent.glob(f"~${MASTER.name}"))
    if locks: raise RuntimeError(f"Excel lock: {locks}")
    tx = json.loads(args.transaction.read_text(encoding="utf-8"))
    wb = load_workbook(MASTER, data_only=False)
    ws = wb["SKILL"]
    h = {s(cell.value): cell.column for cell in ws[1] if cell.value is not None}
    rows = {s(ws.cell(row, h["skill_id"]).value): row for row in range(2, ws.max_row + 1)}
    buff_ws = wb["BUFF_STATUS"]
    bh = {s(cell.value): cell.column for cell in buff_ws[1] if cell.value is not None}
    buff_rows = {s(buff_ws.cell(row, bh["buff_id"]).value): row for row in range(2, buff_ws.max_row + 1)}
    raw_skills = None
    plan, failures, already_applied = [], [], []
    for target in tx["targets"]:
        sid = target["skill_id"]
        if sid not in rows:
            failures.append({"skill_id": sid, "reason": "MISSING_TARGET"}); continue
        row = rows[sid]
        actual = {field: s(ws.cell(row, h[field]).value) for field in ("skill_name_cn", "desc_cn", "status")}
        if actual != target["expected"]:
            translated_now = all(s(ws.cell(row, h[field]).value) == value for field, value in target["translation"].items())
            if actual["skill_name_cn"] == target["expected"]["skill_name_cn"] and actual["desc_cn"] == target["expected"]["desc_cn"] and translated_now:
                already_applied.append(sid); continue
            failures.append({"skill_id": sid, "reason": "SOURCE_OR_STATUS_CHANGED", "actual": actual}); continue
        if token_signature(target["expected"]["desc_cn"]) != token_signature(target["translation"]["desc_vi"]):
            failures.append({"skill_id": sid, "reason": "MARKER_OR_PLACEHOLDER_CHANGED"}); continue
        missing_context = []
        for buff_id in target.get("required_buff_context", []):
            buff_row = buff_rows.get(buff_id)
            if not buff_row or not s(buff_ws.cell(buff_row, bh["buff_name_vi"]).value) or not s(buff_ws.cell(buff_row, bh["buff_desc_vi"]).value):
                missing_context.append(buff_id)
        if missing_context:
            failures.append({"skill_id": sid, "reason": "BLOCKED_CONTEXT", "buff_ids": missing_context}); continue
        raw_display = target.get("raw_display")
        if raw_display:
            if raw_skills is None:
                raw_path = ROOT.parent / "NeoArtifacts" / "MasterData" / "json" / "skillMap.json"
                raw_skills = json.loads(raw_path.read_text(encoding="utf-8"))
            exact = [record for record in raw_skills.values() if s(record.get("GroupId")) == raw_display["group_id"]]
            if not exact or any(s(record.get("NameLanText")) != raw_display["name_cn"] or s(record.get("DescriptionLanText")) != raw_display["desc_cn"] for record in exact):
                failures.append({"skill_id": sid, "reason": "RAW_DISPLAY_SOURCE_CHANGED", "raw_display": raw_display}); continue
        plan.append((row, target))
    report = {"transaction_id": tx["transaction_id"], "planned": len(plan), "already_applied": already_applied, "failures": failures, "applied": []}
    if args.apply:
        if failures: raise RuntimeError("Refusing partial transaction")
        backup = ROOT / "localization" / "backups" / f"localization_master_pre_{tx['transaction_id'].lower()}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        backup.parent.mkdir(parents=True, exist_ok=True)
        before_sha = sha(MASTER); shutil.copy2(MASTER, backup)
        if sha(backup) != before_sha: raise RuntimeError("backup hash mismatch")
        for row, target in plan:
            changes = {}
            for field, value in target["translation"].items():
                old = s(ws.cell(row, h[field]).value); ws.cell(row, h[field]).value = value; changes[field] = {"old": old, "new": value}
            note = s(ws.cell(row, h["notes"]).value).rstrip("; ")
            provenance = target.get("provenance") or f"Imported {tx['transaction_id']}; requires review"
            ws.cell(row, h["notes"]).value = (note + "; " + provenance).strip("; ")
            report["applied"].append({"skill_id": target["skill_id"], "changes": changes})
        wb.save(MASTER); report.update({"backup": str(backup), "backup_sha256": before_sha})
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"report": str(args.report), "planned": len(plan), "already_applied": len(already_applied), "applied": len(report["applied"]), "failures": len(failures)}, ensure_ascii=False))


if __name__ == "__main__": main()
