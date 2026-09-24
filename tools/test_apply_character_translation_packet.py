"""Tests for apply_character_translation_packet generic importer."""
from __future__ import annotations

import gc
import shutil
import tempfile
import unittest
from pathlib import Path
import sys

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))

from apply_character_translation_packet import (
    apply_character_packet,
    plan_packet_apply,
    require_packet_shape,
    validate_source_guards,
)
from safe_workbook_mutation import safe_mutate_workbook


def future_packet(skill_count=2):
    return {
        "BATCH_SUMMARY": [
            {
                "exact_skill_rows": skill_count,
                "unique_buff_count": 1,
                "locked_buff_count": 0,
                "reviewed_reference_buff_count": 0,
                "buff_rows_reviewed": 1,
                "huanzhang_player_rows_translated": 1,
            }
        ],
        "TRANSLATION_QA": [{"check": "sample", "result": "PASS", "details": ""}],
        "SKILLS": [
            {
                "character_id": "A0001",
                "skill_id": f"S{i}",
                "skill_name_cn": "名",
                "desc_cn": "述",
                "proposed_vi_name": "Tên",
                "proposed_vi_desc": "Mô tả",
                "proposal_action": "REVISED",
            }
            for i in range(skill_count)
        ],
        "BUFFS": [
            {
                "buff_id": "Buff_X",
                "buff_name_cn": "增益",
                "buff_desc_cn": "说明",
                "proposed_vi_name": "Tăng Ích",
                "proposed_vi_desc": "Mô tả",
                "proposal_action": "REVISED",
            }
        ],
        "HUANZHANG": [
            {
                "character_id": "A0001",
                "huanzhang_id": "HZ1",
                "title_name_cn": "徽章",
                "full_cn_lore_story": "故事",
                "buff_show_cn": "效果",
                "proposed_vi_title": "Huy Chương",
                "proposed_vi_lore": "Truyện",
                "proposal_action": "REVISED",
            }
        ],
    }


def make_test_master(master_path: Path):
    master_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    
    ws_char = wb.active
    ws_char.title = "CHARACTER"
    ws_char.append(["character_id", "name_cn", "fullname_cn", "name_vi", "fullname_vi"])
    ws_char.append(["A0100", "愿望杯", "愿望杯", "Lotus Chalice", "Lotus Chalice"])
    ws_char.append(["A0001", "人物", "人物", "Nhân Vật", "Nhân Vật"])

    ws_skill = wb.create_sheet("SKILL")
    ws_skill.append(["skill_id", "character_id", "skill_name_cn", "desc_cn", "skill_name_vi", "desc_vi", "translation_review_status", "translation_review_source"])
    ws_skill.append(["S1", "A0001", "名1", "述1", "Tên cũ", "Mô tả cũ", "", ""])

    ws_buff = wb.create_sheet("BUFF_STATUS")
    ws_buff.append(["buff_id", "buff_name_cn", "buff_desc_cn", "buff_name_vi", "buff_desc_vi", "name_authority", "desc_translation_status", "desc_translation_source"])
    ws_buff.append(["Buff_1", "增益1", "说明1", "Tăng Ích", "Mô tả", "OWNER_APPROVED", "", ""])

    ws_hz = wb.create_sheet("HUANZHANG")
    ws_hz.append(["brilliant_id", "character_id", "icon_name_cn", "icon_name_vi", "icon_info_cn", "icon_info_vi", "buff_show_cn", "buff_show_vi", "translation_review_status", "translation_review_source"])
    ws_hz.append(["HZ1", "A0001", "徽章1", "Tên HZ cũ", "故事1", "Truyện HZ cũ", "效果1", "Hiệu ứng cũ", "", ""])
    ws_hz.append(["HZ_EMPTY", "A0001", None, None, None, None, None, None, "", ""])
    ws_hz.append(["HZ_OTHER", "A0001", "徽章Khác", "HZ Khác", "故事Khác", "Truyện Khác", "", "", "", ""])

    wb.save(master_path)
    wb.close()


def make_test_packet(packet_path: Path, hz_records=None):
    packet_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()

    ws_sum = wb.active
    ws_sum.title = "BATCH_SUMMARY"
    ws_sum.append(["exact_skill_rows", "unique_buff_count", "locked_buff_count", "reviewed_reference_buff_count", "review_existing_buff_count", "huanzhang_character_count"])
    ws_sum.append([1, 1, 1, 0, 0, 1])

    ws_qa = wb.create_sheet("TRANSLATION_QA")
    ws_qa.append(["check", "result"])
    ws_qa.append(["sample", "PASS"])

    ws_skills = wb.create_sheet("SKILLS")
    ws_skills.append(["character_id", "skill_id", "skill_name_cn", "desc_cn", "skill_name_vi_current", "desc_vi_current", "proposed_vi_name", "proposed_vi_desc", "translation_action"])
    ws_skills.append(["A0001", "S1", "名1", "述1", "Tên cũ", "Mô tả cũ", "Tên mới", "Mô tả mới", "REVISED"])

    ws_buffs = wb.create_sheet("BUFFS")
    ws_buffs.append(["buff_id", "buff_name_cn", "buff_desc_cn", "translation_action", "proposed_vi_name", "proposed_vi_desc"])
    ws_buffs.append(["Buff_1", "增益1", "说明1", "LOCKED_REFERENCE", "", ""])

    ws_hz = wb.create_sheet("HUANZHANG")
    ws_hz.append(["character_id", "huanzhang_id", "title_name_cn", "title_name_vi_current", "full_cn_lore_story", "current_vi_lore_story", "buff_show_cn", "buff_show_vi_current", "proposed_vi_title", "proposed_vi_lore"])
    if hz_records is None:
        hz_records = [
            ["A0001", "HZ1", "徽章1", "Tên HZ cũ", "故事1", "Truyện HZ cũ", "效果1", "Hiệu ứng cũ", "Tên HZ mới", "Truyện HZ mới"],
            ["A0001", "HZ_EMPTY", None, None, None, None, None, None, None, None],
        ]
    for rec in hz_records:
        ws_hz.append(rec)

    wb.save(packet_path)
    wb.close()


class ApplyCharacterPacketShapeTests(unittest.TestCase):
    def test_accepts_arbitrary_future_row_counts(self):
        require_packet_shape(future_packet(skill_count=3))

    def test_rejects_duplicate_exact_ids(self):
        packet = future_packet(skill_count=2)
        packet["SKILLS"][1]["skill_id"] = packet["SKILLS"][0]["skill_id"]
        with self.assertRaises(SystemExit) as raised:
            require_packet_shape(packet)
        self.assertIn("DUPLICATE_SKILL_ID", str(raised.exception))


class ApplyCharacterPacketGenericHuanzhangTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.root = Path(self.temp_dir)
        self.master_path = self.root / "localization/localization_master.xlsx"
        self.packet_path = self.root / "packet.xlsx"
        make_test_master(self.master_path)
        make_test_packet(self.packet_path)

    def tearDown(self):
        gc.collect()
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_normal_translated_hz_row_apply(self):
        result = apply_character_packet(self.root, self.packet_path)
        self.assertEqual(result["huanzhang_applied"], 1)

        wb = openpyxl.load_workbook(self.master_path, read_only=True)
        hz = wb["HUANZHANG"]
        rows = [list(r) for r in hz.iter_rows(values_only=True)]
        wb.close()

        # Find HZ1 row
        hz1 = next(r for r in rows if r[0] == "HZ1")
        # ['brilliant_id', 'character_id', 'icon_name_cn', 'icon_name_vi', 'icon_info_cn', 'icon_info_vi', 'buff_show_cn', 'buff_show_vi', 'translation_review_status', 'translation_review_source']
        self.assertEqual(hz1[3], "Tên HZ mới")
        self.assertEqual(hz1[5], "Truyện HZ mới")
        self.assertEqual(hz1[8], "CHATGPT_REVIEWED")
        self.assertEqual(hz1[9], self.packet_path.name)
        # Source CN unchanged
        self.assertEqual(hz1[2], "徽章1")
        self.assertEqual(hz1[4], "故事1")

    def test_current_value_guard(self):
        # Mismatch current VI in packet
        bad_packet = self.root / "bad_current_packet.xlsx"
        make_test_packet(
            bad_packet,
            hz_records=[
                ["A0001", "HZ1", "徽章1", "Sai Lệch", "故事1", "Truyện HZ cũ", "效果1", "Hiệu ứng cũ", "Tên mới", "Truyện mới"]
            ]
        )
        with self.assertRaises(SystemExit) as raised:
            apply_character_packet(self.root, bad_packet)
        self.assertIn("SOURCE_CHANGED_HUANZHANG:HZ1:icon_name_vi", str(raised.exception))

    def test_wrong_huanzhang_id_rejection(self):
        bad_id_packet = self.root / "bad_id_packet.xlsx"
        make_test_packet(
            bad_id_packet,
            hz_records=[
                ["A0001", "HZ_UNKNOWN", "徽章", "Cũ", "故事", "Cũ", "", "", "Mới", "Mới"]
            ]
        )
        with self.assertRaises(SystemExit) as raised:
            apply_character_packet(self.root, bad_id_packet)
        self.assertIn("SOURCE_CHANGED_HUANZHANG:HZ_UNKNOWN:missing_id", str(raised.exception))

        # Test wrong character_id
        bad_char_packet = self.root / "bad_char_packet.xlsx"
        make_test_packet(
            bad_char_packet,
            hz_records=[
                ["A9999", "HZ1", "徽章1", "Tên HZ cũ", "故事1", "Truyện HZ cũ", "效果1", "Hiệu ứng cũ", "Tên mới", "Truyện mới"]
            ]
        )
        with self.assertRaises(SystemExit) as raised:
            apply_character_packet(self.root, bad_char_packet)
        self.assertIn("SOURCE_CHANGED_HUANZHANG:HZ1:character_id", str(raised.exception))

    def test_blank_placeholder_no_op(self):
        apply_character_packet(self.root, self.packet_path)
        wb = openpyxl.load_workbook(self.master_path, read_only=True)
        hz = wb["HUANZHANG"]
        rows = [list(r) for r in hz.iter_rows(values_only=True)]
        wb.close()

        empty_row = next(r for r in rows if r[0] == "HZ_EMPTY")
        # Empty row must remain None/empty
        self.assertIsNone(empty_row[2])
        self.assertIsNone(empty_row[3])
        self.assertIsNone(empty_row[4])
        self.assertIsNone(empty_row[5])
        self.assertEqual(empty_row[8] or "", "")

    def test_unrelated_hz_row_immutability(self):
        # Mutator attempting to alter HZ_OTHER without authorization will fail in safe_mutate_workbook
        def illegal_mutate(wb):
            sheet = wb["HUANZHANG"]
            sheet["D4"] = "Tấn Công Trái Phép"  # HZ_OTHER row

        with self.assertRaises(PermissionError):
            safe_mutate_workbook(
                str(self.root),
                illegal_mutate,
                {"huanzhang_ids": {"HZ1"}},
                authorized_new_columns={},
            )

    def test_rollback_safe_write_behavior(self):
        def failing_mutate(wb):
            sheet = wb["HUANZHANG"]
            sheet["D2"] = "Biến đổi dở dang"
            raise RuntimeError("CRASH_DURING_MUTATION")

        with self.assertRaises(RuntimeError):
            safe_mutate_workbook(
                str(self.root),
                failing_mutate,
                {"huanzhang_ids": {"HZ1"}},
                authorized_new_columns={},
            )

        # Ensure master is intact
        wb = openpyxl.load_workbook(self.master_path, read_only=True)
        hz = wb["HUANZHANG"]
        hz1 = next(r for r in hz.iter_rows(values_only=True) if r[0] == "HZ1")
        self.assertEqual(hz1[3], "Tên HZ cũ")
        wb.close()

    def test_character_metadata_update(self):
        spec = {
            "A0001": {
                "expected": {"character_id": "A0001", "name_vi": "Nhân Vật"},
                "updates": {"name_vi": "Nhân Vật Mới"},
            }
        }
        res = apply_character_packet(self.root, self.packet_path, character_updates=spec)
        self.assertEqual(res["character_cells_applied"], 1)

        wb = openpyxl.load_workbook(self.master_path, read_only=True)
        row = next(r for r in wb["CHARACTER"].iter_rows(values_only=True) if r[0] == "A0001")
        self.assertEqual(row[3], "Nhân Vật Mới")
        wb.close()


if __name__ == "__main__":
    unittest.main()
