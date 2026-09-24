"""Standalone migration script to safely consolidate and deduplicate the 45 duplicate Huanzhang skill rows in localization_master.xlsx.

DO NOT RUN AS PART OF TRANSLATION IMPORT.
This script is prepared as an independent migration tool for future execution
when the owner decides to physically deduplicate the workbook rows.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
import hashlib
from pathlib import Path
import sys

import openpyxl

TOOLS_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = TOOLS_DIR.parent
if str(TOOLS_DIR) not in sys.path:
    sys.path.insert(0, str(TOOLS_DIR))

from safe_workbook_mutation import AtomicReplaceLockError, safe_mutate_workbook
from workbook_semantic_fingerprint import workbook_semantic_fingerprint

MASTER_PATH = PROJECT_ROOT / "localization" / "localization_master.xlsx"


def compute_sha256(path: Path | str) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest().upper()


def audit_duplicates(master_path: Path):
    wb = openpyxl.load_workbook(master_path, data_only=True)
    ws = wb["SKILL"]
    headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()

    by_id = defaultdict(list)
    for idx, r in enumerate(rows, start=2):
        r["_excel_row"] = idx
        by_id[r["skill_id"]].append(r)

    duplicates = {k: v for k, v in by_id.items() if len(v) > 1}
    return duplicates, headers


def plan_consolidation(duplicates: dict):
    plans = []
    for skill_id, pair in duplicates.items():
        if len(pair) != 2:
            raise RuntimeError(f"Expected exactly 2 rows for {skill_id}, found {len(pair)}")
        r1, r2 = pair[0], pair[1]

        # Verify parity before permitting consolidation
        if r1.get("character_id") != r2.get("character_id"):
            raise RuntimeError(f"Character mismatch for {skill_id}: {r1.get('character_id')} vs {r2.get('character_id')}")
        if r1.get("skill_name_cn") != r2.get("skill_name_cn"):
            raise RuntimeError(f"CN name mismatch for {skill_id}")
        if r1.get("desc_cn") != r2.get("desc_cn"):
            raise RuntimeError(f"CN desc mismatch for {skill_id}")
        if r1.get("skill_name_vi") != r2.get("skill_name_vi"):
            raise RuntimeError(f"VI name mismatch for {skill_id}")
        if r1.get("desc_vi") != r2.get("desc_vi"):
            raise RuntimeError(f"VI desc mismatch for {skill_id}")

        plans.append({
            "skill_id": skill_id,
            "keep_row": r1["_excel_row"],
            "remove_row": r2["_excel_row"],
            "character_id": r1["character_id"],
        })
    return plans


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dry-run", action="store_true", help="Audit and display planned consolidation without modifying workbook")
    parser.add_argument("--apply", action="store_true", help="Execute consolidation transaction")
    args = parser.parse_args()

    duplicates, headers = audit_duplicates(MASTER_PATH)
    print(f"Total duplicate skill IDs found: {len(duplicates)}")

    plans = plan_consolidation(duplicates)
    print(f"Consolidation planned for {len(plans)} pairs:")
    for p in plans:
        print(f"  {p['skill_id']} ({p['character_id']}): keep row {p['keep_row']}, remove row {p['remove_row']}")

    if args.dry_run or not args.apply:
        print("\nDry run completed. To execute consolidation, specify --apply.")
        return

    print("\nExecuting safe consolidation...")
    rows_to_delete = sorted([p["remove_row"] for p in plans], reverse=True)

    def mutator(wb: openpyxl.Workbook):
        ws = wb["SKILL"]
        for row_idx in rows_to_delete:
            ws.delete_rows(row_idx, 1)

    backup = safe_mutate_workbook(
        base_dir=str(PROJECT_ROOT),
        mutator_fn=mutator,
        authorized_deps={"skill_ids": {p["skill_id"] for p in plans}},
        authorized_new_columns={},
        authorized_cells=set(),
        authorized_deleted_rows={"SKILL": {p["skill_id"] for p in plans}},
    )
    print(f"Consolidation complete. Backup created: {backup}")
    print(f"Master SHA256 after consolidation: {compute_sha256(MASTER_PATH)}")
    print(f"Master semantic fingerprint after consolidation: {workbook_semantic_fingerprint(MASTER_PATH)}")


if __name__ == "__main__":
    main()
