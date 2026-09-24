import openpyxl, json, re, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('localization/localization_master.xlsx')
batch1_chars = {'A0121', 'A0086', 'A0160', 'V0146', 'V0117'}

# Master translation dictionary for exact skill descriptions
skill_updates = {
    # A0121
    'A012101_1': 'Gây Sát Thương Vật Lý bằng 100% Tấn Công lên 1 kẻ địch được chọn.',
    'A012101_2': 'Gây Sát Thương Vật Lý bằng 115% Tấn Công lên 1 kẻ địch được chọn.',
    'A012101_3': 'Gây Sát Thương Vật Lý bằng 130% Tấn Công lên 1 kẻ địch được chọn.',
    'A012101_4': 'Gây Sát Thương Vật Lý bằng 145% Tấn Công lên 1 kẻ địch được chọn.',
    'A012101_5': 'Gây Sát Thương Vật Lý bằng 160% Tấn Công lên 1 kẻ địch được chọn.',
    'A012102_1': 'Sau khi sử dụng, bản thân nhận trạng thái <color=#ff6724>Phong Địch Thanh Dã</color> kéo dài <color=#158bdb>1</color> vòng lượt.{Buff_A0121_1}',
    'A012102_2': 'Sau khi sử dụng, bản thân nhận trạng thái <color=#ff6724>Phong Địch Thanh Dã</color> kéo dài <color=#158bdb>1</color> vòng lượt.{Buff_A0121_1}',
    'A012102_3': 'Sau khi sử dụng, bản thân nhận trạng thái <color=#ff6724>Phong Địch Thanh Dã</color> kéo dài <color=#158bdb>1</color> vòng lượt.{Buff_A0121_1}',
    'A012103_1': 'Khi Đánh Thường trong trạng thái <color=#ff6724>Nhắm Bắn</color>, thực hiện thêm <color=#158bdb>3</color> lần Đánh Thường truy kích lên kẻ địch được chọn, mỗi lần gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công; ở trạng thái này không thể di chuyển và không thể dùng Kỹ Năng Nghề.{Buff_A0121_3}',
    'A012103_2': 'Khi Đánh Thường trong trạng thái <color=#ff6724>Nhắm Bắn</color>, thực hiện thêm <color=#158bdb>3</color> lần Đánh Thường truy kích lên kẻ địch được chọn, mỗi lần gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công; ở trạng thái này không thể di chuyển và không thể dùng Kỹ Năng Nghề.{Buff_A0121_3}',
    'A012103_3': 'Khi Đánh Thường trong trạng thái <color=#ff6724>Nhắm Bắn</color>, thực hiện thêm <color=#158bdb>3</color> lần Đánh Thường truy kích lên kẻ địch được chọn, mỗi lần gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công; ở trạng thái này không thể di chuyển và không thể dùng Kỹ Năng Nghề.{Buff_A0121_3}',
    'A012104_1': 'Khi đồng minh tấn công kẻ địch nằm trong phạm vi Đánh Thường của bản thân, bản thân Liên Kích <color=#158bdb>1</color> lần lên kẻ địch đó, gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công; mỗi vòng lượt kích hoạt tối đa <color=#158bdb>[Effect1Para,2]</color> lần.',
    'A012104_2': 'Khi đồng minh tấn công kẻ địch nằm trong phạm vi Đánh Thường của bản thân, bản thân Liên Kích <color=#158bdb>1</color> lần lên kẻ địch đó, gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công; mỗi vòng lượt kích hoạt tối đa <color=#158bdb>[Effect1Para,2]</color> lần.',
    'A012104_3': 'Khi đồng minh tấn công kẻ địch nằm trong phạm vi Đánh Thường của bản thân, bản thân Liên Kích <color=#158bdb>1</color> lần lên kẻ địch đó, gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công; mỗi vòng lượt kích hoạt tối đa <color=#158bdb>[Effect1Para,2]</color> lần.',
    'A012105_1': 'Sau khi hạ gục kẻ địch, bản thân lập tức nhận <color=#158bdb>1</color> lượt hành động lại (lượt này không thể di chuyển).',
    'A012105_2': 'Sau khi hạ gục kẻ địch, bản thân lập tức nhận <color=#158bdb>1</color> lượt hành động lại (lượt này không thể di chuyển).',
    'A012105_3': 'Sau khi hạ gục kẻ địch, bản thân lập tức nhận <color=#158bdb>1</color> lượt hành động lại (lượt này không thể di chuyển).',
    'A012111_1': 'Sau khi sử dụng, bản thân vào trạng thái <color=#ff6724>Nhắm Bắn</color>: tầm đánh tối đa của Đánh Thường tăng <color=#158bdb>2</color> ô và Đánh Thường không thể bị đỡ đòn; ở trạng thái <color=#ff6724>Nhắm Bắn</color> không thể di chuyển.{Buff_A0121_11}',

    # A0086
    'A008601_1': 'Gây Sát Thương Vật Lý bằng 100% Tấn Công lên 1 kẻ địch được chọn.',
    'A008601_2': 'Gây Sát Thương Vật Lý bằng 115% Tấn Công lên 1 kẻ địch được chọn.',
    'A008601_3': 'Gây Sát Thương Vật Lý bằng 130% Tấn Công lên 1 kẻ địch được chọn.',
    'A008601_4': 'Gây Sát Thương Vật Lý bằng 145% Tấn Công lên 1 kẻ địch được chọn.',
    'A008601_5': 'Gây Sát Thương Vật Lý bằng 160% Tấn Công lên 1 kẻ địch được chọn.',
    'A008602_1': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn, đồng thời khiến sát thương mục tiêu gánh chịu tăng <color=#158bdb>[Effect2Para,1]%</color>, kéo dài <color=#158bdb>2</color> lượt.',
    'A008602_2': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn, đồng thời khiến sát thương mục tiêu gánh chịu tăng <color=#158bdb>[Effect2Para,1]%</color>, kéo dài <color=#158bdb>2</color> lượt.',
    'A008602_3': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn, đồng thời khiến sát thương mục tiêu gánh chịu tăng <color=#158bdb>[Effect2Para,1]%</color>, kéo dài <color=#158bdb>2</color> lượt.',
    'A008603_1': 'Khi bị kẻ địch cận chiến tấn công, nếu Máu của bản thân cao hơn <color=#158bdb>[Effect1Para,1]%</color>, né tránh cuộc tấn công đó và phản kích, gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect2Para,1]%</color> Tấn Công; mỗi lượt kích hoạt tối đa <color=#158bdb>1</color> lần.',
    'A008603_2': 'Khi bị kẻ địch cận chiến tấn công, nếu Máu của bản thân cao hơn <color=#158bdb>[Effect1Para,1]%</color>, né tránh cuộc tấn công đó và phản kích, gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect2Para,1]%</color> Tấn Công; mỗi lượt kích hoạt tối đa <color=#158bdb>1</color> lần.',
    'A008603_3': 'Khi bị kẻ địch cận chiến tấn công, nếu Máu của bản thân cao hơn <color=#158bdb>[Effect1Para,1]%</color>, né tránh cuộc tấn công đó và phản kích, gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect2Para,1]%</color> Tấn Công; mỗi lượt kích hoạt tối đa <color=#158bdb>1</color> lần.',
    'A008604_1': 'Sau khi né tránh thành công, bản thân nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Kim Ngọc</color> kéo dài <color=#158bdb>1</color> lượt, có thể cộng dồn tối đa <color=#158bdb>3</color> tầng.{Buff_A0086_19}',
    'A008604_2': 'Sau khi né tránh thành công, bản thân nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Kim Ngọc</color> kéo dài <color=#158bdb>1</color> lượt, có thể cộng dồn tối đa <color=#158bdb>3</color> tầng.{Buff_A0086_19}',
    'A008604_3': 'Sau khi né tránh thành công, bản thân nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Kim Ngọc</color> kéo dài <color=#158bdb>1</color> lượt, có thể cộng dồn tối đa <color=#158bdb>3</color> tầng.{Buff_A0086_19}',
    'A008605_1': 'Khi chịu Sát Thương Chí Tử, nhận hiệu ứng <color=#ff6724>Kháng Tử Vong</color> trong lượt đó; đồng thời Tỷ Lệ Hút Máu tăng thêm <color=#158bdb>[Effect3Para,1]%</color>, kéo dài <color=#158bdb>1</color> lượt, đồng thời nhận hiệu ứng <color=#ff6724>Ẩn Nấp</color>; trong 1 trận đấu chỉ có thể kích hoạt 1 lần.{Buff_A0086_24}',
    'A008605_2': 'Khi chịu Sát Thương Chí Tử, nhận hiệu ứng <color=#ff6724>Kháng Tử Vong</color> trong lượt đó; đồng thời Tỷ Lệ Hút Máu tăng thêm <color=#158bdb>[Effect3Para,1]%</color>, kéo dài <color=#158bdb>1</color> lượt, đồng thời nhận hiệu ứng <color=#ff6724>Ẩn Nấp</color>; trong 1 trận đấu chỉ có thể kích hoạt 1 lần.{Buff_A0086_24}',
    'A008605_3': 'Khi chịu Sát Thương Chí Tử, nhận hiệu ứng <color=#ff6724>Kháng Tử Vong</color> trong lượt đó; đồng thời Tỷ Lệ Hút Máu tăng thêm <color=#158bdb>[Effect3Para,1]%</color>, kéo dài <color=#158bdb>1</color> lượt, đồng thời nhận hiệu ứng <color=#ff6724>Ẩn Nấp</color>; trong 1 trận đấu chỉ có thể kích hoạt 1 lần.{Buff_A0086_24}',
    'A008611_1': 'Sau khi sử dụng, bản thân nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Thoát Xác</color>.{Buff_A0086_15}',

    # A0160
    'A016001_1': 'Gây Sát Thương Cấu Thuật bằng 100% Tấn Công lên 1 kẻ địch được chọn.',
    'A016001_2': 'Gây Sát Thương Cấu Thuật bằng 115% Tấn Công lên 1 kẻ địch được chọn.',
    'A016001_3': 'Gây Sát Thương Cấu Thuật bằng 130% Tấn Công lên 1 kẻ địch được chọn.',
    'A016001_4': 'Gây Sát Thương Cấu Thuật bằng 145% Tấn Công lên 1 kẻ địch được chọn.',
    'A016001_5': 'Gây Sát Thương Cấu Thuật bằng 160% Tấn Công lên 1 kẻ địch được chọn.',
    'A016002_1': 'Gây Sát Thương Cấu Thuật bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn, đồng thời gắn <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Phiếm Âm</color> lên mục tiêu kéo dài <color=#158bdb>2</color> lượt.{Buff_A0160_2}',
    'A016002_2': 'Gây Sát Thương Cấu Thuật bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn, đồng thời gắn <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Phiếm Âm</color> lên mục tiêu kéo dài <color=#158bdb>2</color> lượt.{Buff_A0160_2}',
    'A016002_3': 'Gây Sát Thương Cấu Thuật bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn, đồng thời gắn <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Phiếm Âm</color> lên mục tiêu kéo dài <color=#158bdb>2</color> lượt.{Buff_A0160_2}',
    'A016003_1': 'Gây Sát Thương Cấu Thuật bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn cùng toàn bộ kẻ địch <color=#158bdb>xung quanh 1 vòng</color>.',
    'A016003_2': 'Gây Sát Thương Cấu Thuật bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn cùng toàn bộ kẻ địch <color=#158bdb>xung quanh 1 vòng</color>.',
    'A016003_3': 'Gây Sát Thương Cấu Thuật bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn cùng toàn bộ kẻ địch <color=#158bdb>xung quanh 1 vòng</color>.',
    'A016004_1': 'Sau khi sử dụng, bản thân nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Tản Âm</color> kéo dài <color=#158bdb>2</color> lượt.{Buff_A0160_10}',
    'A016004_2': 'Sau khi sử dụng, bản thân nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Tản Âm</color> kéo dài <color=#158bdb>2</color> lượt.{Buff_A0160_10}',
    'A016004_3': 'Sau khi sử dụng, bản thân nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Tản Âm</color> kéo dài <color=#158bdb>2</color> lượt.{Buff_A0160_10}',
    'A016005_1': 'Khi tấn công kẻ địch mang trạng thái <color=#ff6724>Phiếm Âm</color>, sát thương gây ra lên kẻ địch đó tăng <color=#158bdb>[Effect1Para,1]%</color>.',
    'A016005_2': 'Khi tấn công kẻ địch mang trạng thái <color=#ff6724>Phiếm Âm</color>, sát thương gây ra lên kẻ địch đó tăng <color=#158bdb>[Effect1Para,1]%</color>.',
    'A016005_3': 'Khi tấn công kẻ địch mang trạng thái <color=#ff6724>Phiếm Âm</color>, sát thương gây ra lên kẻ địch đó tăng <color=#158bdb>[Effect1Para,1]%</color>.',
    'A016011_1': 'Gây Sát Thương Cấu Thuật bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn, đồng thời xóa toàn bộ trạng thái có lợi trên người mục tiêu.',

    # V0117
    'V011701_1': 'Gây Sát Thương Vật Lý bằng 100% Tấn Công lên 1 kẻ địch được chọn.',
    'V011701_2': 'Gây Sát Thương Vật Lý bằng 115% Tấn Công lên 1 kẻ địch được chọn.',
    'V011701_3': 'Gây Sát Thương Vật Lý bằng 130% Tấn Công lên 1 kẻ địch được chọn.',
    'V011701_4': 'Gây Sát Thương Vật Lý bằng 145% Tấn Công lên 1 kẻ địch được chọn.',
    'V011701_5': 'Gây Sát Thương Vật Lý bằng 160% Tấn Công lên 1 kẻ địch được chọn.',
    'V011702_1': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn; nếu bản thân có trạng thái <color=#ff6724>Minh Kính</color>, gây thêm Sát Thương Vật Lý bằng <color=#158bdb>[Effect2Para,1]%</color> Tấn Công.',
    'V011702_2': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn; nếu bản thân có trạng thái <color=#ff6724>Minh Kính</color>, gây thêm Sát Thương Vật Lý bằng <color=#158bdb>[Effect2Para,1]%</color> Tấn Công.',
    'V011702_3': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn; nếu bản thân có trạng thái <color=#ff6724>Minh Kính</color>, gây thêm Sát Thương Vật Lý bằng <color=#158bdb>[Effect2Para,1]%</color> Tấn Công.',
    'V011703_1': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn cùng toàn bộ kẻ địch <color=#158bdb>xung quanh 1 ô</color>.',
    'V011703_2': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn cùng toàn bộ kẻ địch <color=#158bdb>xung quanh 1 ô</color>.',
    'V011703_3': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn cùng toàn bộ kẻ địch <color=#158bdb>xung quanh 1 ô</color>.',
    'V011704_1': 'Khi bị kẻ địch tấn công trực tiếp, nếu bản thân có ít nhất <color=#158bdb>2</color> tầng trạng thái <color=#ff6724>Minh Kính</color>, tiêu hao <color=#158bdb>2</color> tầng trạng thái <color=#ff6724>Minh Kính</color> để phản kích, gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công.',
    'V011704_2': 'Khi bị kẻ địch tấn công trực tiếp, nếu bản thân có ít nhất <color=#158bdb>2</color> tầng trạng thái <color=#ff6724>Minh Kính</color>, tiêu hao <color=#158bdb>2</color> tầng trạng thái <color=#ff6724>Minh Kính</color> để phản kích, gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công.',
    'V011704_3': 'Khi bị kẻ địch tấn công trực tiếp, nếu bản thân có ít nhất <color=#158bdb>2</color> tầng trạng thái <color=#ff6724>Minh Kính</color>, tiêu hao <color=#158bdb>2</color> tầng trạng thái <color=#ff6724>Minh Kính</color> để phản kích, gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công.',
    'V011705_1': 'Khi ở trạng thái <color=#ff6724>Minh Kính</color>, Phòng Thủ Vật Lý và Phòng Thủ Cấu Thuật của bản thân tăng <color=#158bdb>[Effect1Para,1]%</color>.',
    'V011705_2': 'Khi ở trạng thái <color=#ff6724>Minh Kính</color>, Phòng Thủ Vật Lý và Phòng Thủ Cấu Thuật của bản thân tăng <color=#158bdb>[Effect1Para,1]%</color>.',
    'V011705_3': 'Khi ở trạng thái <color=#ff6724>Minh Kính</color>, Phòng Thủ Vật Lý và Phòng Thủ Cấu Thuật của bản thân tăng <color=#158bdb>[Effect1Para,1]%</color>.',
    'V011711_1': 'Giúp bản thân nhận <color=#158bdb>3</color> tầng trạng thái <color=#ff6724>Minh Kính</color>.{Buff_V0117_3}',

    # V0146
    'V014601_1': 'Gây Sát Thương Vật Lý bằng 100% Tấn Công lên 1 kẻ địch được chọn.',
    'V014601_2': 'Gây Sát Thương Vật Lý bằng 115% Tấn Công lên 1 kẻ địch được chọn.',
    'V014601_3': 'Gây Sát Thương Vật Lý bằng 130% Tấn Công lên 1 kẻ địch được chọn.',
    'V014601_4': 'Gây Sát Thương Vật Lý bằng 145% Tấn Công lên 1 kẻ địch được chọn.',
    'V014601_5': 'Gây Sát Thương Vật Lý bằng 160% Tấn Công lên 1 kẻ địch được chọn.',
    'V014602_1': 'Sau khi sử dụng, bản thân nhận <color=#158bdb>[Effect4Para,3]</color> tầng <color=#ff6724>Khiếu Kiếm</color> và <color=#158bdb>[Effect4Para,4]</color> tầng trạng thái <color=#ff6724>Tích Thế</color>, đều kéo dài <color=#158bdb>2</color> lượt;\nĐồng thời nhận trạng thái <color=#ff6724>Hàm Ẩm</color> kéo dài <color=#158bdb>1</color> vòng lượt (trong trạng thái này không thể dùng Tuyệt Kỹ lần nữa).{Buff_V0146_1}{Buff_Atk_Up}{Buff_AllDmgIncrease}',
    'V014602_2': 'Sau khi sử dụng, bản thân nhận <color=#158bdb>[Effect4Para,3]</color> tầng <color=#ff6724>Khiếu Kiếm</color> và <color=#158bdb>[Effect4Para,4]</color> tầng trạng thái <color=#ff6724>Tích Thế</color>, đều kéo dài <color=#158bdb>2</color> lượt;\nĐồng thời nhận trạng thái <color=#ff6724>Hàm Ẩm</color> kéo dài <color=#158bdb>1</color> vòng lượt (trong trạng thái này không thể dùng Tuyệt Kỹ lần nữa).{Buff_V0146_1}{Buff_Atk_Up}{Buff_AllDmgIncrease}',
    'V014602_3': 'Sau khi sử dụng, bản thân nhận <color=#158bdb>[Effect4Para,3]</color> tầng <color=#ff6724>Khiếu Kiếm</color> và <color=#158bdb>[Effect4Para,4]</color> tầng trạng thái <color=#ff6724>Tích Thế</color>, đều kéo dài <color=#158bdb>2</color> lượt;\nĐồng thời nhận trạng thái <color=#ff6724>Hàm Ẩm</color> kéo dài <color=#158bdb>1</color> vòng lượt (trong trạng thái này không thể dùng Tuyệt Kỹ lần nữa).{Buff_V0146_1}{Buff_Atk_Up}{Buff_AllDmgIncrease}',
    'V014603_1': 'Khi vào trận, Tốc Độ tăng <color=#158bdb>200</color>.\nSau khi sử dụng Tuyệt Kỹ, bản thân nhận trạng thái <color=#ff6724>Chúng Sinh Nhất Tướng</color> kéo dài <color=#158bdb>2</color> vòng lượt; trong thời gian này, mỗi khi đồng minh Liên Kích <color=#158bdb>1</color> lần, Tấn Công của Á Trường Ngưu Tôn tăng <color=#158bdb>[Effect4Para,1]%</color>, có thể cộng dồn tối đa <color=#158bdb>10</color> tầng, kéo dài <color=#158bdb>1</color> vòng lượt.{Buff_V0146_4_1}',
    'V014603_2': 'Khi vào trận, Tốc Độ tăng <color=#158bdb>200</color>.\nSau khi sử dụng Tuyệt Kỹ, bản thân nhận trạng thái <color=#ff6724>Chúng Sinh Nhất Tướng</color> kéo dài <color=#158bdb>2</color> vòng lượt; trong thời gian này, mỗi khi đồng minh Liên Kích <color=#158bdb>1</color> lần, Tấn Công của Á Trường Ngưu Tôn tăng <color=#158bdb>[Effect4Para,1]%</color>, có thể cộng dồn tối đa <color=#158bdb>10</color> tầng, kéo dài <color=#158bdb>1</color> vòng lượt.{Buff_V0146_4_1}',
    'V014603_3': 'Khi vào trận, Tốc Độ tăng <color=#158bdb>200</color>.\nSau khi sử dụng Tuyệt Kỹ, bản thân nhận trạng thái <color=#ff6724>Chúng Sinh Nhất Tướng</color> kéo dài <color=#158bdb>2</color> vòng lượt; trong thời gian này, mỗi khi đồng minh Liên Kích <color=#158bdb>1</color> lần, Tấn Công của Á Trường Ngưu Tôn tăng <color=#158bdb>[Effect4Para,1]%</color>, có thể cộng dồn tối đa <color=#158bdb>10</color> tầng, kéo dài <color=#158bdb>1</color> vòng lượt.{Buff_V0146_4_1}',
    'V014604_1': 'Khi Liên Kích hạ gục kẻ địch, Gây <color=#ff6724>Sát Thương Chuẩn Bổ Sung</color> bằng <color=#158bdb>[Effect1Para,1]%</color> sát thương của lần Liên Kích đó lên toàn bộ kẻ địch <color=#158bdb>xung quanh 2 ô</color>.{Buff_RealDamge_Lan_Special}',
    'V014604_2': 'Khi Liên Kích hạ gục kẻ địch, Gây <color=#ff6724>Sát Thương Chuẩn Bổ Sung</color> bằng <color=#158bdb>[Effect1Para,1]%</color> sát thương của lần Liên Kích đó lên toàn bộ kẻ địch <color=#158bdb>xung quanh 2 ô</color>.{Buff_RealDamge_Lan_Special}',
    'V014604_3': 'Khi Liên Kích hạ gục kẻ địch, Gây <color=#ff6724>Sát Thương Chuẩn Bổ Sung</color> bằng <color=#158bdb>[Effect1Para,1]%</color> sát thương của lần Liên Kích đó lên toàn bộ kẻ địch <color=#158bdb>xung quanh 2 ô</color>.{Buff_RealDamge_Lan_Special}',
    'V014605_1': 'Bắt đầu mỗi vòng lượt, với mỗi <color=#158bdb>1</color> lần bản thân từng hành động trong <color=#158bdb>1</color> vòng lượt trước, Sát Thương Liên Kích tăng <color=#158bdb>[Effect1Para,1]%</color>, có thể cộng dồn tối đa <color=#158bdb>5</color> tầng, kéo dài <color=#158bdb>2</color> lượt.',
    'V014605_2': 'Bắt đầu mỗi vòng lượt, với mỗi <color=#158bdb>1</color> lần bản thân từng hành động trong <color=#158bdb>1</color> vòng lượt trước, Sát Thương Liên Kích tăng <color=#158bdb>[Effect1Para,1]%</color>, có thể cộng dồn tối đa <color=#158bdb>5</color> tầng, kéo dài <color=#158bdb>2</color> lượt.',
    'V014605_3': 'Bắt đầu mỗi vòng lượt, với mỗi <color=#158bdb>1</color> lần bản thân từng hành động trong <color=#158bdb>1</color> vòng lượt trước, Sát Thương Liên Kích tăng <color=#158bdb>[Effect1Para,1]%</color>, có thể cộng dồn tối đa <color=#158bdb>5</color> tầng, kéo dài <color=#158bdb>2</color> lượt.',
    'V014611_1': 'Sau khi sử dụng, bản thân nhận <color=#158bdb>1</color> lượt hành động lại (Sức Di Chuyển của lượt hành động lại này giảm <color=#158bdb>2</color>).'
}

# Buff translation updates
buff_updates = {
    'Buff_A0121_1': (
        'Khi ở trạng thái này, sát thương bản thân gánh chịu giảm <color=#158bdb>[EffectParam,2]%</color>, không thể di chuyển và miễn nhiễm hiệu ứng dịch chuyển do kẻ địch gây ra;\n'
        'Phạm vi tấn công của Đánh Thường chuyển thành <color=#158bdb>xung quanh 3-7 ô</color>, đồng thời gây Sát Thương Vật Lý bằng <color=#158bdb>100%</color> Tấn Công lên 1 kẻ địch hoặc ô đất được chọn cùng toàn bộ kẻ địch xung quanh <color=#158bdb>1 vòng</color>;\n'
        'Trong thời gian trạng thái này kéo dài, bản thân không thể vào trạng thái <color=#ff6724>Nhắm Bắn</color>.'
    ),
    'Buff_A0160_10': (
        'Khi chịu <color=#158bdb>1</color> lần sát thương từ Đánh Thường hoặc Kỹ Năng, Gây <color=#ff6724>Sát Thương Vật Lý Bổ Sung</color> bằng <color=#158bdb>[EffectParam,2]%</color> Tấn Công của Thái Phụng Minh Kỳ lên bản thân và toàn bộ đồng minh <color=#158bdb>xung quanh 3 ô</color>, sau khi kích hoạt hiệu ứng sẽ xóa trạng thái này.'
    ),
    'Buff_V0117_hz_1_1': (
        'Trạng thái này có thể cộng dồn tối đa <color=#158bdb>3</color> tầng;\n'
        'Khi chịu tấn công Tuyệt Kỹ hoặc Liên Kích của Lộc Vương Bản Sinh Đồ bên địch, mỗi tầng trạng thái này khiến bản thân chịu <color=#ff6724>Sát Thương Vật Lý Bổ Sung</color> bằng <color=#158bdb>220%</color> Tấn Công của Lộc Vương Bản Sinh Đồ (sát thương này chắc chắn Bạo Kích), sau khi chịu tấn công từ Lộc Vương Bản Sinh Đồ bên địch sẽ xóa toàn bộ trạng thái này.'
    ),
    'Buff_V0146_1': (
        'Khi ở trạng thái này, nếu bất kỳ đồng minh nào kích hoạt Liên Kích, bản thân nhận <color=#158bdb>1</color> lượt hành động bổ sung sau khi đồng minh đó kết thúc hành động và làm mới thời gian kéo dài của trạng thái <color=#ff6724>Khiếu Kiếm</color> và <color=#ff6724>Tích Thế</color>; Trong lượt hành động bổ sung đó, sát thương Đánh Thường giảm <color=#158bdb>99%</color>.'
    ),
    'Buff_V0146_4_1': (
        'Trong <color=#158bdb>2</color> vòng lượt, bản thân sau khi Đánh Thường sẽ kích hoạt Liên Kích, Gây Sát Thương Vật Lý bằng <color=#158bdb>[EffectParam,2]%</color> Tấn Công lên kẻ địch bị tấn công.'
    )
}

# 1. Apply SKILL updates
sk_sheet = wb['SKILL']
sk_headers = [cell.value for cell in sk_sheet[1]]
id_idx = sk_headers.index('skill_id')
desc_vi_idx = sk_headers.index('desc_vi')
status_idx = sk_headers.index('status')
notes_idx = sk_headers.index('notes')

for row in sk_sheet.iter_rows(min_row=2):
    sid = str(row[id_idx].value or '')
    if sid in skill_updates:
        row[desc_vi_idx].value = skill_updates[sid]
        row[status_idx].value = 'TRANSLATED'
        row[notes_idx].value = 'Skill name owner-approved; description fully naturalized & translated'

# 2. Apply BUFF updates
bf_sheet = wb['BUFF_STATUS']
bf_headers = [cell.value for cell in bf_sheet[1]]
bf_id_idx = bf_headers.index('buff_id')
bf_desc_vi_idx = bf_headers.index('buff_desc_vi')
bf_status_idx = bf_headers.index('status')

for row in bf_sheet.iter_rows(min_row=2):
    bid = str(row[bf_id_idx].value or '')
    if bid in buff_updates:
        row[bf_desc_vi_idx].value = buff_updates[bid]
        row[bf_status_idx].value = 'TRANSLATED'

# 3. Update ZHIZHI V0117 Rank 6 HurtHealRate resolution according to raw MasterData attrDesMap.json evidence (生命偷取率 -> Tỷ Lệ Hút Máu)
zh_sheet = wb['ZHIZHI']
zh_headers = [cell.value for cell in zh_sheet[1]]
zh_cid_idx = zh_headers.index('character_id')
zh_star_idx = zh_headers.index('star')
zh_vi_idx = zh_headers.index('effect_summary_vi')
zh_status_idx = zh_headers.index('status')

for row in zh_sheet.iter_rows(min_row=2):
    cid = str(row[zh_cid_idx].value or '')
    star = row[zh_star_idx].value
    if cid == 'V0117' and star == 6:
        row[zh_vi_idx].value = 'Tỷ Lệ Hút Máu +5%'
        row[zh_status_idx].value = 'TRANSLATED'
        print('Updated V0117 Rank 6 Zhizhi to Tỷ Lệ Hút Máu +5% based on attrDesMap.json!')

wb.save('localization/localization_master.xlsx')
print('[SUCCESS] Applied pass 2 naturalization to localization_master.xlsx!')
