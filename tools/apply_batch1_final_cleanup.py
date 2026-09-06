import openpyxl, json, re, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('localization/localization_master.xlsx')
batch1_chars = {'A0121', 'A0086', 'A0160', 'V0146', 'V0117'}

# 1. Update HUANZHANG sheet stories & status
hz_sheet = wb['HUANZHANG']
hz_headers = [cell.value for cell in hz_sheet[1]]
id_idx = hz_headers.index('brilliant_id')
info_vi_idx = hz_headers.index('icon_info_vi')
status_idx = hz_headers.index('status')
notes_idx = hz_headers.index('notes')

v01464_story = (
    '"Trâu nhỏ ơi Trâu nhỏ. Tiếng nước reo nhẹ nhàng bên suối gọi cô bé, uống ngụm nước rồi nghỉ chân đi nào."\n'
    '"Trâu nhỏ à Trâu nhỏ. Lúa dại, bắp ngô uốn eo trước gió khóc lóc than thở: Mau tới giúp chúng tôi với, bị sâu cắn mất rồi!"\n'
    '"Đừng sợ! Cậu ấy là ai chứ? Cậu ấy chính là Á Trưởng Ngưu Tôn đấm đá sâu bọ, uy phong lừng lẫy khắp ruộng đồng đây. Lũ yêu quái — còn không mau đầu hàng —"\n'
    '"Trâu nhỏ Trâu nhỏ, mau về nhà ăn cơm thôi!"\n'
    '"Bóng dáng vác cuốc diệt sâu đang tạo dáng dưới ánh chiều tà bỗng khựng lại, giây tiếp theo liền cắm đầu chạy thục mạng: \\"Cháu về ngay đây!\\""\n\n'
    '"Chị Trâu ơi. Trâu thì ăn gì ạ? Ăn cỏ chứ sao."\n'
    '"Thế cỏ thì ăn gì? Ăn ánh nắng với giọt sương chứ sao."\n'
    '"Thế ánh nắng với giọt sương thì... Ôi! Đừng đánh em! Em tự nghĩ là được chứ gì... Em biết rồi! Ánh nắng với giọt sương ăn chính là tụi mình đó."\n\n'
    '"Lũ trẻ trong làng lúc nào cũng tràn trề năng lượng cùng muôn vàn câu hỏi không bao giờ dứt, phiền phức chết đi được. Tụi nó đinh ninh rằng ánh nắng chiếu trên tấm lưng lăn lộn, hạt mưa tưới lên bàn chân tung tăng chính là bị thiên nhiên nuốt chửng trọn vẹn một lượt."\n\n'
    '"Chị Trâu chị Trâu. Đừng có cười hoài, rốt cuộc có đúng không hả?"\n'
    '"Cái má nhỏ nhắn phải bị ông mặt trời cắn hết lần này đến lần khác mới ửng đỏ lên, đôi chân thần tốc chạy đua với chú chó nhỏ phải được trận mưa rào tưới tắm bao bận mới cứng cáp; cái đầu nhỏ của con người ngoài trồng cấy ra còn nghĩ ra biết bao điều kỳ diệu nữa cơ đấy. Ngưu Ngưu thích lắm. Nhưng Ngưu Ngưu không nói đâu."\n\n'
    '"Bốp —"\n'
    '"Á! Sao lại đánh em nữa!"\n\n'
    '"Đừng có làm phiền chị, mời em lên xe phun nước hiệu Trâu Nhỏ."'
)

v01174_story = (
    'Sau khi Lục Vương Bổn Sinh Đồ bước vào thế giới ấy, câu chuyện liền phát triển đúng như những gì cậu hằng tưởng tượng. Muôn loài tự do tự tại, an nhàn sinh sống trong mảnh thiên địa này. Con người tuy có tranh chấp, nhưng thiện ác đều có báo ứng, mọi thứ diễn ra trật tự đâu vào đấy.\n'
    'Thế nhưng khi những người khác bước chân vào thế giới này, dựa theo từng lựa chọn để viết nên một kết cục khác cho thế giới, mọi chuyện lại chẳng còn như vậy nữa. Họ ít nhiều đều gặp phải khó khăn, chịu đựng thử thách. Có người bảo, chọn sai một bước liền rơi vào vạn kiếp bất phục; cũng có người bảo, thế giới nhỏ bé ấy được thiết kế quá đỗi tinh xảo, khiến người ta chẳng biết phải làm sao.\n'
    'Giờ đây, bạn cùng thiếu niên mang vầng sáng Cửu Sắc Lộc quanh mình tiến vào thế giới nhỏ bé này. Bạn cùng cậu ấy thăm dò đường đi phía trước, trao đổi góc nhìn, trò chuyện vô cùng tâm đắc. Chẳng biết từ lúc nào, hai người đã cùng nhau đi tới tận cùng thế giới...\n'
    'Chúc mừng người chơi đạt thành HE: Thế giới tuyệt đẹp đồng hành cùng Chú Lộc Nhỏ\n'
    'Khi được hỏi về cảm tưởng sau khi phá đảo, bạn trả lời: Thật ra, chỉ cần bạn giữ vững tâm thế giống như Lục Vương Bổn Sinh Đồ, trò chơi này cũng chẳng khó đến thế đâu? Chú hươu chín màu bị thu hút bởi lòng thiện lương kia sẽ dẫn dắt những ai cùng sở hữu tấm lòng nhân hậu cùng bước tới nhạc viện đầm ấm năm xưa.'
)

a01604_story = (
    '"Bác ơi, có tiếng đàn nào vang vọng mãi giữa đất trời, chẳng bị điều gì làm cho dứt đoạn không ạ?"\n'
    'Bên ngoài cửa sổ gỗ, hoa quế âm thầm rụng rơi. Thái Phụng Minh Kỳ tựa nghiêng bên cột hành lang, lắng nghe cuộc trò chuyện giữa một già một trẻ. Chàng mỉm cười ngoái nhìn, cũng lặng lẽ đợi chờ câu trả lời.\n'
    '"Cây đàn này do 雷威 (Lôi Uy) thời Khai Nguyên đại Đường chế tác. Người thời ấy ư, sớm đã rụng rơi thành cát bụi rồi. \'Từng giữ lời người xưa phủ Định, ta ôm đàn này thốt ba tiếng thở dài\'... Rốt cuộc cũng là khúc dứt người xa. Hễ dính dáng đến một chữ \'người\', tất sẽ có lúc đường cùng. Bởi bản thân con người vốn đã có điểm dừng."\n'
    'Thái Phụng Minh Kỳ thầm thở dài. Chàng là cây đàn, những chuyện chàng ghi nhớ nhiều hơn ông lão kia rất nhiều.\n'
    'Những nhà sưu tầm, những nghệ nhân gảy đàn, những người áo vải đi ngang chỉ để nghe chàng cất một tiếng đàn ngân vang.\n'
    'Từ người xa lạ đến chốn tri âm, một khúc hòa vang, gảy mãi tới lúc hừng đông rạng rỡ.\n'
    'Có người gửi gắm tâm sự vào dây đàn, truyền qua mảng tường, bên kia liền cất tiếng họa âm đáp lại. Điều quân tử trao tặng không phải là bản nhạc, mà là một cuộc tương phùng chẳng nỡ buông tay.\n'
    'Chàng từng chứng kiến sự trân trọng sâu sắc nhất: hai người đối mặt không nói một lời, chỉ đẩy qua đẩy lại cây đàn — bạn gảy, rồi tôi lại gảy, như thể ai dừng trước, người đó sẽ phải nói lời từ biệt trước. Nhưng con người rốt cuộc vẫn phải chia ly. Người ta đem những lời nói không cùng, những lá thư chẳng bao giờ tới, gom hết vào tiếng đàn, giao cho chàng cất giữ. Chàng bèn thay họ cất tiếng ngân dài, hết tiếng này đến tiếng khác, cho đến khi đối phương tận hưởng trọn vẹn rồi trở về.\n'
    'Trong mơ, chàng du ngoạn đến Tây Hồ, được mời tới làm khách trong một căn phòng nhỏ sáng bừng.\n'
    'Chiếc hộp nhỏ mở ra, ẩn chứa một bài hát, được những răng lược kim loại gảy đi gảy lại, một lần, rồi lại một lần.\n'
    'Sẽ không bao giờ dừng lại nữa.\n'
    'Thái Phụng Minh Kỳ lắng nghe âm thanh không bao giờ ngưng nghỉ ấy, chợt nhớ về quá khứ xa xăm, chàng từng dưới hành lang ngắm nhìn hoa quế rụng rào rạt, chim trời lướt qua góc mái, mưa đêm tí tách đến tận bình minh. Cưỡi gió mà đi, tận hưởng mà về. Con người, cây đàn, hoa chim, mưa đêm, tất cả đều như vậy.\n'
    'Hãy bay đi, mang theo ta lướt qua bầu trời ngàn năm.\n'
    'Đáp xuống rồi, ta lại đáp từ. Qua lại trao nhau, tiếng đàn bèn ngân vang không bao giờ dứt.'
)

for row in hz_sheet.iter_rows(min_row=2):
    rid = str(row[id_idx].value or '')
    if rid == 'V01464':
        row[info_vi_idx].value = v01464_story
        row[status_idx].value = 'TRANSLATED'
        row[notes_idx].value = 'Full lore translated from CN source'
        print('Updated V01464 story!')
    elif rid == 'V01174':
        row[info_vi_idx].value = v01174_story
        row[status_idx].value = 'TRANSLATED'
        row[notes_idx].value = 'Full lore translated from CN source'
        print('Updated V01174 story!')
    elif rid == 'A01604':
        row[info_vi_idx].value = a01604_story
        row[status_idx].value = 'TRANSLATED'
        row[notes_idx].value = 'Literary prose naturalized from CN source'
        print('Updated A01604 story!')
    elif any(c in rid for c in batch1_chars):
        row[status_idx].value = 'TRANSLATED'

# Helper function for naturalizing Vietnamese gameplay text
def naturalize_text(text, is_skill=False):
    if not text:
        return text
    
    # Fix 致命伤害 semantics
    text = text.replace('sát thương chí mạng', 'Sát Thương Chí Tử')
    text = text.replace('Sát thương chí mạng', 'Sát Thương Chí Tử')
    text = text.replace('Sát Thương Chí Mạng', 'Sát Thương Chí Tử')
    text = text.replace('chịu sát thương chí mạng', 'chịu Sát Thương Chí Tử')
    
    # Fix Damage Constructions
    # 'gây X% sát thương ... thành Sát Thương Chuẩn Bổ Sung'
    text = re.sub(
        r'gây\s+([^\.\,\;\n]+?)\s+thành\s+<color=#ff6724>Sát Thương Chuẩn Bổ Sung</color>',
        r'Gây <color=#ff6724>Sát Thương Chuẩn Bổ Sung</color> bằng \1',
        text,
        flags=re.IGNORECASE
    )
    text = re.sub(
        r'gây\s+([^\.\,\;\n]+?)\s+thành\s+Sát Thương Chuẩn Bổ Sung',
        r'Gây Sát Thương Chuẩn Bổ Sung bằng \1',
        text,
        flags=re.IGNORECASE
    )
    text = re.sub(
        r'gây\s+([^\.\,\;\n]+?)\s+thành\s+<color=#ff6724>Sát Thương Vật Lý Bổ Sung</color>',
        r'Gây <color=#ff6724>Sát Thương Vật Lý Bổ Sung</color> bằng \1',
        text,
        flags=re.IGNORECASE
    )
    text = re.sub(
        r'gây\s+([^\.\,\;\n]+?)\s+thành\s+Sát Thương Vật Lý Bổ Sung',
        r'Gây Sát Thương Vật Lý Bổ Sung bằng \1',
        text,
        flags=re.IGNORECASE
    )
    
    # Normalizations
    text = text.replace('1 kẻ địch đơn thể được chọn', '1 kẻ địch được chọn')
    text = text.replace('đơn vị kẻ địch đơn thể', '1 kẻ địch được chọn')
    text = text.replace('đơn vị kẻ địch', 'kẻ địch')
    text = text.replace('đơn vị đồng minh', 'đồng minh')
    text = text.replace('tiến hành tấn công', 'tấn công')
    text = text.replace('tiến hành di chuyển', 'di chuyển')
    text = text.replace('sau khi vào trận đấu', 'khi vào trận')
    
    # Stack phrases
    text = text.replace('tối đa có thể cộng dồn', 'có thể cộng dồn tối đa')
    text = text.replace('không ít hơn', 'có ít nhất')
    
    # Fix duplicate capitalized phrasing if any
    text = text.replace('Gây Gây', 'Gây')
    return text

# 2. Update SKILL sheet
skill_sheet = wb['SKILL']
sk_headers = [cell.value for cell in skill_sheet[1]]
sk_id_idx = sk_headers.index('skill_id')
sk_cid_idx = sk_headers.index('character_id')
sk_vi_idx = sk_headers.index('desc_vi')
sk_status_idx = sk_headers.index('status')
sk_notes_idx = sk_headers.index('notes')

for row in skill_sheet.iter_rows(min_row=2):
    cid = str(row[sk_cid_idx].value or '')
    if cid in batch1_chars:
        desc_vi = str(row[sk_vi_idx].value or '')
        row[sk_vi_idx].value = naturalize_text(desc_vi, is_skill=True)
        # Ensure status is TRANSLATED for description review, preserving owner name approval in notes
        row[sk_status_idx].value = 'TRANSLATED'
        row[sk_notes_idx].value = 'Skill name owner-approved; description naturalized & translated'

# 3. Update BUFF_STATUS sheet
buff_sheet = wb['BUFF_STATUS']
bf_headers = [cell.value for cell in buff_sheet[1]]
bf_id_idx = bf_headers.index('buff_id')
bf_cid_idx = bf_headers.index('character_id') if 'character_id' in bf_headers else 0
bf_vi_idx = bf_headers.index('buff_desc_vi')
bf_status_idx = bf_headers.index('status')

for row in buff_sheet.iter_rows(min_row=2):
    rid = str(row[bf_id_idx].value or '')
    cid = str(row[bf_cid_idx].value or '') if bf_cid_idx else ''
    if any(c in rid or c in cid for c in batch1_chars):
        desc_vi = str(row[bf_vi_idx].value or '')
        row[bf_vi_idx].value = naturalize_text(desc_vi)
        row[bf_status_idx].value = 'TRANSLATED'

# 4. Update ZHIZHI sheet
zh_sheet = wb['ZHIZHI']
zh_headers = [cell.value for cell in zh_sheet[1]]
zh_cid_idx = zh_headers.index('character_id')
zh_star_idx = zh_headers.index('star')
zh_vi_idx = zh_headers.index('effect_summary_vi')
zh_status_idx = zh_headers.index('status')

for row in zh_sheet.iter_rows(min_row=2):
    cid = str(row[zh_cid_idx].value or '')
    star = row[zh_star_idx].value
    if cid in batch1_chars:
        val = str(row[zh_vi_idx].value or '')
        val = naturalize_text(val)
        # V0117 Rank 6 HurtHealRate resolution
        if cid == 'V0117' and star == 6:
            row[zh_vi_idx].value = 'Hiệu Quả Trị Liệu Nhận Được +5%'
        elif val:
            row[zh_vi_idx].value = val
        row[zh_status_idx].value = 'TRANSLATED'

wb.save('localization/localization_master.xlsx')
print('[SUCCESS] Updated localization_master.xlsx with naturalized descriptions, full lore, and status!')
