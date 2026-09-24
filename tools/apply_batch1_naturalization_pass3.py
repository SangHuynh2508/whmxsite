import openpyxl, json, re, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('localization/localization_master.xlsx')
batch1_chars = {'A0121', 'A0086', 'A0160', 'V0146', 'V0117'}

# Master dictionary matching exact CN placeholder tokens
skill_updates = {
    # Basic Attacks (01)
    'A012101_1': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'A012101_2': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'A012101_3': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'A012101_4': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'A012101_5': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    
    'A008601_1': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'A008601_2': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'A008601_3': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'A008601_4': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'A008601_5': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    
    'A016001_1': 'Gây Sát Thương Cấu Thuật bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'A016001_2': 'Gây Sát Thương Cấu Thuật bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'A016001_3': 'Gây Sát Thương Cấu Thuật bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'A016001_4': 'Gây Sát Thương Cấu Thuật bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'A016001_5': 'Gây Sát Thương Cấu Thuật bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    
    'V011701_1': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'V011701_2': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'V011701_3': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'V011701_4': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'V011701_5': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    
    'V014601_1': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'V014601_2': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'V014601_3': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'V014601_4': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',
    'V014601_5': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn.',

    # V0117 Skills
    'V011702_1': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn cùng toàn bộ kẻ địch <color=#158bdb>xung quanh 1 ô</color>, khiến kẻ địch nhận trạng thái <color=#ff6724>Trầm Nộ</color> và <color=#ff6724>Nghiệp Chướng</color>.{Buff_Taunt}{Buff_V0117_2_1}',
    'V011702_2': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn cùng toàn bộ kẻ địch <color=#158bdb>xung quanh 1 ô</color>, khiến kẻ địch nhận trạng thái <color=#ff6724>Trầm Nộ</color> và <color=#ff6724>Nghiệp Chướng</color>.{Buff_Taunt}{Buff_V0117_2_1}',
    'V011702_3': 'Gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên 1 kẻ địch được chọn cùng toàn bộ kẻ địch <color=#158bdb>xung quanh 1 ô</color>, khiến kẻ địch nhận trạng thái <color=#ff6724>Trầm Nộ</color> và <color=#ff6724>Nghiệp Chướng</color>.{Buff_Taunt}{Buff_V0117_2_1}',
    
    'V011703_1': 'Sau khi chịu tấn công trực tiếp từ kẻ địch, Phản lại <color=#ff6724>Sát Thương Vật Lý Bổ Sung</color> bằng <color=#158bdb><color=#158bdb>[Effect3Para,1]%</color></color> Tấn Công hiện tại và giúp bản thân nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Minh Kính</color>.\nKhi có ít nhất <color=#158bdb>2</color> tầng trạng thái Minh Kính: kẻ địch tấn công trực tiếp bản thân xong sẽ nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Hàn Thiên</color>;\nKhi có ít nhất <color=#158bdb>4</color> tầng trạng thái Minh Kính: Sát Thương Bạo Kích của bản thân tăng <color=#158bdb>50%</color>;\nKhi có ít nhất <color=#158bdb>6</color> tầng trạng thái Minh Kính: kẻ địch tấn công trực tiếp bản thân xong sẽ nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Thiêu Đốt</color>.{Buff_V0117_3}{Buff_Cold}{Buff_Burning}{Buff_SpecialDamgeP_Lan}',
    'V011703_2': 'Sau khi chịu tấn công trực tiếp từ kẻ địch, Phản lại <color=#ff6724>Sát Thương Vật Lý Bổ Sung</color> bằng <color=#158bdb><color=#158bdb>[Effect3Para,1]%</color></color> Tấn Công hiện tại và giúp bản thân nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Minh Kính</color>.\nKhi có ít nhất <color=#158bdb>2</color> tầng trạng thái Minh Kính: kẻ địch tấn công trực tiếp bản thân xong sẽ nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Hàn Thiên</color>;\nKhi có ít nhất <color=#158bdb>4</color> tầng trạng thái Minh Kính: Sát Thương Bạo Kích của bản thân tăng <color=#158bdb>50%</color>;\nKhi có ít nhất <color=#158bdb>6</color> tầng trạng thái Minh Kính: kẻ địch tấn công trực tiếp bản thân xong sẽ nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Thiêu Đốt</color>.{Buff_V0117_3}{Buff_Cold}{Buff_Burning}{Buff_SpecialDamgeP_Lan}',
    'V011703_3': 'Sau khi chịu tấn công trực tiếp từ kẻ địch, Phản lại <color=#ff6724>Sát Thương Vật Lý Bổ Sung</color> bằng <color=#158bdb><color=#158bdb>[Effect3Para,1]%</color></color> Tấn Công hiện tại và giúp bản thân nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Minh Kính</color>.\nKhi có ít nhất <color=#158bdb>2</color> tầng trạng thái Minh Kính: kẻ địch tấn công trực tiếp bản thân xong sẽ nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Hàn Thiên</color>;\nKhi có ít nhất <color=#158bdb>4</color> tầng trạng thái Minh Kính: Sát Thương Bạo Kích của bản thân tăng <color=#158bdb>50%</color>;\nKhi có ít nhất <color=#158bdb>6</color> tầng trạng thái Minh Kính: kẻ địch tấn công trực tiếp bản thân xong sẽ nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Thiêu Đốt</color>.{Buff_V0117_3}{Buff_Cold}{Buff_Burning}{Buff_SpecialDamgeP_Lan}',
    
    'V011704_1': 'Sau khi lũy kế chịu <color=#158bdb>4</color> lần tấn công trực tiếp từ kẻ địch, lập tức nhận thêm <color=#158bdb>1</color> lượt hành động và Đánh Thường tiếp theo sẽ Liên Kích, gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên kẻ địch bị tấn công.',
    'V011704_2': 'Sau khi lũy kế chịu <color=#158bdb>4</color> lần tấn công trực tiếp từ kẻ địch, lập tức nhận thêm <color=#158bdb>1</color> lượt hành động và Đánh Thường tiếp theo sẽ Liên Kích, gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên kẻ địch bị tấn công.',
    'V011704_3': 'Sau khi lũy kế chịu <color=#158bdb>4</color> lần tấn công trực tiếp từ kẻ địch, lập tức nhận thêm <color=#158bdb>1</color> lượt hành động và Đánh Thường tiếp theo sẽ Liên Kích, gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công lên kẻ địch bị tấn công.',
    
    'V011705_1': 'Khi vào trận, Máu tối đa của bản thân tăng <color=#158bdb>[Effect2Para,1]%</color>;\nMỗi lần gây sát thương lên kẻ địch, cướp <color=#158bdb>[Effect1Para,1]%</color> Tấn Công hiện tại của kẻ địch bị tấn công, mỗi lần tối đa <color=#158bdb>200</color> điểm, Tấn Công cướp được có thể cộng dồn vô hạn; khi kẻ địch tử vong, Tấn Công cướp từ kẻ địch đó sẽ bị xóa.',
    'V011705_2': 'Khi vào trận, Máu tối đa của bản thân tăng <color=#158bdb>[Effect2Para,1]%</color>;\nMỗi lần gây sát thương lên kẻ địch, cướp <color=#158bdb>[Effect1Para,1]%</color> Tấn Công hiện tại của kẻ địch bị tấn công, mỗi lần tối đa <color=#158bdb>300</color> điểm, Tấn Công cướp được có thể cộng dồn vô hạn; khi kẻ địch tử vong, Tấn Công cướp từ kẻ địch đó sẽ bị xóa.',
    'V011705_3': 'Khi vào trận, Máu tối đa của bản thân tăng <color=#158bdb>[Effect2Para,1]%</color>;\nMỗi lần gây sát thương lên kẻ địch, cướp <color=#158bdb>[Effect1Para,1]%</color> Tấn Công hiện tại của kẻ địch bị tấn công, mỗi lần tối đa <color=#158bdb>400</color> điểm, Tấn Công cướp được có thể cộng dồn vô hạn; khi kẻ địch tử vong, Tấn Công cướp từ kẻ địch đó sẽ bị xóa.',
    
    'V011711_1': 'Chọn 1 ô trống bất kỳ trong phạm vi <color=#158bdb>xung quanh 5 ô</color> bản thân để dịch chuyển tới, đồng thời nhận <color=#158bdb>1</color> lượt hành động lại (Sức Di Chuyển của lượt này giảm <color=#158bdb>2</color>).'
}

buff_updates = {
    'Buff_A0160_10': 'Khi chịu <color=#158bdb>1</color> lần sát thương từ Đánh Thường hoặc Kỹ Năng, Gây <color=#ff6724>Sát Thương Vật Lý Bổ Sung</color> bằng <color=#158bdb>[EffectParam,2]%</color> Tấn Công của Thái Phụng Minh Kỳ <color=#158bdb>1</color> lần lên bản thân và toàn bộ đồng minh <color=#158bdb>xung quanh 3 ô</color>, sau khi kích hoạt hiệu ứng trên, trạng thái này sẽ bị xóa.',
    'Buff_V0117_hz_1_1': 'Trạng thái này có thể cộng dồn tối đa <color=#158bdb>3</color> tầng;\nKhi chịu tấn công Tuyệt Kỹ hoặc Liên Kích của Lộc Vương Bản Sinh Đồ bên địch, mỗi tầng trạng thái này khiến bản thân chịu Sát Thương Vật Lý Bổ Sung bằng <color=#158bdb>220%</color> Tấn Công của Lộc Vương Bản Sinh Đồ, sát thương này chắc chắn Bạo Kích, sau khi chịu tấn công từ Lộc Vương Bản Sinh Đồ bên địch sẽ xóa toàn bộ trạng thái này.'
}

# Apply to SKILL sheet
sk_sheet = wb['SKILL']
sk_headers = [cell.value for cell in sk_sheet[1]]
id_idx = sk_headers.index('skill_id')
desc_vi_idx = sk_headers.index('desc_vi')
status_idx = sk_headers.index('status')
notes_idx = sk_headers.index('notes')

for row in sk_sheet.iter_rows(min_row=2):
    sid = str(row[id_idx].value or '')
    if sid in skill_updates:
        row[desc_vi_idx].value = skill_updates[sid]
        row[status_idx].value = 'TRANSLATED'
        row[notes_idx].value = 'Skill name owner-approved; description naturalized & translated'

# Apply to BUFF_STATUS sheet
bf_sheet = wb['BUFF_STATUS']
bf_headers = [cell.value for cell in bf_sheet[1]]
bf_id_idx = bf_headers.index('buff_id')
bf_desc_vi_idx = bf_headers.index('buff_desc_vi')
bf_status_idx = bf_headers.index('status')

for row in bf_sheet.iter_rows(min_row=2):
    bid = str(row[bf_id_idx].value or '')
    if bid in buff_updates:
        row[bf_desc_vi_idx].value = buff_updates[bid]
        row[bf_status_idx].value = 'TRANSLATED'

wb.save('localization/localization_master.xlsx')
print('[SUCCESS] Applied pass 3 exact rich-text tag alignment!')
