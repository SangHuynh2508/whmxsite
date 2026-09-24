"""Apply one source-verified CHARACTER localization correction with a backup."""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "localization" / "localization_master.xlsx"


def s(value): return "" if value is None else str(value)
def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--transaction", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    locks = list(MASTER.parent.glob(f"~${MASTER.name}"))
    if locks: raise RuntimeError(f"Excel lock: {locks}")
    tx = json.loads(args.transaction.read_text(encoding="utf-8"))
    wb = load_workbook(MASTER, data_only=False)
    ws = wb["CHARACTER"]
    h = {s(cell.value): cell.column for cell in ws[1] if cell.value is not None}
    rows = {s(ws.cell(row, h["character_id"]).value): row for row in range(2, ws.max_row + 1)}
    plan, failures = [], []
    for target in tx["targets"]:
        cid = target["character_id"]
        if cid not in rows:
            failures.append({"character_id": cid, "reason": "MISSING_TARGET"}); continue
        row = rows[cid]
        actual = {key: s(ws.cell(row, h[key]).value) for key in target["expected"]}
        if actual != target["expected"]:
            failures.append({"character_id": cid, "reason": "SOURCE_OR_OWNER_VALUE_CHANGED", "actual": actual}); continue
        plan.append((row, target))
    report = {"transaction_id": tx["transaction_id"], "planned": len(plan), "failures": failures, "applied": []}
    if args.apply:
        if failures: raise RuntimeError("Refusing partial transaction")
        backup = ROOT / "localization" / "backups" / f"localization_master_pre_{tx['transaction_id'].lower()}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        backup.parent.mkdir(parents=True, exist_ok=True)
        before_sha = sha(MASTER); shutil.copy2(MASTER, backup)
        if sha(backup) != before_sha: raise RuntimeError("backup hash mismatch")
        for row, target in plan:
            changes = {}
            for field, value in target["translation"].items():
                old = s(ws.cell(row, h[field]).value); ws.cell(row, h[field]).value = value; changes[field] = {"old": old, "new": value}
            note = s(ws.cell(row, h["notes"]).value).rstrip("; ")
            ws.cell(row, h["notes"]).value = (note + "; " + target["provenance"]).strip("; ")
            report["applied"].append({"character_id": target["character_id"], "changes": changes})
        wb.save(MASTER); report.update({"backup": str(backup), "backup_sha256": before_sha})
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"report": str(args.report), "planned": len(plan), "applied": len(report["applied"]), "failures": len(failures)}, ensure_ascii=False))


if __name__ == "__main__": main()
