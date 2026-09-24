"""Apply one source-verified group of exact SKILL named-reference migrations."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "localization" / "localization_master.xlsx"


def s(value):
    return "" if value is None else str(value)


def headers(ws):
    return {s(cell.value): cell.column for cell in ws[1] if cell.value is not None}


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def tokens(value):
    return (re.findall(r"<[^>]+>", value), re.findall(r"\[(?:Effect|Condition)\w*,\d+\]", value))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--transaction", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()

    locks = list(MASTER.parent.glob(f"~${MASTER.name}"))
    if locks:
        raise RuntimeError(f"Excel lock: {locks}")
    transaction = json.loads(args.transaction.read_text(encoding="utf-8"))
    wb = load_workbook(MASTER, data_only=False)
    skill_ws, buff_ws = wb["SKILL"], wb["BUFF_STATUS"]
    sh, bh = headers(skill_ws), headers(buff_ws)
    skill_rows = {s(skill_ws.cell(row, sh["skill_id"]).value): row for row in range(2, skill_ws.max_row + 1)}
    buff_rows = {s(buff_ws.cell(row, bh["buff_id"]).value): row for row in range(2, buff_ws.max_row + 1)}
    plan, failures = [], []

    for target in transaction["targets"]:
        sid, bid = target["skill_id"], target["buff_id"]
        if sid not in skill_rows or bid not in buff_rows:
            failures.append({"skill_id": sid, "buff_id": bid, "reason": "MISSING_TARGET"})
            continue
        skill_row, buff_row = skill_rows[sid], buff_rows[bid]
        source = s(skill_ws.cell(skill_row, sh["desc_cn"]).value)
        before = s(skill_ws.cell(skill_row, sh["desc_vi"]).value)
        actual_buff_cn = s(buff_ws.cell(buff_row, bh["buff_name_cn"]).value)
        if actual_buff_cn != target["expected_buff_name_cn"] or target["source_cn"] not in source:
            failures.append({"skill_id": sid, "buff_id": bid, "reason": "SOURCE_CHANGED", "actual_buff_cn": actual_buff_cn})
            continue
        after = before
        for replacement in target["replacements"]:
            old, new, count = replacement["old_vi"], replacement["new_vi"], replacement["count"]
            if before.count(old) != count:
                failures.append({"skill_id": sid, "buff_id": bid, "reason": "OLD_VALUE_CHANGED", "old_vi": old, "expected_count": count, "actual_count": before.count(old)})
                break
            after = after.replace(old, new)
        else:
            if tokens(before) != tokens(after):
                failures.append({"skill_id": sid, "buff_id": bid, "reason": "MARKER_OR_PLACEHOLDER_CHANGED"})
            else:
                plan.append((skill_row, target, before, after))

    report = {"transaction_id": transaction["transaction_id"], "planned": len(plan), "failures": failures, "applied": []}
    if args.apply:
        if failures:
            raise RuntimeError("Refusing partial transaction")
        backup = ROOT / "localization" / "backups" / f"localization_master_pre_{transaction['transaction_id'].lower()}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        backup.parent.mkdir(parents=True, exist_ok=True)
        before_sha = sha256(MASTER)
        shutil.copy2(MASTER, backup)
        if sha256(backup) != before_sha:
            raise RuntimeError("backup hash mismatch")
        for row, target, before, after in plan:
            skill_ws.cell(row, sh["desc_vi"]).value = after
            note = s(skill_ws.cell(row, sh["notes"]).value).rstrip("; ")
            marker = f"Exact named reference migration {transaction['transaction_id']}"
            skill_ws.cell(row, sh["notes"]).value = (note + "; " + marker).strip("; ")
            report["applied"].append({"skill_id": target["skill_id"], "buff_id": target["buff_id"], "before": before, "after": after})
        wb.save(MASTER)
        report.update({"backup": str(backup), "backup_sha256": before_sha})
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"report": str(args.report), "planned": len(plan), "applied": len(report["applied"]), "failures": len(failures)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
