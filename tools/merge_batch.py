"""Safely merge a returned translation batch.

Default mode is dry-run.  Only explicit ``--apply`` can write the master.
"""

import os
import sys

import openpyxl
import pandas as pd


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MASTER_FILE = os.path.join(ROOT, "localization", "localization_master.xlsx")
DEFAULT_INPUT = os.path.join(ROOT, "batch_5_characters_export.xlsx")
PRIMARY_KEYS = {
    "CHARACTER": ("character_id",), "SKILL": ("skill_id",),
    "ZHIZHI": ("character_id", "star"), "HUANZHANG": ("brilliant_id",),
    "SKIN": ("skin_id",), "BUFF_STATUS": ("buff_id",),
}
SHEET_MAP = {"characters": "CHARACTER", "skills": "SKILL", "buffs": "BUFF_STATUS", "zhizhi": "ZHIZHI", "brilliant": "HUANZHANG", "huanzhang": "HUANZHANG", "skins": "SKIN"}
NON_MERGE_SHEETS = {"EXPORT_MANIFEST", "CHARACTER_CHECKLIST", "REFERENCE_CONTEXT", "UNRESOLVED_REFERENCES"}


def canonical(value):
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    return text[:-2].upper() if text.endswith(".0") else text.upper()


def row_key(row, columns):
    return tuple(canonical(row.get(column)) for column in columns)


def duplicates(frame, columns):
    keys, duplicate = set(), []
    for _, row in frame.iterrows():
        identity = row_key(row, columns)
        if not all(identity):
            duplicate.append((identity, "blank_key"))
        elif identity in keys:
            duplicate.append((identity, "duplicate_key"))
        keys.add(identity)
    return duplicate


def merge_batch(input_file=DEFAULT_INPUT, apply=False):
    if not os.path.exists(MASTER_FILE):
        raise FileNotFoundError(f"Missing master workbook: {MASTER_FILE}")
    if not os.path.exists(input_file):
        raise FileNotFoundError(f"Missing batch workbook: {input_file}")

    master = openpyxl.load_workbook(MASTER_FILE)
    batch = pd.ExcelFile(input_file)
    report = {"matched_rows": 0, "changed_cells": 0, "unchanged_cells": 0, "skipped_blank_cells": 0, "missing_ids": [], "unsupported_sheets": [], "duplicate_keys": []}

    for raw_sheet in batch.sheet_names:
        if raw_sheet.upper() in NON_MERGE_SHEETS:
            continue
        sheet = SHEET_MAP.get(raw_sheet.lower(), raw_sheet.upper())
        if sheet not in PRIMARY_KEYS or sheet not in master.sheetnames:
            report["unsupported_sheets"].append(raw_sheet)
            continue
        frame = pd.read_excel(input_file, sheet_name=raw_sheet)
        if frame.empty:
            continue
        keys = PRIMARY_KEYS[sheet]
        if any(column not in frame.columns for column in keys):
            report["unsupported_sheets"].append(f"{raw_sheet}: missing primary key")
            continue
        duplicate = duplicates(frame, keys)
        if duplicate:
            report["duplicate_keys"].extend(f"{sheet}:{entry}" for entry in duplicate)
            continue

        ws = master[sheet]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        if any(column not in headers for column in keys):
            report["unsupported_sheets"].append(f"{sheet}: master primary key mismatch")
            continue
        master_index = {}
        for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            identity = tuple(canonical(values[headers.index(column)]) for column in keys)
            if identity in master_index:
                report["duplicate_keys"].append(f"{sheet}:{identity}: master")
            master_index[identity] = row_number

        for _, batch_row in frame.iterrows():
            identity = row_key(batch_row, keys)
            row_number = master_index.get(identity)
            if not row_number:
                report["missing_ids"].append(f"{sheet}:{identity}")
                continue
            report["matched_rows"] += 1
            for column in frame.columns:
                # Never touch CN/raw/reference columns.  Blank returned cells do
                # not erase existing translation content.
                if not column.endswith("_vi") or column not in headers:
                    continue
                incoming = batch_row[column]
                if incoming is None or pd.isna(incoming) or not str(incoming).strip():
                    report["skipped_blank_cells"] += 1
                    continue
                target = ws.cell(row=row_number, column=headers.index(column) + 1)
                if str(target.value or "") == str(incoming):
                    report["unchanged_cells"] += 1
                else:
                    report["changed_cells"] += 1
                    if apply:
                        target.value = incoming

    if report["duplicate_keys"]:
        print("MERGE_BLOCKED duplicate keys:", report["duplicate_keys"])
        return report
    if apply:
        master.save(MASTER_FILE)
        print(f"MERGE_APPLIED: {MASTER_FILE}")
    else:
        print("DRY_RUN: master workbook was not modified")
    print(report)
    return report


if __name__ == "__main__":
    args = sys.argv[1:]
    apply = "--apply" in args
    args = [arg for arg in args if arg != "--apply"]
    merge_batch(args[0] if args else DEFAULT_INPUT, apply=apply)
