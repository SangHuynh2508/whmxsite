"""Fallback XLSX serializer for the character packet exporter.

Used only when artifact-tool cannot serialize the packet's large rich-text
cells.  It consumes an ephemeral JSON payload and writes a brand-new review
workbook; it never opens or mutates localization_master.xlsx.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def write_packet_from_dict(payload: dict[str, list[dict]], output_path: str | Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    for sheet_name, rows in payload.items():
        worksheet = workbook.create_sheet(sheet_name)
        headers = list(dict.fromkeys(key for row in rows for key in row)) if rows else []
        if headers:
            worksheet.append(headers)
        for row in rows:
            worksheet.append([
                json.dumps(row.get(header, ""), ensure_ascii=False)
                if isinstance(row.get(header, ""), (list, dict))
                else (row.get(header, "") if row.get(header, "") is not None else "")
                for header in headers
            ])
        if headers:
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.font = body_font
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            worksheet.freeze_panes = "A2"
            for index, header in enumerate(headers, 1):
                width = 45 if any(term in header.lower() for term in ("desc", "lore", "story", "notes", "evidence", "info", "path")) else min(55, max(14, len(header) + 3))
                worksheet.column_dimensions[get_column_letter(index)].width = width
            worksheet.row_dimensions[1].height = 30
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main(payload_path: str, output_path: str) -> None:
    payload = json.loads(Path(payload_path).read_text(encoding="utf-8"))
    write_packet_from_dict(payload, output_path)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
