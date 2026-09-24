import os
import sys
import json
import openpyxl
from batch_manifest_helper import load_batch_manifest, collect_batch_dependencies
from validate_localization_batch import validate_localization_batch
from validate_terminology_consistency import validate_terminology_consistency
from check_suspicious_vi import check_suspicious_vi
from zhizhi_resolver import contains_raw_dev_codes, resolve_zhizhi_entry
from generate_localization_review import generate_review

def run_release_gate(base_dir, manifest_path):
    print("==================================================")
    print("PHASE 3 LOCALIZATION RELEASE GATE")
    print("==================================================")
    print(f"Manifest: {manifest_path}\n")

    manifest = load_batch_manifest(manifest_path)
    batch_id = manifest['batch_id']
    deps = collect_batch_dependencies(base_dir, manifest)

    errors = []
    warnings = []

    # 1. Batch Coverage Validation
    print("--- 1. BATCH COVERAGE CHECK ---")
    cov_success, cov_counts, cov_errs = validate_localization_batch(manifest_path, base_dir)
    if not cov_success:
        errors.extend([f"[COVERAGE_ERROR] {e}" for e in cov_errs])
    else:
        print(f"Coverage Passed: SKILL {cov_counts['SKILL']}, BUFF {cov_counts['BUFF_STATUS']}, ZHIZHI {cov_counts['ZHIZHI']}, HUANZHANG {cov_counts['HUANZHANG']}\n")

    # 2. Placeholder & Rich-Text Parity Validation
    print("--- 2. PLACEHOLDER & RICH-TEXT PARITY CHECK ---")
    try:
        master_path = os.path.join(base_dir, 'localization', 'localization_master.xlsx')
        wb = openpyxl.load_workbook(master_path, data_only=True)
        ph_errors = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) <= 1:
                continue
            header = [str(c).strip() if c is not None else '' for c in rows[0]]
            col_pairs = []
            for idx, h in enumerate(header):
                hl = h.lower()
                if 'cn' in hl or 'zh' in hl:
                    vi_col_name = hl.replace('cn', 'vi').replace('zh', 'vi')
                    for vi_idx, vi_h in enumerate(header):
                        if vi_h.lower() == vi_col_name and vi_h.lower() not in ['status', 'notes', 'provenance_vi']:
                            col_pairs.append((idx, h, vi_idx, vi_h))
                            break
            for row_idx, r in enumerate(rows[1:], start=2):
                key_id = str(r[0]) if r[0] is not None else f"Row_{row_idx}"
                for cn_idx, cn_h, vi_idx, vi_h in col_pairs:
                    cn_val = str(r[cn_idx]).strip() if cn_idx < len(r) and r[cn_idx] is not None else ''
                    vi_val = str(r[vi_idx]).strip() if vi_idx < len(r) and r[vi_idx] is not None else ''
                    if not vi_val or vi_val == 'NONE':
                        continue
                    import re
                    cn_ph = sorted(re.findall(r'\[Effect[^\]]+\]', cn_val))
                    vi_ph = sorted(re.findall(r'\[Effect[^\]]+\]', vi_val))
                    if cn_ph != vi_ph:
                        ph_errors.append(f"[PLACEHOLDER_ERROR] {sheet_name}/{key_id}/{vi_h}: CN={cn_ph} != VI={vi_ph}")
                    cn_tags = sorted(re.findall(r'</?color[^>]*>', cn_val))
                    vi_tags = sorted(re.findall(r'</?color[^>]*>', vi_val))
                    if cn_tags != vi_tags:
                        ph_errors.append(f"[RICH_TEXT_ERROR] {sheet_name}/{key_id}/{vi_h}: CN={cn_tags} != VI={vi_tags}")
        if ph_errors:
            errors.extend(ph_errors)
        else:
            print("Placeholder and Rich-Text Parity: 100% Passed\n")
    except Exception as e:
        errors.append(f"[PLACEHOLDER_EXEC_ERROR] {e}")

    # 3. Accidental Han Character Validation
    print("--- 3. ACCIDENTAL HAN-CHARACTER CHECK ---")
    from validate_no_han_characters import APPROVED_WHITELIST
    import re
    cjk_regex = re.compile(r'[\u4e00-\u9fff]')
    han_errors = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = rows[0]
        vi_indices = [i for i, h in enumerate(header) if h and '_vi' in str(h).lower()]
        for row_idx, r in enumerate(rows[1:], start=2):
            rec_id = str(r[0]) if r[0] is not None else f"Row_{row_idx}"
            for idx in vi_indices:
                val = str(r[idx]) if r[idx] is not None else ''
                if val:
                    check_val = val
                    for w in APPROVED_WHITELIST:
                        check_val = check_val.replace(w, '')
                    han_matches = cjk_regex.findall(check_val)
                    if han_matches:
                        han_errors.append({
                            'sheet': sheet_name,
                            'record_id': rec_id,
                            'field': header[idx],
                            'offending_chars': list(set(han_matches))
                        })
    if han_errors:
        for err in han_errors:
            errors.append(f"[HAN_LEAK_ERROR] Sheet: {err['sheet']}, Record: {err['record_id']}, Field: {err['field']}, Chars: {err['offending_chars']}")
    else:
        print("Accidental Han Characters: 0 Leaks Found\n")

    # 4. Terminology Consistency Check
    print("--- 4. TERMINOLOGY CONSISTENCY CHECK ---")
    term_pass, term_warns = validate_terminology_consistency(manifest_path, base_dir, wb_master=wb)
    if term_warns:
        for tw in term_warns:
            warnings.append(f"[TERMINOLOGY_WARNING] CN: {tw['cn_term']} | Skill VI Statuses: {tw['skill_vi_statuses']} | Buff VI: {tw['buff_vi']} (Skill {tw['skill_id']}, Buff {tw['buff_id']})")
    else:
        print("Terminology Consistency: 0 Conflicts Found\n")

    # 5. Suspicious VI Quality Warning Scan
    print("--- 5. SUSPICIOUS VI QUALITY SCAN ---")
    vi_warns = check_suspicious_vi(manifest_path, base_dir, wb_master=wb)
    if vi_warns:
        for vw in vi_warns:
            warnings.append(f"[SUSPICIOUS_VI_WARNING] {vw['sheet']}/{vw['record_id']}: Matched '{vw['pattern']}' -> \"{vw['vi_txt']}\"")
    else:
        print("Suspicious VI Quality Scan: 0 Warnings Found\n")

    # 6. Zhizhi Resolution Sanity
    print("--- 6. ZHIZHI RESOLUTION SANITY CHECK ---")
    ws_zh = wb['ZHIZHI']
    headers_zh = [str(c or '').strip() for c in next(ws_zh.iter_rows(values_only=True))]
    cid_zh_idx = headers_zh.index('character_id')
    star_zh_idx = headers_zh.index('star')
    eff_zh_idx = headers_zh.index('effect_type')
    sum_cn_idx = headers_zh.index('effect_summary_cn')
    base_id_idx = headers_zh.index('skill_up_base_id')

    ws_sk = wb['SKILL']
    headers_sk = [str(c or '').strip() for c in next(ws_sk.iter_rows(values_only=True))]
    sid_sk_idx = headers_sk.index('skill_id')
    name_vi_sk_idx = headers_sk.index('skill_name_vi')
    skill_names_map = {}
    for r in ws_sk.iter_rows(min_row=2, values_only=True):
        sid = str(r[sid_sk_idx] or '').strip()
        nvi = str(r[name_vi_sk_idx] or '').strip()
        if nvi:
            skill_names_map[sid] = nvi
            if len(sid) >= 7:
                skill_names_map[sid[:7]] = nvi

    zh_errors = []
    for r in ws_zh.iter_rows(min_row=2, values_only=True):
        cid = str(r[cid_zh_idx] or '').strip()
        if cid in deps['characters']:
            star = r[star_zh_idx]
            eff_type = r[eff_zh_idx]
            sum_cn = r[sum_cn_idx]
            base_id = r[base_id_idx]
            res_vi = resolve_zhizhi_entry(cid, star, eff_type, sum_cn, base_id, skill_names_map)
            if contains_raw_dev_codes(res_vi):
                zh_errors.append(f"[ZHIZHI_RAW_CODE_ERROR] Character {cid} Rank {star} resolved to raw code: '{res_vi}'")

    if zh_errors:
        errors.extend(zh_errors)
    else:
        print("Zhizhi Resolution: 100% Player-Facing Clean\n")

    # 7. Review Generator Sanity
    print("--- 7. REVIEW GENERATOR SANITY CHECK ---")
    try:
        generate_review(base_dir, manifest_path, wb_master=wb)
        md_path = os.path.join(base_dir, 'localization', 'batch_review', f"{batch_id}_human_review.md")
        if not os.path.exists(md_path):
            errors.append(f"[REVIEW_ERROR] Generated Markdown file missing: {md_path}")
        else:
            with open(md_path, 'r', encoding='utf-8') as f:
                content = f.read()
            if '[MISSING VI' in content:
                errors.append(f"[REVIEW_ERROR] Review package contains missing VI markers!")
            else:
                print("Review Generator Sanity: OK\n")
    except Exception as e:
        errors.append(f"[REVIEW_EXEC_ERROR] {e}")

    # 8. Public Output Leak Validation
    print("--- 8. PUBLIC OUTPUT LEAK CHECK ---")
    try:
        from validate_public_output import validate_public_output
        po_success, po_errs = validate_public_output(manifest_path, os.path.join(base_dir, 'public', 'data.json'))
        if not po_success:
            errors.extend([f"[PUBLIC_OUTPUT_LEAK] {e}" for e in po_errs])
        else:
            print("Public Output Leak Check: 100% Passed\n")
    except Exception as e:
        errors.append(f"[PUBLIC_OUTPUT_EXEC_ERROR] {e}")

    # Summary Output
    print("==================================================")
    print("RELEASE GATE FINAL SUMMARY")
    print("==================================================")
    print(f"Total Errors:   {len(errors)}")
    print(f"Total Warnings: {len(warnings)}")
    print("--------------------------------------------------")

    if warnings:
        print("\n[WARNINGS]")
        for w in warnings:
            print(f"  - {w}")

    if errors:
        print("\n[ERRORS - RELEASE GATE FAILED]")
        for e in errors:
            print(f"  - {e}")
        print("==================================================")
        sys.exit(1)
    else:
        print("\n[PASSED - RELEASE GATE SUCCESSFUL]")
        print("Batch is valid and ready for human/project-owner review.")
        print("==================================================")
        sys.exit(0)

if __name__ == '__main__':
    manifest_arg = sys.argv[1] if len(sys.argv) > 1 else r"localization/batches/phase3_batch1.json"
    base_dir = r"d:\BaiTapCode\WHMX\WhmxCalc"
    run_release_gate(base_dir, manifest_arg)
