import openpyxl, json, os, sys, shutil, time, re, subprocess

sys.stdout.reconfigure(encoding='utf-8')

timestamp = time.strftime('%Y%m%d_%H%M%S')
base_dir = r"d:\BaiTapCode\WHMX\WhmxCalc"
backup_dir = os.path.join(base_dir, 'localization', 'backups')
recovery_dir = os.path.join(base_dir, 'localization', 'recovery')

os.makedirs(backup_dir, exist_ok=True)
os.makedirs(recovery_dir, exist_ok=True)

# ----------------------------------------------------
# 1. CREATE SAFETY BACKUPS FIRST
# ----------------------------------------------------
master_src = os.path.join(base_dir, 'localization', 'localization_master.xlsx')
gen_loc_src = os.path.join(base_dir, 'localization', 'generated_localization.json')
pub_data_src = os.path.join(base_dir, 'public', 'data.json')

backup_master = os.path.join(backup_dir, f'localization_master_backup_{timestamp}.xlsx')
shutil.copy2(master_src, backup_master)
print(f"[STEP 1] Created master backup: {backup_master}")

backup_gen_loc = None
if os.path.exists(gen_loc_src):
    backup_gen_loc = os.path.join(backup_dir, f'generated_localization_backup_{timestamp}.json')
    shutil.copy2(gen_loc_src, backup_gen_loc)
    print(f"[STEP 1] Created generated_localization backup: {backup_gen_loc}")

backup_pub_data = None
if os.path.exists(pub_data_src):
    backup_pub_data = os.path.join(backup_dir, f'data_backup_{timestamp}.json')
    shutil.copy2(pub_data_src, backup_pub_data)
    print(f"[STEP 1] Created public data backup: {backup_pub_data}")

# ----------------------------------------------------
# 2. CREATE TEMPORARY RECOVERY WORKBOOK
# ----------------------------------------------------
temp_master = os.path.join(recovery_dir, 'localization_master_batch1_recovered.xlsx')
shutil.copy2(master_src, temp_master)
print(f"[STEP 2] Created temporary recovery workbook: {temp_master}")

# ----------------------------------------------------
# 3. RECONSTRUCT FROM SURVIVING SOURCES
# ----------------------------------------------------
print(f"[STEP 3] Reconstructing Batch 1 data against temporary recovery workbook...")

wb_temp = openpyxl.load_workbook(temp_master)
batch1_chars = {'A0121', 'A0086', 'A0160', 'V0146', 'V0117'}

# Approved Skill Names Mapping
skill_names_map = {
    'A012101': 'Phao Xạ',
    'A012102': 'Phong Địch Thanh Dã',
    'A012103': 'Huyền Lạc Tiễn Minh',
    'A012104': 'Chiến Hữu Tề Tâm',
    'A012105': 'Tốc Nạp',
    'A012111': 'Vọng Sơn',
    'A008601': 'Dương Chi Mỹ Ngọc',
    'A008602': 'Hạ Thiền Huyên Minh',
    'A008603': 'Kim Thiền Thoát Xác',
    'A008604': 'Động Tĩnh Tương Nghi',
    'A008605': 'Ngọc Diệp Chiết Quang',
    'A008611': 'Ngọa Ngọc',
    'A016001': 'Gảy Đàn',
    'A016002': 'Phụng Minh Kỳ Sơn',
    'A016003': 'Khúc Chấn Lương Trần',
    'A016004': 'Vạn Lại Trầm Tịch',
    'A016005': 'Ti Đồng Khánh Thanh',
    'A016011': 'Lạc Hà',
    'V014601': 'Sam Trác',
    'V014602': 'Phù Ẩm Trường Ca',
    'V014603': 'Vạn Vật Hữu Linh',
    'V014604': 'Tế Dĩ Loại Thương',
    'V014605': 'Kiêu Dũng',
    'V014611': 'Tồi Phong Trảm Kỳ',
    'V011701': 'Lẫm Nhiên Khí Độ',
    'V011702': 'Nhân Quả Vãng Phục',
    'V011703': 'Bồ Đề Vô Ngân',
    'V011704': 'Bách Lộc Trình Tường',
    'V011705': 'Phúc Đức Quả Báo',
    'V011711': 'Lộc Hành Phổ Độ'
}

scratch_dir = r"C:\Users\Legion\.gemini\antigravity-ide\brain\8ce6273a-cafe-4c12-80b2-f6d10ae0f369\scratch"
if scratch_dir not in sys.path:
    sys.path.insert(0, scratch_dir)

import fix_all_skills_exact as fix_skills
TRANSLATIONS = fix_skills.TRANSLATIONS

# Naturalization function
def naturalize_vi(text):
    if not text:
        return text
    text = text.replace('sát thương chí mạng', 'Sát Thương Chí Tử')
    text = text.replace('Sát thương chí mạng', 'Sát Thương Chí Tử')
    text = text.replace('Sát Thương Chí Mạng', 'Sát Thương Chí Tử')
    text = text.replace('Trì Hoãn', 'Trì Trệ')
    text = re.sub(r'gây\s+([^\.\,\;\n]+?)\s+thành\s+<color=#ff6724>Sát Thương Chuẩn Bổ Sung</color>', r'Gây <color=#ff6724>Sát Thương Chuẩn Bổ Sung</color> bằng \1', text, flags=re.IGNORECASE)
    text = re.sub(r'gây\s+([^\.\,\;\n]+?)\s+thành\s+Sát Thương Chuẩn Bổ Sung', r'Gây Sát Thương Chuẩn Bổ Sung bằng \1', text, flags=re.IGNORECASE)
    text = re.sub(r'gây\s+([^\.\,\;\n]+?)\s+thành\s+<color=#ff6724>Sát Thương Vật Lý Bổ Sung</color>', r'Gây <color=#ff6724>Sát Thương Vật Lý Bổ Sung</color> bằng \1', text, flags=re.IGNORECASE)
    text = re.sub(r'gây\s+([^\.\,\;\n]+?)\s+thành\s+Sát Thương Vật Lý Bổ Sung', r'Gây Sát Thương Vật Lý Bổ Sung bằng \1', text, flags=re.IGNORECASE)
    text = re.sub(r'phản lại\s+([^\.\,\;\n]+?)\s+thành\s+<color=#ff6724>Sát Thương Vật Lý Bổ Sung</color>', r'Phản lại <color=#ff6724>Sát Thương Vật Lý Bổ Sung</color> bằng \1', text, flags=re.IGNORECASE)
    text = re.sub(r'phản lại\s+([^\.\,\;\n]+?)\s+thành\s+Sát Thương Vật Lý Bổ Sung', r'Phản lại Sát Thương Vật Lý Bổ Sung bằng \1', text, flags=re.IGNORECASE)
    text = re.sub(r'Tấn\s+Công\s+của\s+bản\s+thân', 'Tấn Công', text, flags=re.IGNORECASE)
    text = text.replace('tiến hành 3 lần Đánh Thường truy kích', 'thực hiện 3 lần Đánh Thường truy kích')
    text = text.replace('tiến hành Đánh Thường', 'Đánh Thường')
    text = text.replace('tiến hành tấn công', 'tấn công')
    text = text.replace('tiến hành di chuyển lại', 'di chuyển lại')
    text = text.replace('tiến hành di chuyển', 'di chuyển')
    text = text.replace('tiến vào trạng thái', 'vào trạng thái')
    text = re.sub(r'tiến\s+hành', 'thực hiện', text, flags=re.IGNORECASE)
    text = text.replace('đơn vị kẻ địch đơn thể', '1 kẻ địch được chọn')
    text = text.replace('1 kẻ địch đơn thể được chọn', '1 kẻ địch được chọn')
    text = text.replace('kẻ địch đơn thể được chọn', '1 kẻ địch được chọn')
    text = text.replace('kẻ địch đơn thể', '1 kẻ địch được chọn')
    text = text.replace('đơn vị kẻ địch chịu đòn', 'kẻ địch bị tấn công')
    text = text.replace('đơn vị chịu đòn', 'kẻ địch bị tấn công')
    text = text.replace('kẻ địch chịu đòn', 'kẻ địch bị tấn công')
    text = text.replace('đơn vị kẻ địch', 'kẻ địch')
    text = text.replace('đơn vị đồng minh', 'đồng minh')
    text = re.sub(r'sau\s+khi\s+vào\s+trận\s+đấu', 'khi vào trận', text, flags=re.IGNORECASE)
    text = text.replace('tối đa có thể cộng dồn', 'có thể cộng dồn tối đa')
    text = text.replace('không ít hơn', 'có ít nhất')
    text = text.replace('Sức Di Chuyển của lần hành động lại đó', 'Sức Di Chuyển của lượt hành động lại này')
    text = text.replace('Gây Gây', 'Gây')
    text = text.replace('Phản lại Phản lại', 'Phản lại')
    return text

# Complete 31 Batch 1 Referenced Buff Map (Name VI, Desc VI)
buff_updates = {
    # Character-specific buffs
    'Buff_A0086_5': ('Giảm Trúng', 'Giảm Tỷ Lệ Trúng <color=#158bdb>[EffectParam,3]%</color>, kéo dài <color=#158bdb>1</color> lượt.'),
    'Buff_A0086_15': ('Thoát Xác', 'Tỷ Lệ Né Tránh tăng <color=#158bdb>[EffectParam,2]%</color>, có thể cộng dồn <color=#158bdb>2</color> tầng, sau khi né tránh thành công sẽ xóa hiệu ứng.'),
    'Buff_A0086_19': ('Kim Ngọc', 'Tỷ Lệ Bạo Kích của Đánh Trả tăng <color=#158bdb>[EffectParam,2]%</color>, kéo dài <color=#158bdb>1</color> lượt.'),
    'Buff_A0086_24': ('Ẩn Nấp', 'Không thể bị chọn làm mục tiêu chính, nhưng sát thương gánh chịu tăng <color=#158bdb>25%</color>, kéo dài <color=#158bdb>1</color> lượt.'),
    'Buff_A0121_1': ('Phong Địch Thanh Dã', 'Khi ở trạng thái này, sát thương bản thân gánh chịu giảm <color=#158bdb>[EffectParam,2]%</color>, không thể di chuyển và miễn nhiễm hiệu ứng dịch chuyển do kẻ địch gây ra;\nPhạm vi tấn công của Đánh Thường chuyển thành <color=#158bdb>xung quanh 3-7 ô</color>, đồng thời gây Sát Thương Vật Lý bằng <color=#158bdb>100%</color> Tấn Công lên 1 kẻ địch hoặc ô đất được chọn cùng toàn bộ kẻ địch xung quanh <color=#158bdb>1 vòng</color>;\nTrong thời gian trạng thái này kéo dài, bản thân không thể vào trạng thái <color=#ff6724>Nhắm Bắn</color>.'),
    'Buff_A0121_3': ('Nhắm Bắn Truy Kích', 'Khi Đánh Thường trong trạng thái <color=#ff6724>Nhắm Bắn</color>, thực hiện thêm <color=#158bdb>3</color> lần Đánh Thường truy kích lên kẻ địch được chọn, mỗi lần gây Sát Thương Vật Lý bằng <color=#158bdb>[Effect1Para,1]%</color> Tấn Công; ở trạng thái này không thể di chuyển và không thể dùng Kỹ Năng Nghề.'),
    'Buff_A0121_8': ('Trường Nỏ', 'Khi ở trạng thái này, sát thương Đánh Thường tiếp theo tăng <color=#158bdb>[EffectParam,2]%</color>, có thể cộng dồn tối đa <color=#158bdb>1</color> tầng.'),
    'Buff_A0121_11': ('Tầm Đánh Mở Rộng', 'Tầm đánh tối đa của Đánh Thường tăng <color=#158bdb>2</color> ô và Đánh Thường không thể bị đỡ đòn; ở trạng thái <color=#ff6724>Nhắm Bắn</color> không thể di chuyển.'),
    'Buff_A0160_2': ('Phiếm Âm', 'Khi chịu tấn công Tuyệt Kỹ của Thái Phụng Minh Kỳ bên địch, sẽ gây hiệu ứng dựa trên số tầng cộng dồn, có thể cộng dồn <color=#158bdb>3</color> tầng, kéo dài <color=#158bdb>2</color> lượt.'),
    'Buff_A0160_10': ('Tản Âm', 'Khi chịu <color=#158bdb>1</color> lần sát thương từ Đánh Thường hoặc Kỹ Năng, Gây <color=#ff6724>Sát Thương Vật Lý Bổ Sung</color> bằng <color=#158bdb>[EffectParam,2]%</color> Tấn Công của Thái Phụng Minh Kỳ <color=#158bdb>1</color> lần lên bản thân và toàn bộ đồng minh <color=#158bdb>xung quanh 3 ô</color>, sau khi kích hoạt hiệu ứng trên, trạng thái này sẽ bị xóa.'),
    'Buff_A0160_hz_1': ('Bộ Giảm Chấn', 'Trạng thái này kéo dài <color=#158bdb>1</color> vòng lượt;\nKhi ở trạng thái này, trong lượt hành động của bất kỳ đơn vị nào, sau mỗi <color=#158bdb>3</color> lần bản thân chịu tấn công trực tiếp, sát thương gánh chịu sau đó tăng <color=#158bdb>50%</color>, kéo dài đến khi bắt đầu <color=#158bdb>1</color> lượt hành động tiếp theo của bất kỳ đơn vị nào.'),
    'Buff_V0117_2_1': ('Suy Nhược', 'Tỷ Lệ Bạo Kích giảm <color=#158bdb>10%</color>, kéo dài <color=#158bdb>2</color> lượt.'),
    'Buff_V0117_3': ('Minh Kính', 'Khi chịu tấn công trực tiếp từ kẻ địch, nhận <color=#158bdb>1</color> tầng trạng thái này.'),
    'Buff_V0117_hz_1_1': ('Trình Sửa Lưu Trữ', 'Trạng thái này có thể cộng dồn tối đa <color=#158bdb>3</color> tầng;\nKhi chịu tấn công Tuyệt Kỹ hoặc Liên Kích của Lộc Vương Bản Sinh Đồ bên địch, mỗi tầng trạng thái này khiến bản thân chịu Sát Thương Vật Lý Bổ Sung bằng <color=#158bdb>220%</color> Tấn Công của Lộc Vương Bản Sinh Đồ, sát thương này chắc chắn Bạo Kích, sau khi chịu tấn công từ Lộc Vương Bản Sinh Đồ bên địch sẽ xóa toàn bộ trạng thái này.'),
    'Buff_V0146_1': ('Hàm Ẩm', 'Khi ở trạng thái này, nếu bất kỳ đồng minh nào kích hoạt Liên Kích, bản thân nhận <color=#158bdb>1</color> lượt hành động bổ sung sau khi đồng minh đó kết thúc hành động và làm mới thời gian kéo dài của trạng thái <color=#ff6724>Khiếu Kiếm</color> và <color=#ff6724>Tích Thế</color>; Trong lượt hành động bổ sung đó, sát thương Đánh Thường giảm <color=#158bdb>99%</color>.'),
    'Buff_V0146_13': ('Trối Hợp', 'Sát Thương Liên Kích tăng <color=#158bdb>[EffectParam,2]%</color>, kéo dài <color=#158bdb>2</color> lượt, có thể cộng dồn <color=#158bdb>5</color> tầng.'),
    'Buff_V0146_15': ('Dũng Nghị', 'Sát Thương Liên Kích tăng <color=#158bdb>[EffectParam,2]%</color>, kéo dài <color=#158bdb>1</color> lượt, có thể cộng dồn <color=#158bdb>5</color> tầng.'),
    'Buff_V0146_4_1': ('Chúng Sinh Nhất Tướng', 'Trong <color=#158bdb>2</color> vòng lượt, bản thân sau khi Đánh Thường sẽ kích hoạt Liên Kích, Gây Sát Thương Vật Lý bằng <color=#158bdb>[EffectParam,2]%</color> Tấn Công lên kẻ địch bị tấn công.'),

    # Generic / Shared buffs
    'Buff_AllDmgIncrease': ('Tích Thế', 'Tăng tất cả sát thương <color=#158bdb>20%</color>, kéo dài <color=#158bdb>[EffectParam,2]</color> lượt, có thể cộng dồn <color=#158bdb>2</color> tầng.'),
    'Buff_Atk_Up': ('Khiếu Kiếm', 'Tấn Công tăng <color=#158bdb>20%</color>, kéo dài <color=#158bdb>[EffectParam,2]</color> lượt, có thể cộng dồn <color=#158bdb>2</color> tầng.'),
    'Buff_Burning': ('Thiêu Đốt', 'Khi bắt đầu lượt, chịu Sát Thương Cấu Thuật bằng <color=#158bdb>20%</color> Tấn Công của người thi triển, có thể cộng dồn <color=#158bdb>5</color> tầng, kéo dài <color=#158bdb>2</color> lượt.'),
    'Buff_CannotMove': ('Cấm Túc', 'Không thể di chuyển.'),
    'Buff_Cold': ('Hàn Thiên', 'Sức Di Chuyển giảm <color=#158bdb>1</color>. Trạng thái này khi cộng dồn đến <color=#158bdb>3</color> tầng, xóa toàn bộ trạng thái <color=#ff6724>Hàn Thiên</color> của bản thân, nhận <color=#158bdb>1</color> tầng trạng thái <color=#ff6724>Sương Giá</color>.'),
    'Buff_Fragile': ('Dễ Vỡ', 'Tăng tất cả sát thương gánh chịu <color=#db2424>30%</color>, mỗi lần chịu 1 đòn tấn công trực tiếp giảm <color=#ff6724>1 tầng</color>, kéo dài <color=#4395d6>2</color> lượt hoặc khi hết tầng.'),
    'Buff_MagicDmgReduce': ('Loạn Thần', 'Khi ở trạng thái này, Sát Thương Cấu Thuật gây ra giảm <color=#158bdb>[EffectParam,3]%</color>, kéo dài <color=#158bdb>[EffectParam,2]</color> lượt.'),
    'Buff_Mov_Down': ('Trì Trệ', 'Giảm Sức Di Chuyển <color=#158bdb>[EffectParam,3]</color>, kéo dài <color=#158bdb>[EffectParam,2]</color> lượt, có thể cộng dồn <color=#158bdb>1</color> tầng.'),
    'Buff_RealDamge_Lan_Special': ('Sát Thương Chuẩn Bổ Sung', 'Sát thương này không kích hoạt hiệu ứng Bạo Kích và Xuyên Thấu, không thể bị Đỡ, bỏ qua Phòng Thủ và hiệu ứng giảm sát thương, đồng thời không chịu ảnh hưởng bởi hiệu ứng tăng sát thương.'),
    'Buff_Reckless': ('Lỗ Mãng', 'Sát thương Đánh Trả gánh chịu tăng <color=#158bdb>[EffectParam,3]%</color>, kéo dài <color=#158bdb>[EffectParam,2]</color> lượt, có thể cộng dồn <color=#158bdb>1</color> tầng.'),
    'Buff_Snipe_Lan': ('Nhắm Bắn', 'Không thể di chuyển.'),
    'Buff_SpecialDamgeP_Lan': ('Sát Thương Vật Lý Bổ Sung', 'Sát thương này không kích hoạt hiệu ứng Bạo Kích và Xuyên Thấu, đồng thời không thể bị Đỡ.'),
    'Buff_Taunt': ('Chấn Nộ', 'Chỉ có thể Đánh Thường đơn vị khiến bản thân nhận trạng thái này, kéo dài <color=#158bdb>[EffectParam,2]</color> lượt.')
}

referenced_buff_ids = set(buff_updates.keys())

v01464_story = (
    '"Trâu nhỏ ơi Trâu nhỏ. Tiếng nước reo nhẹ nhàng bên suối gọi cô bé, uống ngụm nước rồi nghỉ chân đi nào."\n'
    '"Trâu nhỏ à Trâu nhỏ. Lúa dại, bắp ngô uốn eo trước gió khóc lóc than thở: Mau tới giúp chúng tôi với, bị sâu cắn mất rồi!"\n'
    '"Đừng sợ! Cậu ấy là ai chứ? Cậu ấy chính là Á Trưởng Ngưu Tôn đấm đá sâu bọ, uy phong lừng lẫy khắp ruộng đồng đây. Lũ yêu quái — còn không mau đầu hàng —"\n'
    '"Trâu nhỏ Trâu nhỏ, mau về nhà ăn cơm thôi!"\n'
    '"Bóng dáng vác cuốc diệt sâu đang tạo dáng dưới ánh chiều tà bỗng khựng lại, giây tiếp theo liền cắm đầu chạy thục mạng: \\"Cháu về ngay đây!\\""\n\n'
    '"Chị Trâu ơi. Trâu thì ăn gì ạ? Ăn cỏ chứ sao."\n'
    '"Thế cỏ thì ăn gì? Ăn ánh nắng với giọt sương chứ sao."\n'
    '"Thế ánh nắng với giọt sương thì... Ôi! Đừng đánh em! Em tự nghĩ là được chứ gì... Em biết rồi! Ánh nắng với giọt sương ăn chính là tụi mình đó."\n\n'
    '"Lũ trẻ trong làng lúc nào cũng tràn trề năng lượng cùng muôn vàn câu hỏi không bao giờ dứt, phiền phức chết đi được. Tụi nó đinh ninh rằng ánh nắng chiếu trên tấm lưng lăn lộn, hạt mưa tưới lên bàn chân tung tăng chính là bị thiên nhiên nuốt chửng trọn vẹn một lượt."\n\n'
    '"Chị Trâu chị Trâu. Đừng có cười hoài, rốt cuộc có đúng không hả?"\n'
    '"Cái má nhỏ nhắn phải bị ông mặt trời cắn hết lần này đến lần khác mới ửng đỏ lên, đôi chân thần tốc chạy đua với chú chó nhỏ phải được trận mưa rào tưới tắm bao bận mới cứng cáp; cái đầu nhỏ của con người ngoài trồng cấy ra còn nghĩ ra biết bao điều kỳ diệu nữa cơ đấy. Ngưu Ngưu thích lắm. Nhưng Ngưu Ngưu không nói đâu."\n\n'
    '"Bốp —"\n'
    '"Á! Sao lại đánh em nữa!"\n\n'
    '"Đừng có làm phiền chị, mời em lên xe phun nước hiệu Trâu Nhỏ."'
)

v01174_story = (
    'Sau khi Lộc Vương Bản Sinh Đồ bước vào thế giới ấy, câu chuyện liền phát triển đúng như những gì cậu hằng tưởng tượng. Muôn loài tự do tự tại, an nhàn sinh sống trong mảnh thiên địa này. Con người tuy có tranh chấp, nhưng thiện ác đều có báo ứng, mọi thứ diễn ra trật tự đâu vào đấy.\n'
    'Thế nhưng khi những người khác bước chân vào thế giới này, dựa theo từng lựa chọn để viết nên một kết cục khác cho thế giới, mọi chuyện lại chẳng còn như vậy nữa. Họ ít nhiều đều gặp phải khó khăn, chịu đựng thử thách. Có người bảo, chọn sai một bước liền rơi vào vạn kiếp bất phục; cũng có người bảo, thế giới nhỏ bé ấy được thiết kế quá đỗi tinh xảo, khiến người ta chẳng biết phải làm sao.\n'
    'Giờ đây, bạn cùng thiếu niên mang vầng sáng Cửu Sắc Lộc quanh mình tiến vào thế giới nhỏ bé này. Bạn cùng cậu ấy thăm dò đường đi phía trước, trao đổi góc nhìn, trò chuyện vô cùng tâm đắc. Chẳng biết từ lúc nào, hai người đã cùng nhau đi tới tận cùng thế giới...\n'
    'Chúc mừng người chơi đạt thành HE: Thế giới tuyệt đẹp đồng hành cùng Chú Lộc Nhỏ\n'
    'Khi được hỏi về cảm tưởng sau khi phá đảo, bạn trả lời: Thật ra, chỉ cần bạn giữ vững tâm thế giống như Lộc Vương Bản Sinh Đồ, trò chơi này cũng chẳng khó đến thế đâu? Chú hươu chín màu bị thu hút bởi lòng thiện lương kia sẽ dẫn dắt những ai cùng sở hữu tấm lòng nhân hậu cùng bước tới nhạc viện đầm ấm năm xưa.'
)

a01604_story = (
    '"Bác ơi, có tiếng đàn nào vang vọng mãi giữa đất trời, chẳng bị điều gì làm cho dứt đoạn không ạ?"\n'
    'Bên ngoài cửa sổ gỗ, hoa quế âm thầm rụng rơi. Thái Phụng Minh Kỳ tựa nghiêng bên cột hành lang, lắng nghe cuộc trò chuyện giữa một già một trẻ. Chàng mỉm cười ngoái nhìn, cũng lặng lẽ đợi chờ câu trả lời.\n'
    '"Cây đàn này do 雷威 (Lôi Uy) thời Khai Nguyên đại Đường chế tác. Người thời ấy ư, sớm đã rụng rơi thành cát bụi rồi. \'Từng giữ lời người xưa phủ Định, ta ôm đàn này thốt ba tiếng thở dài\'... Rốt cuộc cũng là khúc dứt người xa. Hễ dính dáng đến một chữ \'người\', tất sẽ có lúc đường cùng. Bởi bản thân con người vốn đã có điểm dừng."\n'
    'Thái Phụng Minh Kỳ thầm thở dài. Chàng là cây đàn, những chuyện chàng ghi nhớ nhiều hơn ông lão kia rất nhiều.\n'
    'Những nhà sưu tầm, những nghệ nhân gảy đàn, những người áo vải đi ngang chỉ để nghe chàng cất một tiếng đàn ngân vang.\n'
    'Từ người xa lạ đến chốn tri âm, một khúc hòa vang, gảy mãi tới lúc hừng đông rạng rỡ.\n'
    'Có người gửi gắm tâm sự vào dây đàn, truyền qua mảng tường, bên kia liền cất tiếng họa âm đáp lại. Điều quân tử trao tặng không phải là bản nhạc, mà là một cuộc tương phùng chẳng nỡ buông tay.\n'
    'Chàng từng chứng kiến sự trân trọng sâu sắc nhất: hai người đối mặt không nói một lời, chỉ đẩy qua đẩy lại cây đàn — bạn gảy, rồi tôi lại gảy, như thể ai dừng trước, người đó sẽ phải nói lời từ biệt trước. Nhưng con người rốt cuộc vẫn phải chia ly. Người ta đem những lời nói không cùng, những lá thư chẳng bao giờ tới, gom hết vào tiếng đàn, giao cho chàng cất giữ. Chàng bèn thay họ cất tiếng ngân dài, hết tiếng này đến tiếng khác, cho đến khi đối phương tận hưởng trọn vẹn rồi trở về.\n'
    'Trong mơ, chàng du ngoạn đến Tây Hồ, được mời tới làm khách trong một căn phòng nhỏ sáng bừng.\n'
    'Chiếc hộp nhỏ mở ra, ẩn chứa một bài hát, được những răng lược kim loại gảy đi gảy lại, một lần, rồi lại một lần.\n'
    'Sẽ không bao giờ dừng lại nữa.\n'
    'Thái Phụng Minh Kỳ lắng nghe âm thanh không bao giờ ngưng nghỉ ấy, chợt nhớ về quá khứ xa xăm, chàng từng dưới hành lang ngắm nhìn hoa quế rụng rào rạt, chim trời lướt qua góc mái, mưa đêm tí tách đến tận bình minh. Cưỡi gió mà đi, tận hưởng mà về. Con người, cây đàn, hoa chim, mưa đêm, tất cả đều như vậy.\n'
    'Hãy bay đi, mang theo ta lướt qua bầu trời ngàn năm.\n'
    'Đáp xuống rồi, ta lại đáp từ. Qua lại trao nhau, tiếng đàn bèn ngân vang không bao giờ dứt.'
)

# Apply to SKILL sheet of temp_master
sk_sheet = wb_temp['SKILL']
sk_headers = [cell.value for cell in sk_sheet[1]]
id_idx = sk_headers.index('skill_id')
gid_idx = sk_headers.index('skill_group_id')
cid_idx = sk_headers.index('character_id')
name_vi_idx = sk_headers.index('skill_name_vi')
desc_cn_idx = sk_headers.index('desc_cn')
desc_vi_idx = sk_headers.index('desc_vi')
status_idx = sk_headers.index('status')
notes_idx = sk_headers.index('notes')

skill_count = 0
for row in sk_sheet.iter_rows(min_row=2):
    cid = str(row[cid_idx].value or '')
    if cid in batch1_chars:
        skill_count += 1
        gid = str(row[gid_idx].value or '')
        sid = str(row[id_idx].value or '')
        desc_cn = str(row[desc_cn_idx].value or '')
        
        if gid in skill_names_map:
            row[name_vi_idx].value = skill_names_map[gid]
            
        vi_val = TRANSLATIONS.get(desc_cn)
        if vi_val:
            row[desc_vi_idx].value = naturalize_vi(vi_val)
            
        row[status_idx].value = 'TRANSLATED'
        row[notes_idx].value = 'Skill name owner-approved; description naturalized & translated'

print(f"[STEP 3] Reconstructed {skill_count} SKILL level rows in temporary workbook.")

# Apply to BUFF_STATUS sheet of temp_master
bf_sheet = wb_temp['BUFF_STATUS']
bf_headers = [cell.value for cell in bf_sheet[1]]
bf_id_idx = bf_headers.index('buff_id')
bf_name_vi_idx = bf_headers.index('buff_name_vi')
bf_desc_vi_idx = bf_headers.index('buff_desc_vi')
bf_status_idx = bf_headers.index('status')
bf_notes_idx = bf_headers.index('notes')

buff_count = 0
for row in bf_sheet.iter_rows(min_row=2):
    rid = str(row[bf_id_idx].value or '')
    if rid in referenced_buff_ids:
        buff_count += 1
        name_vi, desc_vi = buff_updates[rid]
        row[bf_name_vi_idx].value = name_vi
        row[bf_desc_vi_idx].value = naturalize_vi(desc_vi)
        row[bf_status_idx].value = 'TRANSLATED'
        row[bf_notes_idx].value = 'Recovered and naturalized for Batch 1'

print(f"[STEP 3] Reconstructed {buff_count} BUFF_STATUS records in temporary workbook.")

# Apply to HUANZHANG sheet of temp_master
hz_sheet = wb_temp['HUANZHANG']
hz_headers = [cell.value for cell in hz_sheet[1]]
hz_id_idx = hz_headers.index('brilliant_id')
hz_info_vi_idx = hz_headers.index('icon_info_vi')
hz_status_idx = hz_headers.index('status')
hz_notes_idx = hz_headers.index('notes')

for row in hz_sheet.iter_rows(min_row=2):
    rid = str(row[hz_id_idx].value or '')
    if rid == 'V01464':
        row[hz_info_vi_idx].value = v01464_story
        row[hz_status_idx].value = 'TRANSLATED'
        row[hz_notes_idx].value = 'Full lore translated from CN source'
    elif rid == 'V01174':
        row[hz_info_vi_idx].value = v01174_story
        row[hz_status_idx].value = 'TRANSLATED'
        row[hz_notes_idx].value = 'Full lore translated from CN source'
    elif rid == 'A01604':
        row[hz_info_vi_idx].value = a01604_story
        row[hz_status_idx].value = 'TRANSLATED'
        row[hz_notes_idx].value = 'Literary prose naturalized from CN source'

print(f"[STEP 3] Reconstructed HUANZHANG lore stories in temporary workbook.")

# Apply to ZHIZHI sheet of temp_master
zh_sheet = wb_temp['ZHIZHI']
zh_headers = [cell.value for cell in zh_sheet[1]]
zh_cid_idx = zh_headers.index('character_id')
zh_star_idx = zh_headers.index('star')
zh_vi_idx = zh_headers.index('effect_summary_vi')
zh_cn_idx = zh_headers.index('effect_summary_cn')
zh_base_idx = zh_headers.index('skill_up_base_id')
zh_type_idx = zh_headers.index('effect_type')
zh_status_idx = zh_headers.index('status')

SLOT_LABELS = {'01': 'Đánh Thường', '11': 'Kỹ Năng Nghề', '02': 'Tuyệt Kỹ', '03': 'Nội Tại 1', '04': 'Nội Tại 2', '05': 'Nội Tại 3'}
STAT_NAME_MAP = {
    'Atk': 'Tấn Công', 'PhysicDef': 'Phòng Thủ Vật Lý', 'Hp': 'Máu', 'MagicDef': 'Phòng Thủ Cấu Thuật',
    'Critical': 'Tỷ Lệ Bạo Kích', 'CritDmg': 'Sát Thương Bạo Kích', 'AllDmgIncrease': 'Tăng Tất Cả Sát Thương',
    'PhysicalDmgIncrease': 'Tăng Sát Thương Vật Lý', 'ConstructDmgIncrease': 'Tăng Sát Thương Cấu Thuật',
    'HurtHealRate': 'Tỷ Lệ Hút Máu', 'Speed': 'Tốc Độ', 'Mov': 'Sức Di Chuyển'
}

def resolve_zhizhi_stat_vi(raw_stat_str):
    if not raw_stat_str or raw_stat_str == 'None':
        return ''
    parts = [p.strip() for p in raw_stat_str.split(';') if p.strip()]
    res_parts = []
    for p in parts:
        if '+' in p: k, v = p.split('+', 1)
        elif '-' in p: k, v = p.split('-', 1); v = f"-{v}"
        else: res_parts.append(p); continue
        sign = '+' if not v.startswith('-') else ''
        if k.endswith('_PERCENT'):
            stat_key = k.replace('_PERCENT', '')
            res_parts.append(f"{STAT_NAME_MAP.get(stat_key, stat_key)} {sign}{v}%")
        elif k.endswith('_FIX'):
            stat_key = k.replace('_FIX', '')
            if stat_key in ['Critical', 'CritDmg', 'AllDmgIncrease', 'PhysicalDmgIncrease', 'ConstructDmgIncrease', 'HurtHealRate']:
                res_parts.append(f"{STAT_NAME_MAP.get(stat_key, stat_key)} {sign}{v}%")
            else:
                res_parts.append(f"{STAT_NAME_MAP.get(stat_key, stat_key)} {sign}{v}")
        else:
            res_parts.append(f"{STAT_NAME_MAP.get(k, k)} {sign}{v}")
    return '; '.join(res_parts)

zhizhi_skill_names = {
    'A008603': 'Kim Thiền Thoát Xác',
    'A012103': 'Huyền Lạc Tiễn Minh',
    'A016004': 'Vạn Lại Trầm Tịch',
    'V011704': 'Bách Lộc Trình Tường',
    'V014605': 'Kiêu Dũng'
}

zhizhi_count = 0
for row in zh_sheet.iter_rows(min_row=2):
    cid = str(row[zh_cid_idx].value or '')
    star = row[zh_star_idx].value
    if cid in batch1_chars:
        zhizhi_count += 1
        effect_type = str(row[zh_type_idx].value or '')
        sum_cn = str(row[zh_cn_idx].value or '')
        base_id = str(row[zh_base_idx].value or '')
        
        if cid == 'V0117' and star == 6:
            row[zh_vi_idx].value = 'Tỷ Lệ Hút Máu +5%'
        elif effect_type == 'SkillUP' or (base_id and base_id != 'None'):
            slot_code = base_id[5:7] if len(base_id) >= 7 else ''
            slot_label = SLOT_LABELS.get(slot_code, 'Kỹ Năng')
            base_skill_vi = zhizhi_skill_names.get(base_id, skill_names_map.get(base_id, ''))
            row[zh_vi_idx].value = f"Cường hóa {slot_label}: {base_skill_vi}"
        else:
            row[zh_vi_idx].value = resolve_zhizhi_stat_vi(sum_cn)
            
        row[zh_status_idx].value = 'TRANSLATED'

print(f"[STEP 3] Reconstructed {zhizhi_count} ZHIZHI rank summaries in temporary workbook.")

wb_temp.save(temp_master)
print(f"[STEP 3 SUCCESS] Temporary recovery workbook saved: {temp_master}")

# ----------------------------------------------------
# 4. VERIFY TEMPORARY WORKBOOK COVERAGE
# ----------------------------------------------------
print(f"[STEP 4] Verifying temporary recovery workbook coverage...")
wb_temp_check = openpyxl.load_workbook(temp_master, data_only=True)

# SKILL check (90 level rows)
sk_unpop = 0
for r in wb_temp_check['SKILL'].iter_rows(min_row=2, values_only=True):
    if r[1] in batch1_chars:
        vi_val = str(r[6] or '').strip()
        if not vi_val or vi_val == 'None':
            sk_unpop += 1
            print(f"  [UNPOPULATED SKILL] Row {r[0]} ({r[1]}): desc_vi is empty!")

# BUFF check (31 referenced buff IDs)
bf_unpop = 0
for r in wb_temp_check['BUFF_STATUS'].iter_rows(min_row=2, values_only=True):
    bid = str(r[0] or '')
    if bid in referenced_buff_ids:
        vi_name = str(r[3] or '').strip()
        vi_desc = str(r[5] or '').strip()
        if not vi_name or vi_name == 'None' or not vi_desc or vi_desc == 'None':
            bf_unpop += 1
            print(f"  [UNPOPULATED BUFF] Row {bid}: name='{vi_name}', desc='{vi_desc}'")

# ZHIZHI check (30 rank rows)
zh_unpop = 0
for r in wb_temp_check['ZHIZHI'].iter_rows(min_row=2, values_only=True):
    if r[0] in batch1_chars:
        vi_summary = str(r[6] or '').strip()  # index 6 is effect_summary_vi
        if not vi_summary or vi_summary == 'None':
            zh_unpop += 1
            print(f"  [UNPOPULATED ZHIZHI] Row {r[0]} Rank {r[1]}: effect_summary_vi is empty!")

# HUANZHANG check (3 stories)
hz_unpop = 0
for r in wb_temp_check['HUANZHANG'].iter_rows(min_row=2, values_only=True):
    if r[0] in {'A01604', 'V01464', 'V01174'}:
        vi_info = str(r[5] or '').strip()
        if not vi_info or vi_info == 'None':
            hz_unpop += 1
            print(f"  [UNPOPULATED HUANZHANG] Row {r[0]}: icon_info_vi is empty!")

print(f"  SKILL unpopulated: {sk_unpop}")
print(f"  BUFF_STATUS unpopulated: {bf_unpop}")
print(f"  ZHIZHI unpopulated: {zh_unpop}")
print(f"  HUANZHANG unpopulated: {hz_unpop}")

if sk_unpop > 0 or bf_unpop > 0 or zh_unpop > 0 or hz_unpop > 0:
    print("[ERROR] Temporary workbook verification failed! Some fields are None or empty.")
    sys.exit(1)

print("[STEP 4 SUCCESS] Temporary workbook coverage is 100% complete!")

# ----------------------------------------------------
# 5. VALIDATE TEMPORARY RECOVERY (Placeholders, Han, Suspicious phrases)
# ----------------------------------------------------
print(f"[STEP 5] Validating temporary recovery workbook...")

def validate_temp_placeholders(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    errors = 0
    audited = 0
    for sheetname in ['SKILL', 'BUFF_STATUS', 'ZHIZHI', 'HUANZHANG']:
        sheet = wb[sheetname]
        headers = [cell.value for cell in sheet[1]]
        cn_col = headers.index('desc_cn' if sheetname == 'SKILL' else ('buff_desc_cn' if sheetname == 'BUFF_STATUS' else ('icon_info_cn' if sheetname == 'HUANZHANG' else 'effect_summary_cn')))
        vi_col = headers.index('desc_vi' if sheetname == 'SKILL' else ('buff_desc_vi' if sheetname == 'BUFF_STATUS' else ('icon_info_vi' if sheetname == 'HUANZHANG' else 'effect_summary_vi')))
        cid_col = headers.index('character_id') if 'character_id' in headers else 0
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            cid = str(row[cid_col] or '')
            rid = str(row[0] or '')
            if (sheetname == 'BUFF_STATUS' and rid in referenced_buff_ids) or any(c in rid or c in cid for c in batch1_chars):
                cn_txt = str(row[cn_col] or '')
                vi_txt = str(row[vi_col] or '')
                if vi_txt and vi_txt != 'None':
                    audited += 1
                    cn_placeholders = re.findall(r'\[Effect[0-9A-Za-z\_,]+\]', cn_txt)
                    vi_placeholders = re.findall(r'\[Effect[0-9A-Za-z\_,]+\]', vi_txt)
                    if cn_placeholders != vi_placeholders:
                        print(f"  [ERROR] Placeholder mismatch in {sheetname} {rid}: CN={cn_placeholders}, VI={vi_placeholders}")
                        errors += 1
    return audited, errors

audited, ph_errors = validate_temp_placeholders(temp_master)
print(f"  Placeholder validation: {audited} fields audited | Errors: {ph_errors}")

def validate_temp_han(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    cjk_regex = re.compile(r'[\u4e00-\u9fff]')
    whitelist = ['《鹿王本生图》', '雷威']
    errors = 0
    audited = 0
    for sheetname in ['SKILL', 'BUFF_STATUS', 'ZHIZHI', 'HUANZHANG']:
        sheet = wb[sheetname]
        headers = [cell.value for cell in sheet[1]]
        vi_col = headers.index('desc_vi' if sheetname == 'SKILL' else ('buff_desc_vi' if sheetname == 'BUFF_STATUS' else ('icon_info_vi' if sheetname == 'HUANZHANG' else 'effect_summary_vi')))
        cid_col = headers.index('character_id') if 'character_id' in headers else 0
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            cid = str(row[cid_col] or '')
            rid = str(row[0] or '')
            if (sheetname == 'BUFF_STATUS' and rid in referenced_buff_ids) or any(c in rid or c in cid for c in batch1_chars):
                vi_txt = str(row[vi_col] or '')
                if vi_txt and vi_txt != 'None':
                    audited += 1
                    test_txt = vi_txt
                    for w in whitelist:
                        test_txt = test_txt.replace(w, '')
                    matches = cjk_regex.findall(test_txt)
                    if matches:
                        print(f"  [ERROR] Han leak in {sheetname} {rid}: {matches}")
                        errors += 1
    return audited, errors

audited_han, han_errors = validate_temp_han(temp_master)
print(f"  Han-character validation: {audited_han} fields audited | Errors: {han_errors}")

suspicious = [
    'tiến hành', 'đơn vị chịu đòn', 'kẻ địch chịu đòn', 'đơn thể',
    'thành Sát Thương', 'không ít hơn', 'tối đa có thể', 'Tấn Công của bản thân',
    'sát thương chí mạng', 'vị ông lão', 'Trì Hoãn'
]
susp_errors = 0
for sheetname in ['SKILL', 'BUFF_STATUS', 'ZHIZHI', 'HUANZHANG']:
    sheet = wb_temp_check[sheetname]
    headers = [cell.value for cell in sheet[1]]
    vi_col = headers.index('desc_vi' if sheetname == 'SKILL' else ('buff_desc_vi' if sheetname == 'BUFF_STATUS' else ('icon_info_vi' if sheetname == 'HUANZHANG' else 'effect_summary_vi')))
    cid_col = headers.index('character_id') if 'character_id' in headers else 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        cid = str(row[cid_col] or '')
        rid = str(row[0] or '')
        if (sheetname == 'BUFF_STATUS' and rid in referenced_buff_ids) or any(c in rid or c in cid for c in batch1_chars):
            vi_txt = str(row[vi_col] or '')
            for s in suspicious:
                if s in vi_txt:
                    print(f"  [SUSPICIOUS PHRASE ERROR] Found '{s}' in {sheetname} {rid}")
                    susp_errors += 1

print(f"  Suspicious phrase self-check errors: {susp_errors}")

if ph_errors > 0 or han_errors > 0 or susp_errors > 0:
    print("[ERROR] Temporary workbook validation failed!")
    sys.exit(1)

print("[STEP 5 SUCCESS] Temporary workbook validation passed 100% cleanly!")

# ----------------------------------------------------
# 6. GENERATE RECOVERY REVIEW FROM TEMP WORKBOOK
# ----------------------------------------------------
temp_review_md = os.path.join(recovery_dir, 'batch1_recovered_review.md')

with open(temp_review_md, 'w', encoding='utf-8') as f:
    f.write(f"# PHASE 3 — BATCH 1 LOCALIZATION RECOVERY REVIEW PACKAGE\n\n")
    f.write(f"Generated directly from temporary recovered workbook: `localization_master_batch1_recovered.xlsx`  \n")
    f.write(f"Scope: A0121, A0086, A0160, V0146, V0117\n\n")
    f.write(f"---\n\n")

    for cid in ['A0121', 'A0086', 'A0160', 'V0146', 'V0117']:
        f.write(f"## Character {cid}\n\n")
        f.write(f"### SKILL\n\n")
        for r in wb_temp_check['SKILL'].iter_rows(min_row=2, values_only=True):
            if r[1] == cid:
                f.write(f"- **ID**: `{r[0]}` | **Name**: {r[4]} ({r[3]})\n")
                f.write(f"  - **CN**: {r[5]}\n")
                f.write(f"  - **VI**: {r[6]}\n\n")

        f.write(f"### BUFF_STATUS\n\n")
        for r in wb_temp_check['BUFF_STATUS'].iter_rows(min_row=2, values_only=True):
            rid = str(r[0] or '')
            if rid in referenced_buff_ids and (cid in rid or r[1] == cid):
                f.write(f"- **ID**: `{r[0]}` | **Name**: {r[3]} ({r[2]})\n")
                f.write(f"  - **CN**: {r[4]}\n")
                f.write(f"  - **VI**: {r[5]}\n\n")

        f.write(f"### ZHIZHI\n\n")
        for r in wb_temp_check['ZHIZHI'].iter_rows(min_row=2, values_only=True):
            if r[0] == cid:
                f.write(f"- **Rank {r[1]}**: {r[3]} | **VI**: {r[6]}\n")

        f.write(f"### HUANZHANG\n\n")
        for r in wb_temp_check['HUANZHANG'].iter_rows(min_row=2, values_only=True):
            if r[1] == cid:
                f.write(f"- **ID**: `{r[0]}` | **Name**: {r[3]} ({r[2]})\n")
                f.write(f"  - **Info CN**: {r[4]}\n")
                f.write(f"  - **Info VI**: {r[5]}\n\n")

print(f"[STEP 6 SUCCESS] Generated temporary recovery review Markdown: {temp_review_md}")

# ----------------------------------------------------
# 7. PROMOTION GATE & CELL-LEVEL PROMOTION
# ----------------------------------------------------
print(f"[STEP 7] PROMOTION GATE PASSED! Promoting ONLY authorized Batch 1 cells into real workbook...")

wb_real = openpyxl.load_workbook(master_src)
promoted_cells = 0

for sheetname in ['SKILL', 'BUFF_STATUS', 'ZHIZHI', 'HUANZHANG']:
    sheet_temp = wb_temp[sheetname]
    sheet_real = wb_real[sheetname]
    
    headers_temp = [cell.value for cell in sheet_temp[1]]
    headers_real = [cell.value for cell in sheet_real[1]]

    if sheetname == 'SKILL':
        id_col_t, cid_col_t = headers_temp.index('skill_id'), headers_temp.index('character_id')
        id_col_r, cid_col_r = headers_real.index('skill_id'), headers_real.index('character_id')
        real_row_map = {str(row[id_col_r].value or ''): r_idx for r_idx, row in enumerate(sheet_real.iter_rows(min_row=2), start=2)}
        
        for row_temp in sheet_temp.iter_rows(min_row=2):
            cid_t = str(row_temp[cid_col_t].value or '')
            rid_t = str(row_temp[id_col_t].value or '')
            if cid_t in batch1_chars and rid_t in real_row_map:
                r_idx = real_row_map[rid_t]
                for c_idx, cell_t in enumerate(row_temp, start=1):
                    val_t = cell_t.value
                    val_r = sheet_real.cell(row=r_idx, column=c_idx).value
                    if val_t != val_r:
                        sheet_real.cell(row=r_idx, column=c_idx).value = val_t
                        promoted_cells += 1

    elif sheetname == 'BUFF_STATUS':
        id_col_t = headers_temp.index('buff_id')
        id_col_r = headers_real.index('buff_id')
        real_row_map = {str(row[id_col_r].value or ''): r_idx for r_idx, row in enumerate(sheet_real.iter_rows(min_row=2), start=2)}
        
        for row_temp in sheet_temp.iter_rows(min_row=2):
            rid_t = str(row_temp[id_col_t].value or '')
            if rid_t in referenced_buff_ids and rid_t in real_row_map:
                r_idx = real_row_map[rid_t]
                for c_idx, cell_t in enumerate(row_temp, start=1):
                    val_t = cell_t.value
                    val_r = sheet_real.cell(row=r_idx, column=c_idx).value
                    if val_t != val_r:
                        sheet_real.cell(row=r_idx, column=c_idx).value = val_t
                        promoted_cells += 1

    elif sheetname == 'ZHIZHI':
        cid_col_t, star_col_t = headers_temp.index('character_id'), headers_temp.index('star')
        cid_col_r, star_col_r = headers_real.index('character_id'), headers_real.index('star')
        real_row_map = {(str(row[cid_col_r].value or ''), row[star_col_r].value): r_idx for r_idx, row in enumerate(sheet_real.iter_rows(min_row=2), start=2)}
        
        for row_temp in sheet_temp.iter_rows(min_row=2):
            cid_t = str(row_temp[cid_col_t].value or '')
            star_t = row_temp[star_col_t].value
            if cid_t in batch1_chars and (cid_t, star_t) in real_row_map:
                r_idx = real_row_map[(cid_t, star_t)]
                for c_idx, cell_t in enumerate(row_temp, start=1):
                    val_t = cell_t.value
                    val_r = sheet_real.cell(row=r_idx, column=c_idx).value
                    if val_t != val_r:
                        sheet_real.cell(row=r_idx, column=c_idx).value = val_t
                        promoted_cells += 1

    elif sheetname == 'HUANZHANG':
        id_col_t = headers_temp.index('brilliant_id')
        id_col_r = headers_real.index('brilliant_id')
        real_row_map = {str(row[id_col_r].value or ''): r_idx for r_idx, row in enumerate(sheet_real.iter_rows(min_row=2), start=2)}
        
        for row_temp in sheet_temp.iter_rows(min_row=2):
            rid_t = str(row_temp[id_col_t].value or '')
            if rid_t in {'A01604', 'V01464', 'V01174'} and rid_t in real_row_map:
                r_idx = real_row_map[rid_t]
                for c_idx, cell_t in enumerate(row_temp, start=1):
                    val_t = cell_t.value
                    val_r = sheet_real.cell(row=r_idx, column=c_idx).value
                    if val_t != val_r:
                        sheet_real.cell(row=r_idx, column=c_idx).value = val_t
                        promoted_cells += 1

wb_real.save(master_src)
print(f"[STEP 7 SUCCESS] Promoted {promoted_cells} cells into real localization_master.xlsx!")

# ----------------------------------------------------
# 8. VERIFY UNRELATED DATA WAS NOT CHANGED
# ----------------------------------------------------
print(f"[STEP 8] Verifying unrelated cells diff = 0 against backup...")

wb_backup = openpyxl.load_workbook(backup_master, data_only=True)
wb_promoted = openpyxl.load_workbook(master_src, data_only=True)

unrelated_diffs = 0

for sheetname in wb_backup.sheetnames:
    ws_b = wb_backup[sheetname]
    ws_p = wb_promoted[sheetname]
    
    headers_b = [cell.value for cell in ws_b[1]]
    
    for r_idx, (r_b, r_p) in enumerate(zip(ws_b.iter_rows(min_row=2, values_only=True), ws_p.iter_rows(min_row=2, values_only=True)), start=2):
        if sheetname == 'SKILL':
            cid = str(r_b[1] or '')
            if cid in batch1_chars: continue
        elif sheetname == 'BUFF_STATUS':
            bid = str(r_b[0] or '')
            if bid in referenced_buff_ids: continue
        elif sheetname == 'ZHIZHI':
            cid = str(r_b[0] or '')
            if cid in batch1_chars: continue
        elif sheetname == 'HUANZHANG':
            bid = str(r_b[0] or '')
            if bid in {'A01604', 'V01464', 'V01174'}: continue
            
        for c_idx, (val_b, val_p) in enumerate(zip(r_b, r_p), start=1):
            if val_b != val_p:
                print(f"  [UNRELATED DIFF ERROR] Sheet '{sheetname}' Row {r_idx} Col {c_idx}: Backup='{val_b}' vs Promoted='{val_p}'")
                unrelated_diffs += 1

print(f"  Unrelated cell diff count: {unrelated_diffs}")
if unrelated_diffs > 0:
    print("[ERROR] Unrelated cells were modified! Aborting pipeline...")
    sys.exit(1)

print("[STEP 8 SUCCESS] Unrelated cell diff = 0! Proof verified 100% perfect.")

# ----------------------------------------------------
# 9. REGENERATE DERIVED ARTIFACTS
# ----------------------------------------------------
print(f"[STEP 9] Regenerating derived JSON and web data artifacts...")

res1 = subprocess.run([sys.executable, 'tools/export_localization_json.py'], cwd=base_dir, capture_output=True, text=True)
print(f"  export_localization_json.py output:\n{res1.stdout}")
if res1.returncode != 0:
    print(f"  [ERROR] export_localization_json.py failed:\n{res1.stderr}")
    sys.exit(1)

res2 = subprocess.run([sys.executable, 'tools/build_web_data.py'], cwd=base_dir, capture_output=True, text=True)
print(f"  build_web_data.py output:\n{res2.stdout}")
if res2.returncode != 0:
    print(f"  [ERROR] build_web_data.py failed:\n{res2.stderr}")
    sys.exit(1)

gen_review_script = os.path.join(scratch_dir, 'generate_batch1_review_package.py')
res3 = subprocess.run([sys.executable, gen_review_script], cwd=base_dir, capture_output=True, text=True)
print(f"  generate_batch1_review_package.py output:\n{res3.stdout}")
if res3.returncode != 0:
    print(f"  [ERROR] generate_batch1_review_package.py failed:\n{res3.stderr}")
    sys.exit(1)

# ----------------------------------------------------
# 10. FINAL VERIFICATION ON REGENERATED HUMAN REVIEW PACKAGE
# ----------------------------------------------------
print(f"[STEP 10] Performing final verification on regenerated human review package...")

final_review_md = os.path.join(base_dir, 'localization', 'batch_review', 'batch1_human_review.md')
with open(final_review_md, 'r', encoding='utf-8') as f:
    md_content = f.read()

none_desc_count = md_content.count("Desc VI**: None")
none_zhizhi_count = md_content.count("VI Player-Facing**: None")
tri_hoan_count = md_content.count("Trì Hoãn")

print(f"  Final Review 'Desc VI**: None' count: {none_desc_count}")
print(f"  Final Review 'VI Player-Facing**: None' count: {none_zhizhi_count}")
print(f"  Final Review 'Trì Hoãn' count: {tri_hoan_count}")

assert_failed = False
if none_desc_count > 0:
    print("[ERROR] Final review package still contains unpopulated Desc VI=None!")
    assert_failed = True
if none_zhizhi_count > 0:
    print("[ERROR] Final review package still contains unpopulated Zhizhi VI=None!")
    assert_failed = True
if tri_hoan_count > 0:
    print("[ERROR] Final review package still contains 'Trì Hoãn' instead of 'Trì Trệ'!")
    assert_failed = True

required_phrases = [
    'Dễ Vỡ', 'Trì Trệ', 'Ẩn Nấp', 'Thiêu Đốt', 'Hàn Thiên', 'Cấm Túc',
    'Sát Thương Chí Tử', 'Tỷ Lệ Hút Máu'
]

for rp in required_phrases:
    if rp not in md_content:
        print(f"  [ERROR] Final review package is missing required terminology: '{rp}'")
        assert_failed = True

if assert_failed:
    print("[ERROR] Final review verification failed!")
    sys.exit(1)

print("[STEP 10 SUCCESS] Final review package verified 100% fully populated and naturalized!")

print("\n==================================================")
print("SAFE DETERMINISTIC RESTORATION COMPLETE SUCCESSFULLY!")
print("==================================================\n")
