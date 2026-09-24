import sys
import hashlib
from collections import Counter
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

wb_path = 'localization/localization_master.xlsx'
with open(wb_path, 'rb') as f:
    sha = hashlib.sha256(f.read()).hexdigest()
print(f"Current SHA-256: {sha}")

wb = openpyxl.load_workbook(wb_path, data_only=True, read_only=True)

# 1. Load CHARACTER
char_ws = wb['CHARACTER']
char_rows = list(char_ws.iter_rows(values_only=True))
char_header = char_rows[0]
cid_idx = char_header.index('character_id')
cname_vi_idx = char_header.index('name_vi')

char_name_map = {}
for r in char_rows[1:]:
    cid = str(r[cid_idx]).strip() if r[cid_idx] is not None else ''
    c_vi = str(r[cname_vi_idx]).strip() if r[cname_vi_idx] is not None else ''
    char_name_map[cid] = c_vi

# 2. Inspect SKIN
skin_ws = wb['SKIN']
skin_rows = list(skin_ws.iter_rows(values_only=True))
header = skin_rows[0]
data = skin_rows[1:]
cols = {name: header.index(name) for name in header}

print(f"Total SKIN rows: {len(data)}")

# Audit obtain_vi
missing_obtain_rows = []
for r in data:
    o_vi = str(r[cols['obtain_vi']]).strip() if r[cols['obtain_vi']] is not None else ''
    o_cn = str(r[cols['obtain_cn']]).strip() if r[cols['obtain_cn']] is not None else ''
    sid = str(r[cols['skin_id']]).strip()
    cid = str(r[cols['character_id']]).strip()
    cname_vi = char_name_map.get(cid, '')
    sname_vi = str(r[cols['skin_name_vi']]).strip() if r[cols['skin_name_vi']] is not None else ''
    sname_cn = str(r[cols['skin_name_cn']]).strip() if r[cols['skin_name_cn']] is not None else ''
    series_id = r[cols['series_id']]
    series_vi = r[cols['series_name_vi']]

    if not o_vi and o_cn:
        missing_obtain_rows.append((sid, cid, cname_vi, sname_cn, sname_vi, series_id, series_vi, o_cn, o_vi))

print(f"\n--- Rows missing obtain_vi ({len(missing_obtain_rows)}) ---")
for row in missing_obtain_rows:
    print(f"ID: {row[0]} | Char: {row[2]} ({row[1]}) | Skin: {row[4]} ({row[3]}) | Series: {row[5]} ({row[6]}) | obtain_cn: '{row[7]}'")

# Audit skin_name_vi
missing_name_rows = []
for r in data:
    sname_vi = str(r[cols['skin_name_vi']]).strip() if r[cols['skin_name_vi']] is not None else ''
    sname_cn = str(r[cols['skin_name_cn']]).strip() if r[cols['skin_name_cn']] is not None else ''
    sid = str(r[cols['skin_id']]).strip()
    cid = str(r[cols['character_id']]).strip()
    cname_vi = char_name_map.get(cid, '')
    if not sname_vi or sname_vi == sid:
        missing_name_rows.append((sid, cid, cname_vi, sname_cn, sname_vi))

print(f"\n--- Rows missing skin_name_vi ({len(missing_name_rows)}) ---")
for row in missing_name_rows:
    print(row)

# Audit duplicate source_cn in desc_cn
desc_cn_list = []
for r in data:
    d_cn = str(r[cols['desc_cn']]).strip() if r[cols['desc_cn']] is not None else ''
    if d_cn:
        desc_cn_list.append(d_cn)

desc_counter = Counter(desc_cn_list)
desc_dups = {k: v for k, v in desc_counter.items() if v > 1}
print(f"\n--- Duplicate desc_cn strings ({len(desc_dups)}) ---")
for k, v in desc_dups.items():
    print(f"Count {v}: '{k[:40]}...'")

# Audit duplicate source_cn in obtain_cn
obtain_cn_list = []
for r in data:
    o_cn = str(r[cols['obtain_cn']]).strip() if r[cols['obtain_cn']] is not None else ''
    if o_cn:
        obtain_cn_list.append(o_cn)

obtain_counter = Counter(obtain_cn_list)
obtain_dups = {k: v for k, v in obtain_counter.items() if v > 1}
print(f"\n--- Duplicate obtain_cn strings ({len(obtain_dups)}) ---")
for k, v in obtain_dups.items():
    print(f"Count {v}: '{k}'")
