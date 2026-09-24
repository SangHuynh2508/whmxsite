"""Apply the owner-approved terminology correction without broad replacements.

The script deliberately edits only canonical BUFF_STATUS rows, the seven shared
GLOSSARY terms, and exact W0164 SKILL named references identified by CN/raw IDs.
"""

from __future__ import annotations

import argparse
import json
import sys
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


TODAY = date.today().isoformat()

SHARED_TERMS = [
    ("蓄势", "Súc Thế", "Tích Thế"),
    ("萧瑟", "Tiêu Sắt", ""),
    ("截招", "Tiệt Chiêu", ""),
    ("滞缓", "Trệ Hoãn", "Trì Hoãn"),
    ("瞄准", "Miêu Chuẩn", "Nhắm Bắn"),
    ("脆弱", "Thúy Nhược", "Dễ Vỡ"),
    ("降低命中率", "Giảm Tỷ Lệ Trúng", "Giảm Trúng"),
]

# The five private names intentionally do not appear in GLOSSARY.
CANONICAL_BUFFS = [
    ("Buff_AllDmgIncrease", "蓄势", "Súc Thế", "Tích Thế", "shared/system"),
    ("Buff_AllDmgReduce", "萧瑟", "Tiêu Sắt", "", "shared/system"),
    ("Buff_HitRate_Down", "截招", "Tiệt Chiêu", "", "shared/system"),
    ("Buff_Mov_Down", "滞缓", "Trệ Hoãn", "Trì Hoãn", "shared/system"),
    ("Buff_HitRate_Up", "瞄准", "Miêu Chuẩn", "Nhắm Bắn", "shared/system"),
    ("Buff_Fragile", "脆弱", "Thúy Nhược", "Dễ Vỡ", "shared/system"),
    ("Buff_A0086_5", "降低命中率", "Giảm Tỷ Lệ Trúng", "Giảm Trúng", "shared/system label"),
    ("Buff_A0084_1", "铜锈", "Đồng Tú", "", "private"),
    ("Buff_A0086_24", "隐蔽", "Ẩn Tế", "Ẩn Nấp", "private"),
    ("Buff_A0090_1", "吉时", "Cát Thời", "Giờ Lành", "private"),
    ("Buff_A0090_3A_1", "避让", "Tị Nhượng", "Nhường Đường", "private"),
    ("Buff_A0090_hz_4_3", "临时干部", "Lâm Thời Cán Bộ", "Cán Bộ Tạm Thời", "private"),
]


def header_map(ws):
    return {cell.value: cell.column for cell in ws[1] if cell.value}


def append_note(existing, note):
    existing = (existing or "").strip()
    return existing if note in existing else (f"{existing} | {note}" if existing else note)


def copy_row_schema(ws, source_row, target_row):
    for col in range(1, ws.max_column + 1):
        src, dst = ws.cell(source_row, col), ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)


def run(path: Path, apply: bool):
    wb = load_workbook(path)
    glossary, buffs, skills = wb["GLOSSARY"], wb["BUFF_STATUS"], wb["SKILL"]
    gh, bh, sh = header_map(glossary), header_map(buffs), header_map(skills)
    report = {
        "dry_run": not apply,
        "glossary_existing": [], "glossary_added": [], "glossary_updated": [],
        "buffs_updated": [], "w0164_named_references_updated": [],
        "private_terms_excluded_from_glossary": [item[1] for item in CANONICAL_BUFFS if item[4] == "private"],
        "errors": [],
    }

    glossary_by_cn = {}
    for row in range(2, glossary.max_row + 1):
        value = glossary.cell(row, gh["term_cn"]).value
        if value:
            glossary_by_cn.setdefault(str(value), []).append(row)

    next_gameplay_id = max(
        int(str(glossary.cell(row, gh["term_id"]).value).rsplit("_", 1)[-1])
        for row in range(2, glossary.max_row + 1)
        if str(glossary.cell(row, gh["term_id"]).value).startswith("OWNER_GAMEPLAY_")
    )
    schema_row = glossary.max_row
    planned_glossary_row = glossary.max_row
    for cn, vi, old_vi in SHARED_TERMS:
        rows = glossary_by_cn.get(cn, [])
        note = f"OWNER_CORRECTED {TODAY}; CN={cn}; VI={vi}; supersedes={old_vi or 'none'}"
        if len(rows) > 1:
            report["errors"].append(f"GLOSSARY duplicate term_cn {cn}: rows {rows}")
            continue
        if rows:
            row = rows[0]
            prior = glossary.cell(row, gh["term_vi"]).value
            report["glossary_existing"].append({"cn": cn, "row": row, "vi": prior})
            if prior != vi or note not in str(glossary.cell(row, gh["notes"]).value or ""):
                report["glossary_updated"].append({"cn": cn, "row": row, "from": prior, "to": vi})
                if apply:
                    glossary.cell(row, gh["term_vi"]).value = vi
                    glossary.cell(row, gh["confidence"]).value = "HIGH"
                    glossary.cell(row, gh["status"]).value = "OWNER_APPROVED"
                    glossary.cell(row, gh["notes"]).value = append_note(glossary.cell(row, gh["notes"]).value, note)
            continue
        next_gameplay_id += 1
        planned_glossary_row += 1
        row = planned_glossary_row
        term_id = f"OWNER_GAMEPLAY_{next_gameplay_id:03d}"
        report["glossary_added"].append({"cn": cn, "vi": vi, "row": row, "term_id": term_id})
        if apply:
            copy_row_schema(glossary, schema_row, row)
            values = {
                "term_id": term_id, "category": "gameplay_terminology", "term_cn": cn,
                "term_vi": vi, "han_viet": None, "confidence": "HIGH",
                "status": "OWNER_APPROVED", "notes": note,
            }
            for field, value in values.items():
                glossary.cell(row, gh[field]).value = value
        glossary_by_cn[cn] = [row]

    buff_rows = {str(buffs.cell(row, bh["buff_id"]).value): row for row in range(2, buffs.max_row + 1)}
    for buff_id, cn, vi, old_vi, scope in CANONICAL_BUFFS:
        row = buff_rows.get(buff_id)
        if not row:
            report["errors"].append(f"Missing canonical BUFF_STATUS row {buff_id}")
            continue
        actual_cn = buffs.cell(row, bh["buff_name_cn"]).value
        if actual_cn != cn:
            report["errors"].append(f"CN mismatch for {buff_id}: expected {cn}, found {actual_cn}")
            continue
        prior = buffs.cell(row, bh["buff_name_vi"]).value
        note = f"OWNER_CORRECTED {TODAY}; {scope}; CN={cn}; VI={vi}; supersedes={old_vi or 'none'}"
        if prior != vi or note not in str(buffs.cell(row, bh["notes"]).value or ""):
            report["buffs_updated"].append({"buff_id": buff_id, "row": row, "from": prior, "to": vi})
            if apply:
                buffs.cell(row, bh["buff_name_vi"]).value = vi
                buffs.cell(row, bh["confidence"]).value = "HIGH"
                buffs.cell(row, bh["status"]).value = "OWNER_APPROVED"
                buffs.cell(row, bh["notes"]).value = append_note(buffs.cell(row, bh["notes"]).value, note)

    # Exact raw/card evidence only: W0164 passive 2 has {Buff_HitRate_Down}
    # and CN 截招. Do not touch coincidental prose in other cards.
    for row in range(2, skills.max_row + 1):
        if skills.cell(row, sh["character_id"]).value != "W0164":
            continue
        if skills.cell(row, sh["skill_group_id"]).value != "W016404":
            continue
        cn = str(skills.cell(row, sh["desc_cn"]).value or "")
        vi = str(skills.cell(row, sh["desc_vi"]).value or "")
        if "截招" not in cn or "{Buff_HitRate_Down}" not in cn:
            report["errors"].append(f"W0164 row {row} lacks raw 截招 / Buff_HitRate_Down evidence")
            continue
        if "Tiết Chiêu" in vi:
            updated = vi.replace("Tiết Chiêu", "Tiệt Chiêu")
            report["w0164_named_references_updated"].append({"row": row, "from": "Tiết Chiêu", "to": "Tiệt Chiêu"})
            if apply:
                skills.cell(row, sh["desc_vi"]).value = updated
                skills.cell(row, sh["notes"]).value = append_note(
                    skills.cell(row, sh["notes"]).value,
                    f"OWNER_CORRECTED {TODAY}; named raw marker 截招 -> Tiệt Chiêu",
                )

    if report["errors"]:
        raise RuntimeError("\n".join(report["errors"]))
    if apply:
        wb.save(path)
    return report


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=Path("localization/localization_master.xlsx"))
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    print(json.dumps(run(args.workbook, args.apply), ensure_ascii=False, indent=2))
