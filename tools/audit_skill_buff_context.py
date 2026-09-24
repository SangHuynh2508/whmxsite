"""Audit exact BUFF markers and source-backed named references in translated skills."""
from __future__ import annotations
import json, re
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
MASTER=ROOT/'localization'/'localization_master.xlsx'
MARKER=re.compile(r'\{(Buff_[^}\s]+)\}',re.I)
TAG=re.compile(r'<[^>]+>')
def s(v):return '' if v is None else str(v)
def plain(v):return TAG.sub('',s(v)).strip()
def header(ws):return {s(c.value):c.column for c in ws[1] if c.value is not None}
def main():
 wb=load_workbook(MASTER,data_only=False);bw=wb['BUFF_STATUS'];bh=header(bw);buffs={s(bw.cell(r,bh['buff_id']).value):{x:s(bw.cell(r,c).value) for x,c in bh.items()} for r in range(2,bw.max_row+1)}
 sw=wb['SKILL'];sh=header(sw);issues=[];audited=0
 for r in range(2,sw.max_row+1):
  status=s(sw.cell(r,sh['status']).value).upper();cn=s(sw.cell(r,sh['desc_cn']).value);vi=s(sw.cell(r,sh['desc_vi']).value)
  if status not in {'TRANSLATED','APPROVED','OWNER_APPROVED','OWNER_CORRECTED'} or not vi:continue
  audited+=1;sid=s(sw.cell(r,sh['skill_id']).value);cm=MARKER.findall(cn);vm=MARKER.findall(vi)
  if cm!=vm:issues.append({'code':'MARKER_ORDER_CHANGED' if Counter(cm)==Counter(vm) else ('MISSING_MARKER_IN_VI' if set(cm)-set(vm) else 'EXTRA_MARKER_IN_VI'),'skill_id':sid,'cn_markers':cm,'vi_markers':vm})
  for marker in cm:
   buff=buffs.get(marker)
   if not buff:issues.append({'code':'AMBIGUOUS_BUFF_CONTEXT','skill_id':sid,'buff_id':marker});continue
   bvi=plain(buff.get('buff_name_vi'));bcn=plain(buff.get('buff_name_cn'))
   if not bvi:issues.append({'code':'MISSING_BUFF_TRANSLATION','skill_id':sid,'buff_id':marker});continue
   if bcn and bcn in plain(cn) and bvi not in plain(vi):issues.append({'code':'SKILL_BUFF_NAME_MISMATCH','skill_id':sid,'buff_id':marker,'expected_name_vi':bvi})
 out=ROOT/'localization'/'batch_reports'/'skill_buff_context_audit_20260910.json';out.write_text(json.dumps({'audited_translated_skills':audited,'counts':dict(Counter(x['code'] for x in issues)),'issues':issues},ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps({'output':str(out),'audited':audited,'counts':dict(Counter(x['code'] for x in issues))},ensure_ascii=False))
if __name__=='__main__':main()
