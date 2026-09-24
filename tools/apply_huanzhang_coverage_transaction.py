"""Safely append the two raw-backed Hoán Chương coverage rows.

This transaction intentionally imports source CN only.  The long-form VI lore
requires editorial review and is therefore left blank/PENDING rather than
machine-authored or marked approved.
"""
from __future__ import annotations

import copy
import hashlib
import json
import sys
from pathlib import Path

import openpyxl

from safe_workbook_mutation import safe_mutate_workbook
from workbook_semantic_fingerprint import workbook_semantic_fingerprint


ROOT = Path(__file__).resolve().parents[2]
PROJECT = ROOT / "WhmxCalc"
NEO = ROOT / "NeoArtifacts"
MASTER = PROJECT / "localization" / "localization_master.xlsx"
RAW = NEO / "MasterData" / "json"
SNAPSHOT_ID = "r3021-20260915T081707511131Z-W0182"
TARGETS = {
    "A01204": "A0120",
    "D00924": "D0092",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_json(name: str) -> dict:
    return json.loads((RAW / name).read_text(encoding="utf-8"))


def snapshot_guard() -> None:
    """Require the mutable mirror to remain identical to the selected snapshot."""
    sys.path.insert(0, str(NEO))
    import character_assets  # Imported only after the NeoArtifacts path is explicit.

    _, manifest = character_assets.load_snapshot(SNAPSHOT_ID)
    expected = manifest["masterdata_json_source_fingerprint"]
    actual = character_assets.fingerprint_tree(RAW)
    if actual != expected:
        raise RuntimeError(
            f"MasterData mirror differs from selected snapshot {SNAPSHOT_ID}: {actual} != {expected}"
        )


def exact_raw_ownership(brilliant: dict, group_to_heroes: dict[str, set[str]]) -> dict[str, list[str]]:
    owners = {}
    for brilliant_id in TARGETS:
        record = brilliant.get(brilliant_id)
        if not isinstance(record, dict):
            raise RuntimeError(f"Missing raw BrilliantMap record: {brilliant_id}")
        groups = []
        for field in ("Buff", "Skill1", "Skill2"):
            value = record.get(field, [])
            groups.extend(value if isinstance(value, list) else [value])
        candidates = sorted({hero for group in groups for hero in group_to_heroes.get(str(group), set())})
        if candidates != [TARGETS[brilliant_id]]:
            raise RuntimeError(
                f"Raw ownership for {brilliant_id} is {candidates}, expected only {TARGETS[brilliant_id]}"
            )
        owners[brilliant_id] = [str(group) for group in groups if group]
    return owners


def build_rows() -> tuple[dict[str, list[object]], dict[str, list[str]]]:
    snapshot_guard()
    brilliant = load_json("BrilliantMap.json")
    skills = load_json("skillMap.json")
    combat = {
        **load_json("characterSkillMap.json"),
        **load_json("characterPassiveSkillMap.json"),
    }
    group_to_heroes: dict[str, set[str]] = {}
    for source in (skills, combat):
        for record in source.values():
            if not isinstance(record, dict):
                continue
            group, hero = str(record.get("GroupId") or ""), str(record.get("HeroId") or "")
            if group and hero:
                group_to_heroes.setdefault(group, set()).add(hero)
    groups_by_id = exact_raw_ownership(brilliant, group_to_heroes)

    icon_dir = PROJECT / "public" / "assets" / "huanzhang"
    exact = {path.name for path in icon_dir.iterdir() if path.is_file()}
    folded = {name.casefold(): name for name in exact}
    rows = {}
    for brilliant_id, character_id in TARGETS.items():
        raw = brilliant[brilliant_id]
        icon = str(raw.get("Icon") or "")
        expected = f"{icon}.png" if icon and not Path(icon).suffix else icon
        actual = folded.get(expected.casefold(), "")
        if not actual:
            raise RuntimeError(f"Missing Hoán Chương icon for {brilliant_id}: {expected!r}")
        rows[brilliant_id] = [
            brilliant_id,
            character_id,
            raw.get("IconNameLanText") or raw.get("IconName") or "",
            "",  # Editorial VI name review remains required.
            raw.get("IconInfoLanText") or raw.get("IconInfo") or "",
            "",  # Never machine-invent long-form VI lore.
            raw.get("BuffShowLanText") or raw.get("BuffShow") or "",
            "",
            "LOW",
            "PENDING",
            (
                f"Raw coverage import; snapshot={SNAPSHOT_ID}; source=BrilliantMap.json; "
                f"groups={','.join(groups_by_id[brilliant_id])}; icon={actual}"
            ),
            "PENDING_RAW_IMPORT",
            f"immutable runtime snapshot {SNAPSHOT_ID}",
        ]
    return rows, groups_by_id


def main() -> int:
    rows, groups_by_id = build_rows()
    before_sha = sha256(MASTER)
    before_semantic = workbook_semantic_fingerprint(MASTER)

    def mutator(workbook):
        ws = workbook["HUANZHANG"]
        headers = [cell.value for cell in ws[1]]
        if headers != [
            "brilliant_id", "character_id", "icon_name_cn", "icon_name_vi",
            "icon_info_cn", "icon_info_vi", "buff_show_cn", "buff_show_vi",
            "confidence", "status", "notes", "translation_review_status",
            "translation_review_source",
        ]:
            raise RuntimeError(f"Unexpected HUANZHANG schema: {headers!r}")
        existing = {str(row[0].value or "").strip() for row in ws.iter_rows(min_row=2, max_col=1)}
        if existing & set(rows):
            raise RuntimeError(f"Refusing to overwrite existing HUANZHANG rows: {sorted(existing & set(rows))}")
        template_row = ws.max_row
        for brilliant_id in sorted(rows):
            destination = ws.max_row + 1
            for column, value in enumerate(rows[brilliant_id], start=1):
                target = ws.cell(destination, column, value)
                source = ws.cell(template_row, column)
                if source.has_style:
                    target._style = copy.copy(source._style)
                if source.number_format:
                    target.number_format = source.number_format
                target.alignment = copy.copy(source.alignment)
            ws.row_dimensions[destination].height = ws.row_dimensions[template_row].height

    backup = safe_mutate_workbook(
        str(PROJECT),
        mutator,
        {"huanzhang_ids": set()},
        authorized_new_columns={},
        authorized_new_rows={"HUANZHANG": set(rows)},
    )
    after_sha = sha256(MASTER)
    after_semantic = workbook_semantic_fingerprint(MASTER)
    print("HUANZHANG_COVERAGE_TRANSACTION=PASS")
    print(f"BACKUP={backup}")
    print(f"SHA256_BEFORE={before_sha}")
    print(f"SHA256_AFTER={after_sha}")
    print(f"SEMANTIC_BEFORE={before_semantic}")
    print(f"SEMANTIC_AFTER={after_semantic}")
    for brilliant_id in sorted(rows):
        print(f"ROW_ADDED=HUANZHANG:{brilliant_id}:owner={TARGETS[brilliant_id]}:groups={','.join(groups_by_id[brilliant_id])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
