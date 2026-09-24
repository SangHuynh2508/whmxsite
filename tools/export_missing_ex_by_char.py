#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
export_missing_ex_by_char.py
Với mỗi nhân vật còn thiếu skill_ex / buff_ex:
- Xuất TOÀN BỘ skill + buff của con đó (đã dịch + chưa dịch)
- Kèm CHARACTER / ZHIZHI / HUANZHANG để có ngữ cảnh
"""

import os
from datetime import datetime
from collections import defaultdict
import openpyxl
from openpyxl import Workbook

MASTER = "localization/localization_master.xlsx"  # sửa đường dẫn nếu cần
OUT = f"localization/missing_ex_context_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

def is_filled(val):
    if val is None:
        return False
    s = str(val).strip()
    return s != "" and s.lower() != "nan"

def get_headers(ws):
    return [c.value for c in ws[1]]

def main():
    if not os.path.exists(MASTER):
        print(f"[-] Không tìm thấy: {MASTER}")
        return

    print(f"[+] Đọc: {MASTER}")
    wb = openpyxl.load_workbook(MASTER, data_only=True)

    # ===== 1. Tìm character_id còn thiếu skill_ex / buff_ex =====
    chars_need = set()

    if "SKILL" in wb.sheetnames:
        ws = wb["SKILL"]
        headers = get_headers(ws)
        col = {h: i for i, h in enumerate(headers)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            sid = row[col.get("skill_id")] if "skill_id" in col else None
            if sid is None or "ex" not in str(sid).lower():
                continue
            name_cn = row[col.get("skill_name_cn")] if "skill_name_cn" in col else None
            name_vi = row[col.get("skill_name_vi")] if "skill_name_vi" in col else None
            desc_cn = row[col.get("desc_cn")] if "desc_cn" in col else None
            desc_vi = row[col.get("desc_vi")] if "desc_vi" in col else None
            if (is_filled(name_cn) and not is_filled(name_vi)) or (is_filled(desc_cn) and not is_filled(desc_vi)):
                cid = row[col.get("character_id")] if "character_id" in col else None
                if cid:
                    chars_need.add(str(cid).strip())

    if "BUFF_STATUS" in wb.sheetnames:
        ws = wb["BUFF_STATUS"]
        headers = get_headers(ws)
        col = {h: i for i, h in enumerate(headers)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            bid = row[col.get("buff_id")] if "buff_id" in col else None
            if bid is None or "ex" not in str(bid).lower():
                continue
            name_cn = row[col.get("buff_name_cn")] if "buff_name_cn" in col else None
            name_vi = row[col.get("buff_name_vi")] if "buff_name_vi" in col else None
            desc_cn = row[col.get("buff_desc_cn")] if "buff_desc_cn" in col else None
            desc_vi = row[col.get("buff_desc_vi")] if "buff_desc_vi" in col else None
            if (is_filled(name_cn) and not is_filled(name_vi)) or (is_filled(desc_cn) and not is_filled(desc_vi)):
                # suy character_id từ Buff_A0144_xxx
                parts = str(bid).replace("Buff_", "").split("_")
                if parts and len(parts[0]) >= 4 and parts[0][0] in "ADSVW":
                    chars_need.add(parts[0])

    print(f"  Nhân vật còn thiếu skill_ex/buff_ex: {len(chars_need)}")
    if not chars_need:
        print("[V] Không còn gì thiếu.")
        return

    # ===== 2. Xuất TOÀN BỘ data của các nhân vật đó =====
    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    def copy_by_chars(sheet_name, id_col):
        if sheet_name not in wb.sheetnames:
            return
        ws = wb[sheet_name]
        headers = get_headers(ws)
        col = {h: i for i, h in enumerate(headers)}
        if id_col not in col:
            print(f"  [-] {sheet_name}: không có cột {id_col}")
            return
        ws_out = wb_out.create_sheet(sheet_name)
        ws_out.append(headers)
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            val = row[col[id_col]]
            if val is None:
                continue
            if str(val).strip() in chars_need:
                ws_out.append(list(row))
                count += 1
        print(f"  → {sheet_name}: {count} dòng")

    # CHARACTER
    copy_by_chars("CHARACTER", "character_id")

    # SKILL: toàn bộ skill của các con đó (đã dịch + chưa dịch + ex)
    copy_by_chars("SKILL", "character_id")

    # ZHIZHI / HUANZHANG
    copy_by_chars("ZHIZHI", "character_id")
    copy_by_chars("HUANZHANG", "character_id")

    # BUFF_STATUS: lấy theo prefix character_id trong buff_id
    if "BUFF_STATUS" in wb.sheetnames:
        ws = wb["BUFF_STATUS"]
        headers = get_headers(ws)
        col = {h: i for i, h in enumerate(headers)}
        ws_out = wb_out.create_sheet("BUFF_STATUS")
        ws_out.append(headers)
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            bid = row[col.get("buff_id")] if "buff_id" in col else None
            if bid is None:
                continue
            bid = str(bid)
            # Buff_A0144_xxx hoặc Buff_A0144ex_xxx
            matched = False
            for cid in chars_need:
                if f"_{cid}_" in bid or bid.startswith(f"Buff_{cid}") or f"Buff_{cid}" in bid:
                    matched = True
                    break
            if matched:
                ws_out.append(list(row))
                count += 1
        print(f"  → BUFF_STATUS: {count} dòng")

    os.makedirs(os.path.dirname(OUT) or ".", exist_ok=True)
    wb_out.save(OUT)
    print(f"\n[V] Đã xuất → {OUT}")
    print("    File gồm toàn bộ skill/buff (đã dịch + chưa dịch) của các nhân vật còn thiếu ex.")
    print("    Gửi file này cho mình để dịch nốt.")

if __name__ == "__main__":
    main()