import openpyxl
import json
import os
import sys
import re

sys.stdout.reconfigure(encoding='utf-8')

from batch_manifest_helper import load_batch_manifest, collect_batch_dependencies

def validate_terminology_consistency(manifest_path, base_dir=None, wb_master=None):
    if base_dir is None:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        
    manifest = load_batch_manifest(manifest_path)
    deps = collect_batch_dependencies(base_dir, manifest)
    
    if wb_master is None:
        master_path = os.path.join(base_dir, 'localization', 'localization_master.xlsx')
        wb = openpyxl.load_workbook(master_path, data_only=True)
    else:
        wb = wb_master
    
    player_buffs = set(deps['player_facing_buff_ids'])
    
    ws_buff = wb['BUFF_STATUS']
    headers_bf = list(next(ws_buff.iter_rows(values_only=True)))
    bid_idx = headers_bf.index('buff_id')
    cn_idx = headers_bf.index('buff_name_cn')
    vi_idx = headers_bf.index('buff_name_vi')

    buff_terms = {}
    for r in ws_buff.iter_rows(min_row=2, values_only=True):
        bid = str(r[bid_idx] or '').strip()
        if bid in player_buffs:
            b_cn = str(r[cn_idx] or '').strip()
            b_vi = str(r[vi_idx] or '').strip()
            if b_cn and b_vi and b_vi != 'None':
                buff_terms[b_cn] = (b_vi, bid)

    ws_skill = wb['SKILL']
    headers_sk = list(next(ws_skill.iter_rows(values_only=True)))
    cid_idx_sk = headers_sk.index('character_id')
    sid_idx_sk = headers_sk.index('skill_id')
    desc_cn_idx_sk = headers_sk.index('desc_cn')
    desc_vi_idx_sk = headers_sk.index('desc_vi')

    batch_chars = set(deps['characters'])
    conflicts = []

    for r in ws_skill.iter_rows(min_row=2, values_only=True):
        cid = str(r[cid_idx_sk] or '').strip()
        if cid in batch_chars:
            sid = str(r[sid_idx_sk] or '').strip()
            d_cn = str(r[desc_cn_idx_sk] or '')
            d_vi = str(r[desc_vi_idx_sk] or '')

            # For each verified named buff in this batch, check if skill text references it in color tags
            for b_cn, (b_vi, bid) in buff_terms.items():
                if b_cn in d_cn:
                    # Find corresponding colored term in VI text
                    # e.g., <color=#ff6724>b_cn</color> in CN
                    cn_pattern = re.compile(r'<color=[^>]+>' + re.escape(b_cn) + r'</color>')
                    if cn_pattern.search(d_cn):
                        # Search for matching colored status in VI
                        vi_color_matches = re.findall(r'<color=[^>]+>([^<]+)</color>', d_vi)
                        # Filter out numbers/percentages
                        status_vi_matches = [m.strip() for m in vi_color_matches if not re.match(r'^[0-9%\-+,\.\s]+$', m.strip())]
                        
                        # Check if buff_vi or equivalent is in VI colored matches
                        if status_vi_matches and b_vi not in status_vi_matches:
                            # Flag conflict if none of the colored status names match buff_vi
                            conflicts.append({
                                'cn_term': b_cn,
                                'skill_vi_statuses': status_vi_matches,
                                'buff_vi': b_vi,
                                'skill_id': sid,
                                'buff_id': bid
                            })

    print(f"=== TERMINOLOGY CONSISTENCY REPORT ({manifest['batch_id']}) ===")
    if conflicts:
        print(f"[WARNING] Found {len(conflicts)} terminology conflict(s):")
        seen_conflicts = set()
        for c in conflicts:
            key = (c['cn_term'], tuple(c['skill_vi_statuses']), c['buff_vi'], c['skill_id'], c['buff_id'])
            if key not in seen_conflicts:
                seen_conflicts.add(key)
                print(f"  [TERM_CONFLICT]")
                print(f"    CN: {c['cn_term']}")
                print(f"    Skill VI Statuses: {c['skill_vi_statuses']} (Skill {c['skill_id']})")
                print(f"    Buff VI: {c['buff_vi']} (Buff {c['buff_id']})")
        return False, conflicts
    else:
        print("[SUCCESS] 0 terminology conflicts found between skills and referenced buffs!")
        return True, []

if __name__ == '__main__':
    manifest_file = sys.argv[1] if len(sys.argv) > 1 else r"d:\BaiTapCode\WHMX\WhmxCalc\localization\batches\phase3_batch1.json"
    validate_terminology_consistency(manifest_file)
