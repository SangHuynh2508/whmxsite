#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
restore_missing_from_backup.py
Chỉ thêm dòng còn thiếu (skill + buff, đặc biệt là ex) từ backup vào master.
Không ghi đè dòng đã có.
"""

import os
import shutil
from datetime import datetime
import openpyxl

MASTER = "localization/localization_master.xlsx"
BACKUP = "localization/backups/localization_master_20260907_125847.xlsx"  # sửa đúng đường dẫn backup của bạn

SHEETS = {
    "SKILL": "skill_id",
    "BUFF_STATUS": "buff_id",
}

def main():
    if not os.path.exists(MASTER) or not os.path.exists(BACKUP):
        print("[-] Không tìm thấy MASTER hoặc BACKUP")
        return

    # Backup master trước
    os.makedirs("localization/backups", exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe = f"localization/backups/master_before_restore_{ts}.xlsx"
    shutil.copyfile(MASTER, safe)
    print(f"[+] Backup master hiện tại → {safe}")

    wb_m = openpyxl.load_workbook(MASTER)
    wb_b = openpyxl.load_workbook(BACKUP)

    total_added = 0

    for sheet, id_col in SHEETS.items():
        if sheet not in wb_m.sheetnames or sheet not in wb_b.sheetnames:
            continue

        ws_m = wb_m[sheet]
        ws_b = wb_b[sheet]

        headers = [c.value for c in ws_m[1]]
        if id_col not in headers:
            print(f"[-] {sheet}: không có cột {id_col}")
            continue
        id_idx = headers.index(id_col)

        # ID đã có trong master
        existing = set()
        for row in ws_m.iter_rows(min_row=2, values_only=True):
            if row[id_idx] is not None:
                existing.add(str(row[id_idx]).strip())

        added = 0
        for row in ws_b.iter_rows(min_row=2, values_only=True):
            sid = row[id_idx]
            if sid is None:
                continue
            sid = str(sid).strip()
            if sid in existing:
                continue  # đã có → bỏ qua

            ws_m.append(list(row))
            existing.add(sid)
            added += 1

        print(f" -> {sheet}: thêm {added} dòng còn thiếu")
        total_added += added

    wb_m.save(MASTER)
    print(f"\n[V] Hoàn tất! Đã thêm tổng {total_added} dòng.")
    print(f"    Bản dịch cũ không bị đụng.")
    print(f"    Backup an toàn: {safe}")

if __name__ == "__main__":
    main()