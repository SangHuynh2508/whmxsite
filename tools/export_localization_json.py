"""
tools/export_localization_json.py
==================================
Reads localization/localization_master.xlsx and generates a deterministic
JSON structure used by data generators (build_web_data.py) and frontend applications.

Source-of-Truth: localization/localization_master.xlsx
Output: localization/generated_localization.json
"""

import openpyxl
import json
import sys
from pathlib import Path

LOC_DIR = Path(__file__).resolve().parent.parent / "localization"
MASTER_WORKBOOK = LOC_DIR / "localization_master.xlsx"
OUT_JSON = LOC_DIR / "generated_localization.json"

def safe_str(val):
    if val is None:
        return ""
    return str(val).strip()

def export():
    if not MASTER_WORKBOOK.exists():
        print(f"[ERROR] Cannot find master workbook: {MASTER_WORKBOOK}")
        sys.exit(1)

    print(f"Loading master workbook: {MASTER_WORKBOOK}...")
    wb = openpyxl.load_workbook(MASTER_WORKBOOK, data_only=True)

    # 1. Export GLOSSARY
    glossary = {}
    if "GLOSSARY" in wb.sheetnames:
        ws = wb["GLOSSARY"]
        for r in range(2, ws.max_row + 1):
            term_id = safe_str(ws.cell(r, 1).value)
            category = safe_str(ws.cell(r, 2).value)
            term_cn = safe_str(ws.cell(r, 3).value)
            term_vi = safe_str(ws.cell(r, 4).value)
            han_viet = safe_str(ws.cell(r, 5).value)
            status = safe_str(ws.cell(r, 7).value)
            if term_cn:
                glossary[term_cn] = {
                    "term_id": term_id,
                    "category": category,
                    "term_cn": term_cn,
                    "term_vi": term_vi,
                    "han_viet": han_viet,
                    "status": status
                }

    # 2. Export CHARACTER
    characters = {}
    if "CHARACTER" in wb.sheetnames:
        ws = wb["CHARACTER"]
        for r in range(2, ws.max_row + 1):
            cid = safe_str(ws.cell(r, 1).value)
            if cid:
                characters[cid] = {
                    "name_cn": safe_str(ws.cell(r, 2).value),
                    "name_vi": safe_str(ws.cell(r, 3).value),
                    "fullname_cn": safe_str(ws.cell(r, 4).value),
                    "fullname_vi": safe_str(ws.cell(r, 5).value),
                    "nickname_vi": safe_str(ws.cell(r, 6).value),
                    "tags_cn": safe_str(ws.cell(r, 7).value),
                    "tags_vi": safe_str(ws.cell(r, 8).value),
                    "status": safe_str(ws.cell(r, 11).value)
                }

    # 3. Export SKILL
    skills = {}
    if "SKILL" in wb.sheetnames:
        ws = wb["SKILL"]
        for r in range(2, ws.max_row + 1):
            sk_id = safe_str(ws.cell(r, 1).value)
            if sk_id:
                skills[sk_id] = {
                    "character_id": safe_str(ws.cell(r, 2).value),
                    "group_id": safe_str(ws.cell(r, 3).value),
                    "skill_name_cn": safe_str(ws.cell(r, 6).value),
                    "skill_name_vi": safe_str(ws.cell(r, 7).value),
                    "desc_cn": safe_str(ws.cell(r, 8).value),
                    "desc_vi": safe_str(ws.cell(r, 9).value),
                    "status": safe_str(ws.cell(r, 11).value)
                }

    # 4. Export BUFF_STATUS
    buffs = {}
    if "BUFF_STATUS" in wb.sheetnames:
        ws = wb["BUFF_STATUS"]
        for r in range(2, ws.max_row + 1):
            b_id = safe_str(ws.cell(r, 1).value)
            if b_id:
                buffs[b_id] = {
                    "group_root_id": safe_str(ws.cell(r, 2).value),
                    "buff_name_cn": safe_str(ws.cell(r, 3).value),
                    "buff_name_vi": safe_str(ws.cell(r, 4).value),
                    "buff_desc_cn": safe_str(ws.cell(r, 5).value),
                    "buff_desc_vi": safe_str(ws.cell(r, 6).value),
                    "classification_scope": safe_str(ws.cell(r, 7).value),
                    "status": safe_str(ws.cell(r, 9).value)
                }

    # 5. Export TALENT
    talents = {}
    if "TALENT" in wb.sheetnames:
        ws = wb["TALENT"]
        for r in range(2, ws.max_row + 1):
            bank_id = safe_str(ws.cell(r, 1).value)
            name_cn = safe_str(ws.cell(r, 2).value)
            entry = {
                "bank_id": bank_id,
                "name_cn": name_cn,
                "name_vi": safe_str(ws.cell(r, 3).value),
                "desc_cn": safe_str(ws.cell(r, 4).value),
                "desc_vi": safe_str(ws.cell(r, 5).value),
                "status": safe_str(ws.cell(r, 8).value)
            }
            if bank_id:
                talents[bank_id] = entry
            if name_cn:
                talents[name_cn] = entry

    # 6. Export ZHIZHI
    zhizhi = {}
    if "ZHIZHI" in wb.sheetnames:
        ws = wb["ZHIZHI"]
        for r in range(2, ws.max_row + 1):
            cid = safe_str(ws.cell(r, 1).value)
            star = ws.cell(r, 2).value
            if cid and star:
                key = f"{cid}_{star}"
                zhizhi[key] = {
                    "star": star,
                    "rank_numeral_cn": safe_str(ws.cell(r, 3).value),
                    "rank_numeral_vi": safe_str(ws.cell(r, 4).value),
                    "effect_type": safe_str(ws.cell(r, 5).value),
                    "effect_summary_cn": safe_str(ws.cell(r, 6).value),
                    "effect_summary_vi": safe_str(ws.cell(r, 7).value),
                    "status": safe_str(ws.cell(r, 11).value)
                }

    # 7. Export HUANZHANG
    huanzhang = {}
    if "HUANZHANG" in wb.sheetnames:
        ws = wb["HUANZHANG"]
        for r in range(2, ws.max_row + 1):
            b_id = safe_str(ws.cell(r, 1).value)
            if b_id:
                huanzhang[b_id] = {
                    "character_id": safe_str(ws.cell(r, 2).value),
                    "icon_name_cn": safe_str(ws.cell(r, 3).value),
                    "icon_name_vi": safe_str(ws.cell(r, 4).value),
                    "icon_info_cn": safe_str(ws.cell(r, 5).value),
                    "icon_info_vi": safe_str(ws.cell(r, 6).value),
                    "buff_show_cn": safe_str(ws.cell(r, 7).value),
                    "buff_show_vi": safe_str(ws.cell(r, 8).value),
                    "status": safe_str(ws.cell(r, 10).value)
                }

    # 8. Export ITEM
    items = {}
    if "ITEM" in wb.sheetnames:
        ws = wb["ITEM"]
        for r in range(2, ws.max_row + 1):
            iid = safe_str(ws.cell(r, 1).value)
            if iid:
                items[iid] = {
                    "category": safe_str(ws.cell(r, 2).value),
                    "name_cn": safe_str(ws.cell(r, 3).value),
                    "name_vi": safe_str(ws.cell(r, 4).value),
                    "desc_cn": safe_str(ws.cell(r, 5).value),
                    "desc_vi": safe_str(ws.cell(r, 6).value),
                    "status": safe_str(ws.cell(r, 8).value)
                }

    # 9. Export PROFILE
    profiles = {}
    if "PROFILE" in wb.sheetnames:
        ws = wb["PROFILE"]
        for r in range(2, ws.max_row + 1):
            pid = safe_str(ws.cell(r, 1).value)
            if pid:
                profiles[pid] = {
                    "character_id": safe_str(ws.cell(r, 2).value),
                    "category": safe_str(ws.cell(r, 3).value),
                    "text_cn": safe_str(ws.cell(r, 6).value),
                    "text_vi": safe_str(ws.cell(r, 7).value),
                    "status": safe_str(ws.cell(r, 9).value)
                }

    # 10. Export SKIN
    skins = {}
    if "SKIN" in wb.sheetnames:
        ws = wb["SKIN"]
        for r in range(2, ws.max_row + 1):
            sid = safe_str(ws.cell(r, 1).value)
            if sid:
                skins[sid] = {
                    "character_id": safe_str(ws.cell(r, 2).value),
                    "skin_name_cn": safe_str(ws.cell(r, 3).value),
                    "skin_name_vi": safe_str(ws.cell(r, 4).value),
                    "desc_cn": safe_str(ws.cell(r, 5).value),
                    "desc_vi": safe_str(ws.cell(r, 6).value),
                    "is_base_skin": safe_str(ws.cell(r, 9).value).lower() == "true",
                    "status": safe_str(ws.cell(r, 11).value)
                }

    # 11. Export UI_SYSTEM
    ui_system = {}
    if "UI_SYSTEM" in wb.sheetnames:
        ws = wb["UI_SYSTEM"]
        for r in range(2, ws.max_row + 1):
            ui_key = safe_str(ws.cell(r, 1).value)
            if ui_key:
                ui_system[ui_key] = {
                    "category": safe_str(ws.cell(r, 2).value),
                    "text_cn": safe_str(ws.cell(r, 3).value),
                    "text_vi": safe_str(ws.cell(r, 4).value),
                    "status": safe_str(ws.cell(r, 7).value)
                }

    out_data = {
        "glossary": glossary,
        "characters": characters,
        "skills": skills,
        "buffs": buffs,
        "talents": talents,
        "zhizhi": zhizhi,
        "huanzhang": huanzhang,
        "items": items,
        "profiles": profiles,
        "skins": skins,
        "ui_system": ui_system
    }

    OUT_JSON.write_text(json.dumps(out_data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[OK] Exported master localization to {OUT_JSON}")
    return out_data

if __name__ == "__main__":
    export()
