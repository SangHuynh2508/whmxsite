"""Validate locked BUFF/STATE terminology in translated SKILL rows."""

from __future__ import annotations

import argparse
from pathlib import Path
import sys

import openpyxl

from term_lock import TermLockRegistry, validate_locked_terms

VALIDATED_TRANSLATION_STATUSES = {"TRANSLATED", "REVIEW", "APPROVED", "OWNER_APPROVED", "OWNER_CORRECTED"}


def validate(workbook_path: Path, buff_ids: set[str] | None = None) -> list[str]:
    registry = TermLockRegistry.from_workbook(workbook_path)
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        ws = wb["SKILL"]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        errors: list[str] = []
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            item = dict(zip(headers, row))
            source = str(item.get("desc_cn") or "")
            target = item.get("desc_vi")
            # PENDING content is frequently a deliberately partial historical
            # draft.  It is not a completed AI batch and must not turn this
            # validator into a broad residual-translation audit.
            if not source or not target or str(item.get("status") or "") not in VALIDATED_TRANSLATION_STATUSES:
                continue
            for issue in validate_locked_terms(source, str(target), registry, protected_phase=False):
                if buff_ids is not None and issue.buff_id not in buff_ids:
                    continue
                errors.append(f"{issue.code} SKILL!{row_no} [{item.get('skill_id')}]: {issue.message}")
        return errors
    finally:
        wb.close()


def main() -> int:
    if getattr(sys.stdout, "encoding", "").lower() != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=Path("localization/localization_master.xlsx"))
    parser.add_argument("--buff-id", action="append", default=[], help="Validate only these exact locked identities")
    args = parser.parse_args()
    errors = validate(args.workbook, set(args.buff_id) or None)
    if errors:
        print("TERM_LOCK_VALIDATION_FAILED")
        print("\n".join(errors))
        return 1
    print("TERM_LOCK_VALIDATION_OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
